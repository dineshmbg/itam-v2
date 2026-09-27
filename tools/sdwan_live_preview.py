"""READ-ONLY preview: reconcile the SD-WAN "Network Inventry" sheet with the live asset register.

  python tools/sdwan_live_preview.py [--sdwan "S D Wan Cisco Detail.xlsx"] [--out samples/SDWAN_Live_Preview_<date>.xlsx]

Reads the sheet and the database; writes ONE workbook and nothing else - the database is opened read-only. The sheet's Login / Password
columns are never read into the output (device credentials do not belong in an inventory system).

Sheets in the preview
  SUMMARY          what was matched, what would change, and cautions
  PROPOSED_FILLS   blank fields in the register that the sheet can fill (hostname, IP address, firmware) - APPROVE column for you to mark
  CI_MISMATCH      sheet rows whose CI number belongs to one register asset but whose serial / IP / hostname belong to another
  CONFLICTS        both sides have a value and they differ - never changed automatically, for you to decide
  SWITCH_MATCH     every sheet row and the register asset it was matched to (and how)
  UNMATCHED_SHEET  sheet rows that match no asset in the register
  NOT_IN_SHEET     switches / routers in the register that the sheet does not list
  SFP_IN_SHEET     every SFP serial in the sheet, with the switch it sits in, and whether the register already has it
  SFP_RECONCILE    SFP lines the register holds (AMC billing lines, no serial) against SFPs the sheet actually shows
"""
import argparse
import datetime as dt
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
import master_db  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
SHEET = "Network Inventry"
IPV4 = re.compile(r"^(\d{1,3}\.){3}\d{1,3}$")


def s(v):
    if v is None or (not isinstance(v, str) and pd.isna(v)):
        return None
    t = re.sub(r"\s+", " ", str(v).replace("\xa0", " ")).strip()
    return t or None


def loose(v):
    t = s(v)
    return re.sub(r"[^A-Z0-9]", "", t.upper()) if t else None


def ip_only(v):
    t = s(v)
    if not t:
        return None
    t = t.split("/")[0].strip()
    return t if IPV4.match(t) else None


def sfp_serial(v):
    t = s(v)
    return None if (not t or t.upper() in ("NO", "N/A", "NA", "-")) else t.upper()


def version_text(v):
    t = s(v)
    if not t:
        return None
    return re.sub(r"^version\s+", "", t, flags=re.I).upper()


def load_sheet(path):
    d = pd.read_excel(path, sheet_name=SHEET, dtype=object, usecols=lambda c: str(c).strip().lower() not in ("login", "pasword", "password"))
    rows = []
    for i, r in d.iterrows():
        sfps = []
        for n, (sc, tc) in enumerate([("SFP-1 S/N", "Type MM/SM"), ("SFP-2 S/N", "Type MM/SM.1"), ("SFP-3 S/N", "Type MM/SM.2"), ("SFP-4  S/N", "Type MM/SM.3")], 1):
            sn = sfp_serial(r.get(sc))
            if sn:
                sfps.append((n, sn, s(r.get(tc))))
        rows.append({"ROW": i + 2, "MAKE": (s(r["Make"]) or "").upper(), "MODEL": s(r["Model"]), "TYPE": s(r["Type"]), "LOCATION": s(r["Locatation"]),
                     "HOSTNAME": s(r["Host Name"]), "CI": s(r["ASSAT ID"]), "SERIAL": s(r["Sr.NO"]), "IP": ip_only(r["SWITECH IP"]), "IP_RAW": s(r["SWITECH IP"]),
                     "VERSION": version_text(r["Version"]), "SHEET_REMARK": s(r.get("Unnamed: 31")), "SFPS": sfps})
    return rows


def load_db():
    con = master_db.connect()
    con.execute("SET default_transaction_read_only = on")
    cur = con.cursor()
    cur.execute("""SELECT asset_key, asset_class, asset_type, record_level, make, model, serial_no, hostname, ip_address, firmware_version, location_code, room, engineer_name, parent_asset_key
                   FROM asset WHERE is_current = 1""")
    cols = [c.name for c in cur.description]
    return [dict(zip(cols, r)) for r in cur.fetchall()]


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--sdwan", default=str(ROOT / "S D Wan Cisco Detail.xlsx"))
    ap.add_argument("--out", default=str(ROOT / "samples" / f"SDWAN_Live_Preview_{dt.date.today():%Y-%m-%d}.xlsx"))
    a = ap.parse_args()

    sd = load_sheet(a.sdwan)
    db_rows = load_db()
    assets = [r for r in db_rows if r["record_level"] == "ASSET"]
    by_key = {r["asset_key"]: r for r in db_rows}
    by_serial, by_ip, by_host = defaultdict(list), defaultdict(list), defaultdict(list)
    for r in assets:
        for idx, k in ((by_serial, loose(r["serial_no"])), (by_ip, r["ip_address"]), (by_host, loose(r["hostname"]))):
            if k:
                idx[k].append(r)

    match_rows, fills, conflicts, unmatched, used, mismatch = [], [], [], [], set(), []
    for x in sd:
        cand = []          # (method, asset)
        if x["CI"] and x["CI"].upper() in by_key and by_key[x["CI"].upper()]["record_level"] == "ASSET":
            cand.append(("CI", by_key[x["CI"].upper()]))
        for method, idx, k in (("SERIAL", by_serial, loose(x["SERIAL"])), ("IP", by_ip, x["IP"]), ("HOSTNAME", by_host, loose(x["HOSTNAME"]))):
            if k and len(idx.get(k, [])) == 1:
                cand.append((method, idx[k][0]))
        keys = {c[1]["asset_key"] for c in cand}
        if not cand:
            unmatched.append({"SHEET_ROW": x["ROW"], "MAKE": x["MAKE"], "MODEL": x["MODEL"], "TYPE": x["TYPE"], "HOSTNAME": x["HOSTNAME"], "CI IN SHEET": x["CI"], "SERIAL": x["SERIAL"], "IP": x["IP_RAW"],
                              "LOCATION": x["LOCATION"], "WHY": "no CI, serial, IP or hostname in the register matches" if not x["CI"] else "CI in sheet is not in the register"})
            continue
        if len(keys) > 1:
            status, asset, how = "AMBIGUOUS", cand[0][1], "+".join(m for m, _ in cand) + " point to different assets: " + ", ".join(sorted(keys))
            found = {m: a_["asset_key"] for m, a_ in cand}
            others = [k for m, k in found.items() if m != "CI"]
            agree = Counter(others).most_common(1)[0] if others else (None, 0)
            mismatch.append({"SHEET_ROW": x["ROW"], "CI IN SHEET": x["CI"], "ASSET FOUND BY CI": found.get("CI"), "BY SERIAL": found.get("SERIAL"), "BY IP": found.get("IP"), "BY HOSTNAME": found.get("HOSTNAME"),
                             "SHEET HOSTNAME": x["HOSTNAME"], "SHEET SERIAL": x["SERIAL"], "SHEET IP": x["IP_RAW"], "SHEET LOCATION": x["LOCATION"],
                             "READING": (f"{agree[1]} of {len(others)} other identifiers (serial / IP / hostname) agree on {agree[0]}, not on {found.get('CI')} - the CI number is probably wrong in the sheet or in the register"
                                         if agree[1] and agree[0] != found.get("CI") else "identifiers disagree with each other - check the device"),
                             "ACTION": "verify on site (Verify button), then correct whichever source is wrong"})
        else:
            asset = cand[0][1]
            how = "+".join(m for m, _ in cand)
            status = "OK"
        used.add(asset["asset_key"])
        row = {"SHEET_ROW": x["ROW"], "MATCH_STATUS": status, "MATCHED_BY": how, "ASSET_KEY": asset["asset_key"], "CI IN SHEET": x["CI"],
               "SHEET_MAKE_MODEL": f'{x["MAKE"]} {x["MODEL"] or ""}'.strip(), "REGISTER_MAKE_MODEL": f'{asset["make"] or ""} {asset["model"] or ""}'.strip(), "SHEET_TYPE": x["TYPE"], "REGISTER_TYPE": asset["asset_type"],
               "SHEET_LOCATION": x["LOCATION"], "REGISTER_LOCATION": asset["location_code"]}
        proposals = []
        for field, sheet_val, reg_val, norm in (("hostname", x["HOSTNAME"], asset["hostname"], loose), ("ip_address", x["IP"], asset["ip_address"], lambda v: v),
                                                ("serial_no", x["SERIAL"], asset["serial_no"], loose), ("firmware_version", x["VERSION"], asset["firmware_version"], lambda v: re.sub(r"^VERSION", "", loose(v) or ""))):
            row[f"SHEET_{field.upper()}"], row[f"REGISTER_{field.upper()}"] = sheet_val, reg_val
            if not sheet_val or status != "OK":
                continue
            if not reg_val:
                if status == "OK" and field != "serial_no":
                    fills.append({"APPROVE (type YES)": "", "ASSET_KEY": asset["asset_key"], "FIELD": field, "CURRENT (blank)": None, "NEW VALUE": sheet_val, "MATCHED_BY": how, "SHEET_ROW": x["ROW"],
                                  "NOTE": "" if "CI" in how else "matched without a CI - check it is the same device"})
                    proposals.append(f"fill {field}")
            elif norm(sheet_val) != norm(reg_val):
                conflicts.append({"ASSET_KEY": asset["asset_key"], "FIELD": field, "REGISTER VALUE": reg_val, "SHEET VALUE": sheet_val, "MATCHED_BY": how, "SHEET_ROW": x["ROW"], "SHEET_REMARK": x["SHEET_REMARK"]})
                proposals.append(f"CHECK {field}")
        row["PROPOSAL"] = "; ".join(proposals) or ("none - already consistent" if status == "OK" else "")
        match_rows.append(row)

    not_in = [{"ASSET_KEY": r["asset_key"], "TYPE": r["asset_type"], "MAKE_MODEL": f'{r["make"] or ""} {r["model"] or ""}'.strip(), "HOSTNAME": r["hostname"], "IP": r["ip_address"],
               "LOCATION": r["location_code"], "SERIAL": r["serial_no"]}
              for r in sorted(assets, key=lambda r: r["asset_key"]) if r["asset_type"] in ("L2 SWITCH", "L3 SWITCH", "ROUTER") and r["asset_key"] not in used]

    # SFPs
    reg_serials = {loose(r["serial_no"]): r for r in db_rows if loose(r["serial_no"])}
    sfp_rows = []
    for x in sd:
        parent = next((m["ASSET_KEY"] for m in match_rows if m["SHEET_ROW"] == x["ROW"] and m["MATCH_STATUS"] == "OK"), None)
        for n, sn, typ in x["SFPS"]:
            have = reg_serials.get(loose(sn))
            sfp_rows.append({"SHEET_ROW": x["ROW"], "SWITCH_CI (matched)": parent, "SWITCH_HOSTNAME": x["HOSTNAME"], "SLOT": f"SFP-{n}", "SFP_SERIAL": sn, "SFP_TYPE": typ,
                             "IN_REGISTER": "yes - " + have["asset_key"] if have else "no", "PROPOSAL": "already recorded" if have else ("add as a component of " + parent if parent else "switch not matched - resolve first")})
    dup_sn = [k for k, v in Counter(loose(r["SFP_SERIAL"]) for r in sfp_rows).items() if v > 1]
    reg_sfp = Counter((r["make"] or "", r["model"] or "") for r in db_rows if r["asset_type"] == "SFP")
    sheet_types = Counter(r["SFP_TYPE"] or "?" for r in sfp_rows)
    recon = [{"SOURCE": "register (AMC billing lines, no serial)", "MAKE": k[0], "MODEL / TYPE": k[1], "COUNT": v} for k, v in reg_sfp.most_common()]
    recon += [{"SOURCE": "sheet (real serials)", "MAKE": "", "MODEL / TYPE": k, "COUNT": v} for k, v in sheet_types.most_common()]

    # cautions
    dup_ci = [k for k, v in Counter((x["CI"] or "").upper() for x in sd if x["CI"]).items() if v > 1]
    dup_ip = [k for k, v in Counter(x["IP"] for x in sd if x["IP"]).items() if v > 1]
    summary = [
        ("Sheet", f"{Path(a.sdwan).name} / {SHEET}: {len(sd)} device rows"),
        ("Register", f"{len(assets):,} current assets; {sum(1 for r in assets if r['asset_type'] in ('L2 SWITCH','L3 SWITCH','ROUTER'))} switches/routers"),
        ("Matched to a register asset", f"{len(match_rows)} of {len(sd)} ({sum(1 for m in match_rows if m['MATCH_STATUS']=='OK')} clean, {sum(1 for m in match_rows if m['MATCH_STATUS']=='AMBIGUOUS')} ambiguous)"),
        ("Sheet rows with no match", len(unmatched)),
        ("Register switches/routers not in the sheet", len(not_in)),
        ("Proposed fills (blank -> value)", f"{len(fills)}  " + ", ".join(f"{k}: {v}" for k, v in Counter(f['FIELD'] for f in fills).items())),
        ("CI numbers that do not agree with the device", f"{len(mismatch)} sheet rows - see CI_MISMATCH (the most important finding)"),
        ("Conflicts to decide (never automatic)", f"{len(conflicts)}  " + ", ".join(f"{k}: {v}" for k, v in Counter(c['FIELD'] for c in conflicts).items())),
        ("SFP serials in the sheet", f"{len(sfp_rows)} on {sum(1 for x in sd if x['SFPS'])} switches; already in register: {sum(1 for r in sfp_rows if r['IN_REGISTER'] != 'no')}; duplicated serials in sheet: {len(dup_sn)}"),
        ("Duplicates inside the sheet", f"CI used twice: {dup_ci or 'none'}; IP used twice: {dup_ip or 'none'}"),
        ("", ""),
        ("IMPORTANT - the 814 unlinked 'components'", "They are AMC billing lines from the inventory's SWITCH sheet (make, model, rate, cover) with NO serial, location or parent, so an individual line cannot be tied to a switch. "
         "Only real SFPs (with serials) can be linked. The sheet shows the ones below; the difference to the register's count is in SFP_RECONCILE."),
        ("Credentials", "The sheet's Login and Password columns were NOT read. Passwords kept in a spreadsheet should be moved to a password vault."),
        ("Nothing was changed", "This workbook is a preview. The database was opened read-only. Mark PROPOSED_FILLS rows with YES and tell me to apply them."),
    ]

    Path(a.out).parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(a.out, engine="openpyxl") as xw:
        pd.DataFrame(summary, columns=["ITEM", "RESULT"]).to_excel(xw, sheet_name="SUMMARY", index=False)
        for name, data in (("PROPOSED_FILLS", fills), ("CI_MISMATCH", mismatch), ("CONFLICTS", conflicts), ("SWITCH_MATCH", match_rows), ("UNMATCHED_SHEET", unmatched), ("NOT_IN_SHEET", not_in), ("SFP_IN_SHEET", sfp_rows), ("SFP_RECONCILE", recon)):
            pd.DataFrame(data).to_excel(xw, sheet_name=name, index=False)
        from openpyxl.styles import Font, PatternFill
        for ws in xw.book.worksheets:
            ws.freeze_panes = "A2"
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1F3A5F")
            for col in ws.columns:
                width = max((len(str(c.value)) for c in col[:200] if c.value is not None), default=8)
                ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 70)
    print(f"Preview written: {a.out}")
    for k, v in summary[:9]:
        print(f"  {k}: {v}")


if __name__ == "__main__":
    main()
