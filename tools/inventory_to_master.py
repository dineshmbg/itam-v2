"""Convert the multi-sheet IT-IMMDSS asset inventory workbook into ONE flat ASSET_MASTER sheet.

  python tools/inventory_to_master.py --raw "IT-IMMDSS ASSET INVENTORY 2026 -Q2.xlsx" [--as-of 2026-09-21] [--out-dir masters] [--no-load-db]

Reads cached values. Resolves user details from the PostgreSQL employee table when reachable.
"""
import argparse
import collections
import datetime as dt
import os
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

MASTER = "ASSET_MASTER"
LOG = collections.Counter()

SHEET_CLASS = {"DESKTOP": "DESKTOP", "LAPTOP": "LAPTOP", "WORK STATION": "WORKSTATION", "SERVER": "SERVER",
               "ROUTER": "ROUTER", "SWITCH": "SWITCH", "PRINTER": "PRINTER", "SCANNER": "SCANNER",
               "UPS": "UPS", "MEDIA CONVERTOR": "MEDIA_CONVERTER"}
CLASS_ABBR = {"DESKTOP": "DT", "LAPTOP": "LT", "WORKSTATION": "WS", "SERVER": "SR", "ROUTER": "RT", "SWITCH": "SW",
              "PRINTER": "PR", "SCANNER": "SC", "UPS": "UP", "MEDIA_CONVERTER": "MC"}

# (target column, group, description)
TARGETS = [
    ("SNAPSHOT_DATE", "TRACE", "Date this snapshot was built (as-of date for cover / PM status)."),
    ("ASSET_KEY", "IDENTITY", "UNIQUE key. = CI_NO when present; otherwise generated (see DQ flag GENERATED_KEY)."),
    ("CI_NO", "IDENTITY", "Configuration Item number / asset tag, e.g. ANKAON00DT189."),
    ("ASSET_CLASS", "IDENTITY", "Asset class = source sheet (DESKTOP, LAPTOP, WORKSTATION, SERVER, ROUTER, SWITCH, PRINTER, SCANNER, UPS, MEDIA_CONVERTER)."),
    ("ASSET_TYPE", "IDENTITY", "H/W type within class (ALL IN ONE, OFFICE LAPTOP, L2 SWITCH, SFP, MONITOR, VM SERVER ...). OFFICE LAPTOP = laptops kept in the DESKTOP class (under AMC + PM)."),
    ("RECORD_LEVEL", "IDENTITY", "ASSET = physical/logical item with its own CI. COMPONENT = child line (monitor, GPU, VM, SFP, fibre module)."),
    ("PARENT_ASSET_KEY", "IDENTITY", "ASSET_KEY of the parent (monitor/GPU -> workstation, VM -> server). Blank for asset rows and for un-linked SFP/fibre lines."),
    ("HOSTNAME", "HARDWARE", "Host / device name."),
    ("ASSET_DESCRIPTION", "HARDWARE", "Server role / purpose (source 'Serve Name')."),
    ("MAKE", "HARDWARE", "Manufacturer, normalised (HP, DELL, ASUS, APPLE ...)."),
    ("MODEL", "HARDWARE", "Model."),
    ("SERIAL_NO", "HARDWARE", "Manufacturer serial number (PC / printer / UPS / scanner / server / router / switch / media-converter Sr No.)."),
    ("ONGC_ASSET_ID", "HARDWARE", "ONGC asset id. Some ids are bulk ids shared by many devices (flag BULK_ASSET_ID)."),
    ("ONGC_CENSUS_NO", "HARDWARE", "ONGC census number."),
    ("SUB_TYPE", "HARDWARE", "Printer type (MFP/mono), scanner type (ADF/A4), UPS type, media-converter mode."),
    ("OS", "HARDWARE", "Operating system as recorded."),
    ("OS_FAMILY", "HARDWARE", "Derived grouping: WINDOWS 11 / WINDOWS 10 / WINDOWS SERVER / MACOS / LINUX / OTHER."),
    ("FIRMWARE_VERSION", "HARDWARE", "JUNOS / IOS version for routers and switches."),
    ("IP_ADDRESS", "HARDWARE", "IP address."),
    ("MGMT_IP", "HARDWARE", "iDRAC / management IP (servers)."),
    ("PROCESSOR", "HARDWARE", "Processor."),
    ("RAM", "HARDWARE", "RAM."),
    ("STORAGE", "HARDWARE", "Disk / storage (source 'HDD')."),
    ("ATTACHED_MONITOR_MODEL", "HARDWARE", "Monitor bundled on a desktop row (source 'Monitor Type')."),
    ("ATTACHED_MONITOR_SN", "HARDWARE", "Serial of the bundled monitor."),
    ("PORTS", "HARDWARE", "Port count (routers/switches)."),
    ("NETWORK_PORT_AVAILABLE", "HARDWARE", "Printer network port availability (YES/NO)."),
    ("CAPACITY", "HARDWARE", "UPS capacity."),
    ("BATTERY_QTY", "HARDWARE", "UPS battery quantity."),
    ("BATTERY_SPEC", "HARDWARE", "UPS battery spec (unlabelled column in source, e.g. 12VA 7AH)."),
    ("INSTALL_DATE", "HARDWARE", "Server install date (unlabelled column in source - assumed)."),
    ("CPF_NO", "OWNERSHIP", "CPF of the ONGC user / owner the asset is assigned to."),
    ("USER_NAME", "OWNERSHIP", "Resolved from the HR master (latest); falls back to source value."),
    ("USER_DESIGNATION", "OWNERSHIP", "Resolved from HR master."),
    ("USER_LEVEL", "OWNERSHIP", "Resolved from HR master."),
    ("USER_MOBILE", "OWNERSHIP", "Resolved from HR master."),
    ("USER_DEPARTMENT", "OWNERSHIP", "Team / department as recorded in the source (not in the HR export)."),
    ("USER_RETIREMENT_DATE", "OWNERSHIP", "Retirement date of the user (HR master). Source laptop column 'RETIRED DATE' held this."),
    ("USER_HR_STATUS", "OWNERSHIP", "ACTIVE / RETIRED / EXITED / NOT_IN_HR_MASTER / NO_CPF."),
    ("USER_CATEGORY", "OWNERSHIP", "Derived from remarks: APPRENTICE / CONTINGENT / AUDITOR / COMMON_USE / DATA_ENTRY."),
    ("ENGINEER_NAME", "LOCATION", "CIPL engineer responsible for the asset."),
    ("LOCATION_CODE", "LOCATION", "Location / network-zone code, e.g. ANK_NEWBDGGF_A."),
    ("FLOOR_AREA", "LOCATION", "Floor or area (typos fixed, see CLEANING_LOG)."),
    ("ROOM", "LOCATION", "Room / office address."),
    ("COVER_TYPE", "CONTRACT", "AMC or WRTY."),
    ("COVER_EXPIRY_DATE", "CONTRACT", "AMC / warranty expiry."),
    ("COVER_STATUS", "CONTRACT", "Derived: ACTIVE / EXPIRING_90D / EXPIRED / REMOVED (removed from AMC) / UNKNOWN."),
    ("RATE_COMPONENT", "CONTRACT", "Rate-card code (see RATE_CARD sheet)."),
    ("RATE_VALUE", "CONTRACT", "Unit rate for the rate component (2 decimals)."),
    ("PM_QUARTER", "PM", "Preventive-maintenance quarter parsed from the source header."),
    ("PM_DATE", "PM", "PM date this quarter."),
    ("PM_STATUS", "PM", "DONE / DONE_OUTSIDE_QUARTER / PENDING / NA (component) / NOT_TRACKED (laptops)."),
    ("PM_DONE_BY", "PM", "Engineer who performed PM."),
    ("PM_SIGNED_BY", "PM", "User who signed the PM."),
    ("PM_TRACKER_DATE", "PM", "PM date looked up from the external PM tracker by CI NO (unlabelled column in source). Equals PM_DATE wherever present; blank = PM not in tracker."),
    ("ASSET_STATUS", "STATUS", "Derived from remarks/location: IN_USE / IN_STORE / SURPLUS / NOT_IN_USE / NOT_ON_NETWORK / STANDBY / REMOVED_FROM_AMC / TRANSFERRED."),
    ("REMARKS", "STATUS", "Original CIPL remark, kept verbatim."),
    ("EXTRA_INFO", "STATUS", "Unlabelled free-text / email column in the desktop sheet, and stray values moved out of the wrong field."),
    ("DQ_FLAGS", "TRACE", "Semicolon-separated data-quality flags for this row."),
    ("SOURCE_SHEET", "TRACE", "Sheet in the source workbook."),
    ("SOURCE_ROW", "TRACE", "Row number in the source sheet."),
]
COLS = [t[0] for t in TARGETS]
GROUP_FILL = {"IDENTITY": "1F3864", "HARDWARE": "2E75B6", "OWNERSHIP": "548235", "LOCATION": "7F6000",
              "CONTRACT": "C55A11", "PM": "7030A0", "STATUS": "C00000", "TRACE": "595959"}

HEADER_MAP = {
    "SR NO": "SR", "CI NO": "CI_NO", "CPF NO.": "CPF", "CPF NO": "CPF", "LEVEL": "U_LEVEL", "USER NAME": "U_NAME",
    "MOBILE NO": "U_MOBILE", "DESIGNATION": "U_DESIG", "DEPARTMENT": "U_DEPT", "ENGINNER NAME": "ENGINEER",
    "ENGINEER NAME": "ENGINEER", "OFFICE ADDRESS": "ROOM", "LOCATION": "LOCATION", "FLOOR": "FLOOR",
    "H/W TYPE": "HW_TYPE", "SCANNER H/W TYPE": "HW_TYPE", "MAKE": "MAKE", "MODEL": "MODEL",
    "OPERATING SYSTEM": "OS", "O.S.": "OS", "ONGC ASSET ID": "ASSET_ID", "ONGC CENSUS NO.": "CENSUS",
    "IP ADDRESS": "IP", "PROCESSOR": "PROCESSOR", "RAM TYPE/SIZE": "RAM", "RAM SIZE": "RAM", "HDD": "STORAGE",
    "MONITOR TYPE": "MON_MODEL", "MONITOR S/N": "MON_SN", "AMC/WRTY": "COVER", "EXPIRED DATE": "COVER_EXP",
    "WRTY EXPIERED DATE": "COVER_EXP", "RATE COMPONENT": "RATE_COMP", "RATE VALUE": "RATE_VAL",
    "PM SIGNED BY": "PM_SIGNED", "PM DONE BY": "PM_DONE", "CIPL REMARKS": "REMARKS", "REMARKS": "REMARKS",
    "HOSTNAME": "HOSTNAME", "PORTS": "PORTS", "JUNOS/IOS DETAILS": "FIRMWARE",
    "NETWORK PORT AVAILABILITY": "NET_PORT", "CAPACITY": "CAPACITY", "BATTERY QTY": "BATT_QTY",
    "SERVE NAME": "DESCRIPTION", "POS NO": "POS_NO", "IDRACK IP": "MGMT_IP", "RETIRED DATE": "U_RETIRE",
}
SERIAL_RE = re.compile(r"^(PC|PRINTER|UPS|SCANNER|SERVER|ROUTER|SWITCH|MEDIA CONVERTOR) SR\.? ?NO\.?$")
KEY_TO_TARGET = {"CI_NO": "CI_NO", "HW_TYPE": "ASSET_TYPE", "DESCRIPTION": "ASSET_DESCRIPTION", "MAKE": "MAKE", "MODEL": "MODEL",
                 "SERIAL": "SERIAL_NO", "ASSET_ID": "ONGC_ASSET_ID", "CENSUS": "ONGC_CENSUS_NO", "SUB_TYPE": "SUB_TYPE",
                 "OS": "OS", "FIRMWARE": "FIRMWARE_VERSION", "IP": "IP_ADDRESS", "MGMT_IP": "MGMT_IP", "HOSTNAME": "HOSTNAME",
                 "PROCESSOR": "PROCESSOR", "RAM": "RAM", "STORAGE": "STORAGE", "MON_MODEL": "ATTACHED_MONITOR_MODEL",
                 "MON_SN": "ATTACHED_MONITOR_SN", "PORTS": "PORTS", "NET_PORT": "NETWORK_PORT_AVAILABLE", "CAPACITY": "CAPACITY",
                 "BATT_QTY": "BATTERY_QTY", "BATT_SPEC": "BATTERY_SPEC", "INSTALL_DATE": "INSTALL_DATE", "CPF": "CPF_NO",
                 "U_NAME": "USER_NAME (HR master first)", "U_DESIG": "USER_DESIGNATION (HR master first)",
                 "U_LEVEL": "USER_LEVEL (HR master first)", "U_MOBILE": "USER_MOBILE (HR master first)", "U_DEPT": "USER_DEPARTMENT",
                 "U_RETIRE": "USER_RETIREMENT_DATE (HR master first)", "ENGINEER": "ENGINEER_NAME", "LOCATION": "LOCATION_CODE",
                 "FLOOR": "FLOOR_AREA", "ROOM": "ROOM", "COVER": "COVER_TYPE", "COVER_EXP": "COVER_EXPIRY_DATE",
                 "RATE_COMP": "RATE_COMPONENT", "RATE_VAL": "RATE_VALUE", "PM_DATE": "PM_DATE", "PM_DONE": "PM_DONE_BY",
                 "PM_SIGNED": "PM_SIGNED_BY", "PREV_PM": "PM_TRACKER_DATE", "REMARKS": "REMARKS", "EXTRA_INFO": "EXTRA_INFO",
                 "CI_NO_TAG": "CI_NO (asset tag; CI NO column then becomes HOSTNAME)"}

MAKE_ALIAS = {"HEWLETT-PACKARD": "HP", "HP (HP)": "HP", "HP (HC)": "HP", "DELL INC": "DELL", "DELL INC.": "DELL",
              "DELL INC. (DL)": "DELL", "ASUSTEK COMPUTER INC": "ASUS", "ASUSTEK COMPUTER INC.": "ASUS",
              "ASUSTEK COMPUTER INC. (AS)": "ASUS", "ASUS (AS)": "ASUS", "APPLE INC": "APPLE", "APPLE INC.": "APPLE",
              "APPLE INC. (AP)": "APPLE", "LENOVO (LN)": "LENOVO", "MICRO-STAR INTERNATIONAL CO., LTD.": "MSI",
              "MICRO-STAR INTERNATIONAL CO., LTD": "MSI", "MICROSOFT CORPORATION": "MICROSOFT",
              "SAMSUNG ELECTRONICS CO., LTD.": "SAMSUNG", "SAMSUNG ELECTRONICS CO., LTD": "SAMSUNG"}
TYPOS = [("SECOUND", "SECOND"), ("CONTOL", "CONTROL"), ("DISPANSARY", "DISPENSARY"), ("DISPNSRY", "DISPENSARY"),
         ("MAINTAINANCE", "MAINTENANCE"), ("MAINTAINACE", "MAINTENANCE"), ("ELECRICAL", "ELECTRICAL"),
         ("ADMINISTATOR", "ADMINISTRATOR"), ("SINGALMODE", "SINGLEMODE"), ("TEBLET", "TABLET"),
         ("TUBULER", "TUBULAR"), ("DISPANSARY", "DISPENSARY")]
IPV4 = re.compile(r"^(25[0-5]|2[0-4]\d|1?\d?\d)(\.(25[0-5]|2[0-4]\d|1?\d?\d)){3}$")
BLANKS = {"", "N/A", "NA", "NIL", "-", "NONE", "#N/A"}


def hnorm(h):
    return re.sub(r"\s+", " ", str(h)).strip().upper()


def isnull(v):
    if v is None:
        return True
    try:
        return bool(pd.isna(v))
    except (TypeError, ValueError):
        return False


def clean_str(v, upper=False, keep_na=False):
    if isnull(v):
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).replace("\xa0", " ")
    s2 = re.sub(r"\s+", " ", s).strip()
    if s2 != s:
        LOG["Whitespace / line-breaks trimmed"] += 1
    if s2.upper() in BLANKS and not keep_na:
        if s2.upper() == "#N/A":
            LOG["#N/A values blanked"] += 1
        return None
    return s2.upper() if upper else s2


def fix_typos(s):
    if s is None:
        return None
    out = s
    for a, b in TYPOS:
        if a in out.upper():
            out = re.sub(a, b, out, flags=re.I)
    if out != s:
        LOG["Spelling typos corrected (floor / room / type)"] += 1
    return out


def norm_make(v):
    s = clean_str(v)
    if s is None:
        return None
    u = s.upper()
    out = MAKE_ALIAS.get(u, u)
    if out != s:
        LOG["MAKE names normalised"] += 1
    return out


def to_date(v):
    if isnull(v) or isinstance(v, (dt.time, str)) and str(v).strip().upper() in BLANKS:
        return None
    if isinstance(v, dt.time):
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return None if v == 0 else None
    if isinstance(v, pd.Timestamp):
        return v.date()
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    d = pd.to_datetime(v, errors="coerce", dayfirst=not re.match(r"\d{4}-\d{2}-\d{2}", str(v)))
    return None if pd.isna(d) else d.date()


def datefrac(series):
    vals = [v for v in series if not isnull(v) and str(v).strip().upper() != "#N/A"]
    if not vals:
        return 0.0
    return sum(isinstance(v, (dt.datetime, dt.date, dt.time, pd.Timestamp)) for v in vals) / len(vals)


def canon_columns(df):
    hn = [hnorm(c) for c in df.columns]
    has_hw = any(h in ("H/W TYPE", "SCANNER H/W TYPE") for h in hn)
    keys, prev = [], None
    extra_used = False
    for c, h in zip(df.columns, hn):
        if h.startswith("UNNAMED"):
            if prev == "CI_NO":
                k = "CI_NO_TAG"
            elif prev == "BATT_QTY":
                k = "BATT_SPEC"
            elif prev == "PM_DONE":
                k = "INSTALL_DATE"
            elif prev in ("REMARKS", "EXTRA_INFO"):
                k = "PREV_PM" if datefrac(df[c]) >= 0.6 else ("EXTRA_INFO" if not extra_used else "IGNORE")
            else:
                k = "IGNORE"
        elif h.startswith("PM DATE"):
            k = "PM_DATE"
        elif SERIAL_RE.match(h):
            k = "SERIAL"
        elif h == "TYPE":
            k = "SUB_TYPE" if (has_hw or "HW_TYPE" in keys) else "HW_TYPE"
        else:
            k = HEADER_MAP.get(h, "IGNORE")
        if k == "REMARKS" and datefrac(df[c]) >= 0.6:
            k = "PREV_PM"
        if k == "EXTRA_INFO":
            extra_used = True
        keys.append(k)
        prev = k if k != "IGNORE" else prev
    return keys


def read_source(path):
    book = pd.read_excel(path, sheet_name=None, dtype=object)
    unknown = [n for n in book if n not in SHEET_CLASS]
    recs, layouts, skipped = [], {}, []
    pm_header = None
    for sheet, cls in SHEET_CLASS.items():
        if sheet not in book:
            continue
        df = book[sheet]
        keys = canon_columns(df)
        layouts[sheet] = list(zip([str(c) for c in df.columns], [get_column_letter(i + 1) for i in range(len(df.columns))], keys))
        for c in df.columns:
            if hnorm(c).startswith("PM DATE") and pm_header is None:
                pm_header = str(c)
        for i, row in df.iterrows():
            nn = [v for v in row.values if not isnull(v)]
            if len(nn) < 3:
                if nn:
                    skipped.append((sheet, i + 2, str(nn[0])))
                continue
            rec = {}
            for k, v in zip(keys, row.values):
                if k != "IGNORE" and k not in rec:
                    rec[k] = v
            rec["_sheet"], rec["_class"], rec["_row"] = sheet, cls, i + 2
            recs.append(rec)
    return recs, layouts, skipped, unknown, pm_header


def parse_quarter(hdr):
    m = re.search(r"Q(\d)\s+(\w+)'(\d\d)\s+TO\s+(\w+)'(\d\d)", hnorm(hdr or ""))
    if not m:
        return None, None, None
    mon = {"JAN": 1, "FEB": 2, "MAR": 3, "APR": 4, "MAY": 5, "JUN": 6, "JUL": 7, "AUG": 8, "SEP": 9, "OCT": 10, "NOV": 11, "DEC": 12}
    m1, m2 = mon[m[2][:3]], mon[m[4][:3]]
    y1, y2 = 2000 + int(m[3]), 2000 + int(m[5])
    start = dt.date(y1, m1, 1)
    end = (dt.date(y2 + (m2 == 12), (m2 % 12) + 1, 1) - dt.timedelta(days=1))
    return f"Q{m[1]} {m[2][:3]}-{m[4][:3]} {y1}", start, end


def load_hr():
    sys.path.insert(0, str(Path(__file__).parent))
    try:
        import master_db
        con = master_db.connect()
    except SystemExit:
        return {}
    try:
        with con.cursor() as cur:
            cur.execute("SELECT cpf_no, employee_name, designation, level, mobile_no, date_of_retirement, record_status FROM employee")
            return {r[0]: r[1:] for r in cur.fetchall()}
    except Exception:
        return {}
    finally:
        con.close()


def os_family(os_):
    if not os_:
        return None
    u = os_.upper()
    if "SERVER" in u and "WIN" in u:
        return "WINDOWS SERVER"
    if "MAC" in u:
        return "MACOS"
    if "11" in u and "WIN" in u:
        return "WINDOWS 11"
    if "10" in u and "WIN" in u:
        return "WINDOWS 10"
    if "LINUX" in u or "CENT" in u:
        return "LINUX"
    return "OTHER"


def asset_id_clean(v, rec_extra):
    s = clean_str(v)
    if s is None:
        return None
    if re.fullmatch(r"\d+(\.0+)?", s):
        if s.endswith(".0"):
            LOG["ONGC asset ids converted 104051019.0 -> 104051019"] += 1
        return s.split(".")[0]
    rec_extra.append(f"Asset ID field held: {s}")
    LOG["Non-numeric value removed from ONGC_ASSET_ID (kept in EXTRA_INFO)"] += 1
    return "__INVALID__"


def status_from(remark, room, floor, loc):
    r = (remark or "").upper()
    cat = None
    if "APPRENTICE" in r:
        cat = "APPRENTICE"
    elif "CONTIGENT" in r or "CONTINGENT" in r:
        cat = "CONTINGENT"
    elif "AUDIT" in r:
        cat = "AUDITOR"
    elif "COMMON USE" in r:
        cat = "COMMON_USE"
    elif "DATA ENTRY" in r:
        cat = "DATA_ENTRY"
    if "STORE ROOM" in r:
        st = "IN_STORE"
    elif "SUR PLUS" in r or "SURPLUS" in r:
        st = "SURPLUS"
    elif "REMOVE" in r and "AMC" in r:
        st = "REMOVED_FROM_AMC"
    elif "NOT IN ONGC NETWORK" in r:
        st = "NOT_ON_NETWORK"
    elif "NOT IN USE" in r or "NOT USED" in r or "NOT IN SITES" in r:
        st = "NOT_IN_USE"
    elif "STAND BY" in r:
        st = "STANDBY"
    elif "TRANSFER" in r:
        st = "TRANSFERRED"
    else:
        st = "IN_USE"
        if any("STORE" in (x or "").upper() or "SPARE" in (x or "").upper() for x in (room, floor, loc)):
            st = "IN_STORE"
    return st, cat


def build_rows(recs, hr, as_of, qlabel, qstart, qend):
    rows = []
    last_ws_parent = None
    server_parent = {}
    counters = collections.Counter()
    for r in recs:
        cls, sheet = r["_class"], r["_sheet"]
        extra = []
        out = {c: None for c in COLS}
        out["SNAPSHOT_DATE"], out["ASSET_CLASS"] = as_of, cls
        out["SOURCE_SHEET"], out["SOURCE_ROW"] = sheet, r["_row"]
        ci = clean_str(r.get("CI_NO"), upper=False)
        hostname = clean_str(r.get("HOSTNAME"))
        if cls == "SERVER":
            tag = clean_str(r.get("CI_NO_TAG"))
            hostname = ci
            ci = tag or ci
        out["CI_NO"], out["HOSTNAME"] = ci, hostname
        hw = fix_typos(clean_str(r.get("HW_TYPE"), upper=True))
        if cls == "DESKTOP" and hw == "LAPTOP":
            hw = "OFFICE LAPTOP"
            LOG["DESKTOP-class LAPTOP renamed to OFFICE LAPTOP"] += 1
        out["ASSET_TYPE"] = hw
        level, parent = "ASSET", None
        sr = clean_str(r.get("SR"))
        if cls == "WORKSTATION":
            if ci:
                last_ws_parent = ci
            elif hw in ("MONITOR", "GRAPHICS CARD"):
                level, parent = "COMPONENT", last_ws_parent
        elif cls == "SERVER":
            base = sr.split(".")[0] if sr else None
            if sr and "." in sr:
                level, parent = "COMPONENT", server_parent.get(base)
                if not hw:
                    out["ASSET_TYPE"] = "SERVER COMPONENT"
            elif base and ci:
                server_parent[base] = ci
        elif cls == "SWITCH" and not ci and hw in ("SFP", "FIBRE MODULE"):
            level = "COMPONENT"
        out["RECORD_LEVEL"], out["PARENT_ASSET_KEY"] = level, parent

        if level == "COMPONENT" and cls == "WORKSTATION":
            out["MODEL"], out["SERIAL_NO"] = clean_str(r.get("MON_MODEL")), clean_str(r.get("MON_SN"))
            out["MAKE"] = None
        else:
            out["MODEL"] = clean_str(r.get("MODEL"))
            out["SERIAL_NO"] = clean_str(r.get("SERIAL"))
            out["ATTACHED_MONITOR_MODEL"] = clean_str(r.get("MON_MODEL"))
            out["ATTACHED_MONITOR_SN"] = clean_str(r.get("MON_SN"))
            out["MAKE"] = norm_make(r.get("MAKE"))
        if level == "COMPONENT" and cls == "WORKSTATION":
            out["MAKE"] = norm_make(r.get("MAKE"))
        out["ASSET_DESCRIPTION"] = clean_str(r.get("DESCRIPTION"))
        aid = asset_id_clean(r.get("ASSET_ID"), extra)
        out["ONGC_ASSET_ID"] = None if aid == "__INVALID__" else aid
        if aid == "__INVALID__":
            out["_bad_aid"] = True
        out["ONGC_CENSUS_NO"] = clean_str(r.get("CENSUS"))
        out["SUB_TYPE"] = fix_typos(clean_str(r.get("SUB_TYPE"), upper=True))
        out["OS"] = clean_str(r.get("OS"))
        out["OS_FAMILY"] = os_family(out["OS"])
        out["FIRMWARE_VERSION"] = clean_str(r.get("FIRMWARE"))
        out["IP_ADDRESS"] = clean_str(r.get("IP"))
        out["MGMT_IP"] = clean_str(r.get("MGMT_IP"))
        out["PROCESSOR"] = clean_str(r.get("PROCESSOR"))
        out["RAM"], out["STORAGE"] = clean_str(r.get("RAM")), clean_str(r.get("STORAGE"))
        out["PORTS"] = clean_str(r.get("PORTS"))
        out["NETWORK_PORT_AVAILABLE"] = clean_str(r.get("NET_PORT"), upper=True)
        out["CAPACITY"], out["BATTERY_QTY"] = clean_str(r.get("CAPACITY")), clean_str(r.get("BATT_QTY"))
        out["BATTERY_SPEC"] = clean_str(r.get("BATT_SPEC"))
        out["INSTALL_DATE"] = to_date(r.get("INSTALL_DATE"))

        try:
            cpf = int(float(r.get("CPF")))
        except (TypeError, ValueError):
            cpf = None
        out["CPF_NO"] = cpf
        h = hr.get(cpf)
        if h:
            out["USER_NAME"], out["USER_DESIGNATION"], out["USER_LEVEL"], out["USER_MOBILE"] = h[0], h[1], h[2], h[3]
            out["USER_RETIREMENT_DATE"] = to_date(h[4])
            out["USER_HR_STATUS"] = h[5]
        else:
            out["USER_NAME"] = clean_str(r.get("U_NAME"))
            out["USER_DESIGNATION"], out["USER_LEVEL"] = clean_str(r.get("U_DESIG")), clean_str(r.get("U_LEVEL"))
            mob = clean_str(r.get("U_MOBILE"))
            out["USER_MOBILE"] = re.sub(r"\D", "", mob)[-10:] if mob else None
            out["USER_RETIREMENT_DATE"] = to_date(r.get("U_RETIRE"))
            out["USER_HR_STATUS"] = "NO_CPF" if cpf is None else "NOT_IN_HR_MASTER"
        out["USER_DEPARTMENT"] = clean_str(r.get("U_DEPT"))
        out["ENGINEER_NAME"] = clean_str(r.get("ENGINEER"), upper=True)
        out["LOCATION_CODE"] = clean_str(r.get("LOCATION"), upper=True)
        out["FLOOR_AREA"] = fix_typos(clean_str(r.get("FLOOR"), upper=True))
        out["ROOM"] = fix_typos(clean_str(r.get("ROOM"), upper=True))

        out["COVER_TYPE"] = clean_str(r.get("COVER"), upper=True)
        out["COVER_EXPIRY_DATE"] = to_date(r.get("COVER_EXP"))
        out["RATE_COMPONENT"] = clean_str(r.get("RATE_COMP"), upper=True)
        try:
            out["RATE_VALUE"] = round(float(r.get("RATE_VAL")), 2)
        except (TypeError, ValueError):
            out["RATE_VALUE"] = None
        out["PM_DATE"] = to_date(r.get("PM_DATE"))
        out["PM_DONE_BY"] = clean_str(r.get("PM_DONE"), upper=True)
        out["PM_SIGNED_BY"] = clean_str(r.get("PM_SIGNED"))
        out["PM_TRACKER_DATE"] = to_date(r.get("PREV_PM"))

        remark = clean_str(r.get("REMARKS"))
        out["REMARKS"] = remark
        ex = clean_str(r.get("EXTRA_INFO"))
        if ex:
            extra.append(ex)
        out["EXTRA_INFO"] = " | ".join(extra) if extra else None
        st, cat = status_from(remark, out["ROOM"], out["FLOOR_AREA"], out["LOCATION_CODE"])
        out["ASSET_STATUS"], out["USER_CATEGORY"] = st, cat

        # keys
        if level == "COMPONENT" and cls == "WORKSTATION":
            tag = "MON" if hw == "MONITOR" else "GPU"
            counters[(parent, tag)] += 1
            out["ASSET_KEY"] = f"{parent}/{tag}{counters[(parent, tag)]}"
        elif level == "COMPONENT" and cls == "SERVER" and not ci:
            counters[(parent, "VM")] += 1
            out["ASSET_KEY"] = f"{parent}/VM{counters[(parent, 'VM')]}"
            out["_gen"] = True
        elif level == "COMPONENT" and cls == "SWITCH":
            t = "SFP" if hw == "SFP" else "FIBRE"
            counters[("SW", t)] += 1
            out["ASSET_KEY"] = f"SWITCH-{t}-{counters[('SW', t)]:04d}"
            out["_gen"] = True
        elif ci:
            out["ASSET_KEY"] = ci
        else:
            counters[(cls, "NOCI")] += 1
            out["ASSET_KEY"] = f"{CLASS_ABBR[cls]}-NOCI-{counters[(cls, 'NOCI')]:03d}"
            out["_gen"] = True
        if level == "COMPONENT" and cls == "WORKSTATION" and not parent:
            out["_orphan"] = True
        rows.append(out)
    return rows


def finish(df, as_of, qlabel, qstart, qend):
    df["PM_QUARTER"] = None
    is_pm = (df["ASSET_CLASS"] != "LAPTOP")
    df.loc[is_pm, "PM_QUARTER"] = qlabel

    def pm_status(r):
        if r["ASSET_CLASS"] == "LAPTOP":
            return "NOT_TRACKED"
        if r["RECORD_LEVEL"] == "COMPONENT":
            return "NA"
        d = r["PM_DATE"]
        if d is None:
            return "PENDING"
        if qstart and qend and not (qstart <= d <= qend):
            return "DONE_OUTSIDE_QUARTER"
        return "DONE"

    df["PM_STATUS"] = df.apply(pm_status, axis=1)

    def cover(r):
        if r["ASSET_STATUS"] == "REMOVED_FROM_AMC":
            return "REMOVED"
        e = r["COVER_EXPIRY_DATE"]
        if e is None:
            return "UNKNOWN"
        if e < as_of:
            return "EXPIRED"
        if e <= as_of + dt.timedelta(days=90):
            return "EXPIRING_90D"
        return "ACTIVE"

    df["COVER_STATUS"] = df.apply(cover, axis=1)

    aid_counts = df["ONGC_ASSET_ID"].value_counts()
    ser_counts = df["SERIAL_NO"].value_counts()
    single_ip = df["IP_ADDRESS"].where(df["IP_ADDRESS"].map(lambda x: bool(x) and bool(IPV4.match(x))))
    ip_counts = single_ip.value_counts()
    flags = []
    for i, r in df.iterrows():
        f = []
        if r.get("_gen") is True:
            f.append("GENERATED_KEY")
        if r.get("_orphan") is True:
            f.append("COMPONENT_NO_PARENT")
        if r["ASSET_CLASS"] == "SWITCH" and r["RECORD_LEVEL"] == "COMPONENT":
            f.append("COMPONENT_NOT_LINKED_TO_SWITCH")
        if r["ASSET_CLASS"] == "MEDIA_CONVERTER" and not r["CI_NO"]:
            f.append("MISSING_CI_NO")
        if r["USER_HR_STATUS"] in ("NOT_IN_HR_MASTER", "NO_CPF"):
            f.append("CPF_NOT_IN_HR_MASTER" if r["USER_HR_STATUS"] == "NOT_IN_HR_MASTER" else "NO_CPF")
        elif r["USER_HR_STATUS"] in ("RETIRED", "EXITED"):
            f.append("USER_RETIRED_OR_EXITED")
        if r["RECORD_LEVEL"] == "ASSET":
            if not r["ONGC_ASSET_ID"] and r.get("_bad_aid") is not True:
                f.append("MISSING_ASSET_ID")
            if not r["SERIAL_NO"] and r["ASSET_CLASS"] != "SERVER":
                f.append("MISSING_SERIAL")
        if r.get("_bad_aid") is True:
            f.append("ASSET_ID_INVALID_MOVED")
        if r["ONGC_ASSET_ID"]:
            n = aid_counts.get(r["ONGC_ASSET_ID"], 0)
            if n > 10:
                f.append("BULK_ASSET_ID")
            elif n > 1:
                f.append("SHARED_ASSET_ID")
        if r["SERIAL_NO"] and ser_counts.get(r["SERIAL_NO"], 0) > 1 and r["RECORD_LEVEL"] == "ASSET":
            f.append("DUPLICATE_SERIAL")
        ip = r["IP_ADDRESS"]
        if ip:
            if not IPV4.match(ip):
                f.append("IP_NOT_STANDARD")
            elif ip_counts.get(ip, 0) > 1:
                f.append("DUPLICATE_IP")
        if r["COVER_STATUS"] == "EXPIRED":
            f.append("COVER_EXPIRED")
        elif r["COVER_STATUS"] == "EXPIRING_90D":
            f.append("COVER_EXPIRING_90D")
        if r["ASSET_STATUS"] == "REMOVED_FROM_AMC" and r["RATE_VALUE"]:
            f.append("REMOVED_FROM_AMC_RATE_PRESENT")
        if r["PM_STATUS"] == "PENDING":
            f.append("PM_PENDING")
        elif r["PM_STATUS"] == "DONE_OUTSIDE_QUARTER":
            f.append("PM_OUTSIDE_QUARTER")
        flags.append("; ".join(f) or None)
    df["DQ_FLAGS"] = flags
    return df


# ------------------------------------------------------------------ workbook
F_BASE = Font(name="Arial", size=10)
F_BOLD = Font(name="Arial", size=10, bold=True)
F_HDR = Font(name="Arial", size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name="Arial", size=14, bold=True)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
DATE_FMT = "DD-MMM-YYYY"
DATE_COLS = {"SNAPSHOT_DATE", "INSTALL_DATE", "USER_RETIREMENT_DATE", "COVER_EXPIRY_DATE", "PM_DATE", "PM_TRACKER_DATE"}
WIDTH = {"ASSET_KEY": 24, "CI_NO": 20, "USER_NAME": 30, "USER_DESIGNATION": 38, "MODEL": 30, "PROCESSOR": 30, "ASSET_DESCRIPTION": 30,
         "HOSTNAME": 24, "ROOM": 24, "LOCATION_CODE": 22, "FLOOR_AREA": 18, "REMARKS": 30, "EXTRA_INFO": 34, "DQ_FLAGS": 44,
         "OS": 24, "RAM": 18, "STORAGE": 22, "SERIAL_NO": 20, "USER_DEPARTMENT": 20, "PARENT_ASSET_KEY": 24, "FIRMWARE_VERSION": 18}


def header_row(ws, names, fills=None, row=1):
    for i, n in enumerate(names, 1):
        c = ws.cell(row=row, column=i, value=n)
        c.font, c.border = F_HDR, BOX
        c.fill = PatternFill("solid", fgColor=(fills[i - 1] if fills else "1F3864"))
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def simple_sheet(wb, title, header, rows, widths):
    ws = wb.create_sheet(title)
    header_row(ws, header)
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.alignment = F_BASE, Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{max(ws.max_row, 2)}"
    return ws


def write_master(ws, df, min_last=0):
    ws.title = MASTER
    header_row(ws, COLS, [GROUP_FILL[t[1]] for t in TARGETS])
    ws.row_dimensions[1].height = 32
    for r, rec in enumerate(df[COLS].itertuples(index=False), 2):
        for i, v in enumerate(rec, 1):
            if v is None or (isinstance(v, float) and pd.isna(v)):
                continue
            ws.cell(row=r, column=i, value=v)
    for i, c in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTH.get(c, 16)
    for row in ws.iter_rows(min_row=2, max_row=max(ws.max_row, min_last)):
        for c in row:
            c.font = F_BASE
            n = COLS[c.column - 1]
            if n in DATE_COLS:
                c.number_format = DATE_FMT
            elif n == "RATE_VALUE":
                c.number_format = "#,##0.00"
            elif n in ("CPF_NO", "SOURCE_ROW"):
                c.number_format = "0"
            elif n in ("USER_MOBILE", "ONGC_ASSET_ID"):
                c.number_format = "@"
    lists = {"COVER_TYPE": '"AMC,WRTY"', "RECORD_LEVEL": '"ASSET,COMPONENT"',
             "ASSET_STATUS": '"IN_USE,IN_STORE,SURPLUS,NOT_IN_USE,NOT_ON_NETWORK,STANDBY,REMOVED_FROM_AMC,TRANSFERRED"',
             "PM_STATUS": '"DONE,DONE_OUTSIDE_QUARTER,PENDING,NA,NOT_TRACKED"',
             "COVER_STATUS": '"ACTIVE,EXPIRING_90D,EXPIRED,REMOVED,UNKNOWN"'}
    last = max(ws.max_row + 500, min_last)
    for name, f in lists.items():
        dv = DataValidation(type="list", formula1=f, allow_blank=True, errorStyle="warning")
        ws.add_data_validation(dv)
        col = get_column_letter(COLS.index(name) + 1)
        dv.add(f"{col}2:{col}{last}")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{max(ws.max_row, 2)}"


def summary_rows(df):
    out = []
    for cls in list(SHEET_CLASS.values()) + ["ALL"]:
        d = df if cls == "ALL" else df[df["ASSET_CLASS"] == cls]
        a = d[d["RECORD_LEVEL"] == "ASSET"]
        out.append((cls, len(d), len(a), len(d) - len(a),
                    int((a["ASSET_STATUS"] == "IN_USE").sum()), int(a["ASSET_STATUS"].isin(["IN_STORE", "SURPLUS", "NOT_IN_USE", "NOT_ON_NETWORK", "STANDBY"]).sum()),
                    int((a["ASSET_STATUS"] == "REMOVED_FROM_AMC").sum()),
                    int((a["PM_STATUS"] == "DONE").sum()), int((a["PM_STATUS"] == "PENDING").sum()),
                    int((a["PM_STATUS"] == "DONE_OUTSIDE_QUARTER").sum()),
                    int((d["COVER_STATUS"] == "EXPIRED").sum()), int((d["COVER_STATUS"] == "EXPIRING_90D").sum()),
                    int(d["DQ_FLAGS"].notna().sum())))
    return out


def write_workbook(path, df, layouts, skipped, unknown, as_of, qlabel, hr_used):
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 110
    ws["A1"] = "IT Asset Master - single-sheet inventory"
    ws["A1"].font = F_TITLE
    n_asset = int((df["RECORD_LEVEL"] == "ASSET").sum())
    lines = [
        ("Built from", "IT-IMMDSS ASSET INVENTORY (10 sheets, one per asset class), merged into the ONE sheet ASSET_MASTER."),
        ("Snapshot / PM quarter", f"{as_of:%d-%b-%Y} / {qlabel}"),
        ("Rows", f"{len(df):,} rows = {n_asset:,} assets + {len(df) - n_asset:,} component lines (monitors, GPUs, VMs, SFPs, fibre modules)."),
        ("Grain", "One row per inventory line. ASSET_KEY is unique. Use ASSET_CLASS / ASSET_TYPE to slice; RECORD_LEVEL = ASSET to count physical items."),
        ("Upload sheet", "ASSET_MASTER only. The other sheets are reference (dictionary, mapping, rate card, quality, cleaning log)."),
        ("User details", "CPF_NO is the link. Name, designation, level, mobile and retirement date come from the HR master (PostgreSQL employee table) so they can no longer go stale or break as external-file VLOOKUPs." if hr_used else "HR master not found: user details fall back to the values stored in the source workbook."),
        ("Derived columns", "OS_FAMILY, USER_CATEGORY, ASSET_STATUS, COVER_STATUS, PM_STATUS and DQ_FLAGS are computed by the converter. Original text is always kept (REMARKS, OS)."),
        ("Header colours", "Dark blue = identity, blue = hardware, green = ownership, brown = location, orange = contract, purple = PM, red = status, grey = trace."),
        ("Quarterly refresh", "Give the new raw inventory workbook; it is converted with tools/inventory_to_master.py. Keep one ASSET_MASTER per quarter for history."),
        ("Assumptions to confirm", "1) Unlabelled server column = INSTALL_DATE.  2) Unlabelled UPS column = BATTERY_SPEC.  3) Unlabelled lookup column = PM_TRACKER_DATE (VLOOKUP into an external PM tracker; matches PM_DATE on every row where present).  4) Desktop unlabelled email/text column kept as EXTRA_INFO - meaning unconfirmed.  5) 'RETIRED DATE' on laptops is the user's retirement date, not the device's."),
    ]
    for i, (a, b) in enumerate(lines, 3):
        ws.cell(row=i, column=1, value=a).font = F_BOLD
        c = ws.cell(row=i, column=2, value=b)
        c.font, c.alignment = F_BASE, Alignment(wrap_text=True, vertical="top")
        ws.cell(row=i, column=1).alignment = Alignment(vertical="top")

    write_master(wb.create_sheet(), df)

    simple_sheet(wb, "SUMMARY", ["ASSET_CLASS", "TOTAL_ROWS", "ASSETS", "COMPONENTS", "IN_USE", "STORE / SURPLUS / NOT USED", "REMOVED_FROM_AMC",
                                 "PM_DONE", "PM_PENDING", "PM_OUTSIDE_QUARTER", "COVER_EXPIRED", "COVER_EXPIRING_90D", "ROWS_WITH_DQ_FLAGS"],
                 summary_rows(df), [18] + [14] * 12)

    dq = collections.Counter()
    for f in df["DQ_FLAGS"].dropna():
        for x in f.split("; "):
            dq[x] += 1
    desc = {
        "GENERATED_KEY": "No CI number in source; key generated.", "COMPONENT_NOT_LINKED_TO_SWITCH": "SFP / fibre-module lines have no parent switch in source.",
        "MISSING_CI_NO": "Asset without CI number.", "CPF_NOT_IN_HR_MASTER": "CPF not found in HR master (retired, transferred or wrong CPF).",
        "NO_CPF": "No CPF on row.", "USER_RETIRED_OR_EXITED": "User marked RETIRED/EXITED in HR master.", "MISSING_ASSET_ID": "ONGC asset id blank.",
        "MISSING_SERIAL": "Serial number blank.", "ASSET_ID_INVALID_MOVED": "ONGC asset id field held a non-numeric value (moved to EXTRA_INFO).",
        "BULK_ASSET_ID": "One asset id shared by >10 devices (contract / batch id).", "SHARED_ASSET_ID": "Asset id shared by 2-10 rows: possible duplicate.",
        "DUPLICATE_SERIAL": "Same serial on more than one asset.", "IP_NOT_STANDARD": "IP is not a single valid IPv4 (e.g. 163 or ranges).",
        "DUPLICATE_IP": "Same IP on more than one row.", "COVER_EXPIRED": "AMC / warranty expiry before snapshot date.",
        "COVER_EXPIRING_90D": "Cover expires within 90 days.", "REMOVED_FROM_AMC_RATE_PRESENT": "Remark says removed from AMC but a rate value is still recorded (billing check).",
        "PM_PENDING": "No PM date this quarter.", "PM_OUTSIDE_QUARTER": "PM date falls outside the PM quarter (stale carry-over).",
        "COMPONENT_NO_PARENT": "Component could not be tied to a parent.",
    }
    rows = [(k, v, desc.get(k, "")) for k, v in dq.most_common()]
    rows.append(("", "", ""))
    rows.append(("Skipped junk rows", len(skipped), "; ".join(f"{s}!{r}='{v}'" for s, r, v in skipped) or "none"))
    rows.append(("Unknown extra sheets", len(unknown), ", ".join(unknown) or "none"))
    simple_sheet(wb, "DATA_QUALITY", ["FLAG / CHECK", "ROWS", "MEANING"], rows, [38, 10, 100])

    rc = df[df["RATE_COMPONENT"].notna()].groupby("RATE_COMPONENT").agg(
        RATE_VALUE=("RATE_VALUE", "first"), CLASS=("ASSET_CLASS", lambda s: ", ".join(sorted(set(s)))),
        TYPES=("ASSET_TYPE", lambda s: ", ".join(sorted({x for x in s if x})[:4])),
        SAMPLE_MODEL=("MODEL", lambda s: next((x for x in s if x), None)), QTY=("ASSET_KEY", "count"),
        VALUE_CONFLICT=("RATE_VALUE", lambda s: "YES" if s.nunique() > 1 else "no")).reset_index()
    rc["_k"] = rc["RATE_COMPONENT"].map(lambda x: (re.sub(r"\d+", "", x), int(re.sub(r"\D", "", x) or 0)))
    rc = rc.sort_values("_k").drop(columns="_k")
    simple_sheet(wb, "RATE_CARD", ["RATE_COMPONENT", "RATE_VALUE", "ASSET_CLASS", "ASSET_TYPES", "SAMPLE_MODEL", "QTY_ROWS", "VALUE_CONFLICT"],
                 rc[["RATE_COMPONENT", "RATE_VALUE", "CLASS", "TYPES", "SAMPLE_MODEL", "QTY", "VALUE_CONFLICT"]].values.tolist(), [16, 12, 20, 40, 34, 10, 14])

    simple_sheet(wb, "FIELD_DICTIONARY", ["COLUMN", "GROUP", "DESCRIPTION"], TARGETS_ROWS(), [28, 14, 120])

    mrows = []
    for sheet, lay in layouts.items():
        for header, letter, key in lay:
            tgt = KEY_TO_TARGET.get(key)
            note = ""
            if key in ("CI_NO_TAG", "BATT_SPEC", "INSTALL_DATE", "PREV_PM", "EXTRA_INFO"):
                note = "UNLABELLED in source - meaning inferred"
            if key == "PREV_PM" and hnorm(header) == "CIPL REMARKS":
                note = "Header says CIPL REMARKS but the cells hold PM-tracker lookup dates"
            if key == "U_RETIRE":
                note = "Really the user's retirement date (VLOOKUP to HR)"
            if key == "SR":
                tgt, note = "(dropped)", "Source serial number restarts / duplicates; replaced by ASSET_KEY"
            if key == "POS_NO":
                tgt, note = "(dropped)", "Always empty"
            if key == "IGNORE":
                tgt, note = "(dropped)", "Empty / unused column"
            if key == "SUB_TYPE":
                tgt = "SUB_TYPE"
            mrows.append((sheet, letter, "(no header)" if header.startswith("Unnamed") else header.replace("\n", " "), tgt or "(dropped)", note))
    simple_sheet(wb, "SOURCE_MAPPING", ["SOURCE SHEET", "COL", "SOURCE HEADER", "TARGET COLUMN", "NOTE"], mrows, [18, 6, 40, 46, 70])

    crow = [(k, v) for k, v in LOG.most_common()]
    simple_sheet(wb, "CLEANING_LOG", ["CLEANING ACTION", "CELLS / ROWS"], crow, [70, 14])
    try:
        wb.save(path)
    except PermissionError:
        sys.exit(f"Cannot write {path} - it is open in Excel. Close it and run again.")


TEMPLATE_NAME = "IT_Asset_Master_Upload_Template.xlsx"
TEMPLATE_DIR = Path(__file__).resolve().parent.parent / "templates"
REQUIRED = {"SNAPSHOT_DATE", "ASSET_KEY", "ASSET_CLASS", "ASSET_TYPE", "RECORD_LEVEL"}
DERIVED = {"OS_FAMILY", "USER_NAME", "USER_DESIGNATION", "USER_LEVEL", "USER_MOBILE", "USER_RETIREMENT_DATE", "USER_HR_STATUS",
           "USER_CATEGORY", "COVER_STATUS", "PM_QUARTER", "PM_STATUS", "ASSET_STATUS", "DQ_FLAGS"}


def write_template(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 110
    ws["A1"] = "IT Asset Master - Standard Bulk Upload Template"
    ws["A1"].font = F_TITLE
    lines = [
        ("Purpose", "One flat sheet (ASSET_MASTER) holding every IT asset and component: desktops, office laptops, laptops, workstations, servers, routers, switches, printers, scanners, UPS and media converters. Structure is identical to IT_Asset_Master_2026-Q2.xlsx."),
        ("Upload sheet", "ASSET_MASTER, headers in row 1. Do not rename, reorder or delete columns."),
        ("Key", "ASSET_KEY - unique per row. Use the CI_NO when the item has one. Components without a CI use PARENT/MON1, PARENT/GPU1, PARENT/VM1 or SWITCH-SFP-0001 style keys."),
        ("Grain", "One row per inventory line. RECORD_LEVEL = ASSET for physical / logical items, COMPONENT for child lines (monitor, GPU, VM, SFP, fibre module). PARENT_ASSET_KEY links a component to its parent asset."),
        ("Required columns", "SNAPSHOT_DATE, ASSET_KEY, ASSET_CLASS, ASSET_TYPE, RECORD_LEVEL. Everything else is optional; leave blank when unknown (never 'NA', 'N/A' or '-')."),
        ("Converter-filled columns", "OS_FAMILY, USER_NAME, USER_DESIGNATION, USER_LEVEL, USER_MOBILE, USER_RETIREMENT_DATE, USER_HR_STATUS, USER_CATEGORY, COVER_STATUS, PM_QUARTER, PM_STATUS, ASSET_STATUS and DQ_FLAGS are calculated by tools/inventory_to_master.py from the raw inventory and the HR master. Only type CPF_NO for the user."),
        ("Formats", "Dates DD-MMM-YYYY. MOBILE and ONGC_ASSET_ID are text. RATE_VALUE 2 decimals. Names in UPPER CASE for make, floor, room, location, engineer."),
        ("Allowed lists", "See the LOOKUPS sheet. Unexpected values show a warning, not a block, so genuinely new values can be entered."),
        ("Header colours", "Dark blue = identity, blue = hardware, green = ownership, brown = location, orange = contract, purple = PM, red = status, grey = trace."),
        ("Quarterly process", "Hand over the raw IT-IMMDSS inventory workbook. It is converted into this layout (IT_Asset_Master_<year>-Q<n>.xlsx); check DATA_QUALITY and DQ_FLAGS, then upload ASSET_MASTER."),
        ("Example row (illustrative)", "2026-09-21 | ANKAON00DT999 | ANKAON00DT999 | DESKTOP | ALL IN ONE | ASSET | | | | HP | PRO ONE 440 G9 24 | 1N13350XXX | 104000000 | 21A138000000 | | WIN 11 PRO | ... | CPF_NO 100001 | ... | HIREN PATEL | ANK_NEWBDGGF_A | GROUND FLOOR | A-003 | AMC | 13-Apr-2029 | | D.4 | 177.30 | Q2 JUL-SEP 2026 | 03-Jul-2026 | | HIREN PATEL"),
    ]
    for i, (a, b) in enumerate(lines, 3):
        ws.cell(row=i, column=1, value=a).font = F_BOLD
        c = ws.cell(row=i, column=2, value=b)
        c.font, c.alignment = F_BASE, Alignment(wrap_text=True, vertical="top")
        ws.cell(row=i, column=1).alignment = Alignment(vertical="top")
    write_master(wb.create_sheet(), pd.DataFrame(columns=COLS), min_last=5000)
    rows = [(t[0], t[1], "REQUIRED" if t[0] in REQUIRED else ("CONVERTER-FILLED" if t[0] in DERIVED else "optional"), t[2]) for t in TARGETS]
    simple_sheet(wb, "FIELD_DICTIONARY", ["COLUMN", "GROUP", "ENTRY", "DESCRIPTION"], rows, [28, 14, 18, 110])
    lists = {
        "ASSET_CLASS": list(SHEET_CLASS.values()),
        "RECORD_LEVEL": ["ASSET", "COMPONENT"],
        "COVER_TYPE": ["AMC", "WRTY"],
        "ASSET_STATUS": ["IN_USE", "IN_STORE", "SURPLUS", "NOT_IN_USE", "NOT_ON_NETWORK", "STANDBY", "REMOVED_FROM_AMC", "TRANSFERRED"],
        "COVER_STATUS": ["ACTIVE", "EXPIRING_90D", "EXPIRED", "REMOVED", "UNKNOWN"],
        "PM_STATUS": ["DONE", "DONE_OUTSIDE_QUARTER", "PENDING", "NA", "NOT_TRACKED"],
        "USER_CATEGORY": ["APPRENTICE", "CONTINGENT", "AUDITOR", "COMMON_USE", "DATA_ENTRY"],
        "OS_FAMILY": ["WINDOWS 11", "WINDOWS 10", "WINDOWS SERVER", "MACOS", "LINUX", "OTHER"],
        "ASSET_TYPE (examples)": ["DESKTOP", "ALL IN ONE", "OFFICE LAPTOP", "LAPTOP", "TABLET", "WORKSTATION", "MONITOR", "GRAPHICS CARD",
                                  "SERVER", "VM SERVER", "ROUTER", "L2 SWITCH", "L3 SWITCH", "HUB", "SFP", "FIBRE MODULE", "PRINTER",
                                  "SCANNER", "UPS", "NETWORK UPS", "MEDIA CONVERTOR"],
    }
    wl = wb.create_sheet("LOOKUPS")
    header_row(wl, list(lists))
    for i, vals in enumerate(lists.values(), 1):
        for r, v in enumerate(vals, 2):
            wl.cell(row=r, column=i, value=v).font = F_BASE
        wl.column_dimensions[get_column_letter(i)].width = 24
    try:
        wb.save(path)
    except PermissionError:
        sys.exit(f"Cannot write {path} - it is open in Excel. Close it and run again.")


def template_columns():
    for f in sorted(TEMPLATE_DIR.glob("*.xlsx")):
        n = f.name.lower()
        if "asset_master" in n and "template" in n and not n.startswith("~$"):
            from openpyxl import load_workbook
            wb = load_workbook(f, read_only=True)
            hdr = [c for c in next(wb["ASSET_MASTER"].iter_rows(min_row=1, max_row=1, values_only=True)) if c]
            wb.close()
            return f.name, hdr
    return None, None


def TARGETS_ROWS():
    return [(t[0], t[1], t[2]) for t in TARGETS]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--raw")
    ap.add_argument("--build-template", action="store_true")
    ap.add_argument("--no-load-db", action="store_true", help="build the workbook only; skip the PostgreSQL load (default: load into database ongc_ank)")
    ap.add_argument("--force", action="store_true")
    ap.add_argument("--dry-run", action="store_true", help="rehearse the database load: everything runs and is reported, nothing is saved")
    ap.add_argument("--as-of")
    ap.add_argument("--out-dir", default=str(Path(__file__).resolve().parent.parent / "masters"))
    a = ap.parse_args()
    if a.dry_run:
        os.environ["ITAM_DRY_RUN"] = "1"
    if a.build_template:
        TEMPLATE_DIR.mkdir(exist_ok=True)
        write_template(TEMPLATE_DIR / TEMPLATE_NAME)
        print("Wrote", TEMPLATE_DIR / TEMPLATE_NAME)
        if not a.raw:
            return
    if not a.raw:
        ap.error("--raw is required")
    as_of = pd.to_datetime(a.as_of).date() if a.as_of else dt.date.today()
    recs, layouts, skipped, unknown, pm_header = read_source(a.raw)
    qlabel, qstart, qend = parse_quarter(pm_header)
    hr = load_hr()
    rows = build_rows(recs, hr, as_of, qlabel, qstart, qend)
    df = pd.DataFrame(rows)
    for c in COLS:
        if c not in df.columns:
            df[c] = None
    df = df.astype(object).where(df.notna(), None)
    df = finish(df, as_of, qlabel, qstart, qend)
    tname, thdr = template_columns()
    if thdr is not None and thdr != COLS:
        sys.exit(f"Output structure differs from template {tname} - stopped. Rebuild the template or fix TARGETS.")
    if tname:
        print("Structure check OK against template:", tname)
    missing_sheets = [n for n in SHEET_CLASS if n not in layouts]
    if missing_sheets and not a.force:
        sys.exit(f"Expected sheets missing from the raw workbook: {missing_sheets}. Stopped (use --force to override).")
    if df["ASSET_KEY"].duplicated().any():
        dup = df.loc[df["ASSET_KEY"].duplicated(keep=False), "ASSET_KEY"].unique()[:5]
        sys.exit(f"ASSET_KEY not unique, e.g. {list(dup)}")
    yq = f"{qstart.year}-Q{re.search(r'Q(\d)', qlabel)[1]}" if qlabel else f"{as_of:%Y-%m}"
    path = Path(a.out_dir) / f"IT_Asset_Master_{yq}.xlsx"
    write_workbook(path, df, layouts, skipped, unknown, as_of, qlabel or "?", bool(hr))
    print(f"Wrote {path}: {len(df)} rows ({int((df.RECORD_LEVEL == 'ASSET').sum())} assets)")
    for r in summary_rows(df):
        print("  ", r)
    print("Skipped junk rows:", skipped, "| unknown sheets:", unknown)
    if a.dry_run or not a.no_load_db:
        sys.path.insert(0, str(Path(__file__).parent))
        import master_db
        master_db.load_assets(path, force=a.force)


if __name__ == "__main__":
    main()
