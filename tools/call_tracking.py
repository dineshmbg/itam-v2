"""Call management: CIPL call tracker (CALLS + SPARE_INWARD + SPARE_OUTWARD) and OEM RMA log (OEM_RMA) -> templates, masters/ workbooks and PostgreSQL.

  python tools/call_tracking.py --build-template
  python tools/call_tracking.py tracker --raw "CIPL 2.0 Open - Closed calls Tracker.xlsx" [--as-of 2026-09-21] [--no-load-db] [--force]
  python tools/call_tracking.py rma     --raw "S D Wan Cisco Detail.xlsx"            [--as-of 2026-09-21] [--no-load-db] [--force]
  python tools/call_tracking.py report

Only the 'Call Log' sheet of the SD-WAN workbook is read. Its 'Password' sheet is never opened.
Outputs go to masters/, templates to templates/, previews to samples/.
"""
import argparse
import collections
import datetime as dt
import difflib
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).parent))
import itam_locks  # noqa: E402
import itam_rules  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE_DIR = ROOT / "templates"
MASTERS_DIR = ROOT / "masters"
TRACKER_TEMPLATE = "Call_Tracker_Upload_Template.xlsx"
RMA_TEMPLATE = "OEM_RMA_Upload_Template.xlsx"

# ------------------------------------------------------------------ table definitions (column, group, description, dtype)
CALLS = [
    ("SNAPSHOT_DATE", "TRACE", "Date of the load.", "date"),
    ("SR_ID", "IDENTITY", "CIPL service request id (UNIQUE KEY), e.g. SR2204202619028.", "text"),
    ("ONGC_TICKET_NO", "IDENTITY", "ONGC ticket number. One ticket can cover several SRs.", "text"),
    ("ONGC_CALL_DATE", "IDENTITY", "Date the call was logged by ONGC.", "date"),
    ("CIPL_CALL_DATE", "IDENTITY", "Date the call was logged in CIPL's system (start of TAT).", "date"),
    ("ZONE", "IDENTITY", "Zone.", "text"),
    ("SITE", "IDENTITY", "Site location.", "text"),
    ("SITE_INCHARGE", "IDENTITY", "CIPL site in-charge.", "text"),
    ("ASSET_KEY", "ASSET", "CI number of the asset (upper case). Links to the asset master.", "text"),
    ("ASSET_CLASS", "ASSET", "From the asset master.", "text"),
    ("ASSET_TYPE", "ASSET", "From the asset master.", "text"),
    ("MAKE", "ASSET", "From the asset master, else source (normalised).", "text"),
    ("MODEL", "ASSET", "From the asset master, else source.", "text"),
    ("SERIAL_NO", "ASSET", "Current serial from the asset master, else source.", "text"),
    ("CPF_NO", "ASSET", "CPF of the user who raised the call.", "int"),
    ("USER_NAME", "ASSET", "From the HR master, else source.", "text"),
    ("ENGINEER", "CALL", "Engineer assigned (spelling corrected against the asset master / roster).", "text"),
    ("PRIORITY", "CALL", "P1 / P2 / P3.", "text"),
    ("PROBLEM_DESCRIPTION", "CALL", "Problem as logged.", "text"),
    ("SPARE_REQUIRED", "CALL", "Spare action needed (source 'Require Spare').", "text"),
    ("PART_REQUIRED", "CALL", "Part required with part number.", "text"),
    ("SPARE_ISSUE_NOTE", "CALL", "DOA / damage / mismatch note.", "text"),
    ("SPARE_STATUS", "CALL", "NO_SPARE_NEEDED / PART_RECEIVED / PART_PENDING (derived from SPARE_INWARD).", "text"),
    ("CALL_STATUS", "LIFECYCLE", "OPEN / CLOSED.", "text"),
    ("CLOSED_DATE", "LIFECYCLE", "Date the call was closed.", "date"),
    ("FAULTY_SPARE_STATUS", "LIFECYCLE", "SENT / NOT_SENT (parsed from the source sentence).", "text"),
    ("FAULTY_SPARE_SENT_DATE", "LIFECYCLE", "Date parsed out of 'Faulty spare sent to warehouse on dd/mm/yyyy'.", "date"),
    ("TAT_DAYS", "METRIC", "CLOSED_DATE - CIPL_CALL_DATE (closed calls; blank if negative).", "int"),
    ("AGEING_DAYS", "METRIC", "Snapshot date - CIPL_CALL_DATE for open calls.", "int"),
    ("LOGGING_LAG_DAYS", "METRIC", "CIPL_CALL_DATE - ONGC_CALL_DATE.", "int"),
    ("REPEAT_CALL_COUNT", "METRIC", "Number of calls on this asset in the file.", "int"),
    ("OEM_RMA_NO", "LINK", "RMA number when the call led to an OEM RMA (fill in manually).", "text"),
    ("DQ_FLAGS", "TRACE", "Semicolon-separated data-quality flags.", "text"),
    ("SOURCE_ROW", "TRACE", "Row in the source sheet.", "int"),
]
SPARE_IN = [
    ("SNAPSHOT_DATE", "TRACE", "Date of the load.", "date"),
    ("INWARD_ID", "IDENTITY", "UNIQUE key: IN-0001. Stable once given; a line is matched by call + part + date, NOT by the sheet SERIAL.", "text"),
    ("INWARD_DATE", "IDENTITY", "Date on the source line (dispatch / bill date).", "date"),
    ("LOCATION", "IDENTITY", "Location.", "text"),
    ("SR_ID", "LINK", "CIPL service request this part is for.", "text"),
    ("ASSET_KEY", "LINK", "CI number, taken from the call.", "text"),
    ("PART_DESCRIPTION", "PART", "Part description.", "text"),
    ("PART_NO", "PART", "Part number if one is embedded in the description.", "text"),
    ("BILL_NO", "LOGISTICS", "Bill / AMC bill number.", "text"),
    ("COURIER_AWB", "LOGISTICS", "Courier AWB.", "text"),
    ("RECEIVED_BY", "LOGISTICS", "Who received the part.", "text"),
    ("RECEIVED_DATE", "LOGISTICS", "Date received.", "date"),
    ("TRANSIT_DAYS", "METRIC", "Received date - inward date.", "int"),
    ("REMARKS", "TRACE", "Remarks.", "text"),
    ("DQ_FLAGS", "TRACE", "Semicolon-separated data-quality flags.", "text"),
    ("SOURCE_ROW", "TRACE", "Row in the Inward sheet.", "int"),
]
SPARE_OUT = [
    ("SNAPSHOT_DATE", "TRACE", "Date of the load.", "date"),
    ("OUTWARD_ID", "IDENTITY", "UNIQUE key: OUT-0001. Stable once given; a line is matched by call + part, NOT by the sheet SERIAL.", "text"),
    ("OUTWARD_DATE", "IDENTITY", "Date on the source line.", "date"),
    ("LOCATION", "IDENTITY", "Location.", "text"),
    ("SR_ID", "LINK", "CIPL service request the faulty part belongs to.", "text"),
    ("ASSET_KEY", "LINK", "CI number of the asset the part came from.", "text"),
    ("DEVICE_SERIAL_NO", "PART", "Serial of the device the part came from.", "text"),
    ("PART_DESCRIPTION", "PART", "Part / item description.", "text"),
    ("PART_SERIAL_NO", "PART", "Serial number of the faulty part.", "text"),
    ("GATEPASS_NO", "LOGISTICS", "Gate pass number.", "text"),
    ("SENT_DATE", "LOGISTICS", "Date sent.", "date"),
    ("SENT_LOCATION", "LOGISTICS", "Where the faulty part was sent.", "text"),
    ("COURIER", "LOGISTICS", "Courier.", "text"),
    ("REMARKS", "TRACE", "Remarks.", "text"),
    ("DQ_FLAGS", "TRACE", "Semicolon-separated data-quality flags.", "text"),
    ("SOURCE_ROW", "TRACE", "Row in the Outward sheet.", "int"),
]
RMA = [
    ("SNAPSHOT_DATE", "TRACE", "Date of the load.", "date"),
    ("RMA_LINE_ID", "IDENTITY", "UNIQUE key: RMA-0001. Stable once given; a line is matched by faulty part serial + call date, NOT by the sheet Sr. no.", "text"),
    ("RMA_NO", "IDENTITY", "OEM RMA number, cleaned (no 'RMA#', spaces or nbsp).", "text"),
    ("VENDOR_CASE_ID", "IDENTITY", "OEM case id (source column 'Cash ID').", "text"),
    ("VENDOR", "DEVICE", "JUNIPER / CISCO.", "text"),
    ("DEVICE_MODEL", "DEVICE", "Device model (EX4300-24P, MX480, 8200 ...).", "text"),
    ("LOCATION", "DEVICE", "Location as written in the log.", "text"),
    ("DEVICE_SERIAL_NO", "DEVICE", "Serial of the switch / router at the time of the call.", "text"),
    ("ASSET_KEY", "DEVICE", "CI number, resolved through the asset master and the replacement chain.", "text"),
    ("LINK_BASIS", "DEVICE", "How ASSET_KEY was found: DEVICE_SERIAL / FAULTY_SERIAL / REPLACEMENT_SERIAL / CHAIN / NONE.", "text"),
    ("FAULT_ITEM", "FAULT", "Faulty part or symptom, cleaned (typos fixed).", "text"),
    ("FAULT_CATEGORY", "FAULT", "WHOLE_UNIT / PSU / POWER_ON / BOOT / PORT / MODULE / OPTIC / SOFTWARE / OTHER.", "text"),
    ("FAULTY_PART_SERIAL", "FAULT", "Serial of the faulty part / unit.", "text"),
    ("CALL_LOG_DATE", "DATES", "Date the case was logged with the OEM.", "date"),
    ("REPLACEMENT_RECEIVED_DATE", "DATES", "Date the replacement arrived.", "date"),
    ("REPLACEMENT_PART_SERIAL", "DATES", "Serial of the replacement part / unit.", "text"),
    ("FAULTY_RETURN_DATE", "DATES", "Date the faulty part was sent back.", "date"),
    ("RETURN_STATUS", "DATES", "RETURNED / PENDING.", "text"),
    ("DC_NO", "LOGISTICS", "Delivery challan number.", "text"),
    ("GATEPASS_NO", "LOGISTICS", "Gate pass number.", "text"),
    ("DAYS_TO_REPLACEMENT", "METRIC", "Replacement received - call logged (blank if negative).", "int"),
    ("DAYS_TO_RETURN", "METRIC", "Faulty return date - replacement received.", "int"),
    ("CALL_SR_ID", "LINK", "CIPL service request that led to this RMA (fill in manually).", "text"),
    ("REMARKS", "TRACE", "Remarks.", "text"),
    ("DQ_FLAGS", "TRACE", "Semicolon-separated data-quality flags.", "text"),
    ("SOURCE_ROW", "TRACE", "Row in the source sheet.", "int"),
]
TABLES = {"CALLS": (CALLS, "SR_ID", "svc_call"), "SPARE_INWARD": (SPARE_IN, "INWARD_ID", "spare_inward"),
          "SPARE_OUTWARD": (SPARE_OUT, "OUTWARD_ID", "spare_outward"), "OEM_RMA": (RMA, "RMA_LINE_ID", "oem_rma")}
VOLATILE = {"SNAPSHOT_DATE", "DQ_FLAGS", "SOURCE_ROW", "AGEING_DAYS", "REPEAT_CALL_COUNT", "SPARE_STATUS", "TAT_DAYS", "TRANSIT_DAYS", "DAYS_TO_REPLACEMENT", "DAYS_TO_RETURN"}
GROUP_FILL = {"IDENTITY": "1F3864", "ASSET": "2E75B6", "CALL": "548235", "LIFECYCLE": "C00000", "METRIC": "7F6000", "LINK": "7030A0", "TRACE": "595959",
              "PART": "2E75B6", "LOGISTICS": "548235", "DEVICE": "2E75B6", "FAULT": "C55A11", "DATES": "C00000"}
LISTS = {"PRIORITY": ["P1", "P2", "P3"], "CALL_STATUS": ["OPEN", "CLOSED"], "SPARE_STATUS": ["NO_SPARE_NEEDED", "PART_RECEIVED", "PART_PENDING"],
         "FAULTY_SPARE_STATUS": ["SENT", "NOT_SENT"], "DIRECTION": ["INWARD", "OUTWARD"], "RETURN_STATUS": ["RETURNED", "PENDING"],
         "FAULT_CATEGORY": ["WHOLE_UNIT", "PSU", "POWER_ON", "BOOT", "PORT", "MODULE", "OPTIC", "SOFTWARE", "OTHER"],
         "VENDOR": ["JUNIPER", "CISCO"]}

F_BASE = Font(name="Arial", size=10)
F_BOLD = Font(name="Arial", size=10, bold=True)
F_HDR = Font(name="Arial", size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name="Arial", size=14, bold=True)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
DATE_FMT = "DD-MMM-YYYY"
LOG = collections.Counter()


# ------------------------------------------------------------------ helpers
def blank(v):
    if v is None:
        return True
    try:
        if pd.isna(v):
            return True
    except (TypeError, ValueError):
        pass
    return isinstance(v, str) and v.strip().upper() in ("", "-", "NA", "N/A", "#N/A", "NIL", "//")


def txt(v, upper=False):
    if blank(v):
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = re.sub(r"\s+", " ", str(v).replace("\xa0", " ")).strip()
    return s.upper() if upper else s


def to_date(v):
    if blank(v):
        return None
    if isinstance(v, (dt.datetime, pd.Timestamp)):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v).strip()
    m = re.fullmatch(r"(\d{2})(\d{2})/(\d{4})", s)            # '0708/2026' typo for 07/08/2026
    if m:
        LOG["Malformed dates repaired (ddmm/yyyy)"] += 1
        return dt.date(int(m[3]), int(m[2]), int(m[1]))
    d = pd.to_datetime(s, errors="coerce", dayfirst=not re.match(r"\d{4}-\d{2}-\d{2}", s))
    return None if pd.isna(d) else d.date()


def hnorm(h):
    return re.sub(r"\s+", " ", str(h or "")).strip().upper()


def days(a, b):
    return (a - b).days if a and b else None


def style_header(ws, names, fills=None):
    for i, n in enumerate(names, 1):
        c = ws.cell(row=1, column=i, value=n)
        c.font, c.border = F_HDR, BOX
        c.fill = PatternFill("solid", fgColor=(fills[i - 1] if fills else "1F3864"))
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def simple_sheet(wb, title, header, rows, widths):
    ws = wb.create_sheet(title)
    style_header(ws, header)
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.alignment = F_BASE, Alignment(wrap_text=True, vertical="top")
            if isinstance(c.value, (dt.date, dt.datetime)):
                c.number_format = DATE_FMT
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{max(ws.max_row, 2)}"
    return ws


def read_rows(path, sheet_name, max_cols, stop_after=200):
    """Stream a sheet (read-only, stops after a run of empty rows - the tracker's used range is a million rows)."""
    wb = load_workbook(path, read_only=True, data_only=True)
    names = {n.upper(): n for n in wb.sheetnames}
    if sheet_name.upper() not in names:
        sys.exit(f"Sheet '{sheet_name}' not found in {Path(path).name}. Found: {wb.sheetnames}")
    ws = wb[names[sheet_name.upper()]]
    header, rows, empty = None, [], 0
    for i, row in enumerate(ws.iter_rows(min_row=1, max_col=max_cols, values_only=True), 1):
        if i == 1:
            header = [hnorm(h) for h in row]
            continue
        if all(blank(v) for v in row):
            empty += 1
            if empty >= stop_after:
                break
            continue
        empty = 0
        rows.append((i, list(row)))
    wb.close()
    return header, rows


def colmap(header, spec):
    """spec: {key: (startswith-or-equal text, occurrence)} -> {key: index}"""
    out = {}
    for key, (name, occ) in spec.items():
        hits = [i for i, h in enumerate(header) if h == name or h.startswith(name)]
        if len(hits) > occ:
            out[key] = hits[occ]
    return out


# ------------------------------------------------------------------ database lookups (read-only)
def lookups():
    try:
        import master_db
        con = master_db.connect()
    except SystemExit:
        return None
    try:
        a = master_db.query_df(con, "SELECT asset_key, asset_class, asset_type, make, model, serial_no, engineer_name FROM asset WHERE is_current = 1")
        h = master_db.query_df(con, "SELECT cpf_no, employee_name, record_status FROM employee")
        try:
            r = master_db.query_df(con, "SELECT employee_name FROM cipl_employee")
        except Exception:
            con.rollback()
            r = pd.DataFrame({"employee_name": []})
    except Exception:
        return None
    finally:
        con.close()
    assets = {str(k).upper(): row for k, row in zip(a["asset_key"], a.to_dict("records"))}
    hr = {int(c): n for c, n in zip(h["cpf_no"], h["employee_name"])}
    eng = {str(x).upper() for x in a["engineer_name"].dropna()} | {str(x).upper() for x in r["employee_name"].dropna()}
    serial_to_asset = {}
    for k, row in assets.items():
        if row["serial_no"]:
            serial_to_asset[str(row["serial_no"]).strip().upper()] = k
    return {"assets": assets, "hr": hr, "engineers": eng, "serials": serial_to_asset}


def norm_engineer(v, known):
    n = txt(v, upper=True)
    if not n:
        return None, None
    if not known or n in known:
        return n, None
    m = difflib.get_close_matches(n, sorted(known), n=1, cutoff=0.88)
    if m:
        LOG["Engineer names corrected to the asset master / roster spelling"] += 1
        return m[0], "ENGINEER_NAME_CORRECTED"
    return n, "ENGINEER_NOT_IN_MASTER_OR_ROSTER"


MAKE_ALIAS = {"HEWLETT-PACKARD": "HP", "HP (HP)": "HP", "DELL INC.": "DELL", "DELL INC": "DELL", "ASUSTEK COMPUTER INC.": "ASUS"}


def norm_make(v):
    s = txt(v, upper=True)
    if not s:
        return None
    out = MAKE_ALIAS.get(s, s)
    if out != txt(v):
        LOG["MAKE names normalised (case / spaces / aliases)"] += 1
    return out


# ------------------------------------------------------------------ tracker extraction
CALL_SPEC = {"zone": ("ZONE", 0), "site": ("SITE LOCATION", 0), "incharge": ("CIPL SITE INCHARGE", 0), "ongc_date": ("ONGC CALL LOGGING DATE", 0),
             "ticket": ("ONGC TICKET NO", 0), "cipl_date": ("CIPL CALL LOGGING DATE", 0), "sr": ("CIPL SERVICE REQ", 0), "asset": ("ASSET ID", 0),
             "cpf": ("CPF NO", 0), "user": ("USER NAME", 0), "make": ("MAKE", 0), "model": ("MODEL", 0), "serial": ("SERIAL NO", 0),
             "engineer": ("ENGINEER ASSIGNED", 0), "prio": ("CALL CATEGORY", 0), "problem": ("PROBLEM DESCRIPTION", 0), "spare": ("REQUIRE SPARE", 0),
             "part": ("PART REQUIRED", 0), "doa": ("DESCRIPTION OF DOA", 0), "status": ("CURRENT STATUS OF CALL", 0), "closed": ("CLOSED DATE", 0),
             "faulty": ("SENT FAULTY SPARE", 0)}
IN_SPEC = {"serial": ("SERIAL", 0), "date": ("DATE", 0), "loc": ("LOCATION", 0), "bill": ("BILL NO", 0), "sr": ("SR CALL NO", 0), "part": ("PART NO", 0),
           "awb": ("COURIER AWB", 0), "by": ("RECEIVED BY", 0), "on": ("ON DATE", 0), "remarks": ("REMARKS", 0)}
OUT_SPEC = {"serial": ("SERIAL", 0), "date": ("DATE", 0), "loc": ("LOCATION", 0), "sr": ("SR CALL NO", 0), "ci": ("CI NO", 0), "dev_serial": ("SERIAL NO", 0),
            "part": ("PART", 0), "desc": ("ITEM DESCRIPTION", 0), "part_serial": ("SERIAL NO", 1), "gp": ("GATEPASS", 0), "sent": ("SENT DATE", 0),
            "sent_loc": ("SENT LOCATION", 0), "courier": ("COURIER", 0), "remarks": ("REMARKS", 0)}


def g(row, cm, k):
    i = cm.get(k)
    return row[i] if i is not None and i < len(row) else None


def build_calls(rows, cm, look, as_of):
    out = []
    for src_row, row in rows:
        sr = txt(g(row, cm, "sr"), upper=True)
        if not sr:
            continue
        r = {c[0]: None for c in CALLS}
        flags = []
        r.update(SNAPSHOT_DATE=as_of, SR_ID=sr, SOURCE_ROW=src_row, ZONE=txt(g(row, cm, "zone"), True), SITE=txt(g(row, cm, "site")), SITE_INCHARGE=txt(g(row, cm, "incharge"), True))
        if not re.fullmatch(r"SR\d{10,16}", sr):
            flags.append("SR_ID_FORMAT")
        tk = g(row, cm, "ticket")
        r["ONGC_TICKET_NO"] = txt(tk)
        if not r["ONGC_TICKET_NO"]:
            flags.append("MISSING_ONGC_TICKET")
        r["ONGC_CALL_DATE"], r["CIPL_CALL_DATE"] = to_date(g(row, cm, "ongc_date")), to_date(g(row, cm, "cipl_date"))
        raw_a = txt(g(row, cm, "asset"))
        ak = raw_a.upper() if raw_a else None
        if raw_a and raw_a != ak:
            LOG["Asset ids upper-cased"] += 1
        r["ASSET_KEY"] = ak
        a = look["assets"].get(ak) if (look and ak) else None
        src_make, src_model, src_serial = norm_make(g(row, cm, "make")), txt(g(row, cm, "model")), txt(g(row, cm, "serial"), True)
        if look is not None:
            if a:
                r.update(ASSET_CLASS=a["asset_class"], ASSET_TYPE=a["asset_type"], MAKE=a["make"] or src_make, MODEL=a["model"] or src_model, SERIAL_NO=a["serial_no"] or src_serial)
                if src_serial and a["serial_no"] and src_serial != str(a["serial_no"]).strip().upper():
                    flags.append("SERIAL_DIFFERS_FROM_MASTER")
            else:
                flags.append("ASSET_NOT_IN_MASTER" if ak else "MISSING_ASSET_ID")
        if not a:
            r.update(MAKE=src_make, MODEL=src_model, SERIAL_NO=src_serial)
        cp = g(row, cm, "cpf")
        try:
            r["CPF_NO"] = int(float(cp))
        except (TypeError, ValueError):
            r["CPF_NO"] = None
        if look is not None and r["CPF_NO"] in look["hr"]:
            r["USER_NAME"] = look["hr"][r["CPF_NO"]]
        else:
            r["USER_NAME"] = txt(g(row, cm, "user"))
            if look is not None and r["CPF_NO"]:
                flags.append("CPF_NOT_IN_HR_MASTER")
        r["ENGINEER"], f = norm_engineer(g(row, cm, "engineer"), look["engineers"] if look else None)
        if f:
            flags.append(f)
        pr = txt(g(row, cm, "prio"), True)
        r["PRIORITY"] = pr
        if pr not in ("P1", "P2", "P3"):
            flags.append("PRIORITY_INVALID")
        r["PROBLEM_DESCRIPTION"], r["SPARE_REQUIRED"], r["PART_REQUIRED"], r["SPARE_ISSUE_NOTE"] = (txt(g(row, cm, k)) for k in ("problem", "spare", "part", "doa"))
        st = txt(g(row, cm, "status"), True)
        r["CALL_STATUS"] = st if st in ("OPEN", "CLOSED") else None
        if st not in ("OPEN", "CLOSED"):
            flags.append("STATUS_INVALID")
        r["CLOSED_DATE"] = to_date(g(row, cm, "closed"))
        fs = txt(g(row, cm, "faulty"), True)
        if fs:
            m = re.search(r"(\d{2})/(\d{2})/(\d{4})", fs)
            if m:
                r["FAULTY_SPARE_STATUS"], r["FAULTY_SPARE_SENT_DATE"] = "SENT", dt.date(int(m[3]), int(m[2]), int(m[1]))
                LOG["Faulty-spare sentences parsed into status + date"] += 1
                if r["FAULTY_SPARE_SENT_DATE"] > as_of:
                    flags.append("FAULTY_SPARE_DATE_IN_FUTURE")
            elif "NOT SENT" in fs:
                r["FAULTY_SPARE_STATUS"] = "NOT_SENT"
        # metrics + lifecycle checks
        cd, od = r["CIPL_CALL_DATE"], r["ONGC_CALL_DATE"]
        if r["CALL_STATUS"] == "CLOSED":
            if not r["CLOSED_DATE"]:
                flags.append("CLOSED_WITHOUT_DATE")
            else:
                t = days(r["CLOSED_DATE"], cd)
                if t is not None and t < 0:
                    flags.append("NEGATIVE_TAT")
                else:
                    r["TAT_DAYS"] = t
        elif r["CALL_STATUS"] == "OPEN":
            if r["CLOSED_DATE"]:
                flags.append("OPEN_WITH_CLOSED_DATE")
            r["AGEING_DAYS"] = days(as_of, cd)
        lag = days(cd, od)
        r["LOGGING_LAG_DAYS"] = lag
        if lag is not None and lag < 0:
            flags.append("CIPL_DATE_BEFORE_ONGC_DATE")
        elif lag is not None and lag >= 5:
            flags.append("SLOW_LOGGING_5D_PLUS")
        r["_flags"] = flags
        out.append(r)
    return out


def build_inward(in_rows, in_cm, calls_by_sr, as_of):
    out, seen, lines = [], collections.Counter(), collections.Counter()
    for src_row, row in in_rows:
        sr = txt(g(row, in_cm, "sr"), True)
        if not sr:
            continue
        r = {c[0]: None for c in SPARE_IN}
        flags = []
        # The sheet's SERIAL column is a label for people, not an identity: it is only used for this provisional ID, which
        # reuse_ids() replaces with the record's real, stable ID whenever the database is reachable (see "record identity" below).
        sn = txt(g(row, in_cm, "serial"))
        mid = f"IN-{int(float(sn)):04d}" if sn and re.fullmatch(r"\d+(\.0)?", sn) else f"IN-R{src_row}"
        seen[mid] += 1
        if seen[mid] > 1:
            mid += f"-{seen[mid]}"
        desc = txt(g(row, in_cm, "part"))
        pn = re.search(r"\b\d{5,7}-\d{3}\b", desc or "")
        r.update(SNAPSHOT_DATE=as_of, INWARD_ID=mid, INWARD_DATE=to_date(g(row, in_cm, "date")), LOCATION=txt(g(row, in_cm, "loc"), True), SR_ID=sr,
                 PART_DESCRIPTION=desc, PART_NO=pn.group() if pn else None, BILL_NO=txt(g(row, in_cm, "bill")), COURIER_AWB=txt(g(row, in_cm, "awb")),
                 RECEIVED_BY=txt(g(row, in_cm, "by"), True), RECEIVED_DATE=to_date(g(row, in_cm, "on")), REMARKS=txt(g(row, in_cm, "remarks")), SOURCE_ROW=src_row)
        lines[line_key("SPARE_INWARD", r)] += 1
        if lines[line_key("SPARE_INWARD", r)] > 1:
            flags.append("DUPLICATE_LINE")           # the very same call + part + date appears more than once in the sheet
        c = calls_by_sr.get(sr)
        r["ASSET_KEY"] = c["ASSET_KEY"] if c else None
        if not c:
            flags.append("SR_NOT_IN_CALLS")
        if not r["RECEIVED_DATE"]:
            flags.append("NOT_RECEIVED_YET")
        else:
            t = days(r["RECEIVED_DATE"], r["INWARD_DATE"])
            r["TRANSIT_DAYS"] = t if t is not None and t >= 0 else None
            if t is not None and t < 0:
                flags.append("RECEIVED_BEFORE_DISPATCH")
        r["_flags"] = flags
        out.append(r)
    return out


def build_outward(out_rows, out_cm, calls_by_sr, look, as_of):
    out, seen, lines = [], collections.Counter(), collections.Counter()
    for src_row, row in out_rows:
        sr = txt(g(row, out_cm, "sr"), True)
        if not sr:
            continue
        r = {c[0]: None for c in SPARE_OUT}
        flags = []
        sn = txt(g(row, out_cm, "serial"))          # a label for people only - see build_inward
        mid = f"OUT-{int(float(sn)):04d}" if sn and re.fullmatch(r"\d+(\.0)?", sn) else f"OUT-R{src_row}"
        seen[mid] += 1
        if seen[mid] > 1:
            mid += f"-{seen[mid]}"
        ak = txt(g(row, out_cm, "ci"), True)
        desc = txt(g(row, out_cm, "desc")) or txt(g(row, out_cm, "part"))
        r.update(SNAPSHOT_DATE=as_of, OUTWARD_ID=mid, OUTWARD_DATE=to_date(g(row, out_cm, "date")), LOCATION=txt(g(row, out_cm, "loc"), True), SR_ID=sr,
                 ASSET_KEY=ak, PART_DESCRIPTION=desc, PART_SERIAL_NO=txt(g(row, out_cm, "part_serial"), True), DEVICE_SERIAL_NO=txt(g(row, out_cm, "dev_serial"), True),
                 COURIER=txt(g(row, out_cm, "courier")), GATEPASS_NO=txt(g(row, out_cm, "gp")), SENT_DATE=to_date(g(row, out_cm, "sent")),
                 SENT_LOCATION=txt(g(row, out_cm, "sent_loc"), True), REMARKS=txt(g(row, out_cm, "remarks")), SOURCE_ROW=src_row)
        lines[line_key("SPARE_OUTWARD", r)] += 1
        if lines[line_key("SPARE_OUTWARD", r)] > 1:
            flags.append("DUPLICATE_LINE")           # the very same call + part appears more than once in the sheet
        if sr not in calls_by_sr:
            flags.append("SR_NOT_IN_CALLS")
        if look is not None and ak and ak not in look["assets"]:
            flags.append("ASSET_NOT_IN_MASTER")
        if not r["GATEPASS_NO"]:
            flags.append("NO_GATEPASS")
        if not r["SENT_DATE"]:
            flags.append("NOT_SENT_YET")
        r["_flags"] = flags
        out.append(r)
    return out


def finish_calls(calls, inward_rows, outward_rows, as_of):
    inward = collections.Counter(s["SR_ID"] for s in inward_rows if s["RECEIVED_DATE"])
    per_asset = collections.Counter(c["ASSET_KEY"] for c in calls if c["ASSET_KEY"])
    for c in calls:
        flags = c.pop("_flags")
        if c["PART_REQUIRED"]:
            c["SPARE_STATUS"] = "PART_RECEIVED" if inward.get(c["SR_ID"]) else "PART_PENDING"
            if c["CALL_STATUS"] == "CLOSED" and c["SPARE_STATUS"] == "PART_PENDING":
                flags.append("CLOSED_PART_NOT_RECEIVED")
            elif c["CALL_STATUS"] == "OPEN" and c["SPARE_STATUS"] == "PART_PENDING":
                flags.append("OPEN_AWAITING_PART")
        else:
            c["SPARE_STATUS"] = "NO_SPARE_NEEDED"
        c["REPEAT_CALL_COUNT"] = per_asset.get(c["ASSET_KEY"]) if c["ASSET_KEY"] else None
        if c["REPEAT_CALL_COUNT"] and c["REPEAT_CALL_COUNT"] >= 3:
            flags.append("REPEAT_FAILURE_3PLUS")
        c["DQ_FLAGS"] = "; ".join(flags) or None
    for s in inward_rows + outward_rows:
        s["DQ_FLAGS"] = "; ".join(s.pop("_flags")) or None


# ------------------------------------------------------------------ RMA extraction
RMA_SPEC = {"sr": ("SR. NO", 0), "loc": ("LOCATATION", 0), "make": ("MAKE", 0), "dev": ("SWITCH S/NO", 0), "fault": ("FULTY PART NAME", 0), "fserial": ("FAULTY PART S/NO", 0),
            "call": ("CALL LOG DATE", 0), "case": ("CASH ID", 0), "rma": ("RMA-CALL NO", 0), "recv": ("NEW RECEIVE PART DATE", 0), "new": ("NEW PART S/NO", 0),
            "sent": ("FAULTY PART SENDING DATE", 0), "dc": ("DC NO", 0), "gp": ("GATPASS NO", 0), "remarks": ("REMARK", 0)}
FAULT_FIX = [("SWITECH", "SWITCH"), ("ENGING", "ENGINE"), ("ISSU", "ISSUE")]


def fault_category(item):
    u = (item or "").upper()
    if re.search(r"XFP|SFP|OPTIC", u):
        return "OPTIC"
    if "PSU" in u or "POWER SUPPLY" in u:
        return "PSU"
    if "POWER ON" in u:
        return "POWER_ON"
    if "BOOT" in u:
        return "BOOT"
    if "PORT" in u or "POE" in u:
        return "PORT"
    if "ENGINE" in u or "CARD" in u or "MODULE" in u:
        return "MODULE"
    if "PASSWORD" in u or "SOFTWARE" in u or "IOS" in u:
        return "SOFTWARE"
    if re.search(r"SWITCH|ROUTER|SRX|MX|L2|L3", u):
        return "WHOLE_UNIT"
    return "OTHER"


def clean_rma_no(v):
    s = txt(v)
    if not s:
        return None
    s2 = re.sub(r"^RMA#\s*", "", s, flags=re.I).strip()
    if s2 != s:
        LOG["RMA numbers cleaned (RMA# prefix / spaces)"] += 1
    return s2


def build_rma(rows, cm, look, as_of):
    out, chain, lines = [], {}, collections.Counter()
    for src_row, row in rows:
        fserial = txt(g(row, cm, "fserial"), True)
        if not fserial:
            continue
        r = {c[0]: None for c in RMA}
        flags = []
        sr = txt(g(row, cm, "sr"))           # the sheet's "Sr. no" is a label for people: this is only a provisional ID, see "record identity"
        if sr and re.fullmatch(r"\d+(\.0)?", sr):
            lid = f"RMA-{int(float(sr)):04d}"
        else:
            lid = f"RMA-R{src_row}"
            flags.append("MISALIGNED_ROW_NO_SR")
        make = txt(g(row, cm, "make")) or ""
        parts = re.split(r"[\s-]+", make.upper(), maxsplit=1)
        vendor = parts[0] if parts and parts[0] in ("JUNIPER", "CISCO") else None
        model = re.sub(r"^ROUTER\s+", "", parts[1].strip()) if vendor and len(parts) > 1 else make
        item = txt(g(row, cm, "fault"), True)
        if item:
            for a, b in FAULT_FIX:
                item = re.sub(r"\b" + a + r"\b", b, item)
        rma = clean_rma_no(g(row, cm, "rma"))
        if rma and re.search(r"[/,]", rma):
            flags.append("MULTIPLE_RMA_NUMBERS")
        cl, rc, sn = to_date(g(row, cm, "call")), to_date(g(row, cm, "recv")), to_date(g(row, cm, "sent"))
        raw_sent = g(row, cm, "sent")
        if raw_sent and not sn:
            flags.append("RETURN_DATE_UNPARSEABLE")
        r.update(SNAPSHOT_DATE=as_of, RMA_LINE_ID=lid, RMA_NO=rma, VENDOR_CASE_ID=txt(g(row, cm, "case")), VENDOR=vendor, DEVICE_MODEL=model, LOCATION=txt(g(row, cm, "loc")),
                 DEVICE_SERIAL_NO=txt(g(row, cm, "dev"), True), FAULT_ITEM=item, FAULT_CATEGORY=fault_category(item), FAULTY_PART_SERIAL=fserial,
                 CALL_LOG_DATE=cl, REPLACEMENT_RECEIVED_DATE=rc, REPLACEMENT_PART_SERIAL=txt(g(row, cm, "new"), True), FAULTY_RETURN_DATE=sn,
                 DC_NO=txt(g(row, cm, "dc")), GATEPASS_NO=txt(g(row, cm, "gp")), REMARKS=txt(g(row, cm, "remarks")), SOURCE_ROW=src_row)
        lines[line_key("OEM_RMA", r)] += 1
        if lines[line_key("OEM_RMA", r)] > 1:
            flags.append("DUPLICATE_LINE")           # the same faulty part, logged on the same date, appears more than once in the sheet
        r["RETURN_STATUS"] = "RETURNED" if sn else "PENDING"
        if not sn:
            flags.append("FAULTY_PART_NOT_RETURNED")
        if not rma:
            flags.append("MISSING_RMA_NO")
        if not cl:
            flags.append("MISSING_CALL_DATE")
        t = days(rc, cl)
        if t is not None and t < 0:
            flags.append("NEGATIVE_DAYS_TO_REPLACEMENT")
        else:
            r["DAYS_TO_REPLACEMENT"] = t
        t2 = days(sn, rc)
        if t2 is not None and t2 < 0:
            flags.append("RETURN_BEFORE_RECEIPT")
        else:
            r["DAYS_TO_RETURN"] = t2
        if r["FAULT_CATEGORY"] == "WHOLE_UNIT" and r["REPLACEMENT_PART_SERIAL"]:
            chain[r["FAULTY_PART_SERIAL"]] = r["REPLACEMENT_PART_SERIAL"]
            if r["DEVICE_SERIAL_NO"]:
                chain.setdefault(r["DEVICE_SERIAL_NO"], r["REPLACEMENT_PART_SERIAL"])
        r["_flags"] = flags
        out.append(r)
    # link to the asset master
    serials = look["serials"] if look else {}
    for r in out:
        basis, key = "NONE", None
        for label, s in (("DEVICE_SERIAL", r["DEVICE_SERIAL_NO"]), ("FAULTY_SERIAL", r["FAULTY_PART_SERIAL"]), ("REPLACEMENT_SERIAL", r["REPLACEMENT_PART_SERIAL"])):
            if s and s in serials:
                basis, key = label, serials[s]
                break
        if not key:
            for s in (r["DEVICE_SERIAL_NO"], r["FAULTY_PART_SERIAL"]):
                hop, cur = 0, s
                while cur and cur in chain and hop < 6:
                    cur, hop = chain[cur], hop + 1
                    if cur in serials:
                        basis, key = "CHAIN", serials[cur]
                        break
                if key:
                    break
        r["ASSET_KEY"], r["LINK_BASIS"] = key, basis
        if look is not None and not key:
            r["_flags"].append("DEVICE_NOT_LINKED_TO_ASSET")
        r["DQ_FLAGS"] = "; ".join(r.pop("_flags")) or None
    return out


def serial_history(rma):
    rows = []
    for r in rma:
        if r["FAULT_CATEGORY"] == "WHOLE_UNIT" and r["ASSET_KEY"] and r["REPLACEMENT_PART_SERIAL"]:
            rows.append((r["ASSET_KEY"], r["FAULTY_PART_SERIAL"], r["REPLACEMENT_PART_SERIAL"], r["RMA_NO"], r["REPLACEMENT_RECEIVED_DATE"]))
    return rows


# ------------------------------------------------------------------ workbooks
def write_table(ws, name, rows, min_last=0):
    cols = TABLES[name][0]
    ws.title = name
    names = [c[0] for c in cols]
    style_header(ws, names, [GROUP_FILL[c[1]] for c in cols])
    ws.row_dimensions[1].height = 32
    dtypes = {c[0]: c[3] for c in cols}
    for r, rec in enumerate(rows or [], 2):
        for i, n in enumerate(names, 1):
            v = rec.get(n)
            if v is not None and not (isinstance(v, float) and pd.isna(v)):
                ws.cell(row=r, column=i, value=v)
    wide = {"PROBLEM_DESCRIPTION": 40, "SPARE_REQUIRED": 34, "PART_REQUIRED": 34, "PART_DESCRIPTION": 40, "DQ_FLAGS": 44, "USER_NAME": 28, "REMARKS": 28,
            "SR_ID": 20, "ASSET_KEY": 20, "FAULT_ITEM": 26, "SPARE_ISSUE_NOTE": 28}
    for i, n in enumerate(names, 1):
        ws.column_dimensions[get_column_letter(i)].width = wide.get(n, 16)
    for row in ws.iter_rows(min_row=2, max_row=max(ws.max_row, min_last)):
        for c in row:
            c.font = F_BASE
            t = dtypes[names[c.column - 1]]
            if t == "date":
                c.number_format = DATE_FMT
            elif t == "int":
                c.number_format = "0"
            elif t == "text":
                c.number_format = "@"
    last = max(ws.max_row + 300, min_last)
    for col in names:
        if col in LISTS:
            dv = DataValidation(type="list", formula1='"' + ",".join(LISTS[col]) + '"', allow_blank=True, errorStyle="warning")
            ws.add_data_validation(dv)
            L = get_column_letter(names.index(col) + 1)
            dv.add(f"{L}2:{L}{last}")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(names))}{max(ws.max_row, 2)}"


def readme(ws, title, lines):
    ws.title = "README"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 112
    ws["A1"] = title
    ws["A1"].font = F_TITLE
    for i, (a, b) in enumerate(lines, 3):
        ws.cell(row=i, column=1, value=a).font = F_BOLD
        c = ws.cell(row=i, column=2, value=b)
        c.font, c.alignment = F_BASE, Alignment(wrap_text=True, vertical="top")
        ws.cell(row=i, column=1).alignment = Alignment(vertical="top")


def dictionary_sheets(wb, tables):
    rows = [(t, c[0], c[1], c[2]) for t in tables for c in TABLES[t][0]]
    simple_sheet(wb, "FIELD_DICTIONARY", ["TABLE", "COLUMN", "GROUP", "DESCRIPTION"], rows, [18, 28, 12, 110])
    lk = wb.create_sheet("LOOKUPS")
    lists = {k: v for k, v in LISTS.items()}
    style_header(lk, list(lists))
    for i, vals in enumerate(lists.values(), 1):
        for r, v in enumerate(vals, 2):
            lk.cell(row=r, column=i, value=v).font = F_BASE
        lk.column_dimensions[get_column_letter(i)].width = 24


TRACKER_README = [
    ("Purpose", "Standard structure for the CIPL call tracker: every service call, and every spare part received or sent."),
    ("Sheets", "CALLS (one row per CIPL service request) | SPARE_INWARD (parts received) | SPARE_OUTWARD (faulty parts sent), one row per line."),
    ("Keys", "CALLS.SR_ID, SPARE_INWARD.INWARD_ID, SPARE_OUTWARD.OUTWARD_ID. Upload = insert new, update existing. SR_ID on a spare line links it to its call; ASSET_KEY links to the asset master."),
    ("Status", "CALL_STATUS OPEN / CLOSED. TAT_DAYS = closed - CIPL call date. AGEING_DAYS counts open calls. SPARE_STATUS is derived from SPARE_INWARD."),
    ("Converter-filled", "ASSET_CLASS/TYPE, MAKE, MODEL, SERIAL_NO, USER_NAME (from the asset and HR masters), SPARE_STATUS, FAULTY_SPARE_*, TAT/AGEING/LAG, REPEAT_CALL_COUNT and DQ_FLAGS."),
    ("Manual link", "OEM_RMA_NO on a call ties it to the OEM RMA table."),
    ("Formats", "Dates DD-MMM-YYYY. Asset ids upper case. Blank = unknown (never '-' or 'NA')."),
    ("Refresh", "Give the raw tracker workbook; it is converted with tools/call_tracking.py tracker and loaded into PostgreSQL with change history."),
]
RMA_README = [
    ("Purpose", "Standard structure for the OEM RMA log of Cisco and Juniper routers and switches."),
    ("Sheet", "OEM_RMA - one row per RMA line (a case can have several lines)."),
    ("Key", "RMA_LINE_ID - stable once given; a line is matched by faulty part serial + call date, not by the sheet Sr. no."),
    ("Asset link", "ASSET_KEY is resolved by serial through the asset master, and through the replacement chain when a whole unit was swapped. LINK_BASIS says how."),
    ("Manual link", "CALL_SR_ID ties an RMA to the CIPL service request that caused it."),
    ("Serial history", "Whole-unit swaps also write oem_serial_history (old -> new serial per asset) in the database."),
    ("Formats", "Dates DD-MMM-YYYY. RMA numbers without 'RMA#' or spaces. Blank = unknown."),
    ("Refresh", "Give the raw SD-WAN workbook (only its 'Call Log' sheet is read); converted with tools/call_tracking.py rma."),
]


def build_templates():
    TEMPLATE_DIR.mkdir(exist_ok=True)
    wb = Workbook()
    readme(wb.active, "CIPL Call Tracker - Standard Bulk Upload Template", TRACKER_README)
    write_table(wb.create_sheet(), "CALLS", [], 1000)
    write_table(wb.create_sheet(), "SPARE_INWARD", [], 1000)
    write_table(wb.create_sheet(), "SPARE_OUTWARD", [], 1000)
    dictionary_sheets(wb, ["CALLS", "SPARE_INWARD", "SPARE_OUTWARD"])
    wb.save(TEMPLATE_DIR / TRACKER_TEMPLATE)
    wb = Workbook()
    readme(wb.active, "OEM RMA Log - Standard Bulk Upload Template", RMA_README)
    write_table(wb.create_sheet(), "OEM_RMA", [], 500)
    dictionary_sheets(wb, ["OEM_RMA"])
    wb.save(TEMPLATE_DIR / RMA_TEMPLATE)
    print("Wrote", TEMPLATE_DIR / TRACKER_TEMPLATE, "and", TEMPLATE_DIR / RMA_TEMPLATE)


def structure_check(fname, sheets):
    f = TEMPLATE_DIR / fname
    if not f.exists():
        print("(no template found to check against)")
        return
    wb = load_workbook(f, read_only=True)
    for s in sheets:
        hdr = [c for c in next(wb[s].iter_rows(min_row=1, max_row=1, values_only=True)) if c]
        if hdr != [c[0] for c in TABLES[s][0]]:
            sys.exit(f"Output structure differs from template {fname} in sheet {s} - stopped.")
    wb.close()
    print("Structure check OK against template:", fname)


def dq_sheet(wb, tables):
    rows = []
    for name, recs in tables.items():
        cnt = collections.Counter(f for r in recs for f in (r["DQ_FLAGS"] or "").split("; ") if f)
        rows += [(name, k, v) for k, v in cnt.most_common()]
    simple_sheet(wb, "DATA_QUALITY", ["TABLE", "FLAG", "ROWS"], rows, [18, 44, 10])
    simple_sheet(wb, "CLEANING_LOG", ["CLEANING ACTION", "COUNT"], LOG.most_common(), [70, 12])


# ------------------------------------------------------------------ database
def col_ddl(cols):
    out = []
    for n, _, _, t in cols:
        out.append(f"{n} {'DATE' if t == 'date' else ('INTEGER' if t == 'int' else 'TEXT')}")
    return ", ".join(out)


def db_setup(con):
    con.execute("DROP TABLE IF EXISTS spare_movement CASCADE")
    for name, (cols, key, table) in TABLES.items():
        con.execute(f"CREATE TABLE IF NOT EXISTS {table} ({col_ddl(cols)}, IS_CURRENT INTEGER, FIRST_SEEN_DATE DATE, LAST_SEEN_DATE DATE, PRIMARY KEY ({key}))")
    con.execute("CREATE TABLE IF NOT EXISTS cm_change_log (TABLE_NAME TEXT, SNAPSHOT_DATE DATE, RECORD_KEY TEXT, CHANGE_TYPE TEXT, FIELD TEXT, OLD_VALUE TEXT, NEW_VALUE TEXT)")
    con.execute("CREATE INDEX IF NOT EXISTS ix_cm_log_key ON cm_change_log (TABLE_NAME, RECORD_KEY)")
    con.execute("CREATE TABLE IF NOT EXISTS oem_serial_history (ASSET_KEY TEXT, OLD_SERIAL TEXT, NEW_SERIAL TEXT, RMA_NO TEXT, CHANGE_DATE DATE, PRIMARY KEY (ASSET_KEY, OLD_SERIAL, NEW_SERIAL))")
    con.execute("CREATE INDEX IF NOT EXISTS ix_svc_call_asset ON svc_call (ASSET_KEY)")
    con.execute("""
        DROP VIEW IF EXISTS v_open_calls;
        CREATE VIEW v_open_calls AS SELECT sr_id, ongc_ticket_no, cipl_call_date, ageing_days, asset_key, asset_class, user_name, engineer, problem_description, spare_status
          FROM svc_call WHERE is_current = 1 AND call_status = 'OPEN' ORDER BY ageing_days DESC;
        DROP VIEW IF EXISTS v_calls_awaiting_part;
        CREATE VIEW v_calls_awaiting_part AS SELECT sr_id, cipl_call_date, ageing_days, asset_key, part_required, call_status FROM svc_call
          WHERE is_current = 1 AND spare_status = 'PART_PENDING' ORDER BY cipl_call_date;
        DROP VIEW IF EXISTS v_call_tat_monthly;
        CREATE VIEW v_call_tat_monthly AS SELECT date_trunc('month', cipl_call_date)::date AS month, COUNT(*) AS calls,
          COUNT(*) FILTER (WHERE call_status = 'CLOSED') AS closed, ROUND(AVG(tat_days)::numeric, 1) AS avg_tat_days
          FROM svc_call WHERE is_current = 1 GROUP BY 1 ORDER BY 1;
        DROP VIEW IF EXISTS v_repeat_failures;
        CREATE VIEW v_repeat_failures AS SELECT asset_key, asset_class, model, COUNT(*) AS calls, MIN(cipl_call_date) AS first_call, MAX(cipl_call_date) AS last_call
          FROM svc_call WHERE is_current = 1 GROUP BY 1, 2, 3 HAVING COUNT(*) >= 3 ORDER BY calls DESC;
        DROP VIEW IF EXISTS v_outward_without_gatepass;
        CREATE VIEW v_outward_without_gatepass AS SELECT outward_id, outward_date, sr_id, asset_key, part_description, sent_date FROM spare_outward
          WHERE is_current = 1 AND gatepass_no IS NULL;
        DROP VIEW IF EXISTS v_spare_movement;
        CREATE VIEW v_spare_movement AS
          SELECT 'INWARD' AS direction, inward_id AS movement_id, inward_date AS movement_date, sr_id, asset_key, part_description, part_no, NULL::text AS part_serial_no,
                 received_date AS done_date, bill_no AS reference_no, dq_flags FROM spare_inward WHERE is_current = 1
          UNION ALL
          SELECT 'OUTWARD', outward_id, outward_date, sr_id, asset_key, part_description, NULL, part_serial_no, sent_date, gatepass_no, dq_flags FROM spare_outward WHERE is_current = 1;
        DROP VIEW IF EXISTS v_rma_pending_return;
        CREATE VIEW v_rma_pending_return AS SELECT rma_line_id, rma_no, vendor, device_model, asset_key, fault_item, call_log_date, replacement_received_date
          FROM oem_rma WHERE is_current = 1 AND return_status = 'PENDING' ORDER BY call_log_date;
        DROP VIEW IF EXISTS v_rma_by_year;
        CREATE VIEW v_rma_by_year AS SELECT EXTRACT(year FROM call_log_date)::int AS year, COUNT(*) AS lines, ROUND(AVG(days_to_replacement)::numeric, 1) AS avg_days_to_replace,
          ROUND(AVG(days_to_return)::numeric, 1) AS avg_days_to_return FROM oem_rma WHERE is_current = 1 GROUP BY 1 ORDER BY 1;
        DROP VIEW IF EXISTS v_asset_rma_count;
        CREATE VIEW v_asset_rma_count AS SELECT asset_key, COUNT(*) AS rma_lines, MIN(call_log_date) AS first_rma, MAX(call_log_date) AS last_rma
          FROM oem_rma WHERE is_current = 1 AND asset_key IS NOT NULL GROUP BY 1 ORDER BY 2 DESC;
    """)
    con.commit()


def _cv(v):
    if v is None or (not isinstance(v, (str, dt.date)) and pd.isna(v)):
        return None
    return v.isoformat() if isinstance(v, dt.date) else str(v)


# ------------------------------------------------------------------ record identity (inward / outward / RMA lines)
# A line is the same line when it is the same thing, not when a hand-typed counter in the sheet says so. The sheet's SERIAL / "Sr. no"
# column is a label for people: it can be blank, retyped or renumbered, and keying on it made a harmless clean-up of the sheet look
# like dozens of deletions. Everything that changes over time (received date, bill, AWB, gate pass, sent date, return date, remarks)
# is deliberately outside the identity so it can be updated without making a "new" record.
# IDs are sticky: a line that matches an existing record keeps that record's ID for good; only a genuinely new line is given a
# new one (the next free number, the same scheme the portal uses when someone adds a line by hand).
#   IDENT[name] = (identity columns, columns that group an "edited" line with its record (same thing, one detail corrected),
#                  columns that pair a lone leftover line on each side)
IDENT = {"SPARE_INWARD": (("SR_ID", "PART_DESCRIPTION", "INWARD_DATE"), ("SR_ID", "PART_DESCRIPTION"), ("SR_ID",)),
         "SPARE_OUTWARD": (("SR_ID", "PART_DESCRIPTION"), ("SR_ID", "PART_DESCRIPTION"), ("SR_ID",)),
         "OEM_RMA": (("FAULTY_PART_SERIAL", "CALL_LOG_DATE"), ("FAULTY_PART_SERIAL",), ("RMA_NO",))}
ID_PREFIX = {"SPARE_INWARD": "IN", "SPARE_OUTWARD": "OUT", "OEM_RMA": "RMA"}


def _ival(v):
    return v.isoformat() if isinstance(v, dt.date) else re.sub(r"\s+", " ", str(v or "")).strip().upper()


def line_key(name, rec, cols=None):
    return tuple(_ival(rec.get(c)) for c in (cols or IDENT[name][0]))


def reuse_ids(con, name, records, lock=False, say=None):
    """Give every incoming line the ID of the existing record it is, or a new ID if it is new. Mutates `records`, returns a summary.
    1. exact match on the identity; repeats of one identity pair up in file order
    2. a line that no longer matches but clearly is an edit of an existing one keeps that record (and is reported, so nothing is
       silently re-pointed): same thing with one detail corrected (e.g. a fixed date), or the only leftover line on both sides
    3. anything left is new: next free number after the highest one in the table (portal-created lines included)."""
    if name not in IDENT:
        return {}
    cols, key, table = TABLES[name]
    ident, edit1, edit2 = IDENT[name]
    pre = ID_PREFIX[name]
    need = list(dict.fromkeys(ident + edit1 + edit2))
    cur = con.cursor()
    if lock:
        cur.execute("SELECT pg_advisory_xact_lock(hashtext(%s))", (table,))      # the same lock the portal takes when it adds a line
    cur.execute(f"SELECT {key}, {','.join(need)}, is_current FROM {table} ORDER BY is_current DESC NULLS LAST, source_row NULLS LAST, {key}")
    by_ident, info, current = collections.defaultdict(list), {}, []
    for row in cur.fetchall():
        rid, vals, cur_flag = row[0], dict(zip(need, (_ival(v) for v in row[1:1 + len(need)]))), row[-1]
        by_ident[tuple(vals[c] for c in ident)].append(rid)
        info[rid] = vals
        if cur_flag == 1:
            current.append(rid)
    used, assigned, nth = set(), {}, collections.Counter()
    for i, rec in enumerate(records):                                                   # 1. same line
        k = line_key(name, rec)
        cand = by_ident.get(k, [])
        if nth[k] < len(cand):
            assigned[i] = cand[nth[k]]
            used.add(cand[nth[k]])
        nth[k] += 1
    edited = []
    left_old = [rid for rid in current if rid not in used]                              # 2. same line, edited
    left_new = [i for i in range(len(records)) if i not in assigned]
    for grp in (edit1, edit2):
        olds, news = collections.defaultdict(list), collections.defaultdict(list)
        for rid in left_old:
            olds[tuple(info[rid][c] for c in grp)].append(rid)
        for i in left_new:
            news[line_key(name, records[i], grp)].append(i)
        for g_, ids in olds.items():
            idx = news.get(g_, [])
            if not any(g_) or (grp is edit2 and grp != edit1 and not (len(ids) == 1 and len(idx) == 1)):
                continue                                                               # a blank value never pairs; the weaker key only one-for-one
            for rid, i in zip(ids, idx):
                assigned[i] = rid
                used.add(rid)
                edited.append((rid, line_key(name, records[i], grp), info[rid], records[i]))
        left_old = [rid for rid in left_old if rid not in used]
        left_new = [i for i in left_new if i not in assigned]
    top = max([int(m[1]) for rid in info if (m := re.fullmatch(rf"{pre}-(\d+)", rid))], default=0)
    for i in left_new:                                                                  # 3. new line
        top += 1
        assigned[i] = f"{pre}-{top:04d}"
    for i, rec in enumerate(records):
        rec[key] = assigned[i]
    n_edit = len(edited)
    out = {"kept": len(records) - len(left_new) - n_edit, "edited": [e[0] for e in edited], "new": len(left_new), "gone": len(left_old)}
    if say:
        say(f"  {name}: {out['kept']} lines matched to existing records, {n_edit} edited lines kept their record, {out['new']} new")
        for rid, k, was, now in edited[:25]:
            changed = [f"{c}: '{was[c]}' -> '{_ival(now.get(c))}'" for c in ident if was.get(c) != _ival(now.get(c))]
            say(f"    {rid} {k[0]}: " + ("; ".join(changed) or "detail corrected"))
    return out


def load_table(con, name, records, as_of, force):
    cols, key, table = TABLES[name]
    names = [c[0] for c in cols]
    records = [itam_rules.upper_record(r) for r in records]        # text rule: upper case everywhere
    reuse_ids(con, name, records, lock=True)                       # same line -> same ID, whatever the sheet's SERIAL column says
    cur = con.cursor()
    cur.execute(f"SELECT {','.join(names)}, first_seen_date FROM {table} WHERE is_current = 1")
    existing = {r[names.index(key)]: r for r in cur.fetchall()}
    new = {r[key]: r for r in records}
    if len(new) != len(records):
        sys.exit(f"{name}: duplicate keys in the file - nothing loaded.")
    itam_locks.ensure_tables(con)
    created, archived = itam_locks.apply_records(con, itam_locks.TABLE_DATASET[table], records, key, as_of)   # keep manual edits made in the portal
    gone = set(existing) - set(new) - created                      # portal-created rows are not in files
    if existing and len(gone) > 0.2 * len(existing) and not force:
        sys.exit(f"Load blocked ({name}): {len(gone)} of {len(existing)} current rows are missing (>20%). Check the file or use --force.")
    log = []
    if existing:
        for k in set(new) - set(existing):
            log.append((name, as_of, k, "ADDED", None, None, None))
    for k in gone:
        log.append((name, as_of, k, "REMOVED", None, None, None))
    for k in set(new) & set(existing):
        for i, n in enumerate(names):
            if n in VOLATILE or n == key:
                continue
            a, b = _cv(existing[k][i]), _cv(new[k].get(n))
            if a != b:
                log.append((name, as_of, k, "CHANGED", n, a, b))
    first = {k: v[-1] for k, v in existing.items()}
    allc = names + ["IS_CURRENT", "FIRST_SEEN_DATE", "LAST_SEEN_DATE"]
    upd = ", ".join(f"{c} = EXCLUDED.{c}" for c in allc if c not in (key, "FIRST_SEEN_DATE"))
    sql = f"INSERT INTO {table} ({','.join(allc)}) VALUES ({','.join(['%s'] * len(allc))}) ON CONFLICT ({key}) DO UPDATE SET {upd}"
    rows = [[new[k].get(n) for n in names] + [1, first.get(k) or as_of, as_of] for k in new]
    cur.executemany(sql, rows)
    if gone:
        cur.execute(f"UPDATE {table} SET is_current = 0 WHERE {key} = ANY(%s)", (list(gone),))
    if archived:
        cur.execute(f"UPDATE {table} SET is_current = 0 WHERE {key} = ANY(%s)", (list(archived),))   # portal-archived rows stay archived
    cur.executemany("INSERT INTO cm_change_log VALUES (%s,%s,%s,%s,%s,%s,%s)", log)
    kinds = collections.Counter(x[3] for x in log)
    print(f"  {name}: {len(new)} rows loaded | changes: {dict(kinds) or 'none'}")


def db_load(tables, as_of, force, history=None):
    import master_db
    con = master_db.connect()
    try:
        db_setup(con)
        for name, recs in tables.items():
            load_table(con, name, recs, as_of, force)
        if history:
            con.cursor().executemany("INSERT INTO oem_serial_history VALUES (%s,%s,%s,%s,%s) ON CONFLICT DO NOTHING", history)
        con.commit()
    finally:
        con.close()
    print(f"Database load OK ({as_of}) -> {master_db.DB_LABEL}")


def report():
    import master_db
    con = master_db.connect()
    try:
        db_setup(con)
        for title, q in (("Calls by status", "SELECT call_status, COUNT(*) n FROM svc_call WHERE is_current = 1 GROUP BY 1"),
                         ("Open calls (oldest 8)", "SELECT sr_id, ageing_days, asset_key, problem_description FROM v_open_calls LIMIT 8"),
                         ("Calls awaiting a part", "SELECT COUNT(*) n FROM v_calls_awaiting_part"),
                         ("Monthly volume / TAT", "SELECT * FROM v_call_tat_monthly"),
                         ("Repeat failures", "SELECT * FROM v_repeat_failures LIMIT 8"),
                         ("RMA by year", "SELECT * FROM v_rma_by_year"),
                         ("RMA pending return", "SELECT COUNT(*) n FROM v_rma_pending_return"),
                         ("Change log", "SELECT table_name, change_type, COUNT(*) n FROM cm_change_log GROUP BY 1,2 ORDER BY 1,2")):
            print(f"\n{title}\n{master_db.query_df(con, q).to_string(index=False)}")
    finally:
        con.close()


# ------------------------------------------------------------------ commands
def run_tracker(raw, as_of, out_dir, load, force):
    look = lookups()
    if look is None:
        print("NOTE: database unreachable - asset / HR / engineer checks skipped.")
    h, rows = read_rows(raw, "Daily Call Tracker", 30)
    calls = build_calls(rows, colmap(h, CALL_SPEC), look, as_of)
    hi, irows = read_rows(raw, "Inward", 12)
    ho, orows = read_rows(raw, "Outward", 16)
    if calls and sum(1 for c in calls if not c["ENGINEER"]) > 0.5 * len(calls):
        sys.exit("STOPPED: most calls have no engineer / lookup values. The tracker's formula columns have no cached values "
                 "(the file was saved by a tool that drops them). Open it in Excel, save it, and run again.")
    by_sr = {c["SR_ID"]: c for c in calls}
    inward = build_inward(irows, colmap(hi, IN_SPEC), by_sr, as_of)
    outward = build_outward(orows, colmap(ho, OUT_SPEC), by_sr, look, as_of)
    finish_calls(calls, inward, outward, as_of)
    if look is not None:                      # database reachable: give each line the ID of the record it already is (read-only here)
        import master_db
        idcon = master_db.connect()
        try:
            for nm, recs in (("SPARE_INWARD", inward), ("SPARE_OUTWARD", outward)):
                reuse_ids(idcon, nm, recs, say=print)
        finally:
            idcon.close()
    for label, recs, key in (("CALLS", calls, "SR_ID"), ("SPARE_INWARD", inward, "INWARD_ID"), ("SPARE_OUTWARD", outward, "OUTWARD_ID")):
        d = [r[key] for r in recs]
        if len(d) != len(set(d)):
            sys.exit(f"{label}: duplicate {key} values, e.g. {[k for k, c in collections.Counter(d).items() if c > 1][:5]}")
    structure_check(TRACKER_TEMPLATE, ["CALLS", "SPARE_INWARD", "SPARE_OUTWARD"])
    Path(out_dir).mkdir(exist_ok=True)
    path = Path(out_dir) / f"Call_Tracker_{as_of:%Y-%m-%d}.xlsx"
    wb = Workbook()
    n_open = sum(1 for c in calls if c["CALL_STATUS"] == "OPEN")
    readme(wb.active, "CIPL Call Tracker", [("Snapshot", as_of.strftime("%d-%b-%Y")), ("CALLS", f"{len(calls)} calls ({n_open} open)"),
                                             ("SPARE_INWARD", f"{len(inward)} parts received lines"), ("SPARE_OUTWARD", f"{len(outward)} faulty parts sent lines"),
                                             ("Check first", "DATA_QUALITY and RECONCILIATION.")])
    write_table(wb.create_sheet(), "CALLS", calls)
    write_table(wb.create_sheet(), "SPARE_INWARD", inward)
    write_table(wb.create_sheet(), "SPARE_OUTWARD", outward)
    dq_sheet(wb, {"CALLS": calls, "SPARE_INWARD": inward, "SPARE_OUTWARD": outward})
    rec = [("PART_CALL_WITHOUT_INWARD", c["SR_ID"], c["CALL_STATUS"], c["PART_REQUIRED"]) for c in calls if c["SPARE_STATUS"] == "PART_PENDING"]
    rec += [("INWARD_SR_NOT_IN_CALLS", s["SR_ID"], s["INWARD_ID"], s["PART_DESCRIPTION"]) for s in inward if s["DQ_FLAGS"] and "SR_NOT_IN_CALLS" in s["DQ_FLAGS"]]
    rec += [("OUTWARD_SR_NOT_IN_CALLS", s["SR_ID"], s["OUTWARD_ID"], s["PART_DESCRIPTION"]) for s in outward if s["DQ_FLAGS"] and "SR_NOT_IN_CALLS" in s["DQ_FLAGS"]]
    simple_sheet(wb, "RECONCILIATION", ["TYPE", "SR_ID", "REFERENCE", "DETAIL"], rec, [34, 20, 20, 60])
    dictionary_sheets(wb, ["CALLS", "SPARE_INWARD", "SPARE_OUTWARD"])
    try:
        wb.save(path)
    except PermissionError:
        sys.exit(f"Cannot write {path} - it is open in Excel. Close it and run again.")
    print(f"Wrote {path}: {len(calls)} calls ({n_open} open), {len(inward)} inward lines, {len(outward)} outward lines")
    if load:
        db_load({"CALLS": calls, "SPARE_INWARD": inward, "SPARE_OUTWARD": outward}, as_of, force)


def run_rma(raw, as_of, out_dir, load, force):
    look = lookups()
    if look is None:
        print("NOTE: database unreachable - asset linking skipped.")
    h, rows = read_rows(raw, "Call Log", 16, stop_after=30)
    rma = build_rma(rows, colmap(h, RMA_SPEC), look, as_of)
    if look is not None:                      # database reachable: give each line the ID of the record it already is (read-only here)
        import master_db
        idcon = master_db.connect()
        try:
            reuse_ids(idcon, "OEM_RMA", rma, say=print)
        finally:
            idcon.close()
    ids = [r["RMA_LINE_ID"] for r in rma]
    if len(ids) != len(set(ids)):
        sys.exit(f"Duplicate RMA_LINE_ID values: {[k for k, c in collections.Counter(ids).items() if c > 1][:5]}")
    structure_check(RMA_TEMPLATE, ["OEM_RMA"])
    Path(out_dir).mkdir(exist_ok=True)
    path = Path(out_dir) / f"OEM_RMA_{as_of:%Y-%m-%d}.xlsx"
    hist = serial_history(rma)
    wb = Workbook()
    linked = sum(1 for r in rma if r["ASSET_KEY"])
    readme(wb.active, "OEM RMA Log", [("Snapshot", as_of.strftime("%d-%b-%Y")), ("OEM_RMA", f"{len(rma)} RMA lines, {linked} linked to an asset"),
                                       ("Not read", "The Password sheet of the source workbook is never opened."), ("Check first", "DATA_QUALITY.")])
    write_table(wb.create_sheet(), "OEM_RMA", rma)
    dq_sheet(wb, {"OEM_RMA": rma})
    simple_sheet(wb, "SERIAL_HISTORY", ["ASSET_KEY", "OLD_SERIAL", "NEW_SERIAL", "RMA_NO", "CHANGE_DATE"], hist, [22, 24, 24, 16, 16])
    dictionary_sheets(wb, ["OEM_RMA"])
    try:
        wb.save(path)
    except PermissionError:
        sys.exit(f"Cannot write {path} - it is open in Excel. Close it and run again.")
    print(f"Wrote {path}: {len(rma)} RMA lines, {linked} linked, {len(hist)} serial-history rows")
    if load:
        db_load({"OEM_RMA": rma}, as_of, force, history=hist)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("cmd", nargs="?", choices=["tracker", "rma", "report"])
    ap.add_argument("--raw")
    ap.add_argument("--as-of")
    ap.add_argument("--out-dir", default=str(MASTERS_DIR))
    ap.add_argument("--build-template", action="store_true")
    ap.add_argument("--no-load-db", action="store_true")
    ap.add_argument("--force", action="store_true")
    a = ap.parse_args()
    if a.build_template:
        build_templates()
    if a.cmd == "report":
        report()
    elif a.cmd in ("tracker", "rma"):
        if not a.raw:
            sys.exit("--raw is required")
        d = pd.to_datetime(a.as_of).date() if a.as_of else dt.date.today()
        (run_tracker if a.cmd == "tracker" else run_rma)(a.raw, d, a.out_dir, not a.no_load_db, a.force)
    elif not a.build_template:
        ap.print_help()
