"""Convert the monthly raw HR/SAP export into the standard Employee Master upload template.

  python tools/hr_export_to_template.py --build-template
  python tools/hr_export_to_template.py --raw Ank-manpower-01.09.2026.xlsx [--prev Employee_Master_Upload_2026-08.xlsx]
"""
import argparse
import datetime as dt
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

TEMPLATE_NAME = "Employee_Master_Upload_Template.xlsx"
TEMPLATE_DIR = Path(__file__).resolve().parent.parent / "templates"
MASTER_SHEET = "EMPLOYEE_MASTER"
MAX_ROWS = 5000

# (template column, raw column or None for system-generated, kind, required, description)
FIELDS = [
    ("SNAPSHOT_DATE", None, "date", True, "Date of the HR export this row came from (taken from file name, dd.mm.yyyy)."),
    ("RECORD_STATUS", None, "status", True, "ACTIVE = present in latest export. RETIRED / EXITED = dropped out of a later export (set by the converter)."),
    ("CPF_NO", "CPF NO", "int", True, "Employee number. UNIQUE KEY - one row per employee; upload = insert new, update existing."),
    ("EMPLOYEE_NAME", "NAME", "text", True, "Full name, upper case, single spaces."),
    ("GENDER", "GENDER KEY", "gender", True, "M or F."),
    ("DATE_OF_BIRTH", "DATE OF BIRTH", "date", True, "Date of birth."),
    ("MOBILE_NO", "Mobile No", "mobile", False, "10-digit mobile number, stored as text."),
    ("HOME_STATE", "HOME STATE", "text", False, "Home state."),
    ("CATEGORY", "CATEGORY", "text", False, "Reservation category (GN / OB / SC / ST / EW)."),
    ("PWD_FLAG", "HANDICAP", "flag", False, "Y if person with disability, else N."),
    ("EX_SERVICEMAN_FLAG", "EXSERVICE MAN", "flag", False, "Y if ex-serviceman, else N."),
    ("NATIVE_FLAG", "NATIVE", "flag", False, "Y if local native, else N."),
    ("DESIGNATION", "DESIGNATION TEXT", "text", True, "Current designation."),
    ("DESIG_CODE", "DESIG CODE", "int", False, "SAP designation code (not 1:1 with designation text)."),
    ("LEVEL", "LEVEL", "text", True, "Pay level / grade, e.g. E3, S1, W7."),
    ("CLASS", "CLASS", "int", False, "Class 1-4."),
    ("DISCIPLINE", "DISCIPLINE TEXT", "text", False, "Discipline."),
    ("SUB_DISCIPLINE", "SUB DISP TEXT", "text", False, "Sub-discipline."),
    ("TECH_NON_TECH", "TNT_TEXT", "text", False, "Tech or Non-Tech."),
    ("ORG_UNIT_CODE", "ORG.UNIT", "int", True, "Organisation unit code."),
    ("ORG_UNIT_NAME", "ORG.UNIT TEXT", "text", False, "Organisation unit name."),
    ("POSITION_CODE", "POSITION", "int", False, "Position code (not unique per employee)."),
    ("POSITION_NAME", "POSITION TEXT", "text", False, "Position name."),
    ("ADMINISTRATOR", "ADMINISTRATOR", "text", False, "HR administrator code."),
    ("LOCATION", "LOCATION", "text", False, "Work location."),
    ("QUALIFICATION", "QUAL_TEXT", "text", False, "Highest qualification."),
    ("QUAL_LEVEL", "QUAL LEVEL", "text", False, "Qualification level code (Q1/Q2/Q3/QL...)."),
    ("DATE_OF_JOIN_ONGC", "DATE OF JOIN ONGC", "date", True, "Date of joining ONGC."),
    ("DATE_OF_JOIN_POST", "DATE OF JOIN POST", "date", False, "Date of joining current post (SAP 'join post')."),
    ("DATE_LAST_PROMOTION", "EFF DATE PROM", "date", False, "Effective date of last promotion."),
    ("DATE_OF_JOIN_LOCATION", "DATE OF JOIN PER AREA", "date", False, "Date of joining current personnel area / location."),
    ("DATE_OF_JOIN_POSITION", "DATE OF JOIN POSITION", "date", False, "Date of joining current position."),
    ("DATE_OF_RETIREMENT", "DATE OF RETIREMENT", "date", True, "Date of superannuation."),
]
COLS = [f[0] for f in FIELDS]
SPEC = {f[0]: f for f in FIELDS}
TRACKED = [c for c in COLS if c not in ("SNAPSHOT_DATE", "RECORD_STATUS")]
KEY = "CPF_NO"

DROPPED = {
    "PERSONAL AREA": "Constant (ANKL); LOCATION already identifies the site.",
    "REGION": "Constant for this file.",
    "PARENT": "Constant 0.",
    "RELIGION": "Sensitive personal data, not needed for employee history. Add back only if a statutory report needs it.",
    "RELIGIOUS DENOMINATION": "Code for RELIGION; also has inconsistent values (Islamic / Muslim / Mohammedan).",
    "ONGC_CODE": "Internal SAP code, ambiguous; DESIG_CODE retained.",
    "JOBID": "Internal SAP job id.",
    "REGULAR/TERM BASE": "Constant 1.",
    "R_P_CD": "Constant MRP80.",
    "VERSION": "SAP record version, no HR meaning.",
    "RELAT": "Completely empty.",
    "MPI-1": "SAP internal index.",
    "MPI-2": "SAP internal index.",
    "DISC_CD": "Code duplicate of DISCIPLINE text.",
    "SUB_DISP_CD": "Code duplicate of SUB_DISCIPLINE text.",
    "ZQUALCODE": "Code duplicate of QUALIFICATION text.",
    "STATE CODE": "Code duplicate of HOME_STATE; mixed data types.",
    "TECH/NON TECHNICAL": "Numeric duplicate of TECH_NON_TECH.",
    "BIRTH YEAR": "Derived from DATE_OF_BIRTH.",
}
RAW_KEPT = {f[1]: f[0] for f in FIELDS if f[1]}

HDR_FILL = PatternFill("solid", fgColor="1F3864")
REQ_FILL = PatternFill("solid", fgColor="C00000")
SYS_FILL = PatternFill("solid", fgColor="7F7F7F")
F_BASE = Font(name="Arial", size=10)
F_BOLD = Font(name="Arial", size=10, bold=True)
F_HDR = Font(name="Arial", size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name="Arial", size=14, bold=True)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
DATE_FMT = "DD-MMM-YYYY"


def blank(v):
    if v is None:
        return True
    try:
        if pd.isna(v):
            return True
    except (TypeError, ValueError):
        pass
    return isinstance(v, str) and v.strip() in ("", "//", "NA", "N/A", "-")


def to_date(v):
    if blank(v):
        return None
    if isinstance(v, pd.Timestamp):
        return v.date()
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    d = pd.to_datetime(v, errors="coerce", dayfirst=True)
    return None if pd.isna(d) else d.date()


def clean(kind, v):
    if kind == "date":
        return to_date(v)
    if kind == "flag":
        return "N" if blank(v) else ("Y" if str(v).strip().lower() in ("yes", "y", "1", "true") else "N")
    if blank(v):
        return None
    if kind == "text":
        return re.sub(r"\s+", " ", str(v)).strip()
    if kind == "int":
        try:
            return int(float(v))
        except (TypeError, ValueError):
            return None
    if kind == "gender":
        return str(v).strip().upper()[:1]
    if kind == "mobile":
        digits = re.sub(r"\D", "", str(int(v)) if isinstance(v, float) else str(v))
        return digits[-10:] if len(digits) >= 10 else None
    if kind == "status":
        return str(v).strip().upper()
    return v


def norm(v):
    """Comparable string form for month-over-month diffs."""
    if blank(v):
        return ""
    if isinstance(v, (pd.Timestamp, dt.datetime, dt.date)):
        return to_date(v).isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def snapshot_from_name(path, override):
    if override:
        return pd.to_datetime(override, dayfirst=True).date()
    m = re.search(r"(\d{2})[.\-_](\d{2})[.\-_](\d{4})", Path(path).stem)
    if not m:
        sys.exit("Cannot read snapshot date from file name; pass --snapshot dd-mm-yyyy")
    return dt.date(int(m[3]), int(m[2]), int(m[1]))


def convert(raw_path, snapshot):
    raw = pd.read_excel(raw_path, sheet_name=0)
    raw.columns = [str(c).strip() for c in raw.columns]
    missing = [c for c in RAW_KEPT if c not in raw.columns]
    if missing:
        sys.exit(f"Raw export is missing expected columns: {missing}")
    unknown = [c for c in raw.columns if c not in RAW_KEPT and c not in DROPPED]
    dq = []
    if unknown:
        dq.append(("New raw columns not in mapping (review)", len(unknown), ", ".join(unknown)))
    rows = []
    placeholders = 0
    for _, r in raw.iterrows():
        row = {"SNAPSHOT_DATE": snapshot, "RECORD_STATUS": "ACTIVE"}
        for tcol, rcol, kind, *_ in FIELDS:
            if rcol:
                v = r[rcol]
                if kind == "date" and isinstance(v, str) and v.strip() in ("//", ""):
                    placeholders += 1
                row[tcol] = clean(kind, v)
        rows.append(row)
    df = pd.DataFrame(rows, columns=COLS)

    dq.append(("Raw rows read", len(raw), Path(raw_path).name))
    expected = len(RAW_KEPT) + len(DROPPED)
    dq.append(("Schema check: raw column count", len(raw.columns), f"expected {expected}" + ("" if len(raw.columns) == expected else "  <-- SCHEMA CHANGED, review")))
    amb = df.groupby("DESIG_CODE")["DESIGNATION"].nunique()
    dq.append(("DESIG_CODE mapping to >1 designation text", int((amb > 1).sum()), "Source ambiguity; both kept"))
    dq.append(("Placeholder dates ('//') blanked", placeholders, "Mostly DATE_LAST_PROMOTION"))
    dq.append(("Duplicate CPF_NO", int(df[KEY].duplicated().sum()), "Must be 0"))
    for tcol in [f[0] for f in FIELDS if f[3] and f[0] not in ("SNAPSHOT_DATE", "RECORD_STATUS")]:
        n = int(df[tcol].isna().sum())
        if n:
            dq.append((f"Missing required: {tcol}", n, "CPF: " + ", ".join(map(str, df.loc[df[tcol].isna(), KEY].head(10)))))
    dq.append(("Invalid/absent mobile", int(df["MOBILE_NO"].isna().sum()), "Blank in template"))
    dq.append(("Retirement date on/before snapshot", int((pd.to_datetime(df["DATE_OF_RETIREMENT"]) <= pd.Timestamp(snapshot)).sum()), "Should be exited"))
    dq.append(("Join date before age 18", int((pd.to_datetime(df["DATE_OF_JOIN_ONGC"]) < pd.to_datetime(df["DATE_OF_BIRTH"]) + pd.DateOffset(years=18)).sum()), "Must be 0"))
    return df, dq


def diff_with_previous(new, prev_path, snapshot):
    prev = pd.read_excel(prev_path, sheet_name=MASTER_SHEET, dtype=object)
    prev = prev[[c for c in COLS if c in prev.columns]].copy()
    prev[KEY] = prev[KEY].astype(int)
    pidx = prev.set_index(KEY)
    nidx = new.set_index(KEY)
    log = []
    for cpf in nidx.index.difference(pidx.index):
        log.append((snapshot, cpf, nidx.at[cpf, "EMPLOYEE_NAME"], "JOINER", "", "", "New in export"))
    carried = []
    for cpf in pidx.index.difference(nidx.index):
        old = pidx.loc[cpf]
        if str(old.get("RECORD_STATUS", "ACTIVE")).upper() != "ACTIVE":
            carried.append(_carry(cpf, old, old["RECORD_STATUS"], old["SNAPSHOT_DATE"]))
            continue
        dor = to_date(old["DATE_OF_RETIREMENT"])
        status = "RETIRED" if dor and dor <= snapshot else "EXITED"
        log.append((snapshot, cpf, old["EMPLOYEE_NAME"], "LEAVER", "RECORD_STATUS", "ACTIVE", status))
        carried.append(_carry(cpf, old, status, snapshot))
    for cpf in nidx.index.intersection(pidx.index):
        for c in TRACKED:
            if c == KEY:
                continue
            a, b = norm(pidx.at[cpf, c]), norm(nidx.at[cpf, c])
            if a != b:
                log.append((snapshot, cpf, nidx.at[cpf, "EMPLOYEE_NAME"], "CHANGE", c, a, b))
        if str(pidx.at[cpf, "RECORD_STATUS"]).upper() != "ACTIVE":
            log.append((snapshot, cpf, nidx.at[cpf, "EMPLOYEE_NAME"], "REJOINED", "RECORD_STATUS", str(pidx.at[cpf, "RECORD_STATUS"]), "ACTIVE"))
    kinds = pd.Series([l[3] for l in log]).value_counts().to_dict()
    dq = [(f"Vs previous month: {k}", v, "") for k, v in kinds.items()]
    if not kinds:
        dq.append(("Vs previous month: no changes", 0, ""))
    combined = pd.concat([new, pd.DataFrame(carried, columns=COLS)], ignore_index=True) if carried else new
    return combined.sort_values(KEY).reset_index(drop=True), log, dq


def _carry(cpf, old, status, snap):
    row = {c: (None if blank(old[c]) else old[c]) for c in COLS if c != KEY}
    row[KEY] = int(cpf)
    for c in COLS:
        if SPEC[c][2] == "date" and row.get(c) is not None:
            row[c] = to_date(row[c])
    row["RECORD_STATUS"] = status
    row["SNAPSHOT_DATE"] = to_date(snap)
    return [row.get(c) for c in COLS]


def style_header(ws, ncols, row=1):
    for i in range(1, ncols + 1):
        c = ws.cell(row=row, column=i)
        c.font, c.alignment, c.border = F_HDR, Alignment(horizontal="center", vertical="center", wrap_text=True), BOX


def build_master_sheet(ws, df):
    ws.title = MASTER_SHEET
    for i, (name, _, kind, req, _) in enumerate(FIELDS, 1):
        c = ws.cell(row=1, column=i, value=name)
        c.fill = SYS_FILL if name in ("SNAPSHOT_DATE", "RECORD_STATUS") else (REQ_FILL if req else HDR_FILL)
        ws.column_dimensions[get_column_letter(i)].width = max(14, min(34, len(name) + 4))
    style_header(ws, len(FIELDS))
    ws.row_dimensions[1].height = 30
    for name, w in (("EMPLOYEE_NAME", 32), ("DESIGNATION", 40), ("ORG_UNIT_NAME", 24), ("POSITION_NAME", 38), ("QUALIFICATION", 32)):
        ws.column_dimensions[get_column_letter(COLS.index(name) + 1)].width = w
    if df is not None:
        for r, rec in enumerate(df[COLS].itertuples(index=False), 2):
            for i, v in enumerate(rec, 1):
                if v is not None and not (isinstance(v, float) and pd.isna(v)):
                    ws.cell(row=r, column=i, value=v.item() if hasattr(v, "item") else v)
    last = max(MAX_ROWS, ws.max_row)
    for i, (name, _, kind, *_) in enumerate(FIELDS, 1):
        fmt = DATE_FMT if kind == "date" else ("@" if kind == "mobile" else ("0" if kind == "int" else None))
        for r in range(2, last + 1):
            c = ws.cell(row=r, column=i)
            c.font = F_BASE
            if fmt:
                c.number_format = fmt
    lists = {"GENDER": '"M,F"', "PWD_FLAG": '"Y,N"', "EX_SERVICEMAN_FLAG": '"Y,N"', "NATIVE_FLAG": '"Y,N"',
             "CATEGORY": '"GN,OB,SC,ST,EW"', "TECH_NON_TECH": '"Tech,Non-Tech"',
             "RECORD_STATUS": '"ACTIVE,RETIRED,EXITED"', "CLASS": '"1,2,3,4"'}
    for name, f in lists.items():
        dv = DataValidation(type="list", formula1=f, allow_blank=True, errorStyle="warning",
                            errorTitle="Unexpected value", error="Value is not in the standard list. Continue only if it is a genuinely new code.")
        col = get_column_letter(COLS.index(name) + 1)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{last}")
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(FIELDS))}{max(ws.max_row, 2)}"


def build_change_log(ws, log):
    ws.title = "CHANGE_LOG"
    heads = ["SNAPSHOT_DATE", "CPF_NO", "EMPLOYEE_NAME", "CHANGE_TYPE", "FIELD", "OLD_VALUE", "NEW_VALUE"]
    ws.append(heads)
    for h, w in zip(range(1, 8), (16, 12, 32, 14, 24, 34, 34)):
        ws.column_dimensions[get_column_letter(h)].width = w
    style_header(ws, 7)
    for h in range(1, 8):
        ws.cell(row=1, column=h).fill = HDR_FILL
    for rec in log:
        ws.append(list(rec))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = F_BASE
        row[0].number_format = DATE_FMT
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{max(ws.max_row, 2)}"


def build_mapping(ws):
    ws.title = "FIELD_MAPPING"
    ws["A1"], ws["A1"].font = "Raw HR/SAP column  ->  Template column", F_TITLE
    ws.append([])
    ws.append(["RAW COLUMN", "STATUS", "TEMPLATE COLUMN", "REASON / DESCRIPTION"])
    style_header(ws, 4, 3)
    order = ["CPF NO", "NAME", "DESIGNATION TEXT", "ONGC_CODE", "LEVEL", "CLASS", "DISCIPLINE TEXT", "SUB DISP TEXT", "PERSONAL AREA",
             "LOCATION", "REGION", "ORG.UNIT TEXT", "PARENT", "POSITION TEXT", "RELIGION", "CATEGORY", "GENDER KEY", "QUAL_TEXT",
             "QUAL LEVEL", "HOME STATE", "DATE OF BIRTH", "DATE OF JOIN ONGC", "DATE OF JOIN POST", "EFF DATE PROM", "DATE OF JOIN PER AREA",
             "DATE OF JOIN POSITION", "DATE OF RETIREMENT", "EXSERVICE MAN", "HANDICAP", "NATIVE", "TNT_TEXT", "REGULAR/TERM BASE",
             "STATE CODE", "RELAT", "DESIG CODE", "R_P_CD", "VERSION", "RELIGIOUS DENOMINATION", "TECH/NON TECHNICAL", "DISC_CD",
             "SUB_DISP_CD", "ZQUALCODE", "MPI-1", "MPI-2", "BIRTH YEAR", "ORG.UNIT", "POSITION", "JOBID", "ADMINISTRATOR", "Mobile No"]
    for raw in order:
        if raw in RAW_KEPT:
            t = RAW_KEPT[raw]
            ws.append([raw, "KEPT", t, SPEC[t][4]])
        else:
            ws.append([raw, "DROPPED", "", DROPPED[raw]])
    for t in ("SNAPSHOT_DATE", "RECORD_STATUS"):
        ws.append(["(generated)", "ADDED", t, SPEC[t][4]])
    for row in ws.iter_rows(min_row=4):
        for c in row:
            c.font, c.alignment, c.border = F_BASE, Alignment(wrap_text=True, vertical="top"), BOX
        if row[1].value == "DROPPED":
            row[1].font = Font(name="Arial", size=10, color="C00000", bold=True)
        elif row[1].value == "KEPT":
            row[1].font = Font(name="Arial", size=10, color="007A33", bold=True)
    for col, w in zip("ABCD", (26, 12, 26, 90)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"


def build_lookups(ws):
    ws.title = "LOOKUPS"
    data = {"GENDER": ["M", "F"], "FLAG (Y/N)": ["Y", "N"], "CATEGORY": ["GN", "OB", "SC", "ST", "EW"],
            "TECH_NON_TECH": ["Tech", "Non-Tech"], "RECORD_STATUS": ["ACTIVE", "RETIRED", "EXITED"],
            "CHANGE_TYPE": ["JOINER", "LEAVER", "CHANGE", "REJOINED"], "CLASS": [1, 2, 3, 4]}
    for i, (h, vals) in enumerate(data.items(), 1):
        ws.cell(row=1, column=i, value=h)
        for r, v in enumerate(vals, 2):
            ws.cell(row=r, column=i, value=v).font = F_BASE
        ws.column_dimensions[get_column_letter(i)].width = 18
    style_header(ws, len(data))
    for i in range(1, len(data) + 1):
        ws.cell(row=1, column=i).fill = HDR_FILL


def build_readme(ws, snapshot, dq):
    ws.title = "README"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 100
    ws["A1"], ws["A1"].font = "Employee Master - Standard Bulk Upload Template", F_TITLE
    lines = [
        ("Purpose", "One standard, cleaned layout for employee master + history, built from the monthly raw HR/SAP export (50 columns -> 33)."),
        ("Snapshot in this file", snapshot.strftime("%d-%b-%Y") if snapshot else "(blank template)"),
        ("Upload sheet", f"{MASTER_SHEET}. Headers are in row 1. Do not rename, reorder or delete columns."),
        ("Key", "CPF_NO. One row per employee. Upload = insert new CPF, update existing CPF."),
        ("History", "CHANGE_LOG lists every joiner / leaver / field change vs the previous month. Leavers stay in EMPLOYEE_MASTER with RECORD_STATUS RETIRED or EXITED."),
        ("Monthly steps", "1) Give the new raw export.  2) Run: python tools/hr_export_to_template.py --raw <raw file> --prev <last month's upload file>.  3) Check DATA_QUALITY.  4) Upload EMPLOYEE_MASTER (and CHANGE_LOG)."),
        ("Formats", "Dates DD-MMM-YYYY; flags Y/N; MOBILE_NO 10-digit text; GENDER M/F; blank = unknown (never '//' or 'NA')."),
        ("Header colours", "Red = required. Blue = optional. Grey = set by the converter, do not type."),
        ("Assumptions to confirm with HR", "1) RELIGION dropped as sensitive - add back only if a statutory report needs it.  2) DATE_OF_JOIN_POST is taken to mean joining the current post (vs position).  3) ONGC_CODE and JOBID dropped as SAP internals - confirm no other system needs them.  4) A CPF missing from the next export is treated as RETIRED (if retirement date passed) else EXITED."),
        ("Extra tools", "tools/master_db.py loads each month into a SQLite history DB with views for age / service / retirement, and checks the IT inventory CPFs against the master."),
        ("Example row (illustrative)", "100001 | RAJ KUMAR SHARMA | M | 15-Jun-1980 | 9800000000 | Gujarat | GN | N | N | N | Manager (Mechanical) | 1234 | E4 | 1 | MECHANICAL | MECHANICAL | Tech | 50001033 | ANK E1400-7 | 70145447 | MECHANICAL INCHARGE | RRS | ANKLESHWAR | B.E. (Mechanical) | Q1 | 01-Sep-2005 | 01-Jan-2021 | 01-Jan-2021 | 12-Aug-2022 | 12-Aug-2022 | 30-Jun-2040"),
    ]
    for i, (a, b) in enumerate(lines, 3):
        ws.cell(row=i, column=1, value=a).font = F_BOLD
        c = ws.cell(row=i, column=2, value=b)
        c.font, c.alignment = F_BASE, Alignment(wrap_text=True, vertical="top")
        ws.cell(row=i, column=1).alignment = Alignment(vertical="top")
    if dq:
        r = len(lines) + 5
        ws.cell(row=r, column=1, value="Data quality summary").font = F_BOLD
        for j, (chk, n, note) in enumerate(dq, r + 1):
            ws.cell(row=j, column=1, value=chk).font = F_BASE
            ws.cell(row=j, column=2, value=f"{n}   {note}").font = F_BASE


def build_dq(ws, dq):
    ws.title = "DATA_QUALITY"
    ws.append(["CHECK", "COUNT", "DETAIL"])
    style_header(ws, 3)
    for i in range(1, 4):
        ws.cell(row=1, column=i).fill = HDR_FILL
    for rec in dq:
        ws.append(list(rec))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = F_BASE
    for col, w in zip("ABC", (52, 10, 80)):
        ws.column_dimensions[col].width = w


def exceptions(df):
    checks = [("MOBILE_NO", "Mobile number missing/invalid"), ("HOME_STATE", "Home state missing"),
              ("QUALIFICATION", "Qualification missing"), ("POSITION_NAME", "Position name missing"),
              ("DATE_LAST_PROMOTION", "Last promotion date missing ('//' in source)")]
    out = []
    act = df[df["RECORD_STATUS"] == "ACTIVE"]
    for col, msg in checks:
        for _, r in act[act[col].isna()].iterrows():
            out.append((r[KEY], r["EMPLOYEE_NAME"], r["LEVEL"], col, msg))
    for _, r in act[act[KEY].duplicated(keep=False)].iterrows():
        out.append((r[KEY], r["EMPLOYEE_NAME"], r["LEVEL"], KEY, "Duplicate CPF_NO"))
    return out


def build_exceptions(ws, rows):
    ws.title = "EXCEPTIONS"
    ws.append(["CPF_NO", "EMPLOYEE_NAME", "LEVEL", "FIELD", "ISSUE"])
    style_header(ws, 5)
    for i in range(1, 6):
        ws.cell(row=1, column=i).fill = HDR_FILL
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = F_BASE
    for col, w in zip("ABCDE", (12, 34, 8, 24, 50)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{max(ws.max_row, 2)}"


def write_workbook(path, df=None, log=None, dq=None, snapshot=None):
    wb = Workbook()
    build_readme(wb.active, snapshot, dq)
    build_master_sheet(wb.create_sheet(), df)
    build_change_log(wb.create_sheet(), log or [])
    if dq:
        build_dq(wb.create_sheet(), dq)
    if df is not None:
        build_exceptions(wb.create_sheet(), exceptions(df))
    build_mapping(wb.create_sheet())
    build_lookups(wb.create_sheet())
    try:
        wb.save(path)
    except PermissionError:
        sys.exit(f"Cannot write {path} - it is open in Excel. Close it and run again.")


def latest_previous(folder, snap):
    found = []
    for f in Path(folder).glob("Employee_Master_Upload_????-??.xlsx"):
        y, m = int(f.stem[-7:-3]), int(f.stem[-2:])
        if (y, m) < (snap.year, snap.month):
            found.append((y, m, f))
    return max(found)[2] if found else None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--build-template", action="store_true")
    ap.add_argument("--raw")
    ap.add_argument("--prev")
    ap.add_argument("--snapshot")
    ap.add_argument("--out-dir", default=str(Path(__file__).resolve().parent.parent / "masters"))
    a = ap.parse_args()
    out = Path(a.out_dir)
    if a.build_template:
        TEMPLATE_DIR.mkdir(exist_ok=True)
        write_workbook(TEMPLATE_DIR / TEMPLATE_NAME)
        print("Wrote", TEMPLATE_DIR / TEMPLATE_NAME)
    if a.raw:
        snap = snapshot_from_name(a.raw, a.snapshot)
        df, dq = convert(a.raw, snap)
        log = []
        prev = a.prev or latest_previous(out, snap)
        if prev:
            print("Comparing with previous month:", prev)
            df, log, dq2 = diff_with_previous(df, prev, snap)
            dq += dq2
        path = out / f"Employee_Master_Upload_{snap:%Y-%m}.xlsx"
        write_workbook(path, df, log, dq, snap)
        print("Wrote", path, f"({len(df)} rows, {len(log)} change-log entries)")
        for chk, n, note in dq:
            print(f"  {chk}: {n} {note}")
    if not (a.build_template or a.raw):
        ap.print_help()


if __name__ == "__main__":
    main()
