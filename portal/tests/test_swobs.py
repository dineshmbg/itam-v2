"""Software/OS observations: inspection date is server-stamped and unforgeable, rows are append-only, latest per source, stale flags. Rolled back (conftest.box)."""
import datetime as dt

import psycopg
import pytest

from portal.app import swobs


def test_inspected_at_is_stamped_by_the_server(box):
    before = dt.datetime.now(dt.timezone.utc)
    oid = swobs.record("A1", "BIGFIX_REPORT", {"os_build": "10.0.26200.9445"}, by="TEST")
    row = box.execute("SELECT inspected_at FROM portal_sw_obs WHERE obs_id = %s", (oid,)).fetchone()
    assert row[0] >= before - dt.timedelta(seconds=5)
    with pytest.raises(TypeError):
        swobs.record("A1", "MANUAL", {}, by="TEST", inspected_at=dt.datetime(2020, 1, 1))
    with pytest.raises(swobs.ObsError):
        swobs.record("A1", "MANUAL", {"inspected_at": "2020-01-01"}, by="TEST")


def test_rejects_unknown_source_and_field(box):
    with pytest.raises(swobs.ObsError):
        swobs.record("A1", "GUESS", {}, by="TEST")
    with pytest.raises(swobs.ObsError):
        swobs.record("A1", "MANUAL", {"colour": "red"}, by="TEST")


def test_rows_cannot_be_updated_or_deleted(box):
    oid = swobs.record("A1", "MANUAL", {"vpn_name": "Zscaler"}, by="TEST")
    for sql in ("UPDATE portal_sw_obs SET vpn_name = 'x' WHERE obs_id = %s", "DELETE FROM portal_sw_obs WHERE obs_id = %s"):
        with pytest.raises(psycopg.errors.RaiseException):
            with box.transaction():
                box.execute(sql, (oid,))


def test_latest_per_source_side_by_side_with_stale_flags(box):
    swobs.record("A2", "BIGFIX_REPORT", {"os_build": "10.0.19045.1"}, by="T")
    new = swobs.record("A2", "BIGFIX_REPORT", {"os_build": "10.0.19045.2"}, by="T", data_as_of=dt.datetime.now(dt.timezone.utc) - dt.timedelta(days=45))
    swobs.record("A2", "TREND_REPORT", {"os_build": "19045"}, by="T", data_as_of=dt.datetime.now(dt.timezone.utc) - dt.timedelta(days=100))
    got = {r["source"]: r for r in swobs.latest_by_source("A2")}
    assert set(got) == {"BIGFIX_REPORT", "TREND_REPORT"}
    assert got["BIGFIX_REPORT"]["obs_id"] == new and got["BIGFIX_REPORT"]["os_build"] == "10.0.19045.2"
    assert got["BIGFIX_REPORT"]["stale"] == "STALE" and got["TREND_REPORT"]["stale"] == "VERY_STALE"
    swobs.record("A2", "MANUAL", {}, by="T")
    assert {r["source"]: r for r in swobs.latest_by_source("A2")}["MANUAL"]["stale"] is None


# ---------------------------------------------------------------- report ingest (task 2)
TODAY = dt.date.today()


def _asset(con, key, cls="DESKTOP"):
    con.execute("""INSERT INTO asset (snapshot_date, asset_key, ci_no, asset_class, asset_type, record_level, asset_status, is_current, first_seen_date, last_seen_date)
                   VALUES (%s,%s,%s,%s,%s,'ASSET','IN_USE',1,%s,%s)""", (TODAY, key, key, cls, cls, TODAY, TODAY))


BIGFIX = ("Computer Name,Device Type,IP Address,OS,CPU,Last Report Time\n"
          "ANKZZ-PC1,Desktop,10.0.0.1,Win11 10.0.26200.9445 (25H2),i7,\"Wed, 14 Oct 2026 21:16:35 +0000\"\n"
          "ankzz-pc2.ongc.local,Desktop,10.0.0.2,Win10 10.0.19045.6456 (22H2),i5,\"Thu, 15 Oct 2026 02:14:30 +0000\"\n"
          "ANKNOPE-9,Desktop,10.0.0.3,Win11 10.0.26200.1 (25H2),i5,\"Thu, 15 Oct 2026 02:14:30 +0000\"\n").encode()
TREND = ("Endpoint name,Object type,Protection Manager,Last agent status reported,Agent runtime protection status,Endpoint GUID,Agent connection status,XDR for Endpoints (EDR),Endpoint group,OS name,OS version,Last logged on user,Agent version status\n"
         "ANKZZ-PC1,Endpoint,Standard,Last 24 hours (2026-10-06 10:39:40),Normal,g1,Communicating,Enabled,G,Windows 11,10.0 (Build 26200),ONGC\\68047,controlledLatestVersion\n"
         "ANKZZ-SRV1,Endpoint,Standard,Last 24 hours (2026-10-06 10:39:40),Normal,g2,Not communicating,Enabled,G,Windows Server 2019,10.0 (Build 17763),ONGC\\141662,olderVersion\n").encode()


def test_bigfix_ingest_matches_our_assets_and_stores_full_build(box):
    _asset(box, "ANKZZ-PC1"), _asset(box, "ANKZZ-PC2")
    s = swobs.ingest_report("bigfix.csv", BIGFIX, by="TEST")
    assert (s["product"], s["rows"], s["matched"], s["unmatched"], s["recorded"]) == ("bigfix", 3, 2, 1, 2) and s["unmatched_names"] == ["ANKNOPE-9"] and s["ignored_other_sites"] == 0
    a = {r["source"]: r for r in swobs.latest_by_source("ANKZZ-PC1")}["BIGFIX_REPORT"]
    assert (a["os_edition"], a["os_build"], a["os_feature_version"]) == ("Windows 11", "10.0.26200.9445", "25H2")
    assert a["data_as_of"].astimezone(dt.timezone.utc).isoformat().startswith("2026-10-14T21:16:35")
    assert {r["source"]: r for r in swobs.latest_by_source("ANKZZ-PC2")}["BIGFIX_REPORT"]["os_edition"] == "Windows 10"
    assert not box.execute("SELECT 1 FROM portal_sw_obs WHERE asset_key = 'ANKNOPE-9'").fetchone()


def test_reimport_of_identical_report_adds_nothing_but_new_data_adds_rows(box):
    _asset(box, "ANKZZ-PC1")
    assert swobs.ingest_report("b.csv", BIGFIX, by="T")["recorded"] == 1
    again = swobs.ingest_report("b.csv", BIGFIX, by="T")
    assert (again["recorded"], again["unchanged"]) == (0, 1)
    newer = BIGFIX.replace(b"14 Oct 2026 21:16:35", b"15 Oct 2026 21:16:35")
    assert swobs.ingest_report("b.csv", newer, by="T")["recorded"] == 1
    assert box.execute("SELECT count(*) FROM portal_sw_obs WHERE asset_key = 'ANKZZ-PC1'").fetchone()[0] == 2


def test_trend_ingest_keeps_status_user_and_server_scope(box):
    _asset(box, "ANKZZ-PC1"), _asset(box, "ANKZZ-SRV1", "SERVER"), _asset(box, "ANKZZ-PR1", "PRINTER")
    s = swobs.ingest_report("trend.csv", TREND, by="T")
    assert (s["product"], s["matched"], s["recorded"]) == ("trendmicro", 2, 2)
    t = {r["source"]: r for r in swobs.latest_by_source("ANKZZ-SRV1")}["TREND_REPORT"]
    assert (t["trend_installed"], t["trend_status"], t["connection"], t["last_user"], t["os_build"]) == (True, "olderVersion", "Not communicating", "ONGC\\141662", "17763")
    assert t["os_edition"] == "Windows Server 2019"


def test_unrecognised_file_is_refused(box):
    with pytest.raises(swobs.ObsError):
        swobs.ingest_report("x.csv", b"a,b\n1,2\n", by="T")


def test_names_not_starting_with_ank_are_discarded_not_reported(box):
    _asset(box, "ANKZZ-PC1")
    other = BIGFIX + b"AAAA-INTWS103,Desktop,10.0.0.9,Win11 10.0.26200.9445 (25H2),i7,\"Wed, 14 Oct 2026 21:16:35 +0000\"\n"
    s = swobs.ingest_report("b.csv", other, by="T")
    assert (s["rows"], s["ignored_other_sites"], s["unmatched_names"]) == (4, 1, ["ankzz-pc2.ongc.local", "ANKNOPE-9"])


# ---------------------------------------------------------------- record view (task 3)
def test_asset_record_carries_software_panel_only_for_in_scope_classes(box):
    from portal.app import queries
    _asset(box, "ANKZZ-PC1"), _asset(box, "ANKZZ-PR1", "PRINTER")
    swobs.record("ANKZZ-PC1", "BIGFIX_REPORT", {"os_edition": "Windows 10", "os_build": "10.0.19045.6456"}, by="T")
    swobs.record("ANKZZ-PC1", "TREND_REPORT", {"os_edition": "Windows 11", "os_build": "19045", "last_user": "ONGC\\68047"}, by="T")
    swobs.record("ANKZZ-PR1", "MANUAL", {"remarks": "x"}, by="T")
    sw = queries.detail("assets", "ANKZZ-PC1", user={"role": "ADMIN"})["related"]["software"]
    assert {r["source"] for r in sw["latest"]} == {"BIGFIX_REPORT", "TREND_REPORT"} and len(sw["history"]) == 2
    big = next(r for r in sw["latest"] if r["source"] == "BIGFIX_REPORT")
    assert "Windows 10 is past end of support" in big["flags"] and big["inspected_at"]
    assert [c["field"] for c in sw["conflicts"]] == ["Windows edition"]            # builds agree (19045 both); editions differ
    assert queries.detail("assets", "ANKZZ-PR1", user={"role": "ADMIN"})["related"]["software"] is None


# ---------------------------------------------------------------- inventory match integration (task 4)
def _ank(box, n=3):
    # machines nobody has recorded anything about yet: the dev database holds real observations (every match run records the report), and a test must not depend on them
    rows = [r[0] for r in box.execute("""SELECT ci_no FROM asset a WHERE is_current = 1 AND record_level = 'ASSET' AND asset_class = 'DESKTOP' AND asset_status = 'IN_USE' AND ci_no ~ '^ANK'
                                         AND NOT EXISTS (SELECT 1 FROM portal_sw_obs o WHERE o.asset_key = a.asset_key) ORDER BY ci_no LIMIT %s""", (n,)).fetchall()]
    if len(rows) < n:
        pytest.skip("not enough ANK desktops in the dev database")
    return rows


def test_match_cross_checks_report_against_script_and_manual_observations(box):
    from test_inventory_match import TM, _go, _tm
    a, b, c = _ank(box)
    swobs.record(a, "SCRIPT", {"trend_installed": False}, by="T")           # report will list a, the machine says no agent
    swobs.record(b, "MANUAL", {"trend_installed": True}, by="T")            # report will not list b, an engineer saw the agent
    summary, res = _go([(TM, [_tm(a), _tm(c)])], {"prefix": "ANK"})
    by = {r["ci"]: r for r in res["assets"]}
    assert "not installed" in by[a]["xcheck"] and "Script" in by[a]["xcheck"]
    assert "installed" in by[b]["xcheck"] and "Entered by hand" in by[b]["xcheck"]
    assert "xcheck" not in by[c] and "observed" not in by[c]                  # no observation: the row is exactly as before
    assert summary["xcheck"] >= 2                                             # at least these two; real observations on other machines may add more


def test_match_without_observations_adds_nothing(box):
    from test_inventory_match import TM, _go, _tm
    a, *_ = _ank(box, 1)
    summary, res = _go([(TM, [_tm(a)])], {"prefix": "ANK"})
    assert all("xcheck" not in r and "observed" not in r for r in res["assets"] if r["ci"] == a)


def test_running_a_match_records_the_report_as_observations(box):
    from test_inventory_match import TM, _csv, _tm
    from portal.app import inventory_match as m
    a, *_ = _ank(box, 1)
    st = m.stage("trend.csv", _csv([_tm(a)], TM), "ADMIN1")
    r = m.run(st["stage_id"], st["mapping"], {"prefix": "ANK"}, "ADMIN1")
    assert r["summary"]["observations"]["recorded"] == 1
    assert {x["source"] for x in swobs.latest_by_source(a)} == {"TREND_REPORT"}


# ---------------------------------------------------------------- check script (task 5)
CHECK = """ANKCHECK=1
COMPUTER=ankzz-pc1
WIN_PRODUCT=Windows 10 Pro
WIN_BUILD=26200
WIN_UBR=9445
WIN_DISPLAY=25H2
WIN_INSTALLTYPE=Client
TREND_NAME=Trend Micro Apex One Security Agent
TREND_VER=14.0.12345
BIGFIX_SVC=1
BIGFIX_NAME=BigFix Client
BIGFIX_VER=11.0.2.130
ZSCALER_NAME=Zscaler
ZSCALER_VER=4.3.0.1
VPN_OTHER=
SAP_VER=8.00 PL 5
O365_ACCOUNT=68047@ongc.co.in
OFFICE_VER=16.0.19029.20136
LAST_USER=ONGC\\68047
DOMAIN=ongc.co.in
EDGE_VER=141.0.3537.71
SECUREBOOT=1
RUN_AT_RAW=15-10-2026 20:11:09.55
"""


def test_parse_check_derives_the_facts():
    c = swobs.parse_check(CHECK)
    f = c["facts"]
    assert c["computer"] == "ANKZZ-PC1"
    assert (f["os_edition"], f["os_build"], f["os_feature_version"]) == ("Windows 11", "10.0.26200.9445", "25H2")       # build >= 22000 is Windows 11 whatever ProductName says
    assert (f["trend_installed"], f["trend_version"], f["bigfix_installed"], f["bigfix_version"]) == (True, "14.0.12345", True, "11.0.2.130")
    assert (f["vpn_name"], f["vpn_version"], f["vpn_other"], f["sap_gui_version"], f["o365_account"], f["last_user"], f["domain_joined"]) == ("Zscaler", "4.3.0.1", False, "8.00 PL 5", "68047@ongc.co.in", "ONGC\\68047", True)


def test_parse_check_server_other_vpn_and_nothing_found():
    c = swobs.parse_check("ANKCHECK=1\nCOMPUTER=ANKSRV1\nWIN_PRODUCT=Windows Server 2019 Standard\nWIN_BUILD=17763\nWIN_UBR=1\nWIN_INSTALLTYPE=Server\nVPN_OTHER=FortiClient\n")
    f = c["facts"]
    assert f["os_edition"] == "Windows Server 2019" and f["trend_installed"] is False and f["bigfix_installed"] is False
    assert f["vpn_other"] is True and f["vpn_name"] == "FortiClient (not Zscaler)" and "FortiClient" in f["remarks"]


def test_parse_check_rejects_junk():
    for bad in ("hello", "ANKCHECK=1\n", "ANKCHECK=1\nCOMPUTER=" + "A" * 80, "x" * 100000):
        with pytest.raises(swobs.ObsError):
            swobs.parse_check(bad)


def test_ingest_check_records_a_script_observation_for_our_assets_only(box):
    _asset(box, "ANKZZ-PC1"), _asset(box, "ANKZZ-PR1", "PRINTER")
    r = swobs.ingest_check(CHECK, by="check script")
    assert (r["matched"], r["recorded"]) == (True, 1)
    s = {x["source"]: x for x in swobs.latest_by_source("ANKZZ-PC1")}["SCRIPT"]
    assert s["trend_version"] == "14.0.12345" and s["recorded_by"] == "check script" and "PC clock" in s["remarks"]
    assert swobs.ingest_check(CHECK.replace("ankzz-pc1", "ankzz-pr1"), by="x")["matched"] is False      # printer: out of scope
    assert swobs.ingest_check(CHECK.replace("ankzz-pc1", "ZZZ-OTHER"), by="x")["matched"] is False


def test_check_token_is_stable_rotatable_and_constant_time_checked(box):
    t1 = swobs.check_token(create=True, by="ADMIN")
    assert len(t1) >= 30 and swobs.check_token() == t1 and swobs.token_ok(t1) and not swobs.token_ok("nope") and not swobs.token_ok("")
    t2 = swobs.check_token(create=True, by="ADMIN", rotate=True)
    assert t2 != t1 and not swobs.token_ok(t1) and swobs.token_ok(t2)


def test_script_text_has_the_address_and_token_and_windows_line_endings():
    s = swobs.script_text("https://10.205.64.25", "TOKEN123")
    assert 'set "URL=https://10.205.64.25"' in s and "/api/software/check" in s and "TOKEN123" in s
    # one file only, next to the script (the share); Documents only when that folder cannot be written; no report folder, no scratch files
    assert 'set "OUT=%HERE%ankcheck_%COMPUTERNAME%.txt"' in s and "%~dp0" in s and "\\Documents\\ankcheck_" in s
    assert "REPORTDIR" not in s and "Details.txt" not in s and "Software_Raw" not in s and "mkdir" not in s.lower() and "__DROP__" not in s
    assert s.count("\r\n") > 40 and "\n" not in s.replace("\r\n", "")
    assert "powershell" not in s.lower().replace("no powershell", "")


def test_check_endpoints(box, monkeypatch):
    from starlette.testclient import TestClient
    from conftest import HDR, fake_user
    from portal.app.main import app
    _asset(box, "ANKZZ-PC1")
    tok = swobs.check_token(create=True, by="T")
    c = TestClient(app)
    assert c.post("/api/software/check", content=CHECK).status_code == 403
    assert c.post("/api/software/check", content=CHECK, headers={"X-Check-Token": "wrong"}).status_code == 403
    ok = c.post("/api/software/check", content=CHECK, headers={"X-Check-Token": tok})
    assert ok.status_code == 200 and ok.json()["recorded"] == 1
    assert c.post("/api/software/check", content=CHECK.replace("ankzz-pc1", "ANKNONE"), headers={"X-Check-Token": tok}).status_code == 422
    assert c.post("/api/software/check", content=b"x" * 500000, headers={"X-Check-Token": tok}).status_code == 413
    assert c.post("/api/software/check", content="junk", headers={"X-Check-Token": tok}).status_code == 400
    from portal.app import routes_swobs
    u = fake_user(monkeypatch, role="ADMIN", username="ADM")

    async def me(request):                       # the upload route reads the user itself (raw body), so it needs the same stand-in
        return u
    monkeypatch.setattr(routes_swobs, "current_user", me)
    up = c.post("/api/software/check-upload", content=CHECK, headers={"X-Requested-With": "itam-portal"})
    assert up.status_code == 200 and up.json()["matched"] is True
    dl = c.post("/api/software/script", json={}, headers=HDR)
    assert dl.status_code == 200 and tok in dl.text and "attachment" in dl.headers["content-disposition"]
    u = fake_user(monkeypatch, role="USER", username="U1")
    assert c.post("/api/software/script", json={}, headers=HDR).status_code == 403
    assert c.post("/api/software/check-upload", content=CHECK, headers={"X-Requested-With": "itam-portal"}).status_code == 403


# ---------------------------------------------------------------- manual entry + Excel template (task 6)
def _wb(rows, header=None):
    import io
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(header or swobs.template_headers())
    for r in rows:
        ws.append(r)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_manual_entry_cleans_values_and_records_a_manual_observation(box):
    _asset(box, "ANKZZ-PC1"), _asset(box, "ANKZZ-PR1", "PRINTER")
    oid = swobs.record_manual("ANKZZ-PC1", {"os_edition": "Windows 11", "trend_installed": "Yes", "bigfix_installed": "no", "vpn_name": " Zscaler ", "sap_gui_version": "", "remarks": "checked at desk"}, by="ENG1")
    row = box.execute("SELECT source, recorded_by, trend_installed, bigfix_installed, vpn_name, sap_gui_version FROM portal_sw_obs WHERE obs_id = %s", (oid,)).fetchone()
    assert tuple(row) == ("MANUAL", "ENG1", True, False, "Zscaler", None)
    for bad in ({}, {"colour": "red"}, {"trend_installed": "maybe"}, {"os_build": "x" * 300}):
        with pytest.raises(swobs.ObsError):
            swobs.record_manual("ANKZZ-PC1", bad, by="E")
    with pytest.raises(swobs.ObsError):
        swobs.record_manual("ANKZZ-PR1", {"remarks": "x"}, by="E")          # printer: out of scope


def test_template_headers_roundtrip_and_workbook_all_or_nothing(box):
    _asset(box, "ANKZZ-PC1"), _asset(box, "ANKZZ-PC2")
    h = swobs.template_headers()
    assert h[0] == "Asset (CI)" and "Windows edition" in h and "Trend Micro installed (Yes/No)" in h
    def line(ci, **kv):
        r = [""] * len(h)
        r[0] = ci
        for k, v in kv.items():
            r[h.index(k)] = v
        return r
    good = swobs.ingest_workbook(_wb([line("ANKZZ-PC1", **{"Windows edition": "Windows 11", "Trend Micro installed (Yes/No)": "Yes"}), line("ankzz-pc2", **{"VPN": "Zscaler"})]), by="ENG")
    assert good == {"recorded": 2, "errors": []}
    assert {x["source"] for x in swobs.latest_by_source("ANKZZ-PC2")} == {"MANUAL"}
    before = box.execute("SELECT count(*) FROM portal_sw_obs").fetchone()[0]
    bad = swobs.ingest_workbook(_wb([line("ANKZZ-PC1", **{"Windows edition": "Windows 10"}), line("ANKNOPE", **{"Windows edition": "Windows 10"}), line("ANKZZ-PC2", **{"Trend Micro installed (Yes/No)": "perhaps"})]), by="ENG")
    assert bad["recorded"] == 0 and len(bad["errors"]) == 2 and "Row 3" in bad["errors"][0] and "Row 4" in bad["errors"][1]
    assert box.execute("SELECT count(*) FROM portal_sw_obs").fetchone()[0] == before          # one bad row: nothing recorded
    with pytest.raises(swobs.ObsError):
        swobs.ingest_workbook(_wb([["x"]], header=["Something else"]), by="E")
    with pytest.raises(swobs.ObsError):
        swobs.ingest_workbook(b"not an excel file", by="E")


def test_manual_and_bulk_endpoints_and_permissions(box, monkeypatch):
    from starlette.testclient import TestClient
    from conftest import HDR, fake_user
    from portal.app import routes_swobs
    from portal.app.main import app
    _asset(box, "ANKZZ-PC1")
    box.execute("UPDATE asset SET engineer_name = 'ZZ_ENG' WHERE asset_key = 'ANKZZ-PC1'")
    c = TestClient(app)
    u = fake_user(monkeypatch, role="USER", username="E1", engineer_key="ZZ_ENG")
    ok = c.post("/api/software/manual", json={"key": "ANKZZ-PC1", "facts": {"trend_installed": "Yes"}}, headers=HDR)
    assert ok.status_code == 200 and ok.json()["software"]["latest"][0]["source"] == "MANUAL"
    u2 = fake_user(monkeypatch, role="USER", username="E2", engineer_key="SOMEONE_ELSE")
    assert c.post("/api/software/manual", json={"key": "ANKZZ-PC1", "facts": {"trend_installed": "Yes"}}, headers=HDR).status_code == 403
    fake_user(monkeypatch, role="ADMIN", username="ADM")
    assert c.post("/api/software/manual", json={"key": "ANKZZ-PC1", "facts": {"remarks": "by admin"}}, headers=HDR).status_code == 200
    tpl = c.post("/api/software/template", json={}, headers=HDR)
    assert tpl.status_code == 200 and tpl.content[:2] == b"PK"
    admin = fake_user(monkeypatch, role="ADMIN", username="ADM")

    async def me(request):
        return admin
    monkeypatch.setattr(routes_swobs, "current_user", me)
    h = swobs.template_headers()
    r = [""] * len(h)
    r[0], r[h.index("VPN")] = "ANKZZ-PC1", "Zscaler"
    up = c.post("/api/software/bulk", content=_wb([r]), headers={"X-Requested-With": "itam-portal"})
    assert up.status_code == 200 and up.json() == {"recorded": 1, "errors": []}


def test_engineer_bulk_upload_is_limited_to_their_own_assets(box, monkeypatch):
    from starlette.testclient import TestClient
    from conftest import HDR, fake_user
    from portal.app import routes_swobs
    from portal.app.main import app
    _asset(box, "ANKZZ-PC1"), _asset(box, "ANKZZ-PC2")
    box.execute("UPDATE asset SET engineer_name = 'ZZ_ENG' WHERE asset_key = 'ANKZZ-PC1'")
    box.execute("UPDATE asset SET engineer_name = 'ZZ_OTHER' WHERE asset_key = 'ANKZZ-PC2'")
    h = swobs.template_headers()

    def sheet(*cis):
        rows = []
        for ci in cis:
            r = [""] * len(h)
            r[0], r[h.index("VPN")] = ci, "Zscaler"
            rows.append(r)
        return _wb(rows)
    c = TestClient(app)
    u = fake_user(monkeypatch, role="USER", username="E1", engineer_key="ZZ_ENG")

    async def me(request):
        return u
    monkeypatch.setattr(routes_swobs, "current_user", me)
    H = {"X-Requested-With": "itam-portal"}
    assert c.post("/api/software/template", json={}, headers=HDR).status_code == 200            # any signed-in user may fetch the template
    ok = c.post("/api/software/bulk", content=sheet("ANKZZ-PC1"), headers=H)
    assert ok.status_code == 200 and ok.json()["recorded"] == 1
    bad = c.post("/api/software/bulk", content=sheet("ANKZZ-PC1", "ANKZZ-PC2"), headers=H).json()
    assert bad["recorded"] == 0 and "not assigned to you" in bad["errors"][0] and "Row 3" in bad["errors"][0]
    u = fake_user(monkeypatch, role="USER", username="NOBODY", engineer_key=None)
    monkeypatch.setattr(routes_swobs, "current_user", me)
    assert c.post("/api/software/bulk", content=sheet("ANKZZ-PC1"), headers=H).status_code == 403


# ---------------------------------------------------------------- richer check script: services, updates, installed software
CHECK2 = CHECK.replace("TREND_NAME=Trend Micro Apex One Security Agent\nTREND_VER=14.0.12345\n", "TREND_STATUS=DETECTED\nTREND_NAME=Trend Micro - detected by TmListen service\nTREND_VER=\nTREND_LISTENER_STATE=4\nTREND_SCAN_STATE=1\n") + """WIN_PRODUCT_RAW=Windows 10 Pro
ZSCALER_SERVICE_STATE=4
BIGFIX_SERVICE_STATE=4
O365_STATUS=REGISTRY_IDENTITY_FOUND_NOT_VERIFIED_ACTIVE
O365_ACCOUNT_SOURCE=HKCU\\Software\\Microsoft\\Office\\16.0\\Common\\Identity\\Identities / EmailAddress
TREND_CANDIDATE_NAME=Some Security Agent
TREND_CANDIDATE_VER=1.2
[UPDATES]
Event[0]
  Log Name: System
  Date: 2026-10-17T11:37:16.3800000Z
  Event ID: 19
  Description:
Installation Successful: Windows successfully installed the following update: Security Intelligence Update for Microsoft Defender Antivirus - KB2267602 (Version 1.461.65.0) - Current Channel (Broad)

Event[1]
  Log Name: System
  Date: 2026-10-14T03:10:00.0000000Z
  Event ID: 19
  Description:
Installation Successful: Windows successfully installed the following update: 2026-10 Cumulative Update for Windows 11 Version 25H2 for x64-based Systems (KB5066835)

Event[2]
  Date: not a date
  Description:
nothing useful here
[SOFTWARE]
Source: HKLM\\SOFTWARE\\Microsoft\\Windows\\CurrentVersion\\Uninstall
HKEY_LOCAL_MACHINE\\SOFTWARE\\Microsoft\\Windows\\CurrentVersion\\Uninstall\\{AAA}
    DisplayName    REG_SZ    BigFix Client
    DisplayVersion    REG_SZ    11.0.6.137
    Publisher    REG_SZ    HCL
HKEY_LOCAL_MACHINE\\SOFTWARE\\Microsoft\\Windows\\CurrentVersion\\Uninstall\\AddressBook
Source: HKLM\\SOFTWARE\\WOW6432Node\\Microsoft\\Windows\\CurrentVersion\\Uninstall
HKEY_LOCAL_MACHINE\\SOFTWARE\\WOW6432Node\\Microsoft\\Windows\\CurrentVersion\\Uninstall\\SAP
    DisplayName    REG_SZ    SAP GUI for Windows 8.00
    Publisher    REG_SZ    SAP SE
Source: HKCU\\SOFTWARE\\Microsoft\\Windows\\CurrentVersion\\Uninstall
HKEY_CURRENT_USER\\SOFTWARE\\Microsoft\\Windows\\CurrentVersion\\Uninstall\\Vowen
    DisplayName    REG_SZ    Vowen 0.5.9
    DisplayVersion    REG_SZ    0.5.9
"""


def test_parse_check_reads_services_candidates_updates_and_software():
    c = swobs.parse_check(CHECK2)
    f = c["facts"]
    assert f["trend_installed"] is True and f["trend_services"] == "TmListen running; ntrtscan stopped"
    assert f["vpn_service_state"] == "ZSAService running" and f["bigfix_service_state"] == "BESClient running" and f["os_product_raw"] == "Windows 10 Pro"
    assert f["o365_status"].startswith("REGISTRY_IDENTITY_FOUND") and "EmailAddress" in f["o365_account_source"] and f["trend_candidate"] == "Some Security Agent 1.2"
    assert f["os_update_title"].startswith("2026-10 Cumulative Update") and f["os_update_at"] == "2026-10-14 03:10"            # the OS update, not the newest Defender one
    assert (f["defender_signature"], f["defender_signature_at"]) == ("1.461.65.0", "2026-10-17 11:37")
    assert [(u["kb"], u["at"].day) for u in c["updates"]] == [("KB2267602", 17), ("KB5066835", 14)]                           # the unreadable event is skipped
    assert c["apps"] == [("BigFix Client", "11.0.6.137", "HCL", "HKLM"), ("SAP GUI for Windows 8.00", None, "SAP SE", "HKLM (32-bit)"), ("Vowen 0.5.9", "0.5.9", None, "HKCU")]


def test_ingest_check_stores_updates_once_and_software_only_when_it_changed(box):
    _asset(box, "ANKZZ-PC1")
    r1 = swobs.ingest_check(CHECK2, by="check script")
    assert (r1["recorded"], r1["updates_seen"], r1["software"]) == (1, 2, 3)
    r2 = swobs.ingest_check(CHECK2, by="check script")                                   # same machine, nothing changed
    assert r2["recorded"] == 1
    n = lambda sql: box.execute(sql).fetchone()[0]
    assert n("SELECT count(*) FROM portal_sw_update WHERE asset_key = 'ANKZZ-PC1'") == 2                                  # each update event stored once
    assert n("SELECT count(*) FROM portal_sw_obs_app a JOIN portal_sw_obs o USING (obs_id) WHERE o.asset_key = 'ANKZZ-PC1'") == 3
    rows = box.execute("SELECT obs_id, apps_obs_id, apps_count FROM portal_sw_obs WHERE asset_key = 'ANKZZ-PC1' ORDER BY obs_id").fetchall()
    assert rows[0][1] is None and rows[1][1] == rows[0][0] and rows[1][2] == 3                                           # the second run points at the first run's list
    swobs.ingest_check(CHECK2.replace("Vowen 0.5.9", "Vowen 0.6.0"), by="check script")
    assert n("SELECT count(*) FROM portal_sw_obs_app a JOIN portal_sw_obs o USING (obs_id) WHERE o.asset_key = 'ANKZZ-PC1'") == 6      # changed list: stored again
    for sql in ("UPDATE portal_sw_update SET title = 'x'", "DELETE FROM portal_sw_obs_app"):
        with pytest.raises(psycopg.errors.RaiseException):
            with box.transaction():
                box.execute(sql)


def test_panel_shows_latest_software_and_updates(box):
    _asset(box, "ANKZZ-PC1")
    swobs.ingest_check(CHECK2, by="check script")
    swobs.ingest_check(CHECK2, by="check script")
    p = swobs.panel("ANKZZ-PC1", "DESKTOP")
    assert [a["name"] for a in p["apps"]] == ["BigFix Client", "SAP GUI for Windows 8.00", "Vowen 0.5.9"]                    # found through the pointer of the latest run
    assert [u["kb"] for u in p["updates"]] == ["KB2267602", "KB5066835"]
    assert swobs.panel("ANKZZ-PC1", "PRINTER") is None


# ---------------------------------------------------------------- Chrome / Firefox versions and the software report
def test_parse_check_reads_chrome_and_firefox_and_falls_back_for_edge():
    f = swobs.parse_check(CHECK + "CHROME_VER=155.0.8059.84\nFIREFOX_VER=150.0.1\n")["facts"]
    assert (f["chrome_version"], f["firefox_version"], f["browser_version"]) == ("155.0.8059.84", "150.0.1", "141.0.3537.71")
    g = swobs.parse_check(CHECK.replace("CHROME_VER", "X").replace("EDGE_VER=141.0.3537.71\n", ""))["facts"]
    assert "chrome_version" not in g and "firefox_version" not in g and "browser_version" not in g


def _report_fixture(box):
    _asset(box, "ANKZZ-PC1"), _asset(box, "ANKZZ-PC2"), _asset(box, "ANKZZ-PC3"), _asset(box, "ANKZZ-PR1", "PRINTER")
    box.execute("UPDATE asset SET engineer_name = 'ZZ_ENG', user_name = 'A USER' WHERE asset_key IN ('ANKZZ-PC1', 'ANKZZ-PC2')")
    box.execute("UPDATE asset SET engineer_name = 'ZZ_OTHER' WHERE asset_key = 'ANKZZ-PC3'")
    swobs.ingest_check(CHECK2.replace("COMPUTER=ankzz-pc1", "COMPUTER=ANKZZ-PC1") + "", by="check script")
    swobs.ingest_check(CHECK2.replace("ankzz-pc1", "ankzz-pc3").replace("Vowen 0.5.9", "Vowen 0.6.0"), by="check script")
    swobs.record("ANKZZ-PC2", "MANUAL", {"trend_installed": True, "chrome_version": "154.0.1"}, by="ENG")


def test_software_report_has_machines_installed_software_and_a_version_summary(box):
    _report_fixture(box)
    t = {x["title"]: x for x in swobs.report_tables()}
    assert set(t) == {"Machines", "Installed software", "Software and versions"}
    machines = [r for r in t["Machines"]["rows"] if r["ci"].startswith("ANKZZ-")]
    by = {(r["ci"], r["source"]): r for r in machines}
    assert ("ANKZZ-PC1", "Check script") in by and ("ANKZZ-PC2", "Entered by hand") in by and ("ANKZZ-PC3", "Check script") in by
    assert by[("ANKZZ-PC1", "Check script")]["trend_installed"] == "Yes" and by[("ANKZZ-PC2", "Entered by hand")]["chrome_version"] == "154.0.1"
    assert "ANKZZ-PR1" not in {r["ci"] for r in machines}                                                          # a printer is out of scope
    soft = [r for r in t["Installed software"]["rows"] if r["ci"].startswith("ANKZZ-")]
    assert {(r["ci"], r["program"], r["version"]) for r in soft} >= {("ANKZZ-PC1", "BigFix Client", "11.0.6.137"), ("ANKZZ-PC3", "Vowen 0.6.0", "0.5.9")}
    vow = [r for r in t["Software and versions"]["rows"] if r["program"].startswith("Vowen")]
    assert {(r["program"], r["machines"]) for r in vow} >= {("Vowen 0.5.9", 1), ("Vowen 0.6.0", 1)}
    mine = {x["title"]: x for x in swobs.report_tables("ZZ_ENG")}
    assert {r["ci"] for r in mine["Machines"]["rows"]} == {"ANKZZ-PC1", "ANKZZ-PC2"}                              # an engineer sees their own machines only
    assert {r["ci"] for r in mine["Installed software"]["rows"]} == {"ANKZZ-PC1"}


def test_software_report_endpoint_and_permissions(box, monkeypatch):
    import io
    from openpyxl import load_workbook
    from starlette.testclient import TestClient
    from conftest import HDR, fake_user
    from portal.app.main import app
    _report_fixture(box)
    c = TestClient(app)
    fake_user(monkeypatch, role="ADMIN", username="ADM")
    r = c.post("/api/software/report", json={}, headers=HDR)
    assert r.status_code == 200 and "Software_Report_" in r.headers["content-disposition"]
    names = load_workbook(io.BytesIO(r.content)).sheetnames
    assert {"MACHINES", "INSTALLED SOFTWARE", "SOFTWARE AND VERSIONS"} <= set(names)
    fake_user(monkeypatch, role="USER", username="E1", engineer_key="ZZ_ENG")
    assert c.post("/api/software/report", json={}, headers=HDR).status_code == 200
    fake_user(monkeypatch, role="USER", username="NOBODY", engineer_key=None)
    assert c.post("/api/software/report", json={}, headers=HDR).status_code == 403


def test_control_characters_from_a_pc_are_stripped():
    c = swobs.parse_check(CHECK.replace("EDGE_VER=141.0.3537.71", "EDGE_VER=141.0" + chr(8) + ".35" + chr(0) + "37"))
    assert c["facts"]["browser_version"] == "141.0.3537"


def test_engineer_uploads_result_files_for_their_own_pcs_only(box, monkeypatch):
    from starlette.testclient import TestClient
    from conftest import fake_user
    from portal.app import routes_swobs
    from portal.app.main import app
    _asset(box, "ANKZZ-PC1"), _asset(box, "ANKZZ-PC2")
    box.execute("UPDATE asset SET engineer_name = 'ZZ_ENG' WHERE asset_key = 'ANKZZ-PC1'")
    box.execute("UPDATE asset SET engineer_name = 'ZZ_OTHER' WHERE asset_key = 'ANKZZ-PC2'")
    u = fake_user(monkeypatch, role="USER", username="E1", engineer_key="ZZ_ENG")

    async def me(request):
        return u
    monkeypatch.setattr(routes_swobs, "current_user", me)
    H, c = {"X-Requested-With": "itam-portal"}, TestClient(app)
    ok = c.post("/api/software/check-upload", content=CHECK, headers=H)
    assert ok.status_code == 200 and ok.json()["matched"] is True and ok.json()["recorded"] == 1
    other = c.post("/api/software/check-upload", content=CHECK.replace("ankzz-pc1", "ankzz-pc2"), headers=H).json()
    assert other["matched"] is False and "not assigned to you" in other["reason"]
    assert not box.execute("SELECT 1 FROM portal_sw_obs WHERE asset_key = 'ANKZZ-PC2'").fetchone()
    assert "Not a desktop" in swobs.ingest_check(CHECK.replace("ankzz-pc1", "ZZZ-OTHER"), by="x", engineer_key="ZZ_ENG")["reason"]
    u = fake_user(monkeypatch, role="USER", username="NOBODY", engineer_key=None)
    monkeypatch.setattr(routes_swobs, "current_user", me)
    assert c.post("/api/software/check-upload", content=CHECK, headers=H).status_code == 403
