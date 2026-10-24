"""Detailed Trend Micro investigation: built from the newest analysed reports + PC checks, classified per machine, scoped per engineer, downloadable as Excel / PDF,
mailed automatically after a Trend Micro related analysis. Machines come from the dev database; nothing here depends on specific data. Rolled back (see conftest.box)."""
import io
from collections import Counter

import pytest
from openpyxl import load_workbook
from starlette.testclient import TestClient

from conftest import HDR, fake_user
from portal.app import inventory_match as m
from portal.app import mailer, swobs
from portal.app import trend_investigation as ti
from portal.app.main import app
from test_inventory_match import BF, TM, _bf, _csv, _tm

EXTRA = ("Sensor connectivity", "Endpoint sensor version", "Package", "Protection module version")
EDR, AM = TM.index("XDR for Endpoints (EDR)"), TM.index("Anti-malware")


def _vision(name, am, edr, sensor="Connected", sensor_ver="1.2.0.7045"):
    r = _tm(name, av=am)
    r[EDR] = edr
    return r + [sensor, sensor_ver, "Endpoint Security Essentials|Sandbox Analysis", "14.0.20958"]


@pytest.fixture()
def fleet(box, tmp_path, monkeypatch):
    """Eight desktops of ONE engineer, each given a known place in the centre reports and PC checks."""
    monkeypatch.setattr(m, "STAGE", tmp_path)
    eng = box.execute("""SELECT engineer_name FROM asset WHERE is_current = 1 AND record_level = 'ASSET' AND asset_class = 'DESKTOP' AND asset_status = 'IN_USE' AND engineer_name IS NOT NULL
                         AND ci_no ~ '^ANK' GROUP BY engineer_name HAVING count(*) >= 8 ORDER BY count(*) DESC LIMIT 1""").fetchone()
    if not eng:
        pytest.skip("needs an engineer with at least eight in-use desktops")
    rows = box.execute("""SELECT ci_no, asset_key FROM asset WHERE is_current = 1 AND record_level = 'ASSET' AND asset_class = 'DESKTOP' AND asset_status = 'IN_USE' AND engineer_name = %s
                          AND ci_no ~ '^ANK' ORDER BY ci_no LIMIT 8""", (eng[0],)).fetchall()
    ci = [r[0] for r in rows]
    box.execute("DELETE FROM portal_match_run")                         # only the reports of this test count (rolled back)
    run_ids = {}

    def go(label, name, data, minute):
        st = m.stage(name, data, "ADMIN1")
        r = m.run([st["stage_id"]], st["mapping"], {"prefix": "ANK", "classes": ["DESKTOP", "WORKSTATION", "SERVER"]}, "ADMIN1")
        box.execute("UPDATE portal_match_run SET at = now() + %s * interval '1 minute' WHERE run_id = %s", (minute, r["run_id"]))
        run_ids[label] = r["run_id"]
        return r

    vision = [_vision(ci[0], "Enabled - recommended features", "Enabled"),                       # both
              _vision(ci[1], "Enabled - recommended features", "Disabled", sensor="", sensor_ver=""),   # anti-malware only, no sensor installed
              _vision(ci[2], "Disabled", "Enabled"),                                              # sensor only
              _vision(ci[3], "Disabled", "Disabled", sensor="", sensor_ver=""),                  # neither
              _vision(ci[4], "Enabled - recommended features", "Enabled", sensor="Disconnected")]   # both, but the sensor is disconnected
    go("vision", "Endpoint Inventory_20261006.csv", _csv(vision, TM + EXTRA), 1)
    go("list", "Old_Apex One agent Inventory.csv", _csv([[ci[5] + "$", "Desktop", "Windows 11"]], ("Endpoint name", "Object type", "OS name")), 2)          # a name written with a trailing $
    go("bigfix", "BIGFIX_DESKTOP.csv", _csv([_bf(ci[6])]), 3)
    for key, c, installed in ((rows[6][1], ci[6], False), (rows[7][1], ci[7], True)):
        swobs.record(key, "SCRIPT", {"trend_installed": installed, "trend_version": "14.0.20958" if installed else None, "trend_status": "Apex One Security Agent" if installed else None}, "T", con=box)
    return {"ci": ci, "eng": eng[0], "runs": run_ids}


def _by(d):
    return {x["ci"]: x for x in d["machines"]}


def test_which_analyses_count(fleet):
    assert ti.is_trend_run(fleet["runs"]["vision"]) and ti.is_trend_run(fleet["runs"]["list"])
    assert not ti.is_trend_run(fleet["runs"]["bigfix"])               # BigFix is not a Trend Micro report
    assert not ti.is_trend_run("nosuchrun")


def test_every_machine_gets_its_status_and_protection_state(fleet):
    ci = fleet["ci"]
    by = _by(ti.build())
    want = [("Installed - per Vision One console, not checked on the PC", ti.PROT[0]), ("Installed - per Vision One console, not checked on the PC", ti.PROT[1]),
            (ti.CATS[1], ti.PROT[2]), (ti.CATS[2], ti.PROT[3]), ("Installed - per Vision One console, not checked on the PC", ti.PROT[0]),
            (ti.CATS[5], ti.PROT[4]),                                  # only in the older list - found although the name ends with $
            (ti.CATS[0], ti.PROT[4]),                                  # the PC check found no agent
            (ti.CATS[3], ti.PROT[4])]                                  # the PC check found the agent although no centre report lists it
    for c, (cat, prot) in zip(ci, want):
        assert (by[c]["final"], by[c]["prot"]) == (cat, prot), c
    assert by[ci[4]]["sensor_disc"] and not by[ci[0]]["sensor_disc"]
    assert by[ci[1]]["missing"] == "EDR sensor (not installed)" and by[ci[2]]["missing"] == "Anti-malware" and by[ci[3]]["missing"] == "Anti-malware and EDR sensor"
    assert "no centre report" in by[ci[7]]["notes"] and by[ci[7]]["pc_version"] == "14.0.20958" and by[ci[6]]["notes"] == ""       # a "no agent" note is only for a machine a centre report lists
    assert by[ci[0]]["pc"] == "Not checked" and by[ci[6]]["pc"] == "Not installed" and by[ci[7]]["pc"] == "Installed"


def test_numbers_add_up_and_one_engineer_sees_only_their_own(fleet):
    d = ti.build()
    n = len(d["machines"])
    stat = next(t for t in ti.tables(d) if t["title"] == "Where machines stand")["rows"]
    assert sum(r["machines"] for r in stat[:-1]) == n == stat[-1]["machines"]
    assert sum(1 for r in next(t for t in ti.tables(d) if t["title"] == "Sensor catch-up")["rows"] if r["prot"] == ti.PROT[3]) == Counter(x["prot"] for x in d["machines"])[ti.PROT[3]]
    mine = ti.build(fleet["eng"])
    assert mine["machines"] and all(x["engineer"] == fleet["eng"] for x in mine["machines"]) and set(fleet["ci"]) <= {x["ci"] for x in mine["machines"]}
    titles = [t["title"] for t in ti.tables(mine)]
    assert "By engineer" not in titles and "Centre names not in register" not in titles and "Findings" in titles          # fleet-wide parts are for administrators only
    assert "By engineer" in [t["title"] for t in ti.tables(d)]
    with pytest.raises(ti.InvError) as e:
        ti.build("NO SUCH ENGINEER")
    assert e.value.status == 404


def test_excel_and_pdf_open(fleet):
    d = ti.build(fleet["eng"])
    x, mime, ext = ti.render(d, "xlsx", "T")
    wb = load_workbook(io.BytesIO(x), data_only=True)
    assert ext == "xlsx" and "MACHINE STATUS" in wb.sheetnames and "SENSOR CATCH-UP" in wb.sheetnames and wb["SUMMARY"]["B5"].value == len(d["machines"])
    ci = {r[0] for r in wb["MACHINE STATUS"].iter_rows(min_row=2, values_only=True)}
    assert set(fleet["ci"]) <= ci and len(ci) == len(d["machines"])
    p, mime, ext = ti.render(d, "pdf", "T")
    assert ext == "pdf" and p.startswith(b"%PDF")


def test_download_route_and_who_may_use_it(monkeypatch, fleet):
    u = fake_user(monkeypatch, role="USER", username="E1", engineer_key=fleet["eng"])
    c = TestClient(app)
    post = lambda fmt="xlsx": c.post("/api/investigation/trend", json={"format": fmt}, headers=HDR)      # noqa: E731
    r = post()                                                                                          # an engineer: their own machines
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"] and fleet["eng"].replace(" ", "_") in r.headers["content-disposition"]
    sheet = load_workbook(io.BytesIO(r.content), data_only=True)["MACHINE STATUS"]
    assert {x[3] for x in sheet.iter_rows(min_row=2, values_only=True)} == {fleet["eng"]}
    assert post("pdf").status_code == 200 and post("pdf").content.startswith(b"%PDF") and post("exe").status_code == 400
    u["engineer_key"] = None
    assert post().status_code == 403                                                                    # not linked to an engineer
    u.update(role="ADMIN", group="ADMIN")
    full = post()
    assert full.status_code == 200 and len(load_workbook(io.BytesIO(full.content), data_only=True)["MACHINE STATUS"]["A"]) > 100     # everyone's
    u["read_only"] = True                                                                               # the public-password demo account
    assert post().status_code == 403


def test_no_report_yet_says_so(monkeypatch, box):
    fake_user(monkeypatch, role="ADMIN")
    box.execute("DELETE FROM portal_match_run")
    r = TestClient(app).post("/api/investigation/trend", json={"format": "xlsx"}, headers=HDR)
    assert r.status_code == 404 and "No Trend Micro report" in r.json()["error"]


def test_automatic_mail_after_a_trend_analysis(monkeypatch, fleet):
    sent = []
    monkeypatch.setattr(mailer, "get_smtp", lambda: {**mailer.DEFAULT_SMTP, "enabled": True, "portal_url": ""})
    monkeypatch.setattr(mailer, "send", lambda to, subject, text, html=None, att=(), cfg=None: sent.append((to, subject, [a[0] for a in att], dict((a[0], a[1]) for a in att))))
    monkeypatch.setattr(mailer, "engineer_email", lambda key: "eng@example.com" if key == fleet["eng"] else None)         # everyone else has no address: skipped, not an error
    user = {"username": "ADMIN1", "email": "boss@example.com"}
    assert ti.on_run(fleet["runs"]["bigfix"], user) is None and not sent                              # not a Trend Micro report: nothing happens
    out = ti.on_run(fleet["runs"]["vision"], user)
    assert out["built"] and out["mail_enabled"]
    assert {s[0] for s in sent} == {"eng@example.com", "boss@example.com"}
    for to, subject, names, files in sent:
        assert len(names) == 2 and names[0].endswith(".xlsx") and names[1].endswith(".pdf") and "Trend Micro detailed investigation" in subject
    mine = next(s for s in sent if s[0] == "eng@example.com")
    sheet = load_workbook(io.BytesIO(mine[3][mine[2][0]]), data_only=True)["MACHINE STATUS"]
    assert {x[3] for x in sheet.iter_rows(min_row=2, values_only=True)} == {fleet["eng"]}               # the engineer's copy holds only their machines
    boss = next(s for s in sent if s[0] == "boss@example.com")
    assert len({x[3] for x in load_workbook(io.BytesIO(boss[3][boss[2][0]]), data_only=True)["MACHINE STATUS"].iter_rows(min_row=2, values_only=True)}) > 1          # the person who ran it gets everyone's
    assert m.get(fleet["runs"]["vision"])["investigation"]["built"]                                    # kept with the analysis, so reopening it still says what was done
    n = len(sent)
    ti.on_run(fleet["runs"]["vision"], user)
    assert len(sent) == n                                                                              # a second call mails nobody twice
    sent.clear()
    monkeypatch.setattr(mailer, "get_smtp", lambda: {**mailer.DEFAULT_SMTP, "enabled": False, "portal_url": ""})
    out = ti.on_run(fleet["runs"]["list"], {"username": "ADMIN2", "email": "boss2@example.com"})
    assert out["built"] and not sent and {p["status"] for p in out["plan"]} <= {"MAIL_OFF", "NO_ADDRESS"}      # e-mail switched off: still built, nothing sent


def test_the_analysis_route_makes_the_investigation_by_itself(monkeypatch, fleet):
    from portal.app import routes_tools
    user = fake_user(monkeypatch, role="ADMIN", username="ADMIN1")

    async def me(request):
        return user
    monkeypatch.setattr(routes_tools, "current_user", me)                # the upload route reads the user itself (it takes a raw body)
    monkeypatch.setattr(mailer, "get_smtp", lambda: {**mailer.DEFAULT_SMTP, "enabled": False, "portal_url": ""})
    c = TestClient(app)
    up = lambda data, name: c.post("/api/match/upload", content=data, headers={"X-Requested-With": "itam-portal", "X-Filename": name, "Content-Type": "application/octet-stream"}).json()        # noqa: E731
    run = lambda st: c.post("/api/match/run", json={"stage_ids": [st["stage_id"]], "mapping": st["mapping"], "params": {"prefix": "ANK"}}, headers=HDR)      # noqa: E731
    r = run(up(_csv([_vision(fleet["ci"][0], "Enabled - recommended features", "Enabled")], TM + EXTRA), "Endpoint Inventory_20261007.csv"))
    assert r.status_code == 200 and r.json()["investigation"]["built"] and r.json()["investigation"]["mail_enabled"] is False
    assert m.get(r.json()["run_id"])["investigation"]["built"]
    plain = run(up(_csv([_bf(fleet["ci"][0])]), "PAN INDIA BIGFIX.csv"))                                  # a BigFix report is not a Trend Micro report
    assert plain.status_code == 200 and plain.json().get("investigation") is None


def test_analysis_keeps_what_the_investigation_needs(fleet, box):
    run = m.get(fleet["runs"]["vision"])
    row = next(a for a in run["assets"] if a["ci"] == fleet["ci"][2])
    assert (row["t_am"], row["t_edr"], row["t_pm_ver"], row["t_pkg"]) == ("Disabled", "Enabled", "14.0.20958", "Endpoint Security Essentials")
    assert m.key_of("ANKAON01DT076$") == "ANKAON01DT076" and m.key_of("ankaon01dt076.ongc.net") == "ANKAON01DT076"
    assert m.profile_file(["Domain", "Endpoint", "IP Address"])["mapping"]["name"] == "Endpoint"


def test_version_mismatch_and_duplicate_console_records(fleet, box):
    ci = fleet["ci"]
    key = {r[0]: r[1] for r in box.execute("SELECT ci_no, asset_key FROM asset WHERE is_current = 1 AND record_level = 'ASSET' AND ci_no = ANY(%s)", (ci,)).fetchall()}
    for c, ver in ((ci[0], "14.0.21008"), (ci[1], "14.0.20000"), (ci[2], "14.0.20958")):         # newer than, older than, and equal to the console's 14.0.20958
        swobs.record(key[c], "SCRIPT", {"trend_installed": True, "trend_version": ver, "trend_status": "Apex One Security Agent"}, "T", con=box)
    assert [ti.version_gap(a, b) for a, b in (("14.0.21008", "14.0.20958"), ("14.0.20000", "14.0.20958"), ("14.0.20958", "14.0.20958"), ("", "14.0.20958"), ("14.0.x", "14.0.1"))] == ["newer", "older", "", "", ""]
    g = ti.gather()
    g["dups"][m.key_of(ci[0])] = {"name": ci[0], "rows": 2, "detail": "two records"}
    by = {x["ci"]: x for x in ti.machines(g)}
    assert [by[c]["ver_gap"] for c in ci[:3]] == ["newer", "older", ""]
    assert "Agent version on PC (14.0.21008) differs from console (14.0.20958" in by[ci[0]]["notes"] and "differs" not in by[ci[2]]["notes"]
    assert by[ci[0]]["v1_dup"] and "two records" in by[ci[0]]["notes"] and not by[ci[1]]["v1_dup"] and "two records" not in by[ci[1]]["notes"]
    rows = next(t for t in ti.tables({**ti.build(), "machines": list(by.values())}) if t["title"] == "PC vs centre")["rows"]
    vr = {r["ci"]: r for r in rows if r["t"].startswith("Agent version")}
    assert set(vr) == {ci[0], ci[1]} and "upgraded" in vr[ci[0]]["hint"] and "older agent" in vr[ci[1]]["hint"]
