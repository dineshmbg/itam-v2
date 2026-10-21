"""Remove from AMC. Every test runs in one rolled-back transaction on synthetic assets (TST-*): real data is never touched."""
import datetime as dt

import pytest
from starlette.testclient import TestClient

from portal.app import amcremoval, dashboards, pm
from portal.app.edit import Conflict, Invalid
from portal.app.main import app
from tests.conftest import HDR, fake_user
from tests.test_lifecycle import make_asset, n, one, rows

ED, IP = "TEST EDITOR", "127.0.0.1"
Q = pm.quarter(dt.date.today())


@pytest.fixture()
def fleet(box):
    """Three whole assets of different classes under AMC, plus one laptop and one warranty asset (both ineligible)."""
    box.execute("INSERT INTO portal_user (username, display_name, role, password_hash) VALUES ('TST-APPROVER','TST APPROVER','ADMIN','x') ON CONFLICT DO NOTHING")
    make_asset(box, "TST-D-1", rate_value=1500)
    make_asset(box, "TST-S-1", asset_class="SWITCH", asset_type="CORE SWITCH", rate_value=9000)
    make_asset(box, "TST-U-1", asset_class="UPS", asset_type="UPS", rate_value=400)
    make_asset(box, "TST-L-1", asset_class="LAPTOP", asset_type="LAPTOP")
    make_asset(box, "TST-W-1", cover_type="WRTY")
    return box


def go(keys, **body):
    return amcremoval.remove(keys, {"reason": "SURPLUS", "approved_by": "TST-APPROVER", **body}, ED, IP)


def test_removes_one_asset_by_archiving_it_with_a_snapshot_and_audit(fleet):
    r = go(["TST-D-1"], note="verbal OK from SI")
    assert r["assets"] == ["TST-D-1"] and r["removal_no"].startswith("AMCR-")
    a = one(fleet, "SELECT * FROM asset WHERE asset_key = 'TST-D-1'")
    assert (a["is_current"], a["asset_status"], a["rate_value"]) == (0, "REMOVED_FROM_AMC", None)      # archived, rate cleared, row still there
    assert n(fleet, "SELECT count(*) FROM portal_lock WHERE dataset='assets' AND record_key='TST-D-1' AND field='__archived__'") == 1
    line = one(fleet, "SELECT * FROM amc_removal_line WHERE asset_key = 'TST-D-1'")
    assert line["reason"] == "SURPLUS" and line["snapshot"]["cover_type"] == "AMC" and float(line["snapshot"]["rate_value"]) == 1500
    head = one(fleet, "SELECT * FROM amc_removal WHERE removal_no = %s", (r["removal_no"],))
    assert (head["effective_date"], head["approved_by"], head["status"], head["requested_by"]) == (Q["start"], "TST-APPROVER", "APPLIED", ED)
    assert n(fleet, "SELECT count(*) FROM portal_audit WHERE dataset='assets' AND record_key='TST-D-1' AND action='AMC_REMOVE'") == 1


def test_removed_asset_leaves_the_asset_counts(fleet):
    key = "kpi" if "kpi" in dashboards.assets() else "k"
    before = dashboards.assets()[key]["assets"]
    go(["TST-D-1"])
    assert dashboards.assets()[key]["assets"] == before - 1


def test_one_request_can_mix_classes_and_types(fleet):
    r = go(["TST-D-1", "TST-S-1", "TST-U-1"], reason={"TST-D-1": "SURPLUS", "TST-S-1": "END_OF_LIFE", "TST-U-1": "REPLACED"})
    assert n(fleet, "SELECT count(*) FROM amc_removal_line WHERE removal_no = %s", (r["removal_no"],)) == 3
    assert n(fleet, "SELECT count(*) FROM asset WHERE asset_key LIKE 'TST-%%-1' AND is_current = 0") == 3
    assert one(fleet, "SELECT reason FROM amc_removal_line WHERE asset_key = 'TST-S-1'")["reason"] == "END_OF_LIFE"


@pytest.mark.parametrize("bad,why", [("TST-L-1", "laptops"), ("TST-W-1", "not covered by AMC"), ("TST-NOPE", "not found")])
def test_ineligible_assets_block_the_whole_request_and_nothing_changes(fleet, bad, why):
    with pytest.raises(Conflict, match=why):
        go(["TST-D-1", bad])
    assert one(fleet, "SELECT is_current FROM asset WHERE asset_key = 'TST-D-1'")["is_current"] == 1       # the good one was not removed either
    assert n(fleet, "SELECT count(*) FROM amc_removal WHERE requested_by = %s", (ED,)) == 0


def test_cannot_remove_twice(fleet):
    go(["TST-D-1"])
    with pytest.raises(Conflict, match="already removed"):
        go(["TST-D-1"])


def test_open_call_needs_an_override_note(fleet):
    fleet.execute("INSERT INTO svc_call (sr_id, asset_key, call_status, is_current, cipl_call_date) VALUES ('TST-SR-9','TST-D-1','OPEN',1,%s)", (dt.date.today(),))
    with pytest.raises(Conflict, match="open call"):
        go(["TST-D-1"])
    go(["TST-D-1"], override_note="Call is for a printer fault, not AMC work")
    assert one(fleet, "SELECT is_current FROM asset WHERE asset_key = 'TST-D-1'")["is_current"] == 0


# ---------------------------------------------------------------- T5: component lines leave with their parent
def with_parts(box):
    make_asset(box, "TST-D-1/MON1", record_level="COMPONENT", parent_asset_key="TST-D-1", asset_type="MONITOR", rate_value=189.12)
    make_asset(box, "TST-D-1/GPU1", record_level="COMPONENT", parent_asset_key="TST-D-1", asset_type="GRAPHICS CARD", rate_value=141.84, cover_type="WRTY")


def test_components_leave_with_their_parent_and_their_rates_are_credited(fleet):
    with_parts(fleet)
    r = go(["TST-D-1"])
    assert r["assets"] == ["TST-D-1"] and r["components"] == 2
    got = {x["asset_key"]: x for x in rows(fleet, "SELECT a.asset_key, a.is_current, l.via_parent, l.credit_amount FROM asset a JOIN amc_removal_line l USING (asset_key) WHERE a.asset_key LIKE 'TST-D-1%%'")}
    assert {k: (v["is_current"], v["via_parent"]) for k, v in got.items()} == {"TST-D-1": (0, None), "TST-D-1/MON1": (0, "TST-D-1"), "TST-D-1/GPU1": (0, "TST-D-1")}
    assert str(got["TST-D-1/MON1"]["credit_amount"]) == "189.12"
    assert n(fleet, "SELECT count(*) FROM asset WHERE parent_asset_key = 'TST-D-1' AND is_current = 1") == 0       # no orphan lines left in the register


def test_impact_adds_the_component_rates_to_the_parent(fleet):
    with_parts(fleet)
    a = amcremoval.impact(["TST-D-1"], Q["start"].isoformat())["assets"][0]
    assert (a["components"], str(a["component_amount"]), str(a["amount"])) == (2, "330.96", "1830.96")


def test_a_component_can_be_removed_on_its_own_and_the_parent_stays(fleet):
    with_parts(fleet)
    go(["TST-D-1/MON1"])
    assert one(fleet, "SELECT is_current FROM asset WHERE asset_key = 'TST-D-1'")["is_current"] == 1
    assert one(fleet, "SELECT is_current FROM asset WHERE asset_key = 'TST-D-1/MON1'")["is_current"] == 0
    assert one(fleet, "SELECT is_current FROM asset WHERE asset_key = 'TST-D-1/GPU1'")["is_current"] == 1


def test_choosing_a_component_and_its_parent_together_gives_one_line_each(fleet):
    with_parts(fleet)
    r = go(["TST-D-1", "TST-D-1/MON1"])
    assert n(fleet, "SELECT count(*) FROM amc_removal_line WHERE asset_key = 'TST-D-1/MON1' AND removal_no = %s", (r["removal_no"],)) == 1


def test_a_component_added_after_a_pending_request_still_leaves_and_a_blocked_parent_keeps_its_parts(fleet, monkeypatch):
    d = dt.date.today() + dt.timedelta(days=1)
    if d > Q["end"]:
        pytest.skip("today is the last day of the quarter")
    with_parts(fleet)
    first = go(["TST-D-1"], effective_date=d.isoformat())["removal_no"]
    second = go(["TST-S-1"], effective_date=d.isoformat())["removal_no"]
    make_asset(fleet, "TST-D-1/UPS9", record_level="COMPONENT", parent_asset_key="TST-D-1", asset_type="UPS", rate_value=10)       # arrives after the request
    make_asset(fleet, "TST-S-1/SFP1", record_level="COMPONENT", parent_asset_key="TST-S-1", asset_type="SFP")
    fleet.execute("UPDATE asset SET cover_type = 'WRTY' WHERE asset_key = 'TST-S-1'")                  # the second parent stops being eligible
    monkeypatch.setattr(amcremoval, "TODAY", lambda: d)
    assert amcremoval.apply_due() == [(first, 4, 0), (second, 0, 1)]
    assert one(fleet, "SELECT is_current FROM asset WHERE asset_key = 'TST-D-1/UPS9'")["is_current"] == 0
    assert one(fleet, "SELECT is_current FROM asset WHERE asset_key = 'TST-S-1/SFP1'")["is_current"] == 1       # a parent that stays keeps its parts


# ---------------------------------------------------------------- T6: finding assets to pick
def test_candidates_list_only_eligible_assets_across_classes_with_facets(fleet):
    c = amcremoval.candidates({"q": "TST-"})
    keys = {r["asset_key"] for r in c["rows"]}
    assert keys == {"TST-D-1", "TST-S-1", "TST-U-1"}                                                # laptop and warranty asset are not offered
    assert c["total"] == 3 and set(c["keys"]) == keys and not c["capped"]
    assert {x["value"] for x in c["facets"]["classes"]} >= {"DESKTOP", "SWITCH", "UPS"}
    assert "LAPTOP" not in {x["value"] for x in c["facets"]["classes"]}


def test_candidates_filter_by_class_type_engineer_and_location(fleet):
    fleet.execute("UPDATE asset SET engineer_name = 'TST ENG', location_code = 'TST_LOC' WHERE asset_key IN ('TST-D-1','TST-S-1')")
    pick = lambda **f: {r["asset_key"] for r in amcremoval.candidates({"q": "TST-", **f})["rows"]}  # noqa: E731
    assert pick(classes=["SWITCH", "UPS"]) == {"TST-S-1", "TST-U-1"}                                # several classes at once
    assert pick(types=["CORE SWITCH"]) == {"TST-S-1"}
    assert pick(engineer="tst eng") == {"TST-D-1", "TST-S-1"}
    assert pick(location="TST_LOC", classes=["DESKTOP"]) == {"TST-D-1"}


def test_candidates_hide_component_lines_unless_asked_and_drop_removed_assets(fleet):
    with_parts(fleet)
    ids = lambda **f: {r["asset_key"] for r in amcremoval.candidates({"q": "TST-D-1", **f})["rows"]}  # noqa: E731
    assert ids() == {"TST-D-1"}
    assert ids(components=True) == {"TST-D-1", "TST-D-1/MON1"}                                       # the warranty component is not AMC
    assert amcremoval.candidates({"q": "TST-D-1"})["rows"][0]["components"] == 2
    go(["TST-D-1"])
    assert ids() == set()


def test_candidates_route_is_limited_and_serialises(fleet, monkeypatch):
    fake_user(monkeypatch, "USER")
    with TestClient(app) as c:
        assert c.get("/api/amc-removal/candidates?q=TST-", headers=HDR).status_code == 403
    fake_user(monkeypatch, "ADMIN")
    with TestClient(app) as c:
        r = c.get("/api/amc-removal/candidates?q=TST-&classes=SWITCH,UPS", headers=HDR)
        assert r.status_code == 200 and {x["asset_key"] for x in r.json()["rows"]} == {"TST-S-1", "TST-U-1"}


def test_effective_date_must_fall_in_the_running_quarter(fleet):
    with pytest.raises(Invalid, match="running quarter"):
        go(["TST-D-1"], effective_date=(Q["start"] - dt.timedelta(days=1)).isoformat())
    with pytest.raises(Invalid, match="running quarter"):
        go(["TST-D-1"], effective_date=(Q["end"] + dt.timedelta(days=1)).isoformat())
    no = go(["TST-D-1"], effective_date=dt.date.today().isoformat())["removal_no"]
    assert one(fleet, "SELECT effective_date FROM amc_removal WHERE removal_no = %s", (no,))["effective_date"] == dt.date.today()


def test_approver_must_be_an_active_portal_user_and_reason_is_required(fleet):
    with pytest.raises(Invalid, match="verbal approval"):
        go(["TST-D-1"], approved_by="NOBODY")
    fleet.execute("UPDATE portal_user SET active = FALSE WHERE username = 'TST-APPROVER'")
    with pytest.raises(Invalid, match="verbal approval"):
        go(["TST-D-1"])
    fleet.execute("UPDATE portal_user SET active = TRUE WHERE username = 'TST-APPROVER'")
    with pytest.raises(Invalid, match="reason"):
        go(["TST-D-1"], reason="")
    with pytest.raises(Invalid, match="at least one"):
        go([])


def test_request_numbers_run_on_within_the_quarter(fleet):
    a = go(["TST-D-1"])["removal_no"]
    b = go(["TST-S-1"])["removal_no"]
    assert a[:-3] == b[:-3] and int(b[-3:]) == int(a[-3:]) + 1


def test_route_is_for_administrators_and_team_leaders_only(fleet, monkeypatch):
    body = {"keys": ["TST-D-1"], "reason": "SURPLUS", "approved_by": "TST-APPROVER"}
    fake_user(monkeypatch, "USER")
    with TestClient(app) as c:
        assert c.post("/api/amc-removal/remove", json=body, headers=HDR).status_code == 403
    fake_user(monkeypatch, "ADMIN")
    with TestClient(app) as c:
        r = c.post("/api/amc-removal/remove", json=body, headers=HDR)
        assert r.status_code == 200, r.text
        assert r.json()["assets"] == ["TST-D-1"]


# ---------------------------------------------------------------- T3: billing and future-dated removals
def test_billing_is_pro_rata_on_the_quarterly_rate():
    q = pm.quarter(dt.date(2026, 10, 21))                                   # Q3: 1 Oct - 31 Dec 2026, 92 days
    assert (q["start"], q["end"]) == (dt.date(2026, 10, 1), dt.date(2026, 12, 31))
    full = amcremoval.billing(1500, dt.date(2026, 10, 1))
    assert (full["days"], full["quarter_days"], str(full["amount"])) == (92, 92, "1500.00")      # removed on day 1 = the whole quarter is credited
    assert (amcremoval.billing(1500, dt.date(2026, 11, 1))["days"], str(amcremoval.billing(1500, dt.date(2026, 11, 1))["amount"])) == (61, "994.57")
    assert str(amcremoval.billing(1500, dt.date(2026, 12, 31))["amount"]) == "16.30"                # last day only: 1 of 92 days
    assert amcremoval.billing(None, dt.date(2026, 10, 1))["amount"] is None


def test_credit_is_stored_on_the_line_when_applied(fleet):
    r = go(["TST-D-1", "TST-S-1"], effective_date=Q["start"].isoformat())
    lines = {x["asset_key"]: x for x in rows(fleet, "SELECT * FROM amc_removal_line WHERE removal_no = %s", (r["removal_no"],))}
    qdays = (Q["end"] - Q["start"]).days + 1
    assert (lines["TST-D-1"]["credit_days"], str(lines["TST-D-1"]["credit_amount"])) == (qdays, "1500.00")
    assert str(lines["TST-S-1"]["credit_amount"]) == "9000.00"
    assert [x for x in amcremoval.requests() if x["removal_no"] == r["removal_no"]][0]["credit_total"] == 10500


def test_impact_shows_days_credit_and_total_before_anything_is_changed(fleet):
    out = amcremoval.impact(["TST-D-1", "TST-S-1", "TST-L-1"], Q["start"].isoformat())
    by = {a["asset_key"]: a for a in out["assets"]}
    assert str(by["TST-D-1"]["amount"]) == "1500.00" and by["TST-L-1"]["blocked"] == "laptops are outside AMC scope"
    assert str(out["total"]) == "10500.00" and out["quarter"]["days"] == (Q["end"] - Q["start"]).days + 1
    assert one(fleet, "SELECT is_current FROM asset WHERE asset_key = 'TST-D-1'")["is_current"] == 1        # read only


@pytest.fixture()
def tomorrow():
    d = dt.date.today() + dt.timedelta(days=1)
    if d > Q["end"]:
        pytest.skip("today is the last day of the quarter - no later day to test with")
    return d


def test_a_later_date_waits_and_is_applied_when_it_arrives(fleet, tomorrow, monkeypatch):
    r = go(["TST-D-1"], effective_date=tomorrow.isoformat())
    assert r["status"] == "PENDING"
    assert one(fleet, "SELECT is_current, rate_value FROM asset WHERE asset_key = 'TST-D-1'") == {"is_current": 1, "rate_value": 1500}      # untouched
    assert amcremoval.apply_due() == []                                                                  # its day has not come
    monkeypatch.setattr(amcremoval, "TODAY", lambda: tomorrow)
    assert amcremoval.apply_due() == [(r["removal_no"], 1, 0)]
    assert one(fleet, "SELECT is_current, asset_status FROM asset WHERE asset_key = 'TST-D-1'") == {"is_current": 0, "asset_status": "REMOVED_FROM_AMC"}
    assert one(fleet, "SELECT status FROM amc_removal WHERE removal_no = %s", (r["removal_no"],))["status"] == "APPLIED"
    assert amcremoval.apply_due() == []                                                                  # never applied twice


def test_a_pending_asset_that_became_ineligible_is_skipped_with_the_reason(fleet, tomorrow, monkeypatch):
    r = go(["TST-D-1", "TST-S-1"], effective_date=tomorrow.isoformat())
    fleet.execute("UPDATE asset SET cover_type = 'WRTY' WHERE asset_key = 'TST-S-1'")                   # cover changed in the meantime
    monkeypatch.setattr(amcremoval, "TODAY", lambda: tomorrow)
    assert amcremoval.apply_due() == [(r["removal_no"], 1, 1)]
    assert one(fleet, "SELECT skipped_reason FROM amc_removal_line WHERE asset_key = 'TST-S-1'")["skipped_reason"] == "it is not covered by AMC"
    assert one(fleet, "SELECT is_current FROM asset WHERE asset_key = 'TST-S-1'")["is_current"] == 1


def test_a_pending_request_can_be_cancelled_but_an_applied_one_cannot(fleet, tomorrow):
    r = go(["TST-D-1"], effective_date=tomorrow.isoformat())
    amcremoval.cancel(r["removal_no"], ED, IP)
    assert amcremoval.apply_due() == [] and one(fleet, "SELECT is_current FROM asset WHERE asset_key = 'TST-D-1'")["is_current"] == 1
    done = go(["TST-S-1"])
    with pytest.raises(Conflict, match="not been applied"):
        amcremoval.cancel(done["removal_no"], ED, IP)


# ---------------------------------------------------------------- T4: what an archived asset drops out of
def test_removed_asset_drops_out_of_pm_scope_and_cover_alerts(fleet):
    from portal.app import mailer
    scope = f"SELECT count(*) FROM asset WHERE {pm.IN_SCOPE} AND asset_key = 'TST-D-1'"
    fleet.execute("UPDATE asset SET pm_status = 'PENDING', engineer_name = 'TST ENG', cover_expiry_date = current_date + 10, cover_status = 'EXPIRING_90D' WHERE asset_key = 'TST-D-1'")
    assert n(fleet, scope) == 1
    go(["TST-D-1"])
    assert n(fleet, scope) == 0                                                                           # no longer in the PM list or PM counts
    for rule in ("AMC_EXPIRED", "AMC_EXPIRING"):
        got = mailer.detect(rule, {"lookback_days": 3650, "within_days": 90}, today=dt.date(2026, 10, 19))      # a Monday
        assert not [i for d in got.values() for i in d["rows"] if i["asset_key"] == "TST-D-1"]


def test_the_engineer_is_told_the_asset_left_their_list(fleet):
    from portal.app import mailer
    fleet.execute("UPDATE asset SET engineer_name = 'TST ENG' WHERE asset_key = 'TST-D-1'")
    go(["TST-D-1"])
    got = mailer.detect("ASSET_REMOVED", {"lookback_days": 7}, today=dt.date.today())
    assert [i["asset_key"] for i in got["TST ENG"]["rows"]] == ["TST-D-1"]


def test_the_archive_lock_the_inventory_loader_reads_is_set(fleet):
    import itam_locks
    go(["TST-D-1"])
    _, _, archived = itam_locks.load_overrides(fleet, "assets")
    assert "TST-D-1" in archived                                  # master_db.py keeps these rows is_current = 0 on every re-import


def test_read_routes_serialise_money_and_dates(fleet, monkeypatch):
    fake_user(monkeypatch, "ADMIN")
    with TestClient(app) as c:
        r = c.get(f"/api/amc-removal/impact?keys=TST-D-1,TST-S-1&date={Q['start']}", headers=HDR)
        assert r.status_code == 200, r.text
        assert float(r.json()["total"]) == 10500 and r.json()["assets"][0]["days"] == (Q["end"] - Q["start"]).days + 1
        assert c.post("/api/amc-removal/remove", json={"keys": ["TST-D-1"], "reason": "SURPLUS", "approved_by": "TST-APPROVER"}, headers=HDR).status_code == 200
        rows = c.get("/api/amc-removal/requests", headers=HDR).json()["rows"]
        assert rows[0]["status"] == "APPLIED" and float(rows[0]["credit_total"]) == 1500


# ---------------------------------------------------------------- T7: a pasted or uploaded list
def test_parse_list_splits_on_lines_commas_tabs_and_drops_a_heading():
    assert amcremoval.parse_list("Asset (CI)\n tst-d-1 ,TST-S-1;\tTST-U-1\r\n\nTST-D-1\n'TST-L-1'") == ["TST-D-1", "TST-S-1", "TST-U-1", "TST-L-1"]
    assert amcremoval.parse_list("") == [] and amcremoval.parse_list(None) == []


def test_match_list_finds_by_key_hostname_serial_and_ongc_id_and_explains_the_rest(fleet):
    fleet.execute("UPDATE asset SET hostname = 'TSTHOST-S' WHERE asset_key = 'TST-S-1'")
    fleet.execute("UPDATE asset SET serial_no = 'TSTSER-U' WHERE asset_key = 'TST-U-1'")
    fleet.execute("UPDATE asset SET ongc_asset_id = 'TST-ONGC-D' WHERE asset_key = 'TST-D-1'")
    m = amcremoval.match_list(["tsthost-s", "TSTSER-U", "TST-ONGC-D", "TST-L-1", "TST-W-1", "TST-NOPE", "TST-S-1"])
    by = {r["input"]: r for r in m["rows"]}
    assert by["TSTHOST-S"]["row"]["asset_key"] == "TST-S-1" and by["TSTSER-U"]["row"]["asset_key"] == "TST-U-1" and by["TST-ONGC-D"]["row"]["asset_key"] == "TST-D-1"
    assert by["TST-L-1"]["result"] == "BLOCKED" and "laptops" in by["TST-L-1"]["why"] and by["TST-W-1"]["why"] == "it is not covered by AMC"
    assert by["TST-NOPE"]["result"] == "NOT_FOUND"
    assert by["TST-S-1"]["result"] == "BLOCKED" and "same asset as TSTHOST-S" in by["TST-S-1"]["why"]       # two names for one asset: counted once
    assert (m["ok"], m["blocked"], m["not_found"]) == (3, 3, 1)


def test_match_list_recognises_an_asset_that_was_already_removed_and_ambiguous_names(fleet):
    go(["TST-D-1"])
    assert amcremoval.match_list(["TST-D-1"])["rows"][0]["why"] == "it is already removed from AMC"
    fleet.execute("UPDATE asset SET hostname = 'TSTSAME' WHERE asset_key IN ('TST-S-1','TST-U-1')")
    r = amcremoval.match_list(["TSTSAME"])["rows"][0]
    assert r["result"] == "AMBIGUOUS" and "TST-S-1" in r["why"] and "TST-U-1" in r["why"]


def test_list_from_file_reads_text_and_the_first_column_of_a_workbook(fleet):
    import io
    from openpyxl import Workbook
    assert amcremoval.list_from_file(b"CI\nTST-D-1\nTST-S-1\n", "a.csv") == ["TST-D-1", "TST-S-1"]
    wb = Workbook()
    wb.active.append(["Asset (CI)", "Remarks"])
    wb.active.append(["TST-U-1", "x"])
    wb.active.append(["TST-D-1", "y"])
    buf = io.BytesIO()
    wb.save(buf)
    assert amcremoval.list_from_file(buf.getvalue(), "a.xlsx") == ["TST-U-1", "TST-D-1"]


# ---------------------------------------------------------------- T8: engineers request, deciders approve
ENG = "TST ENG"


@pytest.fixture()
def mine_(fleet):
    fleet.execute("UPDATE asset SET engineer_name = %s WHERE asset_key IN ('TST-D-1','TST-S-1')", (ENG,))
    return fleet


def ask(keys=("TST-D-1",), **body):
    return amcremoval.request(list(keys), {"reason": "SURPLUS", "note": "machine is in the store", **body}, ENG, "TST-ENG-USER", IP)


def test_an_engineer_request_changes_nothing_until_it_is_approved(mine_):
    r = ask(["TST-D-1", "TST-S-1"])
    assert r["status"] == "REQUESTED" and r["assets"] == ["TST-D-1", "TST-S-1"]
    head = one(mine_, "SELECT * FROM amc_removal WHERE removal_no = %s", (r["removal_no"],))
    assert (head["status"], head["requested_by"], head["approved_by"], head["approved_on"]) == ("REQUESTED", "TST-ENG-USER", None, None)
    assert n(mine_, "SELECT count(*) FROM asset WHERE asset_key IN ('TST-D-1','TST-S-1') AND is_current = 1") == 2         # still in the register
    assert amcremoval.apply_due() == []                                                                                # and the scheduler ignores it


def test_an_engineer_can_only_request_their_own_assets(mine_):
    with pytest.raises(Conflict, match="TST-U-1: it is not assigned to you"):
        ask(["TST-D-1", "TST-U-1"])
    with pytest.raises(Invalid, match="engineer"):
        amcremoval.request(["TST-D-1"], {"reason": "SURPLUS"}, None, "SOMEONE", IP)
    mine_.execute("UPDATE asset SET engineer_name = %s WHERE asset_key = 'TST-L-1'", (ENG,))
    with pytest.raises(Conflict, match="laptops"):
        ask(["TST-L-1"])
    assert n(mine_, "SELECT count(*) FROM amc_removal WHERE requested_by = 'TST-ENG-USER'") == 0


def test_approving_records_the_verbal_approver_and_applies_it(mine_):
    no = ask()["removal_no"]
    with pytest.raises(Invalid, match="verbal approval"):
        amcremoval.approve(no, {}, "TST-LEAD", IP)
    r = amcremoval.approve(no, {"approved_by": "TST-APPROVER"}, "TST-LEAD", IP)
    assert r["status"] == "APPLIED"
    head = one(mine_, "SELECT * FROM amc_removal WHERE removal_no = %s", (no,))
    assert (head["status"], head["approved_by"], head["decided_by"], head["requested_by"], head["effective_date"]) == ("APPLIED", "TST-APPROVER", "TST-LEAD", "TST-ENG-USER", Q["start"])
    assert one(mine_, "SELECT is_current, asset_status FROM asset WHERE asset_key = 'TST-D-1'") == {"is_current": 0, "asset_status": "REMOVED_FROM_AMC"}
    with pytest.raises(Conflict, match="not waiting"):
        amcremoval.approve(no, {"approved_by": "TST-APPROVER"}, "TST-LEAD", IP)                                   # decided once


def test_the_decider_can_set_a_later_date_and_open_calls_need_an_override(mine_):
    mine_.execute("INSERT INTO svc_call (sr_id, asset_key, call_status, is_current, cipl_call_date) VALUES ('TST-SR-8','TST-D-1','OPEN',1,%s)", (dt.date.today(),))
    no = ask()["removal_no"]                                                                                          # the engineer is not asked about calls
    with pytest.raises(Conflict, match="open call"):
        amcremoval.approve(no, {"approved_by": "TST-APPROVER"}, "TST-LEAD", IP)
    assert one(mine_, "SELECT status FROM amc_removal WHERE removal_no = %s", (no,))["status"] == "REQUESTED"          # a refused approval leaves it waiting
    d = dt.date.today() + dt.timedelta(days=1)
    if d <= Q["end"]:
        assert amcremoval.approve(no, {"approved_by": "TST-APPROVER", "override_note": "printer call, not AMC", "effective_date": d.isoformat()}, "TST-LEAD", IP)["status"] == "PENDING"


def test_approval_checks_the_assets_again(mine_):
    no = ask()["removal_no"]
    mine_.execute("UPDATE asset SET cover_type = 'WRTY' WHERE asset_key = 'TST-D-1'")                                  # changed while it waited
    with pytest.raises(Conflict, match="not covered by AMC"):
        amcremoval.approve(no, {"approved_by": "TST-APPROVER"}, "TST-LEAD", IP)


def test_rejecting_needs_a_reason_the_engineer_can_read(mine_):
    no = ask()["removal_no"]
    with pytest.raises(Invalid, match="reason"):
        amcremoval.reject(no, {"reason": "no"}, "TST-LEAD", IP)
    assert amcremoval.reject(no, {"reason": "Still in use by the control room"}, "TST-LEAD", IP)["status"] == "REJECTED"
    assert one(mine_, "SELECT is_current FROM asset WHERE asset_key = 'TST-D-1'")["is_current"] == 1
    assert amcremoval.mine(ENG, "TST-ENG-USER")["requests"][0]["decision_note"] == "Still in use by the control room"
    with pytest.raises(Conflict, match="not waiting"):
        amcremoval.reject(no, {"reason": "again and again"}, "TST-LEAD", IP)


def test_an_engineer_can_withdraw_only_their_own_waiting_request(mine_):
    no = ask()["removal_no"]
    with pytest.raises(Conflict, match="only by the person"):
        amcremoval.cancel(no, "TST-OTHER", IP, only_own=True)
    assert amcremoval.cancel(no, "TST-ENG-USER", IP, only_own=True)["status"] == "CANCELLED"
    done = ask()["removal_no"]
    amcremoval.approve(done, {"approved_by": "TST-APPROVER"}, "TST-LEAD", IP)
    with pytest.raises(Conflict):
        amcremoval.cancel(done, "TST-ENG-USER", IP, only_own=True)


def test_the_engineers_page_lists_their_assets_requests_and_removed_assets_outside_the_counts(mine_):
    before = amcremoval.mine(ENG, "TST-ENG-USER")
    assert {a["asset_key"] for a in before["assets"]} == {"TST-D-1", "TST-S-1"} and before["removed"] == []
    no = ask(["TST-D-1"])["removal_no"]
    amcremoval.approve(no, {"approved_by": "TST-APPROVER"}, "TST-LEAD", IP)
    after = amcremoval.mine(ENG, "TST-ENG-USER")
    assert {a["asset_key"] for a in after["assets"]} == {"TST-S-1"}                                                    # no longer one of their assets...
    assert [(r["asset_key"], r["removal_no"]) for r in after["removed"]] == [("TST-D-1", no)]                           # ...but they can still see it was removed
    assert after["requests"][0]["status"] == "APPLIED"
    assert amcremoval.mine("TST OTHER ENG", "x")["removed"] == []                                                       # and nobody else's


def test_engineer_routes_and_permissions(mine_, monkeypatch):
    fake_user(monkeypatch, "USER", username="TST-ENG-USER", engineer_key=ENG)
    with TestClient(app) as c:
        r = c.post("/api/amc-removal/request", json={"keys": ["TST-D-1"], "reason": "SURPLUS"}, headers=HDR)
        assert r.status_code == 200 and r.json()["status"] == "REQUESTED", r.text
        no = r.json()["removal_no"]
        assert c.get("/api/amc-removal/mine", headers=HDR).json()["requests"][0]["removal_no"] == no
        assert c.post("/api/amc-removal/approve", json={"removal_no": no, "approved_by": "TST-APPROVER"}, headers=HDR).status_code == 403      # engineers cannot decide
        assert c.post("/api/amc-removal/reject", json={"removal_no": no, "reason": "no reason"}, headers=HDR).status_code == 403
        assert c.get("/api/amc-removal/requests", headers=HDR).status_code == 403
        m = c.post("/api/amc-removal/match", json={"text": "TST-D-1\nTST-S-1\nTST-U-1"}, headers=HDR).json()
        assert [x["result"] for x in m["rows"]] == ["OK", "OK", "NOT_FOUND"]                                                              # TST-U-1 is not theirs: invisible, not "found"
        assert c.post("/api/amc-removal/withdraw", json={"removal_no": no}, headers=HDR).status_code == 200
    fake_user(monkeypatch, "USER", username="NOBODY")                                                                          # a user with no engineer link
    with TestClient(app) as c:
        assert c.post("/api/amc-removal/request", json={"keys": ["TST-D-1"], "reason": "SURPLUS"}, headers=HDR).status_code in (400, 403, 422)
        assert c.post("/api/amc-removal/match", json={"text": "TST-D-1"}, headers=HDR).status_code in (400, 403, 422)
    lead = fake_user(monkeypatch, "USER", username="TST-LEAD")
    lead["lead_tools"] = True
    no2 = ask()["removal_no"]
    with TestClient(app) as c:
        assert c.get("/api/amc-removal/requests?status=REQUESTED", headers=HDR).status_code == 200
        r = c.post("/api/amc-removal/approve", json={"removal_no": no2, "approved_by": "TST-APPROVER"}, headers=HDR)
        assert r.status_code == 200 and r.json()["status"] == "APPLIED", r.text
        assert len(c.post("/api/amc-removal/match", json={"text": "TST-U-1"}, headers=HDR).json()["rows"]) == 1                 # a team leader sees every asset


# ---------------------------------------------------------------- T9: reinstatement
def back(keys, **body):
    return amcremoval.reinstate(keys, {"approved_by": "TST-APPROVER", "reason": "removed by mistake", **body}, ED, IP)


def test_reinstating_restores_cover_rate_pm_and_unarchives(fleet):
    fleet.execute("UPDATE asset SET pm_status = 'PENDING' WHERE asset_key = 'TST-D-1'")
    no = go(["TST-D-1"])["removal_no"]
    r = back(["TST-D-1"])
    assert r == {"reinstated": ["TST-D-1"], "removals": [no]}
    a = one(fleet, "SELECT * FROM asset WHERE asset_key = 'TST-D-1'")
    assert (a["is_current"], a["asset_status"], a["cover_type"], float(a["rate_value"]), a["pm_status"]) == (1, "IN_USE", "AMC", 1500.0, "PENDING")
    assert n(fleet, "SELECT count(*) FROM portal_lock WHERE dataset='assets' AND record_key='TST-D-1' AND field='__archived__'") == 0
    assert one(fleet, "SELECT status FROM amc_removal WHERE removal_no = %s", (no,))["status"] == "REVERSED"             # the only asset came back
    assert n(fleet, "SELECT count(*) FROM portal_audit WHERE record_key='TST-D-1' AND action='AMC_REINSTATE'") == 1
    assert n(fleet, "SELECT count(*) FROM asset WHERE asset_key = 'TST-D-1' AND " + pm.IN_SCOPE) == 1                       # back in the PM list


def test_reinstating_brings_its_parts_back_and_a_removal_with_other_assets_stays_applied(fleet):
    with_parts(fleet)
    no = go(["TST-D-1", "TST-S-1"])["removal_no"]
    back(["TST-D-1"])
    assert n(fleet, "SELECT count(*) FROM asset WHERE asset_key LIKE 'TST-D-1%%' AND is_current = 1") == 3
    assert one(fleet, "SELECT status FROM amc_removal WHERE removal_no = %s", (no,))["status"] == "APPLIED"               # TST-S-1 is still out
    assert one(fleet, "SELECT is_current FROM asset WHERE asset_key = 'TST-S-1'")["is_current"] == 0


def test_net_credit_after_reinstatement(fleet):
    d = Q["start"]
    go(["TST-D-1"], effective_date=d.isoformat())
    back(["TST-D-1"])                                                       # back today, same quarter
    line = one(fleet, "SELECT * FROM amc_removal_line WHERE asset_key = 'TST-D-1'")
    days = (dt.date.today() - d).days
    assert (line["net_credit_days"], line["reinstated_approved_by"]) == (days, "TST-APPROVER")
    assert str(line["net_credit_amount"]) == str((amcremoval.Decimal("1500") * days / 92).quantize(amcremoval.Decimal("0.01"), amcremoval.ROUND_HALF_UP))
    assert [x for x in amcremoval.requests() if x["removal_no"] == line["removal_no"]][0]["credit_total"] == line["net_credit_amount"]      # the figure that stands


def test_reinstating_needs_approval_a_reason_and_a_removed_asset(fleet):
    go(["TST-D-1"])
    with pytest.raises(Invalid, match="verbal approval"):
        back(["TST-D-1"], approved_by="NOBODY")
    with pytest.raises(Invalid, match="reason"):
        back(["TST-D-1"], reason="x")
    with pytest.raises(Conflict, match="TST-S-1: it was not removed"):
        back(["TST-D-1", "TST-S-1"])
    assert one(fleet, "SELECT is_current FROM asset WHERE asset_key = 'TST-D-1'")["is_current"] == 0                       # all or none
    back(["TST-D-1"])
    with pytest.raises(Conflict, match="already been reinstated"):
        back(["TST-D-1"])
    with pytest.raises(Invalid, match="between"):
        back(["TST-D-1"], effective_date=(dt.date.today() + dt.timedelta(days=1)).isoformat())


def test_it_can_be_removed_again_after_a_reinstatement(fleet):
    go(["TST-D-1"])
    back(["TST-D-1"])
    r2 = go(["TST-D-1"])["removal_no"]
    assert n(fleet, "SELECT count(*) FROM amc_removal_line WHERE asset_key = 'TST-D-1' AND reinstated_at IS NULL") == 1
    back(["TST-D-1"])
    assert one(fleet, "SELECT status FROM amc_removal WHERE removal_no = %s", (r2,))["status"] == "REVERSED"


def test_plain_restore_is_refused_for_an_asset_removed_from_amc(fleet):
    from portal.app import edit
    go(["TST-D-1"])
    with pytest.raises(Conflict, match="Reinstate"):
        edit.restore("assets", "TST-D-1", ED, IP, "test")


# ---------------------------------------------------------------- T10: reports and alerts
def test_summary_counts_assets_parts_reinstated_and_credit_per_quarter(fleet):
    with_parts(fleet)
    go(["TST-D-1", "TST-S-1"])
    back(["TST-S-1"])
    s = amcremoval.summary()
    q = [x for x in s["quarters"] if x["quarter_label"] == Q["label"]][0]
    assert q["assets"] >= 2 and q["parts"] >= 2 and q["reinstated"] >= 1
    assert s["waiting"] >= 0


def test_the_dashboard_removed_figure_counts_requests_not_asset_rows(fleet):
    from portal.app import dashboards
    base = dashboards.assets()["kpi"]["removed"]
    go(["TST-D-1"])
    assert dashboards.assets()["kpi"]["removed"] == base + 1
    back(["TST-D-1"])
    assert dashboards.assets()["kpi"]["removed"] == base
    fleet.execute("UPDATE asset SET engineer_name = 'TST ENG' WHERE asset_key = 'TST-S-1'")
    go(["TST-S-1"])
    assert dashboards.assets(eng="TST ENG")["kpi"]["removed"] == 1 and dashboards.assets(eng="NOBODY ELSE")["kpi"]["removed"] == 0


def test_excel_report_has_a_row_per_line_with_credit_and_people(fleet):
    import io
    from openpyxl import load_workbook
    with_parts(fleet)
    no = go(["TST-D-1"], note="x")["removal_no"]
    wb = load_workbook(io.BytesIO(amcremoval.report_bytes(Q["label"])))
    rows = [r for r in wb["AMC removals"].iter_rows(values_only=True) if r[0] == no]
    assert [r[3] for r in rows] == ["TST-D-1", "TST-D-1/GPU1", "TST-D-1/MON1"] and rows[1][4] == "TST-D-1"
    assert (rows[0][13], rows[0][15], rows[0][22], rows[0][23]) == (1500, 1500, ED, "TST-APPROVER")
    assert wb["By quarter"].max_row >= 2


def test_report_and_reinstate_routes(fleet, monkeypatch):
    fake_user(monkeypatch, "USER", username="TST-ENG-USER", engineer_key="TST ENG")
    with TestClient(app) as c:
        assert c.get("/api/amc-removal/export", headers=HDR).status_code == 403
        assert c.get("/api/amc-removal/summary", headers=HDR).status_code == 403
        assert c.post("/api/amc-removal/reinstate", json={"keys": ["TST-D-1"]}, headers=HDR).status_code == 403
    fake_user(monkeypatch, "ADMIN")
    go(["TST-D-1"])
    with TestClient(app) as c:
        r = c.get("/api/amc-removal/export", headers=HDR)
        assert r.status_code == 200 and r.content[:2] == b"PK" and "AMC_removals.xlsx" in r.headers["content-disposition"]
        assert c.get("/api/amc-removal/summary", headers=HDR).status_code == 200
        r = c.post("/api/amc-removal/reinstate", json={"keys": ["TST-D-1"], "approved_by": "TST-APPROVER", "reason": "needed again"}, headers=HDR)
        assert r.status_code == 200 and r.json()["reinstated"] == ["TST-D-1"], r.text


def test_waiting_requests_alert_administrators_and_team_leaders(mine_):
    from portal.app import mailer
    eng = one(mine_, "SELECT engineer_key FROM portal_engineer ORDER BY engineer_key LIMIT 1")["engineer_key"]
    mine_.execute("INSERT INTO portal_user (username, display_name, role, password_hash, engineer_key) VALUES ('TST-BOSS','TST BOSS','ADMIN','x',%s) ON CONFLICT DO NOTHING", (eng,))
    no = ask()["removal_no"]
    got = mailer.detect("AMC_REQUEST", {})
    assert eng in got and no in {r["removal_no"] for r in got[eng]["rows"]} and "waiting for your decision" in got[eng]["subject"]
    amcremoval.reject(no, {"reason": "not now at all"}, "TST-LEAD", IP)
    assert no not in {r["removal_no"] for r in mailer.detect("AMC_REQUEST", {}).get(eng, {"rows": []})["rows"]}              # decided: no longer announced


# ---------------------------------------------------------------- only administrators and Team Leader/SI give the verbal approval; the upload template
def test_only_administrators_and_team_leaders_can_give_the_approval(fleet, monkeypatch):
    from portal.app import auth
    fleet.execute("INSERT INTO portal_user (username, display_name, role, password_hash) VALUES ('TST-PLAIN','TST PLAIN','USER','x'), ('TST-RO','TST RO','ADMIN','x') ON CONFLICT DO NOTHING")
    fleet.execute("UPDATE portal_user SET read_only = TRUE WHERE username = 'TST-RO'")
    names = {a["username"] for a in amcremoval.meta()["approvers"]}
    assert "TST-APPROVER" in names and "TST-PLAIN" not in names and "TST-RO" not in names             # an ordinary user and a read-only account are not offered
    with pytest.raises(Invalid, match="administrator or Team Leader"):
        go(["TST-D-1"], approved_by="TST-PLAIN")
    # a plain user whose designation is Team Leader/SI is offered
    eng = one(fleet, "SELECT engineer_key FROM portal_engineer ORDER BY engineer_key LIMIT 1")["engineer_key"]
    fleet.execute("UPDATE portal_user SET engineer_key = %s WHERE username = 'TST-PLAIN'", (eng,))
    monkeypatch.setattr(auth, "is_lead_designation", lambda d: True)
    assert "TST-PLAIN" in {a["username"] for a in amcremoval.meta()["approvers"]}
    assert go(["TST-D-1"], approved_by="TST-PLAIN")["status"] == "APPLIED"


def test_template_has_the_columns_and_a_reason_list():
    import io
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(amcremoval.template_bytes()))
    assert [c.value for c in wb["Assets"][1]] == ["Asset (CI)", "Serial no.", "Reason"]
    assert [r[0].value for r in wb["Lists"].iter_rows()] == list(amcremoval.REASONS.values())
    assert wb["Assets"].data_validations.dataValidation[0].formula1 == "=Lists!$A$1:$A$9"


def test_a_filled_template_gives_names_serials_and_a_reason_per_row(fleet):
    import io
    from openpyxl import load_workbook
    fleet.execute("UPDATE asset SET serial_no = 'TSTSER-S' WHERE asset_key = 'TST-S-1'")
    wb = load_workbook(io.BytesIO(amcremoval.template_bytes()))
    ws = wb["Assets"]
    ws.append(["TST-D-1", None, "Surplus"])
    ws.append([None, "TSTSER-S", "End of life / support"])
    ws.append(["TST-U-1", None, "Because I said so"])
    ws.append(["TST-W-1", None, None])
    buf = io.BytesIO()
    wb.save(buf)
    names, reasons = amcremoval.read_list(buf.getvalue())
    assert names == ["TST-D-1", "TSTSER-S", "TST-U-1", "TST-W-1"]
    assert reasons == {"TST-D-1": "SURPLUS", "TSTSER-S": "END_OF_LIFE", "TST-U-1": "?Because I said so"}
    m = {r["input"]: r for r in amcremoval.match_list(names, None, reasons)["rows"]}
    assert m["TSTSER-S"]["row"]["asset_key"] == "TST-S-1" and m["TSTSER-S"]["reason"] == "END_OF_LIFE" and m["TST-D-1"]["reason"] == "SURPLUS"
    assert "not one of the listed reasons" in m["TST-U-1"]["reason_note"] and "reason" not in m["TST-U-1"]
    assert m["TST-W-1"]["result"] == "BLOCKED" and "reason" not in m["TST-W-1"]


def test_an_exact_asset_key_beats_a_hostname_that_reads_the_same(fleet):
    fleet.execute("UPDATE asset SET hostname = 'TST-D-1' WHERE asset_key = 'TST-S-1'")
    r = amcremoval.match_list(["TST-D-1"])["rows"][0]
    assert r["result"] == "OK" and r["row"]["asset_key"] == "TST-D-1"


def test_template_route_is_open_to_signed_in_users(fleet, monkeypatch):
    fake_user(monkeypatch, "USER", username="TST-ENG-USER", engineer_key="TST ENG")
    with TestClient(app) as c:
        r = c.get("/api/amc-removal/template", headers=HDR)
        assert r.status_code == 200 and r.content[:2] == b"PK" and "AMC_Removal_Upload_Template.xlsx" in r.headers["content-disposition"]


# ---------------------------------------------------------------- assets marked "Out of AMC" by hand before the process existed
def legacy(box):
    make_asset(box, "TST-OLD-1", status="REMOVED_FROM_AMC", rate_value=118.20, remarks="REMOVE FROM AMC 01/10/2026")
    make_asset(box, "TST-OLD-W", status="REMOVED_FROM_AMC", cover_type="WRTY")


def test_an_asset_marked_out_of_amc_by_hand_can_go_through_the_process(fleet):
    legacy(fleet)
    m = {r["input"]: r for r in amcremoval.match_list(["TST-OLD-1", "TST-OLD-W"])["rows"]}
    assert m["TST-OLD-1"]["result"] == "OK" and m["TST-OLD-W"]["why"] == "it is not covered by AMC"                      # a warranty asset still is not an AMC matter
    assert "TST-OLD-1" in {r["asset_key"] for r in amcremoval.candidates({"q": "TST-OLD"})["rows"]}
    no = go(["TST-OLD-1"], reason="SCOPE_CHANGE")["removal_no"]
    a = one(fleet, "SELECT is_current, rate_value, asset_status FROM asset WHERE asset_key = 'TST-OLD-1'")
    assert (a["is_current"], a["rate_value"], a["asset_status"]) == (0, None, "REMOVED_FROM_AMC")                      # archived, rate cleared
    line = one(fleet, "SELECT * FROM amc_removal_line WHERE removal_no = %s", (no,))
    assert str(line["credit_amount"]) == "118.20" and float(line["snapshot"]["rate_value"]) == 118.2                      # the credit comes from the rate it still carried
    with pytest.raises(Conflict, match="already removed"):
        go(["TST-OLD-1"])                                                                                                # now it really is done


def test_reinstating_a_hand_marked_asset_puts_it_back_as_deployed(fleet):
    legacy(fleet)
    go(["TST-OLD-1"])
    back(["TST-OLD-1"])
    a = one(fleet, "SELECT is_current, asset_status, rate_value FROM asset WHERE asset_key = 'TST-OLD-1'")
    assert (a["is_current"], a["asset_status"], float(a["rate_value"])) == (1, "IN_USE", 118.2)
