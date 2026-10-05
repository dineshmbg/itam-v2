"""Tools: hover cards, reports and exports, PM cycles, e-mail rules, backup, import, dashboards packs. Transaction-isolated (see conftest.box)."""
import datetime as dt
import io
import zipfile

import pytest
from starlette.testclient import TestClient

from conftest import HDR, fake_user
from portal.app import auth, backup, export, importer, mailer, packs, pm, reports
from portal.app.main import app


def one(con, sql, params=()):
    return con.execute(sql, params).fetchone()


# ---------------------------------------------------------------- quarters
@pytest.mark.parametrize("day,label,start,end", [
    ("2026-09-22", "Q2 JUL-SEP 2026", "2026-07-01", "2026-09-30"), ("2026-04-01", "Q1 APR-JUN 2026", "2026-04-01", "2026-06-30"),
    ("2026-12-31", "Q3 OCT-DEC 2026", "2026-10-01", "2026-12-31"), ("2027-02-10", "Q4 JAN-MAR 2027", "2027-01-01", "2027-03-31"), ("2027-04-01", "Q1 APR-JUN 2027", "2027-04-01", "2027-06-30")])
def test_financial_year_quarters(day, label, start, end):
    q = pm.quarter(dt.date.fromisoformat(day))
    assert (q["label"], q["start"].isoformat(), q["end"].isoformat()) == (label, start, end)
    assert q["kickoff"] == q["start"] + dt.timedelta(days=60)


# ---------------------------------------------------------------- reports
ADMIN = {"username": "T", "role": "ADMIN"}
PLAIN = {"username": "T", "role": "USER", "engineer_key": None, "call_parts_access": "NONE", "asset_access": "NONE"}


def test_report_builder_filters_groups_and_rejects_bad_input(box):
    # personal-field gating is catalog metadata only - no rows involved, so a real non-admin user is fine here
    d = reports.describe("assets", PLAIN)
    keys = {f["key"] for f in d["fields"]}
    assert {"asset_key", "make", "cover_expiry_date", "cpf_no"} <= keys and "user_mobile" not in keys and "is_current" not in keys
    assert "mobile_no" in {f["key"] for f in reports.describe("employees", ADMIN)["fields"]} and "mobile_no" not in {f["key"] for f in reports.describe("employees", PLAIN)["fields"]}
    # everything below is about report-building correctness (filters, grouping, injection safety), not scoping - run unscoped as admin
    r = reports.run({"dataset": "assets", "columns": ["asset_key", "make", "model"], "filters": [{"field": "make", "op": "in", "values": ["hp"]}, {"field": "record_level", "op": "eq", "value": "asset"}],
                     "sort": [{"field": "asset_key", "dir": "desc"}]}, ADMIN, limit=5)
    assert r["total"] > 100 and len(r["rows"]) == 5 and all(x["make"] == "HP" for x in r["rows"]) and r["rows"][0]["asset_key"] >= r["rows"][1]["asset_key"]
    g = reports.run({"dataset": "assets", "group_by": ["asset_class", "install_date:year"], "metrics": [{"fn": "count_distinct", "field": "make"}]}, ADMIN, limit=50)
    assert [c["key"] for c in g["columns"]] == ["asset_class", "install_date_year", "n", "m0"] and g["rows"][0]["n"] >= g["rows"][-1]["n"]
    any_ = reports.run({"dataset": "assets", "match": "any", "columns": ["asset_key"], "filters": [{"field": "make", "op": "eq", "value": "DELL"}, {"field": "make", "op": "eq", "value": "APPLE"}]}, ADMIN)
    only = {reports.run({"dataset": "assets", "columns": ["asset_key"], "filters": [{"field": "make", "op": "eq", "value": m}]}, ADMIN)["total"] for m in ("DELL", "APPLE")}
    assert any_["total"] == sum(only)
    dated = reports.run({"dataset": "assets", "columns": ["asset_key"], "filters": [{"field": "cover_expiry_date", "op": "between", "value": "2020-01-01", "value2": "2099-01-01"}]}, ADMIN)
    assert dated["total"] > 0
    for bad in ({"dataset": "nope"}, {"dataset": "assets", "columns": ["password_hash"]}, {"dataset": "assets", "filters": [{"field": "make", "op": "gt", "value": 1}]},
                {"dataset": "assets", "columns": ["asset_key"], "filters": [{"field": "make; DROP TABLE asset", "op": "eq", "value": "x"}]}, {"dataset": "assets", "group_by": ["make:month"]},
                {"dataset": "assets", "group_by": ["make"], "metrics": [{"fn": "sum", "field": "make"}]}):
        with pytest.raises(reports.ReportError):
            reports.run(bad, ADMIN)
    with pytest.raises(reports.ReportError):     # a non-admin cannot select a personal field as a report column, same as catalog() hides it
        reports.run({"dataset": "employees", "columns": ["mobile_no"]}, PLAIN)
    hostile = reports.run({"dataset": "assets", "columns": ["asset_key"], "filters": [{"field": "asset_key", "op": "contains", "value": "'; DROP TABLE asset; --"}]}, ADMIN)
    assert hostile["total"] == 0 and one(box, "SELECT count(*) FROM asset")[0] > 4000        # values are parameters, never SQL


def test_reports_apply_the_same_access_policy_as_the_registers(box):
    """The permissions fix (2026-09-30): reports used to see everything regardless of role - closing that so a report can never
    show more than the equivalent register/dashboard would."""
    # a User with no engineer_key and no grant sees nothing in an Assets report, same as the Assets register would show them
    empty = reports.run({"dataset": "assets", "columns": ["asset_key"]}, PLAIN)
    assert empty["total"] == 0
    # Calls (call_parts_access=NONE) is not even an openable dataset for reports, same as the register
    with pytest.raises(auth.AuthError):
        reports.describe("calls", PLAIN)
    with pytest.raises(auth.AuthError):
        reports.run({"dataset": "calls", "columns": ["sr_id"]}, PLAIN)
    # asset_access=READ or FULL un-scopes Assets for reports, same as it does for the register and dashboard
    for grant in ("READ", "FULL"):
        r = reports.run({"dataset": "assets", "columns": ["asset_key"]}, {**PLAIN, "asset_access": grant})
        assert r["total"] > 100
    # a real engineer with assigned assets only ever sees their own rows through a report, exactly like the register - the
    # comparison query matches reports.SETS["assets"]'s own base condition (is_current = 1, no record_level filter: components included)
    eng = one(box, "SELECT engineer_name FROM asset WHERE is_current = 1 AND record_level = 'ASSET' AND engineer_name IS NOT NULL LIMIT 1")[0]
    scoped_user = {**PLAIN, "engineer_key": eng}
    from_report = reports.run({"dataset": "assets", "columns": ["asset_key"]}, scoped_user, limit=10000)["total"]
    from_table = one(box, "SELECT count(*) FROM asset WHERE is_current = 1 AND engineer_name = %s", (eng,))[0]
    assert from_report == from_table > 0


def test_report_save_share_and_delete(box):
    u = {"username": "TESTER", "role": "USER"}
    reports.save(u, "hp laptops", {"dataset": "assets", "columns": ["asset_key"], "filters": [{"field": "make", "op": "eq", "value": "HP"}]}, shared=True)
    saved = reports.list_saved({"username": "OTHER", "role": "USER"})
    assert saved[0]["name"] == "HP LAPTOPS" and saved[0]["shared"]
    with pytest.raises(reports.ReportError):
        reports.delete({"username": "OTHER", "role": "USER"}, saved[0]["id"])
    reports.delete(u, saved[0]["id"])
    assert reports.list_saved(u) == []
    with pytest.raises(reports.ReportError):
        reports.save(u, "x", {"dataset": "assets", "columns": ["nope"]}, False)


# ---------------------------------------------------------------- files
def test_export_formats_are_real_files(box):
    t, total = reports.table_for_export({"dataset": "assets", "columns": ["asset_key", "make", "cover_expiry_date"], "filters": [{"field": "make", "op": "eq", "value": "HP"}]}, ADMIN, "Assets")
    x, _, ext = export.render("xlsx", "Assets", "test", [t])
    assert ext == "xlsx" and zipfile.ZipFile(io.BytesIO(x)).testzip() is None
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(x))["ASSETS"]
    assert ws.max_row == total + 1 and [c.value for c in ws[1]] == ["ASSET KEY", "MAKE", "COVER EXPIRY DATE"]
    c, _, _ = export.render("csv", "Assets", "test", [t])
    assert c.startswith(b"\xef\xbb\xbf") and c.decode("utf-8-sig").splitlines()[0] == "ASSET KEY,MAKE,COVER EXPIRY DATE" and len(c.decode("utf-8-sig").splitlines()) == total + 1
    p, mime, _ = export.render("pdf", "Assets", "test", [t], kpis=[("Assets", total)], charts=[{"title": "x", "rows": [{"label": "HP", "n": 3}]}])
    assert p[:5] == b"%PDF-" and mime == "application/pdf" and len(p) > 2000
    with pytest.raises(ValueError):
        export.render("doc", "x", "x", [t])


@pytest.mark.parametrize("name", ["assets", "calls", "engineers", "pm"])
def test_every_dashboard_pack_renders_to_pdf_and_excel(box, name):
    p = packs.PACKS[name]()
    assert p["kpis"] and p["tables"] and p["charts"] is not None
    pdf, _, _ = export.render("pdf", p["title"], p["subtitle"], p["tables"], p["kpis"], p["charts"])
    xl, _, _ = export.render("xlsx", p["title"], p["subtitle"], p["tables"], p["kpis"], p["charts"])
    assert pdf[:5] == b"%PDF-" and zipfile.ZipFile(io.BytesIO(xl)).testzip() is None


def test_export_and_share_endpoints(box, monkeypatch):
    fake_user(monkeypatch, "USER")
    before = one(box, "SELECT count(*) FROM portal_activity WHERE action IN ('EXPORT_REPORT','EXPORT_DASHBOARD','SHARE_EMAIL')")[0]
    sent = []
    monkeypatch.setattr(mailer, "send", lambda to, subject, text, html_body=None, attachments=(), cfg=None: sent.append((to, subject, [a[0] for a in attachments])))
    with TestClient(app) as c:
        r = c.post("/api/reports/export", json={"definition": {"dataset": "assets", "columns": ["asset_key"]}, "format": "pdf"}, headers=HDR)
        assert r.status_code == 200 and r.content[:5] == b"%PDF-" and "attachment" in r.headers["content-disposition"] and ".pdf" in r.headers["content-disposition"]
        assert c.post("/api/reports/export", json={"definition": {"dataset": "assets", "columns": ["asset_key"]}, "format": "csv"}, headers=HDR).content.startswith(b"\xef\xbb\xbf")
        r = c.post("/api/dashboard/export", json={"pack": "calls", "format": "xlsx"}, headers=HDR)
        assert r.status_code == 200 and r.content[:2] == b"PK"
        assert c.post("/api/dashboard/export", json={"pack": "calls", "format": "csv"}, headers=HDR).status_code == 400
        assert c.post("/api/dashboard/export", json={"pack": "nope", "format": "pdf"}, headers=HDR).status_code == 400
        r = c.post("/api/dashboard/share", json={"pack": "assets", "to": ["boss@example.com"], "formats": ["pdf", "xlsx"]}, headers=HDR)
        assert r.status_code == 400 and "not switched on" in r.json()["error"]               # nothing is mailed until e-mail is switched on
        box.execute("INSERT INTO portal_setting (key, value) VALUES ('smtp', %s::jsonb) ON CONFLICT (key) DO UPDATE SET value = EXCLUDED.value",
                    ('{"enabled": true, "host": "smtp.example", "from_addr": "portal@example.com"}',))
        assert c.post("/api/dashboard/share", json={"pack": "assets", "to": ["not-an-address"], "formats": ["pdf"]}, headers=HDR).status_code == 400
        r = c.post("/api/dashboard/share", json={"pack": "assets", "to": ["boss@example.com", "a@b.co"], "formats": ["pdf", "xlsx"], "message": "for review"}, headers=HDR)
        assert r.status_code == 200 and sent[0][0] == ["boss@example.com", "a@b.co"] and sorted(x.rsplit(".", 1)[1] for x in sent[0][2]) == ["pdf", "xlsx"]
        assert one(box, "SELECT count(*) FROM portal_activity WHERE action IN ('EXPORT_REPORT','EXPORT_DASHBOARD','SHARE_EMAIL')")[0] - before == 4


# ---------------------------------------------------------------- PM
def _put_assets_in(box, q):
    """Pin the quarter the in-scope assets carry, so these tests do not depend on where the real data happens to be today."""
    box.execute(f"UPDATE asset SET pm_quarter = %s WHERE {pm.IN_SCOPE}", (q["label"],))


def test_recording_pm_updates_asset_history_and_status(box):
    q = pm.quarter(dt.date.today())
    _put_assets_in(box, q)
    a = one(box, "SELECT asset_key FROM asset WHERE is_current = 1 AND record_level = 'ASSET' AND pm_status = 'PENDING' LIMIT 2")
    keys = [r[0] for r in box.execute("SELECT asset_key FROM asset WHERE is_current = 1 AND record_level = 'ASSET' AND pm_status = 'PENDING' ORDER BY asset_key LIMIT 2").fetchall()]
    assert a and len(keys) == 2
    user = {"username": "ENG1", "role": "USER"}
    out = pm.record(keys, dt.date.today().isoformat(), "dinesh gadaria", "signer one", "ok", user, "127.0.0.1")
    assert out == {"recorded": 2, "quarter": q["label"]}
    rows = box.execute("SELECT pm_status, pm_done_by, pm_signed_by, pm_date FROM asset WHERE asset_key = ANY(%s)", (keys,)).fetchall()
    assert all(r[0] == "DONE" and r[1] == "DINESH GADARIA" and r[2] == "SIGNER ONE" and r[3] == dt.date.today() for r in rows)
    assert one(box, "SELECT count(*) FROM pm_record WHERE asset_key = ANY(%s)", (keys,))[0] == 2
    assert one(box, "SELECT count(*) FROM portal_lock WHERE dataset = 'assets' AND record_key = ANY(%s) AND field = 'pm_date'", (keys,))[0] == 2      # survives loader runs
    with pytest.raises(pm.PmError):
        pm.record(keys, (q["start"] - dt.timedelta(days=1)).isoformat(), None, None, None, user, "x")            # outside the cycle
    with pytest.raises(pm.PmError):
        pm.record(keys, (dt.date.today() + dt.timedelta(days=1)).isoformat(), None, None, None, user, "x")       # future
    with pytest.raises(pm.PmError):
        pm.record(["NO-SUCH-ASSET"], dt.date.today().isoformat(), None, None, None, user, "x")
    lap = one(box, "SELECT asset_key FROM asset WHERE is_current = 1 AND asset_class = 'LAPTOP' LIMIT 1")[0]
    with pytest.raises(pm.PmError):
        pm.record([lap], dt.date.today().isoformat(), None, None, None, user, "x")                                # laptops are not tracked
    d = pm.dashboard()
    assert d["kpi"]["scope"] == d["kpi"]["done"] + d["kpi"]["stale"] + d["kpi"]["pending"] and d["cycle"]["label"] == q["label"]


def test_rollover_snapshots_then_resets_the_quarter(box):
    user = {"username": "ADMIN1", "role": "ADMIN"}
    _put_assets_in(box, pm.quarter(dt.date.today()))
    audited = one(box, "SELECT count(*) FROM portal_audit WHERE record_key = 'PM-ROLLOVER'")[0]      # real earlier roll-overs exist
    before = pm.rollover_preview()
    assert before["closing"] == pm.quarter(dt.date.today())["label"] and before["opening"] == pm.next_quarter(dt.date.today())["label"] and not before["overdue"]
    with pytest.raises(pm.PmError):
        pm.rollover(user, "127.0.0.1")                                    # the quarter has not ended yet
    out = pm.rollover(user, "127.0.0.1", force=True)
    assert out["opened"] == before["opening"] and out["assets"] == before["scope"]
    snap = one(box, "SELECT count(*) FROM pm_snapshot WHERE quarter_label = %s AND as_of = %s", (before["closing"], dt.date.today()))[0]      # today's snapshot only: real earlier ones may exist
    assert snap == before["scope"]
    k = one(box, "SELECT count(*), count(*) FILTER (WHERE pm_status = 'PENDING'), count(*) FILTER (WHERE pm_date IS NOT NULL), count(DISTINCT pm_quarter) FROM asset WHERE " + pm.IN_SCOPE)
    assert k[0] == before["scope"] and k[1] == k[0] and k[2] == 0 and k[3] == 1
    assert one(box, "SELECT status FROM pm_cycle WHERE quarter_label = %s", (before["closing"],))[0] == "CLOSED"
    assert one(box, "SELECT count(*) FROM portal_audit WHERE record_key = 'PM-ROLLOVER'")[0] == audited + 1


def test_closing_a_quarter_after_it_has_ended_closes_that_quarter_not_the_new_one(box):
    """The 2026-10-01 bug: on the first day of Q3 the roll-over offered to close Q3 and open Q4, and the dashboard showed Q2's
    done/pending under a Q3 heading. The quarter to close is the one the assets carry; the one to open is the one containing today."""
    user = {"username": "ADMIN1", "role": "ADMIN"}
    today = dt.date.today()
    cur = pm.quarter(today)
    old = pm.quarter(cur["start"] - dt.timedelta(days=1))
    _put_assets_in(box, old)
    # Start from a clean slate: the dev database may already hold PM completed this quarter (it did, 179 assets, on 2026-10-04). Those would
    # rightly survive a roll-over and make the "exactly one kept" figures below depend on today's data. Rolled back with everything else.
    box.execute(f"UPDATE asset SET pm_date = NULL, pm_done_by = NULL, pm_signed_by = NULL, pm_status = 'PENDING' WHERE {pm.IN_SCOPE}")
    box.execute("DELETE FROM pm_cycle WHERE quarter_label = %s", (cur["label"],))
    kept, wiped = [r[0] for r in box.execute(f"SELECT asset_key FROM asset WHERE {pm.IN_SCOPE} ORDER BY asset_key LIMIT 2").fetchall()]
    box.execute("UPDATE asset SET pm_date = %s, pm_done_by = 'SOMEONE' WHERE asset_key = %s", (today, kept))                 # done already, in the new quarter
    box.execute("UPDATE asset SET pm_date = %s, pm_done_by = 'SOMEONE' WHERE asset_key = %s", (old["end"], wiped))           # done in the old quarter

    d = pm.dashboard()["cycle"]
    assert d["label"] == old["label"] and d["overdue"] and d["calendar_label"] == cur["label"]       # old figures are labelled as old
    assert pm.snapshot_now(user, "x")["quarter"] == old["label"]                                       # ... and so is a snapshot of them
    before = pm.rollover_preview()
    assert (before["closing"], before["opening"], before["overdue"]) == (old["label"], cur["label"], True) and before["starts_in_days"] <= 0

    out = pm.rollover(user, "127.0.0.1")                                                               # no early cut-over needed
    assert (out["closed"], out["opened"]) == (old["label"], cur["label"])
    assert one(box, "SELECT status FROM pm_cycle WHERE quarter_label = %s", (old["label"],))[0] == "CLOSED"
    assert one(box, "SELECT status FROM pm_cycle WHERE quarter_label = %s", (cur["label"],))[0] == "OPEN"
    assert one(box, "SELECT count(*) FROM pm_snapshot WHERE quarter_label = %s AND as_of = %s", (old["label"], old["end"]))[0] == before["scope"]
    assert one(box, "SELECT pm_status, pm_date, pm_done_by FROM asset WHERE asset_key = %s", (kept,)) == ("DONE", today, "SOMEONE")
    assert one(box, "SELECT pm_status, pm_date, pm_done_by FROM asset WHERE asset_key = %s", (wiped,)) == ("PENDING", None, None)
    k = pm.dashboard()
    assert k["cycle"]["label"] == cur["label"] and not k["cycle"]["overdue"]
    assert (k["kpi"]["done"], k["kpi"]["stale"], k["kpi"]["pending"]) == (1, 0, before["scope"] - 1)


# ---------------------------------------------------------------- e-mail
def test_notification_rules_detect_and_send_once(box, monkeypatch):
    box.execute("INSERT INTO portal_setting (key, value) VALUES ('smtp', %s::jsonb) ON CONFLICT (key) DO UPDATE SET value = EXCLUDED.value", ('{"enabled": true, "host": "h", "from_addr": "p@example.com"}',))
    box.execute("UPDATE cipl_employee SET company_email = 'eng@example.com' WHERE ecode = (SELECT ecode FROM portal_engineer WHERE engineer_key = 'DINESH GADARIA')")
    mailer.save_rule("CALL_OVERDUE", True, "lead@example.com", {"older_than_days": 1}, "T")
    mailer.save_rule("PM_KICKOFF", True, "", {}, "T")
    sent = []
    monkeypatch.setattr(mailer, "send", lambda to, subject, text, html_body=None, attachments=(), cfg=None: sent.append((to, subject)))
    monday = dt.date(2026, 9, 21)
    monkeypatch.setattr(mailer.dt, "date", type("D", (dt.date,), {"today": classmethod(lambda cls: monday)}))
    d = mailer.detect("CALL_OVERDUE", {"older_than_days": 1}, today=monday)
    assert d and all(v["rows"] for v in d.values())
    assert mailer.detect("CALL_OVERDUE", {"older_than_days": 1}, today=dt.date(2026, 9, 22)) == {}      # weekly digest: Mondays only
    pmk = mailer.detect("PM_KICKOFF", {}, today=pm.quarter(monday)["kickoff"])
    assert pmk and mailer.detect("PM_KICKOFF", {}, today=pm.quarter(monday)["kickoff"] - dt.timedelta(days=1)) == {}    # nothing before day 60
    preview = mailer.run_rules(dry_run=True, only="CALL_OVERDUE")
    assert preview and all(p["status"] in ("DRY_RUN", "NO_ADDRESS") for p in preview) and sent == []          # a preview never sends
    first = mailer.run_rules(only="CALL_OVERDUE")
    n = len(sent)
    assert n >= 1 and {"eng@example.com", "lead@example.com"} & {a for s in sent for a in ([s[0]] if isinstance(s[0], str) else s[0])}
    assert all(x["status"] == "SENT" for x in first)
    mailer.run_rules(only="CALL_OVERDUE")
    assert len(sent) == n, "the same event must never be mailed twice"
    assert one(box, "SELECT count(*) FROM notify_log WHERE rule_key = 'CALL_OVERDUE' AND status = 'SENT'")[0] == n


def test_pm_close_quarter_reminder_only_fires_while_overdue_and_open(box):
    """Finds the quarter by pm_cycle.status/end_date, not by pm.quarter(today) - the label of the quarter *containing* today is
    the wrong question once today has already rolled past the boundary. Never closes the quarter itself - detect() only ever
    produces a reminder, nothing in this rule calls pm.rollover()."""
    today = dt.date(2026, 10, 3)      # a few days into Q3, Q2 (ends 2026-09-30) never rolled over
    box.execute("INSERT INTO pm_cycle (quarter_label, start_date, end_date, kickoff_date, status) VALUES ('Q2 JUL-SEP 2026', '2026-07-01', '2026-09-30', '2026-08-30', 'OPEN') "
                "ON CONFLICT (quarter_label) DO UPDATE SET status = 'OPEN', end_date = '2026-09-30'")
    d = mailer.detect("PM_CLOSE_QUARTER", {}, today=today)
    admin_keys = {r[0] for r in box.execute("SELECT engineer_key FROM portal_user WHERE role = 'ADMIN' AND active AND engineer_key IS NOT NULL").fetchall()}
    assert admin_keys, "need at least one active administrator with a linked engineer for this test to mean anything"
    assert set(d.keys()) == admin_keys
    assert all(v["event_key"] == f"Q2 JUL-SEP 2026:{today.isoformat()}" and "Q2 JUL-SEP 2026" in v["subject"] for v in d.values())
    assert mailer.detect("PM_CLOSE_QUARTER", {}, today=dt.date(2026, 9, 15)) == {}       # well before the quarter even ends
    box.execute("UPDATE pm_cycle SET status = 'CLOSED' WHERE quarter_label = 'Q2 JUL-SEP 2026'")
    assert mailer.detect("PM_CLOSE_QUARTER", {}, today=today) == {}                       # rolled over -> stops reminding


def test_mail_failures_are_logged_and_retried(box, monkeypatch):
    box.execute("INSERT INTO portal_setting (key, value) VALUES ('smtp', %s::jsonb) ON CONFLICT (key) DO UPDATE SET value = EXCLUDED.value", ('{"enabled": true, "host": "h", "from_addr": "p@example.com"}',))
    box.execute("UPDATE cipl_employee SET company_email = 'eng@example.com' WHERE company_email IS NOT NULL")
    mailer.save_rule("CALL_ASSIGNED", True, "", {"lookback_days": 365}, "T")
    calls = {"n": 0}

    def flaky(*a, **k):
        calls["n"] += 1
        raise mailer.MailError("server down")
    monkeypatch.setattr(mailer, "send", flaky)
    out = mailer.run_rules(only="CALL_ASSIGNED")
    assert out and all(x["status"] == "FAILED" for x in out) and one(box, "SELECT count(*) FROM notify_log WHERE status = 'FAILED'")[0] == len(out)
    monkeypatch.setattr(mailer, "send", lambda *a, **k: None)
    again = mailer.run_rules(only="CALL_ASSIGNED")
    assert all(x["status"] == "SENT" for x in again) and one(box, "SELECT max(attempts) FROM notify_log")[0] == 2


def test_smtp_settings_validation_and_secret_handling(box, tmp_path, monkeypatch):
    monkeypatch.setattr(mailer, "SECRET_FILE", tmp_path / "secret")
    for bad in ({"enabled": True, "host": "", "from_addr": ""}, {"host": "h", "port": 0}, {"host": "h", "port": "x"}, {"host": "h", "security": "weird"}, {"host": "h", "from_addr": "nope"}):
        with pytest.raises(mailer.MailError):
            mailer.save_smtp(bad, None, "T")
    cfg = mailer.save_smtp({"enabled": True, "host": "smtp.example", "port": "587", "security": "starttls", "from_addr": "p@example.com", "user": "u"}, "s3cret!", "T")
    assert cfg["has_password"] and cfg["port"] == 587 and "s3cret!" not in str(cfg) and (tmp_path / "secret").read_text() == "s3cret!"
    assert "s3cret!" not in str(one(box, "SELECT value FROM portal_setting WHERE key = 'smtp'"))       # never stored in the database
    with pytest.raises(mailer.MailError):
        mailer.save_rule("PM_KICKOFF", True, "not-an-address", {}, "T")
    with pytest.raises(mailer.MailError):
        mailer.save_rule("NOPE", True, "", {}, "T")


# ---------------------------------------------------------------- backup
def test_backup_verify_and_restore_guards(box, tmp_path, monkeypatch):
    monkeypatch.setattr(backup, "BACKUP_DIR", tmp_path)
    out = backup.run_backup("MANUAL", "TEST", "unit test")
    assert (tmp_path / out["file"]).stat().st_size > 10000 and len(out["sha256"]) == 64
    v = backup.verify(out["file"])
    assert v["ok"] and v["tables"] >= 10
    (tmp_path / out["file"]).write_bytes((tmp_path / out["file"]).read_bytes()[:-50] + b"x" * 50)          # damaged file
    with pytest.raises(backup.BackupError):
        backup.verify(out["file"])
    for bad in ("../secret.dump", "a/b.dump", "..\\x.dump", "notadump.txt", "", "missing.dump"):
        with pytest.raises(backup.BackupError):
            backup._resolve(bad)
    assert backup.list_backups()[0]["file"] == out["file"]
    pre = one(box, "SELECT count(*) FROM portal_backup WHERE kind = 'AUTO' AND ok")[0]         # the live scheduler may already have made some
    for i in range(3):
        backup.run_backup("AUTO", "scheduler")
    assert backup.prune(1) == pre + 2 and one(box, "SELECT count(*) FROM portal_backup WHERE kind = 'AUTO' AND note LIKE '%%retention%%'")[0] == pre + 2
    assert one(box, "SELECT count(*) FROM portal_backup WHERE kind = 'MANUAL' AND made_by = 'TEST'")[0] == 1      # manual backups are never pruned


# ---------------------------------------------------------------- import
def _xlsx_bytes():
    from openpyxl import Workbook
    b = io.BytesIO()
    Workbook().save(b)
    return b.getvalue()


def test_import_upload_rules(box, tmp_path, monkeypatch):
    monkeypatch.setattr(importer, "UPLOADS", tmp_path)
    ok = importer.save_upload("calls", "../../evil name.xlsx", _xlsx_bytes(), "ADMIN1")
    assert ok["filename"] == "evil name.xlsx" and (tmp_path / ok["job_id"] / "evil name.xlsx").exists()
    for kind, name, data in (("nope", "a.xlsx", _xlsx_bytes()), ("calls", "a.csv", b"PKxx"), ("calls", "a.xlsx", b"not a zip"), ("calls", "a.xlsx", b""), ("calls", "a.xlsx", b"PK" + b"0" * (importer.MAX_BYTES + 1))):
        with pytest.raises(importer.ImportError_):
            importer.save_upload(kind, name, data, "ADMIN1")
    with pytest.raises(importer.ImportError_) as e:
        importer.load(ok["job_id"], "ADMIN1")                                 # cannot load a file that was never checked
    assert e.value.status == 409
    for bad in ("../x", "zzzz", "0" * 11):
        with pytest.raises(importer.ImportError_):
            importer.detail(bad)


def test_import_check_rehearses_the_load_without_touching_the_database(box, tmp_path, monkeypatch):
    """The check runs the converter AND rehearses the load (--dry-run: rolled back), so it can warn about a refused load beforehand - and saves nothing."""
    monkeypatch.setattr(importer, "UPLOADS", tmp_path)
    src = tmp_path / "IT-IMMDSS ASSET INVENTORY 2026 -Q2.xlsx"
    import shutil
    from pathlib import Path
    real = Path(__file__).resolve().parents[2] / "IT-IMMDSS ASSET INVENTORY 2026 -Q2.xlsx"
    if not real.exists():
        pytest.skip("source workbook not present")
    shutil.copy(real, src)
    job = importer.save_upload("assets", src.name, src.read_bytes(), "ADMIN1")
    before = one(box, "SELECT count(*), max(snapshot_date) FROM asset WHERE is_current = 1")
    latest = one(box, "SELECT max(snapshot_date) FROM asset_snapshot")[0]
    res = importer.check(job["job_id"], (latest or dt.date.today()).isoformat())
    # Whether this old workbook PASSES depends on what the dev database holds today (a replacement already made in the portal, a newer snapshot,
    # a mass removal): the rehearsal reports those honestly, which is its job. What must hold whatever the data is: it ran, and it saved nothing.
    assert any(n.endswith(".xlsx") for n in res["outputs"]) and "ASSET" in res["log"].upper()
    assert "nothing was saved" in res["log"]                                           # it was a rehearsal
    assert one(box, "SELECT count(*), max(snapshot_date) FROM asset WHERE is_current = 1") == before
    assert one(box, "SELECT status FROM portal_import WHERE job_id = %s", (job["job_id"],))[0] == ("CHECKED" if res["ok"] else "CHECK_FAILED")


def test_import_endpoints_need_an_administrator(box, monkeypatch):
    fake_user(monkeypatch, "USER")
    with TestClient(app) as c:
        h = {"X-Requested-With": "itam-portal", "X-Kind": "calls", "X-Filename": "a.xlsx", "Content-Type": "application/octet-stream"}
        assert c.post("/api/admin/import/upload", content=_xlsx_bytes(), headers=h).status_code == 403
        assert c.get("/api/admin/import").status_code == 403
        assert c.post("/api/admin/backups/restore", json={"file": "x.dump", "confirm": "RESTORE"}, headers=HDR).status_code == 403
        assert c.post("/api/admin/email/run", json={"dry_run": True}, headers=HDR).status_code == 403
    fake_user(monkeypatch, "ADMIN")
    with TestClient(app) as c:
        assert c.post("/api/admin/backups/restore", json={"file": "x.dump"}, headers=HDR).status_code == 400            # typed confirmation is required
        assert c.post("/api/pm/rollover", json={"confirm": "no"}, headers=HDR).status_code == 400
        assert c.post("/api/admin/import/upload", content=b"not a workbook", headers={**h, "X-Requested-With": "x"}).status_code == 403


def test_pm_rollover_confirm_is_case_insensitive_and_trimmed(box, monkeypatch):
    """Every text input on this page displays upper-case via CSS regardless of what was actually typed (app.css's
    body { text-transform: uppercase }) - an exact-match confirm check would silently reject a correctly-typed
    lower-case "roll over" that LOOKS right on screen. Only checking that this gets past the confirm gate itself -
    whatever happens next (success, or a different PM error like "the quarter has not ended") is not what's under test."""
    fake_user(monkeypatch, "ADMIN")
    with TestClient(app) as c:
        for typed in ("roll over", "  ROLL OVER  ", "Roll Over"):
            r = c.post("/api/pm/rollover", json={"confirm": typed}, headers=HDR)
            assert r.status_code != 400 or "type roll over" not in r.json().get("error", "").lower()


# ---------------------------------------------------------------- cards
def test_hover_cards_show_personal_data_to_administrators_only(box, monkeypatch):
    a = one(box, "SELECT asset_key, cpf_no FROM asset WHERE is_current = 1 AND record_level = 'ASSET' AND cpf_no IS NOT NULL LIMIT 1")
    c1 = one(box, "SELECT sr_id FROM svc_call WHERE is_current = 1 LIMIT 1")[0]
    for role in ("USER", "ADMIN"):
        fake_user(monkeypatch, role, engineer_key="DINESH GADARIA" if role == "ADMIN" else None)
        with TestClient(app) as c:
            asset = c.get(f"/api/card/asset/{a[0]}").json()
            assert asset["title"] == a[0] and asset["row"]["asset_key"] == a[0] and "calls" in asset["extra"]
            assert c.get(f"/api/card/call/{c1}").json()["row"]["sr_id"] == c1
            emp = c.get(f"/api/card/cpf/{a[1]}").json()
            assert ("mobile_no" in emp["row"]) == (role == "ADMIN") and ("date_of_birth" in emp["row"]) == (role == "ADMIN")
            assert c.get("/api/card/asset/NO-SUCH").status_code == 404 and c.get("/api/card/other/1").status_code == 400


def test_engineer_hover_card_admin_full_self_own_other_hidden(box, monkeypatch):
    """A non-admin never sees another engineer's floating-window details - only their own; an admin sees anyone's."""
    fake_user(monkeypatch, "ADMIN")
    with TestClient(app) as c:
        eng = c.get("/api/card/engineer/DINESH GADARIA").json()
        assert "mobile_no" in eng["row"] and eng["extra"]["checklist"] and eng["subtitle"].startswith("ECODE")
        assert c.get("/api/card/ecode/" + eng["subtitle"].split()[1]).json()["id"] == "DINESH GADARIA"
    fake_user(monkeypatch, "USER", engineer_key="DINESH GADARIA")
    with TestClient(app) as c:
        mine = c.get("/api/card/engineer/DINESH GADARIA")
        assert mine.status_code == 200 and "mobile_no" in mine.json()["row"]
    fake_user(monkeypatch, "USER", engineer_key="SOMEONE ELSE ENTIRELY")
    with TestClient(app) as c:
        other = c.get("/api/card/engineer/DINESH GADARIA")
        assert other.status_code == 404


def test_registers_hide_personal_data_and_people_registers_work(box, monkeypatch):
    fake_user(monkeypatch, "ADMIN")
    with TestClient(app) as c:
        for name in ("employees", "engineers"):
            r = c.get(f"/api/registers/{name}?facets=1&limit=5").json()
            assert r["total"] > 5 and r["rows"] and r["facets"]
            assert not ({"mobile_no", "date_of_birth", "personal_email"} & set(r["rows"][0]))
        g = c.get("/api/registers/employees?f.gender=F").json()
        assert 0 < g["total"] < c.get("/api/registers/employees").json()["total"]
        eid = c.get("/api/registers/employees?limit=1").json()["rows"][0]["id"]
        d = c.get(f"/api/registers/employees/{eid}").json()
        assert d["row"]["employee_name"] and "mobile_no" in d["row"] and d["archived"] is False    # admin sees it here too, matching the hover card
        owner_cpf, owner_eng = one(box, "SELECT cpf_no, engineer_name FROM asset WHERE is_current=1 AND cpf_no IS NOT NULL AND engineer_name IS NOT NULL LIMIT 1")
        # a plain user linked to the engineer who owns this employee's asset (so row-level scoping still lets them open it) must not see personal data;
        # a different username from the ADMIN call above, since the per-user response cache is keyed by username, not role
        fake_user(monkeypatch, "USER", username="PLAINUSER2", engineer_key=owner_eng)
        d2 = c.get(f"/api/registers/employees/{owner_cpf}").json()
        assert "mobile_no" not in d2["row"]
        fake_user(monkeypatch, "ADMIN")
        q = c.get("/api/search?q=dinesh").json()
        assert "employees" in {g["dataset"] for g in q["groups"]} or "engineers" in {g["dataset"] for g in q["groups"]}


# ---------------------------------------------------------------- past quarters
def _frozen_quarter(box, label="Q9 TST-TST 2099", as_of=dt.date(2099, 1, 1)):
    """Two engineers' assets frozen in a made-up quarter, so the assertions never depend on what the real snapshots hold."""
    keys = [r[0] for r in box.execute(f"SELECT asset_key FROM asset WHERE {pm.IN_SCOPE} ORDER BY asset_key LIMIT 4").fetchall()]
    box.execute("DELETE FROM pm_snapshot WHERE quarter_label = %s", (label,))
    for key, eng, st in zip(keys, ["ENG-A", "ENG-A", "ENG-B", None], ["DONE", "PENDING", "DONE_OUTSIDE_QUARTER", "PENDING"]):
        box.execute("INSERT INTO pm_snapshot (quarter_label, as_of, asset_key, asset_class, engineer_name, location_code, pm_status) VALUES (%s,%s,%s,'DESKTOP',%s,'LOC1',%s)", (label, as_of, key, eng, st))
    return label, as_of, keys


def test_past_quarters_list_and_detail(box):
    label, as_of, keys = _frozen_quarter(box)
    row = next(r for r in pm.history()["quarters"] if r["label"] == label)
    assert (row["scope"], row["done"], row["stale"], row["pending"], row["pct_done"]) == (4, 1, 1, 2, 50.0)
    d = pm.history_detail(label)
    assert d["as_of"] == as_of and d["kpi"]["unassigned"] == 1 and {r["label"] for r in d["by_engineer"]} == {"ENG-A", "ENG-B", "UNASSIGNED"}
    assert d["assets"]["total"] == 4 and [r["asset_key"] for r in d["assets"]["rows"]] == sorted(keys)
    assert pm.history_detail(label, status="PENDING")["assets"]["total"] == 2
    assert pm.history_detail(label, q=keys[0].lower())["assets"]["total"] >= 1
    assert pm.history_detail(label, q="100%_")["assets"]["total"] == 0                         # wildcard characters are searched literally
    mine = pm.history_detail(label, eng="ENG-A")
    assert mine["kpi"]["scope"] == 2 and {r["asset_key"] for r in mine["assets"]["rows"]} <= set(keys)
    assert pm.history_detail(label, eng="~no-engineer~")["kpi"]["scope"] == 0                  # a user with no linked engineer sees nothing
    for bad in (lambda: pm.history_detail("Q1 NOT-REAL 1999"), lambda: pm.history_detail(label, as_of=dt.date(1999, 1, 1)), lambda: pm.history_detail(label, status="BOGUS")):
        with pytest.raises(pm.PmError):
            bad()


def test_past_quarters_api_scoping_and_export(box, monkeypatch):
    label, as_of, keys = _frozen_quarter(box)
    fake_user(monkeypatch, role="USER", username="E1", engineer_key="ENG-A")
    with TestClient(app) as c:
        q = c.get("/api/pm/history").json()["quarters"]
        assert next(r for r in q if r["label"] == label)["scope"] == 2                         # only their own two assets
        d = c.get("/api/pm/history/detail", params={"quarter": label}).json()
        assert d["kpi"]["scope"] == 2 and {r["engineer_name"] for r in d["assets"]["rows"]} == {"ENG-A"}
        assert c.get("/api/pm/history/detail", params={"quarter": "nope"}).status_code == 404
        r = c.post("/api/pm/history/export", json={"quarter": label, "format": "csv"}, headers=HDR)
        assert r.status_code == 200 and "Completed" in r.text and "ENG-B" not in r.text and "DONE_OUTSIDE" not in r.text
        r = c.post("/api/pm/history/export", json={"quarter": label, "format": "xlsx"}, headers=HDR)
        assert r.status_code == 200 and zipfile.is_zipfile(io.BytesIO(r.content))
