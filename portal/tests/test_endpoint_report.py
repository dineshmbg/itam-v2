"""Detailed endpoint report: each machine of an analysed endpoint list compared with the PC check, Sensor only / Full suite with its basis, the second-route self check, the per-engineer
copies, the e-mail attachment and the download. Synthetic machines (ANKEZ-*) in the rolled-back test transaction: nothing here depends on what the dev database holds."""
import datetime as dt
import io

import pytest
from openpyxl import Workbook, load_workbook
from starlette.testclient import TestClient

from conftest import HDR, fake_user
from portal.app import endpoint_report as er, inventory_match as m, mailer, swobs
from portal.app import trend_investigation as ti
from portal.app.main import app
from test_inventory_match import APEX, TM, _csv, _tm

NAMES = {k: f"ANKEZ-{k}" for k in ("A1", "A2", "A3", "A4", "A5", "A6")}
ENG = {"A1": "ENG_ONE", "A2": "ENG_ONE", "A3": "ENG_TWO", "A4": "ENG_TWO", "A5": "ENG_ONE", "A6": "ENG_ONE"}


def _asset(con, k):
    ci = NAMES[k]
    con.execute("""INSERT INTO asset (snapshot_date, asset_key, ci_no, asset_class, asset_type, record_level, asset_status, is_current, first_seen_date, last_seen_date, engineer_name, user_name, location_code)
                   VALUES (current_date, %s, %s, 'DESKTOP', 'DESKTOP', 'ASSET', 'IN_USE', 1, current_date, current_date, %s, %s, 'LOC1')""", (ci, ci, ENG[k], "USER " + k))


def _pc(k, installed=True, services="TmListen running; ntrtscan running"):
    facts = {"trend_installed": installed}
    if installed:
        facts.update(trend_status="Apex One Security Agent", trend_version="14.0.1", trend_services=services)
    swobs.record(NAMES[k], "SCRIPT", facts, "ankcheck")


def _xlsx(rows):
    wb = Workbook()
    ws = wb.active
    ws.append(["Endpoint name", "Object type", "Protection Manager", "OS name", "OS type"])
    for n in rows:
        ws.append([n, "Desktop", "Active Directory (on-premises)", "Windows 11 Pro 10.0 (26200)", "Windows"])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def _later(con, run_id, minutes):
    """Every run of a test shares one transaction time; give each its own, as separate real analyses have."""
    con.execute("UPDATE portal_match_run SET at = now() + %s * interval '1 minute' WHERE run_id = %s", (minutes, run_id))


def _setup(box, tmp_path, monkeypatch):
    monkeypatch.setattr(m, "STAGE", tmp_path)
    for k in NAMES:
        _asset(box, k)
    _pc("A1")
    _pc("A3", installed=False)
    _pc("A5")
    P = {"prefix": "ANKEZ", "classes": ["DESKTOP"]}
    sv = m.stage("Apex One Security Agent List1(1.2.3.4).csv", _csv([["Workgroup\\", NAMES["A1"], "10.1.1.1", "Online", "Windows 11", "x64", "ONGC\\1"],
                                                                    ["Workgroup\\", NAMES["A4"], "10.1.1.4", "Offline", "Windows 11", "x64", "ONGC\\4"]], APEX), "ADMIN1")
    _later(box, m.run([sv["stage_id"]], sv["mapping"], P, "ADMIN1")["run_id"], 1)
    vc = m.stage("Endpoint Inventory_x.csv", _csv([_tm(NAMES["A5"], av="Disabled")], TM), "ADMIN1")
    _later(box, m.run([vc["stage_id"]], vc["mapping"], P, "ADMIN1")["run_id"], 2)
    ad = m.stage("Old_Apex One agent Inventory.xlsx", _xlsx([NAMES["A1"], NAMES["A2"] + "$", NAMES["A3"], NAMES["A4"], "ANKEZ-NOPE$"]), "ADMIN1")
    r = m.run([ad["stage_id"]], ad["mapping"], P, "ADMIN1")
    _later(box, r["run_id"], 3)
    return r["run_id"]


def test_each_machine_is_classified_from_its_own_evidence(box, tmp_path, monkeypatch):
    run = _setup(box, tmp_path, monkeypatch)
    d = er.build(run)
    by = {x["ci"]: x for x in d["all"]}
    assert {x["ci"] for x in d["listed"]} == {NAMES[k] for k in ("A1", "A2", "A3", "A4")}          # the primary analysis' own machines
    a1, a2, a3, a4, a5, a6 = (by[NAMES[k]] for k in ("A1", "A2", "A3", "A4", "A5", "A6"))
    assert (a1["cmp"], a1["final"], a1["conf"]) == (er.CMP[0], "Full suite", "High") and "ntrtscan" in a1["basis"]
    assert (a2["cmp"], a2["final"], a2["conf"]) == (er.CMP[3], "Not determined", "-") and a2["file_name"] == NAMES["A2"] + "$" and "ends with $" in a2["notes"]
    assert (a3["cmp"], a3["final"]) == (er.CMP[1], "Not installed") and "no Trend Micro" in a3["basis"]
    assert (a4["cmp"], a4["final"], a4["conf"]) == (er.CMP[2], "Full suite", "Medium") and "Offline" in a4["notes"] and a4["serverlist"] == "Yes - Offline"
    assert (a5["final"], a5["conf"]) == ("Full suite", "High") and a5["view"].startswith("EDR sensor only") and "DISABLED" in a5["notes"]       # installed != switched on
    assert a5["listed"] is False and a6["final"] == "Not determined" and a6["pc_checked"] == "Not checked"
    assert d["verified"] and d["checked"] == 4 and d["bad"] == []


def test_sheets_verify_first_and_check_these_pcs(box, tmp_path, monkeypatch):
    run = _setup(box, tmp_path, monkeypatch)
    d = er.build(run)
    t, info = er.tables(d)
    ti_ = {x["title"]: x for x in t}
    assert list(ti_) == ["About this report", "Endpoint report", "Verify first", "Check these PCs", "Not in register", "Accuracy checks", "All register machines"]
    assert [r["ci"] for r in ti_["Check these PCs"]["rows"]] == [NAMES["A2"]]                      # the only listed machine nothing is known about
    issues = {(r["ci"], r["issue"][:24]) for r in ti_["Verify first"]["rows"]}
    assert (NAMES["A3"], "Listed by the centre, bu") in issues and (NAMES["A5"], "Console shows anti-malwa") in issues
    assert [r["name"] for r in ti_["Not in register"]["rows"]] == ["ANKEZ-NOPE$"]
    assert info["n"] == 4 and info["determined"] == 3 and info["need85"] == 1 and len(ti_["All register machines"]["rows"]) == 6
    data, mime, ext = er.render(d, "tester")
    assert ext == "xlsx" and data[:2] == b"PK"
    names = load_workbook(io.BytesIO(data), read_only=True).sheetnames
    assert "SUMMARY" in names and "ENDPOINT REPORT" in names and "VERIFY FIRST" in names


def test_engineer_copy_holds_only_their_machines(box, tmp_path, monkeypatch):
    run = _setup(box, tmp_path, monkeypatch)
    d = er.build(run)
    mine = er.scoped(d, "ENG_TWO")
    assert {x["ci"] for x in mine["machines"]} == {NAMES["A3"], NAMES["A4"]} and mine["scope"] == "ENG_TWO"
    t, _ = er.tables(mine)
    titles = [x["title"] for x in t]
    assert titles == ["About this report", "Endpoint report", "Verify first", "Check these PCs"]        # no fleet-wide sheets in an engineer's copy
    assert {r["ci"] for r in t[1]["rows"]} == {NAMES["A3"], NAMES["A4"]}
    assert er.scoped(d, "NOBODY") is None and er.file(d, "NOBODY", "x") is None
    name, data, mime = er.file(d, "ENG_ONE", "x")
    assert name.startswith("Endpoint_Report_ENG_ONE_") and data[:2] == b"PK"


def test_a_wrong_conclusion_is_caught_and_the_report_is_marked_do_not_circulate(box, tmp_path, monkeypatch):
    run = _setup(box, tmp_path, monkeypatch)
    real = er._decide
    monkeypatch.setattr(er, "_decide", lambda pcf, console, server, listed: ("Full suite", "forced", "High") if pcf["agent"] == "No" else real(pcf, console, server, listed))
    d = er.build(run)
    assert not d["verified"] and len(d["bad"]) == 1 and NAMES["A3"] in d["bad"][0]
    t, _ = er.tables(d)
    about = {r["k"]: r["v"] for r in t[0]["rows"]}
    assert about["Checked before release"].startswith("DO NOT CIRCULATE")


def test_the_investigation_mail_carries_the_report_for_each_engineer_only_when_verified(box, tmp_path, monkeypatch):
    run = _setup(box, tmp_path, monkeypatch)
    sent = []
    monkeypatch.setattr(mailer, "get_smtp", lambda: {"enabled": True, "portal_url": "", "host": "x"})
    monkeypatch.setattr(mailer, "engineer_email", lambda e: f"{e.lower()}@test.invalid")
    monkeypatch.setattr(mailer, "send", lambda to, subject, text, body, files, cfg=None: sent.append((to, [f[0] for f in files])))
    out = ti.on_run(run, {"username": "ADMIN1", "email": "admin@test.invalid"})
    assert out["endpoint_report"]["built"] and out["endpoint_report"]["verified"] and out["endpoint_report"]["attached"] >= 2
    by = dict(sent)
    assert any(f.startswith("Endpoint_Report_ENG_ONE_") for f in by["eng_one@test.invalid"]) and any(f.startswith("Trend_Micro_Investigation") for f in by["eng_one@test.invalid"])
    assert any(f.startswith("Endpoint_Report_ENG_TWO_") for f in by["eng_two@test.invalid"]) and not any("ENG_ONE" in f for f in by["eng_two@test.invalid"])
    assert any(f.startswith("Endpoint_Report_2") for f in by["admin@test.invalid"])                  # the person who ran it gets everything
    sent.clear()
    box.execute("DELETE FROM notify_log WHERE rule_key = %s", (ti.RULE,))
    real = er._decide
    monkeypatch.setattr(er, "_decide", lambda pcf, console, server, listed: ("Full suite", "forced", "High") if pcf["agent"] == "No" else real(pcf, console, server, listed))
    out = ti.on_run(run, {"username": "ADMIN1", "email": "admin@test.invalid"})
    assert out["endpoint_report"]["verified"] is False and out["endpoint_report"]["attached"] == 0 and out["endpoint_report"]["differences"]
    assert sent and all(not any(f.startswith("Endpoint_Report") for f in files) for _, files in sent)         # the investigation still goes; the doubtful report does not


def test_download_is_scoped_to_the_signed_in_person(box, tmp_path, monkeypatch):
    run = _setup(box, tmp_path, monkeypatch)
    box.execute("UPDATE portal_match_run SET published_at = now() WHERE run_id = %s", (run,))
    c = TestClient(app)
    body = {"kind": "endpoint", "format": "xlsx"}

    def sheet_rows(resp):
        wb = load_workbook(io.BytesIO(resp.content), read_only=True)
        ws = wb["ENDPOINT REPORT"]
        return [r for r in ws.iter_rows(values_only=True) if r and any(str(x or "").startswith("ANKEZ-") for x in r)]
    fake_user(monkeypatch, role="ADMIN", username="ADM")
    r = c.post("/api/match/pack", json=body, headers=HDR)
    assert r.status_code == 200 and "Endpoint_Report_" in r.headers["content-disposition"] and len(sheet_rows(r)) >= 4
    fake_user(monkeypatch, role="USER", username="E2", engineer_key="ENG_TWO")
    r = c.post("/api/match/pack", json=body, headers=HDR)
    assert r.status_code == 200 and len(sheet_rows(r)) == 2 and "ENG_TWO" in r.headers["content-disposition"]
    fake_user(monkeypatch, role="USER", username="E9", engineer_key="NOBODY")
    assert c.post("/api/match/pack", json=body, headers=HDR).status_code == 404
    fake_user(monkeypatch, role="USER", username="NOLINK", engineer_key=None)
    assert c.post("/api/match/pack", json=body, headers=HDR).status_code == 403
    assert c.post("/api/match/pack", json={"kind": "endpoint", "format": "pdf"}, headers=HDR).status_code == 400


def test_which_analyses_count_as_endpoint_reports():
    run = lambda product, name: {"summary": {"product": product, "files": [{"name": name}]}, "tool": "", "filename": name}      # noqa: E731
    assert er.is_endpoint_run(run("trendmicro", "x.csv")) and er.is_endpoint_run(run("apex", "x.csv")) and er.is_endpoint_run(run("generic", "Endpoint report.xlsx"))
    assert er.is_endpoint_run(run("generic", "ank-Old_Apex One agent Inventory.xlsx")) and er.is_endpoint_run(run("generic", "Vision One list.csv"))
    assert not er.is_endpoint_run(run("bigfix", "PAN INDIA BIGFIX DESKTOP_26th_OCT.xlsx")) and not er.is_endpoint_run(run("generic", "assets.xlsx"))
    assert not er.is_endpoint_run(run("bigfix", "endpoint bigfix.xlsx"))                          # BigFix is never one, whatever it is called
