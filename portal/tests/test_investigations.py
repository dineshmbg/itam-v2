"""Investigations: workbook parsing (values as Excel shows them, styles, merges, header detection), file checks, who may use it, nothing is erased. Rolled back (see conftest.box)."""
import datetime as dt
import io

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from starlette.testclient import TestClient

from conftest import HDR, fake_user
from portal.app import investigations as inv
from portal.app.main import app

UP = {"X-Requested-With": "itam-portal"}
PDF = b"%PDF-1.4\n" + b"x" * 100
NAVY = PatternFill("solid", fgColor="1F3864")


def book(rows=40):
    wb = Workbook()
    s = wb.active
    s.title = "Summary"
    s["A1"] = "Coverage report"; s["A1"].font = Font(bold=True)
    s.merge_cells("A2:D2"); s["A2"] = "A merged paragraph of text"
    s["A4"], s["B4"] = "Share", 0.2366; s["B4"].number_format = "0.0%"
    s["A5"], s["B5"] = "Delta", 33; s["B5"].number_format = "+#,##0;-#,##0;0"
    s["A6"], s["B6"] = "Count", 1175
    s["A7"], s["B7"] = "Formula", "=SUM(B5:B6)"                       # no stored result: must not show as text
    s["A8"], s["B8"] = "When", dt.datetime(2026, 10, 6, 10, 39)
    s["A9"], s["B9"] = "Flag", True
    d = wb.create_sheet("Machine status")
    d.append(["Asset (CI)", "Status", "Count"])
    for c in d[1]:
        c.fill = NAVY; c.font = Font(bold=True, color="FFFFFF")
    for i in range(rows):
        d.append([f"ANK{i:03d}", "Installed" if i % 2 else "Missing", i])
    d["B2"].fill = PatternFill("solid", fgColor="C6EFCE")
    h = wb.create_sheet("Hidden")
    h["A1"] = "secret"; h.sheet_state = "hidden"
    out = io.BytesIO(); wb.save(out)
    return out.getvalue()


def up(c, data=None, title="Trend Micro coverage", **h):
    return c.post("/api/investigations/upload", content=book() if data is None else data, headers={**UP, "X-Title": title, "X-Filename": "report.xlsx", **h})


def test_parse_shows_values_as_excel_does():
    s = {x["name"]: x for x in inv.parse(book())}
    assert list(s) == ["Summary", "Machine status"]                      # the hidden sheet is skipped
    doc = s["Summary"]
    assert doc["kind"] == "document" and doc["header"] is None and doc["styles"]["0:0"][0] == 1       # bold is kept
    cells = {(r, c): v for r, row in enumerate(doc["rows"]) for c, v in enumerate(row)}
    assert cells[(3, 1)] == "23.7%" and cells[(4, 1)] == "+33" and cells[(5, 1)] == 1175 and cells[(7, 1)] == "06-Oct-2026 10:39" and cells[(8, 1)] == "Yes"
    assert cells.get((6, 1)) is None                                      # formula without a stored result
    assert inv._show("ctrl\x01char\x0b", "") == "ctrlchar"               # control characters never reach the page (openpyxl itself refuses to write them, other tools may not)
    assert doc["merges"] == [[1, 0, 1, 1]]                                # A2:D2, cut to the two columns that hold data
    tab = s["Machine status"]
    assert tab["kind"] == "table" and tab["header"] == 0 and len(tab["rows"]) == 41 and tab["styles"]["1:1"][1] == "C6EFCE"


def test_parse_refuses_what_is_not_a_workbook(monkeypatch):
    for bad in (b"hello", b"PK\x03\x04 not really a zip", b""):
        with pytest.raises(inv.InvError):
            inv.parse(bad)
    monkeypatch.setattr(inv, "MAX_UNZIPPED", 100)                         # a zip that expands to more than the cap is refused before it is opened
    with pytest.raises(inv.InvError) as e:
        inv.parse(book())
    assert e.value.status == 413
    monkeypatch.setattr(inv, "MAX_UNZIPPED", 150 * 1024 * 1024)
    monkeypatch.setattr(inv, "MAX_CELLS", 10)
    with pytest.raises(inv.InvError):
        inv.parse(book())


def test_only_administrators_and_team_leads(monkeypatch, box):
    u = fake_user(monkeypatch, role="USER", username="E1", engineer_key="X")
    c = TestClient(app)
    assert c.get("/api/investigations").status_code == 403
    assert up(c).status_code == 403
    u["lead_tools"] = True
    assert up(c).status_code == 200 and c.get("/api/investigations").status_code == 200
    u.update(lead_tools=False, role="ADMIN", group="USER", extended_access=True)         # extended access is not enough
    assert c.get("/api/investigations").status_code == 403
    u.update(group="ADMIN", extended_access=False)
    assert c.get("/api/investigations").status_code == 200
    u["read_only"] = True                                                                 # the demo account has a public password
    assert c.get("/api/investigations").status_code == 403 and up(c).status_code == 403


def test_upload_list_read_download_remove(monkeypatch, box):
    fake_user(monkeypatch, role="ADMIN", username="A1")
    c = TestClient(app)
    data = book()
    r = up(c, data, title="  Trend   Micro coverage  ", **{"X-Note": "check %26 act"})
    assert r.status_code == 200 and r.json()["sheets"] == 2
    iid = r.json()["inv_id"]
    it = next(x for x in c.get("/api/investigations").json()["items"] if x["inv_id"] == iid)
    assert it["title"] == "Trend Micro coverage" and it["note"] == "check & act" and it["uploaded_by"] == "A1" and not it["superseded_at"] and it["pdf_name"] is None
    assert [s["name"] for s in it["sheet_index"]] == ["Summary", "Machine status"] and it["sheet_index"][1]["rows"] == 40 and it["sheet_index"][1]["kind"] == "table"
    sh = c.get(f"/api/investigations/{iid}/sheet/1").json()
    assert sh["name"] == "Machine status" and sh["rows"][0] == ["Asset (CI)", "Status", "Count"]
    assert c.get(f"/api/investigations/{iid}/sheet/9").status_code == 404 and c.get("/api/investigations/999999999/sheet/0").status_code == 404
    f = c.get(f"/api/investigations/{iid}/file/xlsx")
    assert f.status_code == 200 and f.content == data and "attachment" in f.headers["content-disposition"] and "nosniff" in f.headers["x-content-type-options"]
    assert c.get(f"/api/investigations/{iid}/file/pdf").status_code == 404 and c.get(f"/api/investigations/{iid}/file/zip").status_code == 404
    # the PDF can be added once and is never overwritten
    p = lambda body: c.post(f"/api/investigations/{iid}/pdf", content=body, headers={**UP, "X-Filename": "summary.pdf"})
    assert p(b"not a pdf").status_code == 415
    assert p(PDF).status_code == 200 and c.get(f"/api/investigations/{iid}/file/pdf").content == PDF
    assert p(PDF + b"2").status_code == 409 and c.get(f"/api/investigations/{iid}/file/pdf").content == PDF
    # remove hides it from the current list; the files and the sheets stay
    assert c.post("/api/investigations/remove", json={"inv_id": iid}, headers=HDR).status_code == 200
    it = next(x for x in c.get("/api/investigations").json()["items"] if x["inv_id"] == iid)
    assert it["superseded_at"] and it["superseded_by"] == "A1"
    assert c.get(f"/api/investigations/{iid}/file/xlsx").content == data and c.get(f"/api/investigations/{iid}/sheet/0").status_code == 200
    assert c.post("/api/investigations/remove", json={"inv_id": iid}, headers=HDR).status_code == 404         # already removed
    assert box.execute("SELECT count(*) FROM portal_investigation WHERE inv_id = %s AND xlsx IS NOT NULL AND pdf IS NOT NULL", (iid,)).fetchone()[0] == 1       # removal erased nothing


def test_upload_checks(monkeypatch, box):
    fake_user(monkeypatch, role="ADMIN")
    c = TestClient(app)
    before = box.execute("SELECT count(*) FROM portal_investigation").fetchone()[0]          # the real database may already hold investigations
    assert up(c, title="ab").status_code == 400 and up(c, title="").status_code == 400
    assert up(c, b"").status_code == 400
    assert up(c, b"plain text, not a workbook").status_code == 415
    assert up(c, book() + b"z" * inv.MAX_XLSX).status_code in (400, 413)
    assert c.post("/api/investigations/upload", content=book(), headers={"X-Title": "No marker header"}).status_code == 403      # not a portal request
    assert c.post("/api/investigations/upload", content=book(), headers={**UP, "X-Title": "Other site", "Origin": "http://evil.example"}).status_code == 403
    assert c.post("/api/investigations/999999999/pdf", content=PDF, headers=UP).status_code == 404
    assert box.execute("SELECT count(*) FROM portal_investigation").fetchone()[0] == before          # none of the refused uploads was stored
