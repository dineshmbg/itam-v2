"""Inventory match: the centre's report(s) matched against the asset register by Asset (CI) = computer name.
Reports are built from whatever machines the dev database holds, so nothing here depends on specific data."""
import csv
import datetime as dt
import io

import pytest
from starlette.testclient import TestClient

from conftest import HDR, fake_user
from portal.app import inventory_match as m
from portal.app.main import app

NOW = dt.datetime(2026, 10, 6, 12, 0)
BF = ("Computer Name", "Device Type", "IP Address", "OS", "CPU", "Last Report Time")
FRESH, OLD = "Tue, 06 Oct 2026 01:00:00 +0000", "Mon, 01 Jun 2026 01:00:00 +0000"


def _machines(box, n=6):
    rows = box.execute("""SELECT ci_no, ip_address FROM asset WHERE is_current = 1 AND record_level = 'ASSET' AND asset_class = 'DESKTOP' AND asset_status = 'IN_USE'
                          AND ci_no ~ '^[A-Z]{3,4}' ORDER BY ci_no LIMIT %s""", (n,)).fetchall()
    if len(rows) < n:
        pytest.skip("not enough desktops in the dev database")
    return rows


def _prefix(ci):
    return ci[:3]


def _csv(rows, header=BF):
    b = io.StringIO()
    w = csv.writer(b)
    w.writerow(header)
    w.writerows(rows)
    return b.getvalue().encode()


def _bf(name, ip="10.9.9.9", seen=FRESH, os="Win11 10.0.26200"):
    return [name, "Desktop", ip, os, "i7", seen]


TM = ("Endpoint name", "Object type", "Recommended actions", "Last agent status reported", "Protection Manager", "Agent connection status", "Agent runtime protection status", "XDR for Endpoints (EDR)",
      "Endpoint group", "IP address", "OS name", "OS version", "Anti-malware", "Agent version status", "Last scanned", "Endpoint GUID", "Protection module last connected")


def _tm(name, group="ANK", conn="Communicating", action="", av="Enabled - recommended features", seen="Last 24 hours (2026-10-06 10:00:00)", ip="10.9.9.9", agent="controlledLatestVersion"):
    return [name, "Desktop", action, seen, "Standard Endpoint Protection Manager", conn, "Protected", "Enabled", group, ip, "Windows 11", "10.0 (Build 26200)", av, agent, "2026-10-05 10:00:00", "g-" + name, "2026-10-06 10:00:00"]


def _files(*parts):
    out = []
    for i, (header, rows) in enumerate(parts):
        h, r = m.read_file(f"p{i}.csv", _csv(rows, header))
        out.append({"name": f"p{i}.csv", "headers": h, "rows": r, "product": m.profile_file(h)["product"]})
    return out


def _go(parts, params, today=NOW):
    files = _files(*parts)
    return m.analyse(files, m.profile_file(files[0]["headers"])["mapping"], params, today=today)


def test_read_file_and_profile_the_product_from_the_columns_not_the_name():
    h, rows = m.read_file("whatever.csv", _csv([_bf("A-1")]))
    p = m.profile_file(h)
    assert h == list(BF) and p["product"] == "bigfix" and p["mapping"]["name"] == "Computer Name" and p["mapping"]["seen"] == "Last Report Time"
    h2, _ = m.read_file("Endpoint Inventory_x.csv", _csv([_tm("A-1")], TM))
    p2 = m.profile_file(h2)
    assert p2["product"] == "trendmicro" and p2["label"] == "Trend Micro Vision One" and p2["mapping"]["group"] == "Endpoint group" and p2["confidence"] > 0.5
    h3, _ = m.read_file("x.csv", _csv([["a", "1"]], ("Hostname", "Foo")))
    assert m.profile_file(h3)["product"] == "generic" and m.profile_file(h3)["mapping"]["name"] == "Hostname"
    with pytest.raises(m.MatchError):
        m.read_file("r.pdf", b"%PDF")
    with pytest.raises(m.MatchError):
        m.read_file("r.xlsx", b"not a workbook")


def test_key_time_and_typo_helpers():
    assert m.key_of(" ankaon26lt001.ongc.local ") == "ANKAON26LT001" and m.key_of("10.1.2.3") == "10.1.2.3"
    assert m._when("Tue, 06 Oct 2026 01:35:28 +0000") == dt.datetime(2026, 10, 6, 1, 35, 28) and m._when("garbage") is None
    assert m._seen("Last 24 hours (2026-10-06 10:39:40)", "trendmicro") == dt.datetime(2026, 10, 6, 10, 39, 40)
    assert m._typo("ANKAON26OL143", "ANKAON260L143") and m._typo("ANKAON26OL178", "ANKAON26OL78")              # O/0 look-alike; a character missing
    assert m._typo("ANKAON00DT140", "ANKAON00DT147") is None                                                   # a different digit is a different machine
    assert m._install_reason("Installation package is missing one or more required files.") == "installation package incomplete"


def test_bigfix_match_health_dormant_mismatch_and_unregistered(box):
    ms = _machines(box)
    pre = _prefix(ms[0][0])
    report = [_bf(ms[0][0]), _bf(ms[1][0].lower() + ".ongc.local", seen=OLD), _bf(pre + "-NOTINREGISTER-1"), _bf(ms[0][0], seen=OLD)]
    if ms[2][1]:
        report.append(_bf(pre + "-RENAMED-1", ip=ms[2][1].strip()))                                           # a different name on the third machine's own IP
    s, res = _go([(BF, report)], {"prefix": pre, "classes": ["DESKTOP"]})
    by = {r["ci"]: r for r in res["assets"]}
    assert s["tool"] == "BigFix" and s["product"] == "bigfix"                                                  # tool name comes from the file, not from anything typed
    assert by[ms[0][0]]["installed"] and by[ms[0][0]]["health"] == "Healthy" and by[ms[0][0]]["reporting"] == "Active"
    assert by[ms[1][0]]["installed"] and by[ms[1][0]]["reporting"] == "Dormant" and by[ms[1][0]]["health"] == "Needs attention" and "Not reported for" in by[ms[1][0]]["issues"]
    assert not by[ms[3][0]]["installed"] and by[ms[3][0]]["priority"] == "High" and by[ms[3][0]]["health"] == "Not installed"
    if ms[2][1]:
        assert by[ms[2][0]]["possible"].endswith("RENAMED-1") and by[ms[2][0]]["possible_by"] == "same IP address" and s["name_mismatch"] == 1
    assert any(u["name"].endswith("NOTINREGISTER-1") and u["kind"] == "unregistered" for u in res["unregistered"])
    assert s["duplicate_names"] == 1 and res["duplicates"][0]["rows"] == 2
    assert s["installed"] + s["missing"] == s["assets"] and s["healthy"] + s["attention"] == s["installed"] and s["reconciled"] is True
    assert s["coverage"] == round(100 * s["installed"] / s["assets"], 1) and s["findings"][0]["finding"].startswith("BigFix:")
    assert s["rag"] in ("red", "amber", "green") and {c["check"] for c in s["checks"]} >= {"What the file is", "Report date"}
    assert s["report_as_of"].startswith("2026-10-06")                                                          # ages are counted from the report's own date


def test_trend_micro_parts_are_combined_and_installed_does_not_mean_healthy(box):
    ms = _machines(box, 6)
    pre = _prefix(ms[0][0])
    ank_part = [_tm(ms[0][0]), _tm(ms[1][0], conn="Not communicating"), _tm(ms[2][0], action="Installation package is missing one or more required files."), _tm(ms[3][0], av="Disabled"),
                _tm("CAUXYZ01PSAV"), _tm(ms[0][0])]                                                          # a foreign device in our group, and a repeat of a machine
    workgroup_part = [_tm(ms[4][0], group="Workgroup"), _tm("DESKTOP-ABC1234", group="Workgroup", ip="10.9.9.99")]
    s, res = _go([(TM, ank_part), (TM, workgroup_part)], {"prefix": pre, "classes": ["DESKTOP"]})
    by = {r["ci"]: r for r in res["assets"]}
    assert s["tool"] == "Trend Micro Vision One" and len(s["files"]) == 2
    assert by[ms[0][0]]["health"] == "Healthy"
    assert by[ms[1][0]]["installed"] and by[ms[1][0]]["health"] == "Needs attention" and "not communicating" in by[ms[1][0]]["issues"]
    assert "installation failed: installation package incomplete" in by[ms[2][0]]["issues"]
    assert "Anti-malware disabled" in by[ms[3][0]]["issues"]
    assert by[ms[4][0]]["installed"] and by[ms[4][0]]["r_group"] == "Workgroup"                                  # the second file was read as part of the same report
    assert s["group_issues"] >= 1 and s["home_group"] == "ANK" and s["foreign"] == 1
    assert any("outside the usual group" in c["check"] for c in s["checks"]) and any("Other sites" in c["check"] for c in s["checks"])
    assert s["installed"] > s["healthy"] and s["attention"] >= 3 and s["reconciled"] is True


def test_files_of_different_products_are_refused(box):
    files = _files((BF, [_bf("ANKA-1")]), (TM, [_tm("ANKA-2")]))
    with pytest.raises(m.MatchError) as e:
        m.analyse(files, {"name": "Computer Name"}, {"prefix": "ANK"}, today=NOW)
    assert "different products" in str(e.value)


def test_needs_a_name_column_and_a_matching_prefix(box):
    ms = _machines(box, 2)
    files = _files((BF, [_bf(ms[0][0])]))
    with pytest.raises(m.MatchError):
        m.analyse(files, {}, {}, today=NOW)
    with pytest.raises(m.MatchError):
        m.analyse(files, m.profile_file(files[0]["headers"])["mapping"], {"prefix": "ZZQ"}, today=NOW)           # the report has no such site: say so, do not report 0%


def test_run_is_stored_reopened_and_exported(box, tmp_path, monkeypatch):
    monkeypatch.setattr(m, "STAGE", tmp_path)
    ms = _machines(box)
    st = m.stage("centre.csv", _csv([_bf(ms[0][0])]), "ADMIN1")
    assert st["product"] == "bigfix" and st["mapping"]["name"] == "Computer Name" and st["rows"] == 1 and (tmp_path / f"{st['stage_id']}__centre.csv").exists()
    r = m.run([st["stage_id"]], st["mapping"], {"prefix": _prefix(ms[0][0]), "classes": ["DESKTOP"]}, "ADMIN1")
    back = m.get(r["run_id"])
    assert back["summary"]["installed"] == r["summary"]["installed"] and len(back["assets"]) == len(r["assets"]) and m.history()[0]["run_id"] == r["run_id"]
    xlsx, _, ext = m.render(back, "xlsx", {"username": "ADMIN1"})
    pdf, _, pext = m.render(back, "pdf", {"username": "ADMIN1"})
    assert ext == "xlsx" and xlsx[:2] == b"PK" and pext == "pdf" and pdf[:4] == b"%PDF"
    from openpyxl import load_workbook
    names = {n.upper() for n in load_workbook(io.BytesIO(xlsx)).sheetnames}
    assert {"ABOUT THIS REPORT", "KEY FINDINGS", "DATA CHECKS ON THE REPORT", "COVERAGE BY ENGINEER", "FULL MATCH"} <= names
    for bad in ("../x", "zzzz", ""):
        with pytest.raises(m.MatchError):
            m._staged(bad)
    with pytest.raises(m.MatchError):
        m.get("nope")
    with pytest.raises(m.MatchError):
        m.run([], {}, {}, "ADMIN1")


def test_endpoints_need_a_full_administrator(box, monkeypatch, tmp_path):
    monkeypatch.setattr(m, "STAGE", tmp_path)
    fake_user(monkeypatch, "USER")
    h = {"X-Requested-With": "itam-portal", "X-Filename": "a.csv", "Content-Type": "application/octet-stream"}
    with TestClient(app) as c:
        assert c.post("/api/match/upload", content=b"a,b\n1,2\n", headers=h).status_code == 403
        assert c.get("/api/match/meta").status_code == 403
        assert c.post("/api/match/run", json={}, headers=HDR).status_code == 403
    user = fake_user(monkeypatch, "ADMIN")
    from portal.app import mailer, routes_tools
    monkeypatch.setattr(mailer, "get_smtp", lambda: {**mailer.DEFAULT_SMTP, "enabled": False})        # a test must never send a real e-mail

    async def me(request):
        return user
    monkeypatch.setattr(routes_tools, "current_user", me)          # the upload route reads the user itself (it takes a raw body), so it needs the same stand-in
    ms = _machines(box)
    with TestClient(app) as c:
        up = c.post("/api/match/upload", content=_csv([_bf(ms[0][0])]), headers=h)
        assert up.status_code == 200 and up.json()["product"] == "bigfix"
        run = c.post("/api/match/run", json={"stage_ids": [up.json()["stage_id"]], "mapping": up.json()["mapping"], "params": {"prefix": _prefix(ms[0][0]), "classes": ["DESKTOP"]}}, headers=HDR)
        assert run.status_code == 200 and run.json()["summary"]["installed"] >= 1
        rid = run.json()["run_id"]
        assert c.get(f"/api/match/runs/{rid}").status_code == 200 and c.get("/api/match/runs/nope").status_code == 404
        assert c.post("/api/match/export", json={"run_id": rid, "format": "pdf"}, headers=HDR).headers["content-type"] == "application/pdf"
        assert c.post("/api/match/export", json={"run_id": rid, "format": "csv"}, headers=HDR).status_code == 400
        assert c.post("/api/match/run", json={"stage_ids": ["bad"], "mapping": {"name": "x"}}, headers=HDR).status_code == 400


def test_each_engineer_gets_only_their_own_list_once(box, tmp_path, monkeypatch):
    from portal.app import mailer
    monkeypatch.setattr(m, "STAGE", tmp_path)
    rows = box.execute("""SELECT ci_no, engineer_name FROM asset WHERE is_current = 1 AND record_level = 'ASSET' AND asset_class = 'DESKTOP' AND asset_status = 'IN_USE' AND engineer_name IS NOT NULL
                          AND ci_no ~ '^[A-Z]{3,4}' ORDER BY ci_no LIMIT 400""").fetchall()
    if len({r[1] for r in rows}) < 2:
        pytest.skip("needs desktops with at least two engineers")
    st = m.stage("c.csv", _csv([_bf(rows[0][0])]), "ADMIN1")                    # only one machine is in the report: everyone else has gaps
    r = m.run([st["stage_id"]], st["mapping"], {"prefix": _prefix(rows[0][0]), "classes": ["DESKTOP"]}, "ADMIN1")
    sent = []
    monkeypatch.setattr(mailer, "engineer_email", lambda key: f"{key.replace(' ', '.').lower()}@example.com")
    monkeypatch.setattr(mailer, "send", lambda to, subject, text, html=None, att=(), cfg=None: sent.append((to, subject, att)))
    monkeypatch.setattr(mailer, "get_smtp", lambda: {**mailer.DEFAULT_SMTP, "enabled": True, "host": "h", "from_addr": "a@b.c", "has_password": False})
    user = {"username": "ADMIN1"}
    dry = m.send_to_engineers(r["run_id"], user, dry_run=True)
    assert not sent and all(p["status"] in ("READY", "NO_ENGINEER", "NOTHING_URGENT") for p in dry["plan"]) and any(p["status"] == "READY" for p in dry["plan"])
    out = m.send_to_engineers(r["run_id"], user)
    ready = [p for p in out["plan"] if p["status"] == "SENT"]
    assert ready and len(sent) == len(ready) and len({s[0] for s in sent}) == len(sent)
    from openpyxl import load_workbook
    for to, subject, att in sent:                                            # each attachment holds only that engineer's machines
        eng = next(p["engineer"] for p in ready if p["email"] == to)
        wb = load_workbook(io.BytesIO(att[0][1]))
        ws = wb[next(n for n in wb.sheetnames if n.startswith("YOUR WORK LIST"))]      # the attachment is the field report; sheet names are upper-cased and cut to 31 characters
        cis = [row[0] for row in ws.iter_rows(min_row=2, values_only=True) if row[0]]
        assert cis and all(next(a for a in r["assets"] if a["ci"] == c)["engineer"] == eng for c in cis) and "no agent" in subject
    again = m.send_to_engineers(r["run_id"], user)
    assert len(sent) == len(ready) and all(p["status"] in ("ALREADY_SENT", "NO_ENGINEER", "NOTHING_URGENT") for p in again["plan"])     # a second send mails nobody twice
    monkeypatch.setattr(mailer, "get_smtp", lambda: {**mailer.DEFAULT_SMTP, "enabled": False})
    with pytest.raises(mailer.MailError):
        m.send_to_engineers(r["run_id"], user)


def test_engineers_see_only_their_own_machines_and_only_once_shared(box, tmp_path, monkeypatch):
    monkeypatch.setattr(m, "STAGE", tmp_path)
    rows = box.execute("""SELECT ci_no, engineer_name FROM asset WHERE is_current = 1 AND record_level = 'ASSET' AND asset_class = 'DESKTOP' AND asset_status = 'IN_USE' AND engineer_name IS NOT NULL
                          AND ci_no ~ '^[A-Z]{3,4}' ORDER BY ci_no LIMIT 400""").fetchall()
    engs = sorted({r[1] for r in rows})
    if len(engs) < 2:
        pytest.skip("needs desktops with at least two engineers")
    me, other = engs[0], engs[1]
    st = m.stage("c.csv", _csv([_bf(rows[0][0])]), "ADMIN1")
    r = m.run([st["stage_id"]], st["mapping"], {"prefix": _prefix(rows[0][0]), "classes": ["DESKTOP"]}, "ADMIN1")
    rid = r["run_id"]
    m.publish(rid, "ADMIN1", on=False)                                    # a run is shared on creation now; start from the unshared state this test is about
    mine = {a["ci"] for a in r["assets"] if a["engineer"] == me}
    theirs = {a["ci"] for a in r["assets"] if a["engineer"] == other}
    assert mine and theirs
    sc = m.scoped(m.get(rid), me)
    assert {a["ci"] for a in sc["assets"]} == mine and sc["summary"]["machines"] == len(mine) and sc["summary"]["installed"] + sc["summary"]["missing"] == len(mine)
    # as the engineer, through the API
    fake_user(monkeypatch, "USER", username="ENG1", engineer_key=me)
    with TestClient(app) as c:
        assert c.get("/api/match/mine").json()["runs"] == [] or all(x["run_id"] != rid for x in c.get("/api/match/mine").json()["runs"])      # not shared yet
        assert c.get(f"/api/match/mine/{rid}").status_code == 403
        assert c.post("/api/match/publish", json={"run_id": rid}, headers=HDR).status_code in (403, 404)                                 # no manual sharing route: it happens when the tool runs
        m.publish(rid, "ADMIN1")
        assert any(x["run_id"] == rid for x in c.get("/api/match/mine").json()["runs"])
        got = c.get(f"/api/match/mine/{rid}", params={"as": other})                                                                               # asking for someone else's list is ignored
        assert got.status_code == 200 and {a["ci"] for a in got.json()["assets"]} == mine and not ({a["ci"] for a in got.json()["assets"]} & theirs)
        x = c.post("/api/match/mine/export", json={"run_id": rid, "as": other}, headers=HDR)
        assert x.status_code == 200 and x.content[:2] == b"PK"
        m.publish(rid, "ADMIN1", on=False)
        assert c.get(f"/api/match/mine/{rid}").status_code == 403                                                                                   # unshared again
    fake_user(monkeypatch, "USER", username="NOLINK", engineer_key=None)
    with TestClient(app) as c:
        assert c.get("/api/match/mine").status_code == 403                                                                                          # not linked to an engineer: nothing to show
    fake_user(monkeypatch, "ADMIN")
    with TestClient(app) as c:
        assert c.get(f"/api/match/mine/{rid}").status_code == 400                                                                                   # an administrator must choose whose view
        ok = c.get(f"/api/match/mine/{rid}", params={"as": other})
        assert ok.status_code == 200 and {a["ci"] for a in ok.json()["assets"]} == theirs


def test_engineer_gets_a_sign_in_notice_for_a_shared_run_until_opened(box, tmp_path, monkeypatch):
    monkeypatch.setattr(m, "STAGE", tmp_path)
    rows = box.execute("""SELECT ci_no, engineer_name FROM asset WHERE is_current = 1 AND record_level = 'ASSET' AND asset_class = 'DESKTOP' AND asset_status = 'IN_USE' AND engineer_name IS NOT NULL
                          AND ci_no ~ '^[A-Z]{3,4}' ORDER BY ci_no LIMIT 400""").fetchall()
    if not rows:
        pytest.skip("needs desktops with an engineer")
    me = rows[0][1]
    st = m.stage("c.csv", _csv([_bf(rows[0][0])]), "ADMIN1")
    rid = m.run([st["stage_id"]], st["mapping"], {"prefix": _prefix(rows[0][0]), "classes": ["DESKTOP"]}, "ADMIN1")["run_id"]
    m.publish(rid, "ADMIN1", on=False)
    fake_user(monkeypatch, "USER", username="ENG1", engineer_key=me)
    with TestClient(app) as c:
        assert c.get("/api/match/notice").json()["notice"] in (None,) or c.get("/api/match/notice").json()["notice"]["run_id"] != rid      # not shared yet
        m.publish(rid, "ADMIN1")
        n = c.get("/api/match/notice").json()["notice"]
        assert n["run_id"] == rid and n["summary"]["machines"] > 0 and n["summary"]["machines"] == m.scoped(m.get(rid), me)["summary"]["machines"]
        assert c.post("/api/match/notice/seen", json={"run_id": rid}, headers=HDR).status_code == 200
        again = c.get("/api/match/notice").json()["notice"]
        assert again is None or again["run_id"] != rid                                                                                        # opened: not shown again
    fake_user(monkeypatch, "ADMIN")
    with TestClient(app) as c:
        assert c.get("/api/match/notice").json()["notice"] is None                                                                           # administrators get no engineer notice


def test_every_run_is_shared_and_mailed_automatically(box, tmp_path, monkeypatch):
    from portal.app import mailer
    monkeypatch.setattr(m, "STAGE", tmp_path)
    rows = box.execute("""SELECT ci_no FROM asset WHERE is_current = 1 AND record_level = 'ASSET' AND asset_class = 'DESKTOP' AND asset_status = 'IN_USE' AND engineer_name IS NOT NULL
                          AND ci_no ~ '^[A-Z]{3,4}' ORDER BY ci_no LIMIT 5""").fetchall()
    if len(rows) < 5:
        pytest.skip("not enough desktops")
    st = m.stage("c.csv", _csv([_bf(rows[0][0])]), "ADMIN1")
    r = m.run([st["stage_id"]], st["mapping"], {"prefix": _prefix(rows[0][0]), "classes": ["DESKTOP"]}, "ADMIN1")
    sent = []
    monkeypatch.setattr(mailer, "engineer_email", lambda key: f"{key.replace(' ', '.').lower()}@example.com")
    monkeypatch.setattr(mailer, "send", lambda to, subject, text, html=None, att=(), cfg=None: sent.append(to))
    user = {"username": "ADMIN1"}
    # e-mail switched off: still shared in the portal, nothing sent, the reason and the would-be recipients are kept with the analysis
    monkeypatch.setattr(mailer, "get_smtp", lambda: {**mailer.DEFAULT_SMTP, "enabled": False})
    off = m.distribute(r["run_id"], user)
    assert not sent and "not switched on" in off["error"] and any(p["status"] == "READY" for p in off["plan"])
    assert any(x["run_id"] == r["run_id"] for x in m.published()) and m.get(r["run_id"])["mail"]["error"] == off["error"]
    # e-mail on: one message per engineer with something to do, outcome stored
    monkeypatch.setattr(mailer, "get_smtp", lambda: {**mailer.DEFAULT_SMTP, "enabled": True, "host": "h", "from_addr": "a@b.c", "has_password": False})
    on = m.distribute(r["run_id"], user)
    assert on["error"] is None and sent and len(sent) == len({p["email"] for p in on["plan"] if p["status"] == "SENT"}) and m.get(r["run_id"])["mail"]["error"] is None


APEX = ("Domain", "Endpoint", "IP Address", "Connection Status", "Platform", "Architecture", "Logon User")


def test_apex_files_sensor_full_and_work_lists(box, tmp_path, monkeypatch):
    """Apex One CSV + old inventory workbook as one report: a machine in both keeps the details of whichever file has them, the check script decides Sensor only / Full suite,
    and a machine the PC says has an agent is a 'check' job, never a 'no agent' one."""
    from openpyxl import Workbook
    from portal.app import swobs
    monkeypatch.setattr(m, "STAGE", tmp_path)
    rows = box.execute("""SELECT ci_no, engineer_name FROM asset WHERE is_current = 1 AND record_level = 'ASSET' AND asset_class = 'DESKTOP' AND asset_status = 'IN_USE' AND engineer_name IS NOT NULL
                         AND ci_no ~ '^[A-Z]{3,4}' AND NOT EXISTS (SELECT 1 FROM portal_sw_obs o WHERE o.asset_key = asset.asset_key)
                         ORDER BY ci_no LIMIT 6""").fetchall()  # real PC-check observations would turn 'unknown'/'both' into Full suite
    if len(rows) < 6:
        pytest.skip("needs six desktops")
    full, sensor, unknown, both, offline, pc_only = [r[0] for r in rows]
    c = io.StringIO()
    w = csv.writer(c)
    w.writerow(APEX)
    for ci, st in ((full, "Online"), (sensor, "Online"), (unknown, "Online"), (both, "Online"), (offline, "Offline")):
        w.writerow(["Workgroup_x\\", ci, "10.1.1.9", st, "Windows 11  10.0.26100", "x64", "ONGC\\123"])
    wb = Workbook()
    ws = wb.active
    ws.append(["Endpoint name", "Object type", "Protection Manager", "OS name", "OS type"])
    ws.append([both + "$", "Desktop", "Active Directory (on-premises)", "Windows 11 Pro 10.0 (26200)", "Windows"])
    ws.append([both + "$", "Desktop", "Active Directory (on-premises)", "Windows 11 Pro 10.0 (26200)", "Windows"])
    buf = io.BytesIO()
    wb.save(buf)
    st1 = m.stage("list.csv", c.getvalue().encode(), "ADMIN1")
    st2 = m.stage("old.xlsx", buf.getvalue(), "ADMIN1")
    assert st1["product"] == st2["product"] == "apex"
    swobs.record(full, "SCRIPT", {"trend_services": "ntrtscan; TmListen"}, "t")
    swobs.record(sensor, "SCRIPT", {"trend_services": "TmListen running; EndpointBasecamp running"}, "t")
    swobs.record(pc_only, "SCRIPT", {"trend_services": "ntrtscan"}, "t")
    r = m.run([st2["stage_id"], st1["stage_id"]], st1["mapping"], {"prefix": _prefix(full), "classes": ["DESKTOP"]}, "ADMIN1")
    by = {a["ci"]: a for a in m._with_state(r["assets"])}
    assert [by[x]["state"] for x in (full, sensor, unknown, both)] == ["Full suite", "Sensor only", "Not determined", "Not determined"]
    assert by[pc_only]["state"] == "Not installed" and "Do not reinstall" in by[pc_only]["todo"] and "do not reinstall" in by[pc_only]["action"]
    assert by[both]["r_ip"] == "10.1.1.9" and by[both]["r_conn"] == "Online" and by[offline]["r_conn"] == "Offline"          # the file that has the details wins, in either upload order
    assert r["summary"]["suite_counts"]["Sensor only"] >= 1 and r["summary"]["reconciled"]
    eng = by[pc_only]["engineer"].upper()
    assert pc_only not in {x["ci"] for x in m._work(r, eng)["miss"]} and pc_only in {x["ci"] for x in m._work(r, eng)["quiet"]}
    d = m.get(r["run_id"])
    merged, absent = m._merged_rows(d)
    assert {x["endpoint"] for x in merged} >= {full, sensor, both} and any(x["endpoint"] == pc_only for x in absent)
    for fmt in ("xlsx", "csv", "pdf"):
        assert m.consolidated(d, fmt, {"username": "t"})[0]


def test_pack_buttons_scope_by_role(box, tmp_path, monkeypatch):
    """The Option 2 buttons: an engineer gets only their own machines and only from a shared Trend analysis; administrators/team leads get everything."""
    from openpyxl import load_workbook
    monkeypatch.setattr(m, "STAGE", tmp_path)
    rows = box.execute("""SELECT ci_no, engineer_name FROM asset WHERE is_current = 1 AND record_level = 'ASSET' AND asset_class = 'DESKTOP' AND asset_status = 'IN_USE' AND engineer_name IS NOT NULL
                         AND ci_no ~ '^[A-Z]{3,4}' ORDER BY ci_no LIMIT 400""").fetchall()
    engs = sorted({r[1] for r in rows})
    if len(engs) < 2:
        pytest.skip("needs desktops with at least two engineers")
    me, other = engs[0], engs[1]
    c_ = io.StringIO()
    w = csv.writer(c_)
    w.writerow(APEX)
    for ci in {x[0] for x in rows}:
        w.writerow(["Workgroup_x\\", ci, "10.1.1.9", "Online", "Windows 11  10.0.26100", "x64", "ONGC\\123"])
    st = m.stage("list.csv", c_.getvalue().encode(), "ADMIN1")
    r = m.run([st["stage_id"]], st["mapping"], {"prefix": _prefix(rows[0][0]), "classes": ["DESKTOP"]}, "ADMIN1")
    ci_of = lambda who: {a["ci"] for a in r["assets"] if a["engineer"] == who}      # noqa: E731
    sheet = lambda body: {x[0] for x in load_workbook(io.BytesIO(body))["MERGED LIST"].iter_rows(min_row=2, values_only=True) if x[0]}      # noqa: E731
    fake_user(monkeypatch, "USER", username="ENG1", engineer_key=me)
    with TestClient(app) as c:
        m.publish(r["run_id"], "ADMIN1")
        mine = c.post("/api/match/pack", json={"kind": "consolidated", "format": "xlsx"}, headers=HDR)
        assert mine.status_code == 200 and sheet(mine.content) <= ci_of(me) and sheet(mine.content)
        for kind, fmt in (("consolidated", "csv"), ("consolidated", "pdf"), ("investigation", "pdf"), ("field", "xlsx")):
            assert c.post("/api/match/pack", json={"kind": kind, "format": fmt, "as": other}, headers=HDR).status_code == 200
        assert c.post("/api/match/pack", json={"kind": "field", "format": "pdf"}, headers=HDR).status_code == 400
    fake_user(monkeypatch, "ADMIN")
    with TestClient(app) as c:
        allx = c.post("/api/match/pack", json={"kind": "consolidated", "format": "xlsx"}, headers=HDR)
        assert allx.status_code == 200 and sheet(allx.content) >= (ci_of(me) | ci_of(other)) & {x[0] for x in rows}          # every machine that is in the file, whoever owns it


def test_a_run_is_shared_the_moment_it_exists_and_can_be_shared_again(box, tmp_path, monkeypatch):
    """Engineers must never find 'nothing shared' after an analysis: run() publishes it itself; the share route repairs an older unshared run."""
    monkeypatch.setattr(m, "STAGE", tmp_path)
    ci = box.execute("SELECT ci_no FROM asset WHERE is_current = 1 AND record_level = 'ASSET' AND asset_class = 'DESKTOP' AND ci_no ~ '^[A-Z]{3,4}' LIMIT 1").fetchone()[0]
    st = m.stage("c.csv", _csv([_bf(ci)]), "ADMIN1")
    r = m.run([st["stage_id"]], st["mapping"], {"prefix": _prefix(ci), "classes": ["DESKTOP"]}, "ADMIN1")
    assert any(x["run_id"] == r["run_id"] for x in m.published(50))
    m.publish(r["run_id"], "ADMIN1", on=False)                                            # how an older / unshared run looks
    assert not any(x["run_id"] == r["run_id"] for x in m.published(50))
    fake_user(monkeypatch, "ADMIN")
    with TestClient(app) as c:
        assert c.post("/api/match/share", json={"run_id": r["run_id"]}, headers=HDR).status_code == 200
    assert any(x["run_id"] == r["run_id"] for x in m.published(50))
    fake_user(monkeypatch, "USER", username="ENG1", engineer_key="X")
    with TestClient(app) as c:
        assert c.post("/api/match/share", json={"run_id": r["run_id"]}, headers=HDR).status_code == 403          # an engineer cannot share
