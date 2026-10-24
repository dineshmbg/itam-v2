"""Inventory match: upload the report(s) received from the centre (BigFix, Trend Micro, antivirus, patch, EDR...) and match them against the asset
register by Asset (CI) = computer name. A CI found in the report means the tool is installed; a CI missing from it means not.

The portal does not trust the file name or a typed tool name. It reads the file first: which product it is (from its columns), what each column
means, how old the data is, whether the same fleet is split over several files, and what is wrong with the data itself (duplicates, devices parked in
the wrong group, install failures, agents not communicating, names that do not match any CI). Then it matches, grades each installed machine as healthy
or needing attention, and writes findings that name the machines and say what to do. Every total is reconciled before it is shown.

Nothing in the asset data is changed. Each result is stored as a run so it can be reopened and downloaded again - an Excel action pack for the team, a PDF
for management, and a personal list for each engineer.
"""
import csv
import datetime as dt
import io
import json
import re
import uuid
from collections import Counter
from email.utils import parsedate_to_datetime
from pathlib import Path

from . import db, export, swobs
from .importer import UPLOADS

MAX_BYTES = 40 * 1024 * 1024
MAX_FILES = 6
STAGE = UPLOADS / "inventory-match"
DEFAULT_CLASSES = ["DESKTOP", "WORKSTATION", "SERVER"]      # plain laptops are left out by default; Office Laptops are class DESKTOP, so they stay in
ALL_CLASSES = ["DESKTOP", "LAPTOP", "WORKSTATION", "SERVER", "PRINTER", "SCANNER", "SWITCH", "ROUTER", "UPS", "MEDIA_CONVERTER"]
HIGH, MEDIUM, LOW = "High", "Medium", "Low"
PRIORITY_BY_STATUS = {"IN_USE": HIGH, "NOT_ON_NETWORK": MEDIUM, "NOT_IN_USE": MEDIUM, "STANDBY": MEDIUM}      # anything else (store, surplus, transferred...) is Low

DDL = ["""CREATE TABLE IF NOT EXISTS portal_match_run (
          run_id TEXT PRIMARY KEY, at TIMESTAMPTZ NOT NULL DEFAULT now(), username TEXT NOT NULL, tool TEXT NOT NULL, filename TEXT NOT NULL,
          params JSONB NOT NULL, summary JSONB NOT NULL, result JSONB NOT NULL)""",
       "ALTER TABLE portal_match_run ADD COLUMN IF NOT EXISTS published_at TIMESTAMPTZ", "ALTER TABLE portal_match_run ADD COLUMN IF NOT EXISTS published_by TEXT",
       "CREATE TABLE IF NOT EXISTS portal_match_seen (username TEXT NOT NULL, run_id TEXT NOT NULL, at TIMESTAMPTZ NOT NULL DEFAULT now(), PRIMARY KEY (username, run_id))"]

# Products the portal recognises from the columns of a file (never from its name). A file with at least `min` of the signature columns is that product.
PRODUCTS = {
    "trendmicro": {"label": "Trend Micro Vision One", "min": 4,
                   "sig": ("endpoint name", "protection manager", "agent runtime protection status", "endpoint guid", "agent connection status", "xdr for endpoints (edr)", "endpoint group"),
                   "map": {"name": "Endpoint name", "ip": "IP address", "os": "OS name", "seen": "Last agent status reported", "type": "Object type", "group": "Endpoint group"},
                   "extra": {"conn": "Agent connection status", "action": "Recommended actions", "av": "Anti-malware", "edr": "XDR for Endpoints (EDR)", "agent": "Agent version status",
                             "scanned": "Last scanned", "user": "Last logged on user", "seen2": "Protection module last connected", "osver": "OS version", "runtime": "Agent runtime protection status"}},
    # Apex One console exports (no sensor/full column in them: that comes from the check script). The CSV "Security Agent List" and the old inventory workbook are
    # given the same headings by `_canon` so that both are read as parts of one report.
    "apex": {"label": "Trend Micro Apex One", "min": 4,
             "sig": ("endpoint name", "os name", "ip address", "connection status", "last logged on user", "object type", "os type"),
             "map": {"name": "Endpoint name", "ip": "IP address", "os": "OS name", "type": "Object type"},
             "extra": {"conn": "Connection status", "user": "Last logged on user"}},
    "bigfix": {"label": "BigFix", "min": 3, "sig": ("computer name", "device type", "last report time", "cpu", "os", "ip address"),
               "map": {"name": "Computer Name", "ip": "IP Address", "os": "OS", "seen": "Last Report Time", "type": "Device Type"}, "extra": {}},
}
# generic guesses for a product the portal does not know: header (lower-case, letters and digits only) -> field
GUESS = {
    "name": ("computername", "hostname", "devicename", "machinename", "endpointname", "assetci", "cino", "ci", "assetname", "computer", "nodename", "systemname", "name", "device", "asset"),
    "ip": ("ipaddress", "ip", "ipv4", "ipv4address", "lastipaddress", "lastipused"),
    "os": ("os", "osname", "operatingsystem", "osversion", "platform"),
    "seen": ("lastreporttime", "lastreport", "lastseen", "lastcheckin", "lastcontact", "lastcommunication", "lastupdate", "lastscan", "lastconnected", "lastactivity", "lastagentstatusreported"),
    "type": ("devicetype", "objecttype", "type", "computertype", "category"),
    "group": ("endpointgroup", "group", "domain", "domainhierarchy", "ou", "site"),
}
BASE_FIELDS = ("name", "ip", "os", "seen", "type", "group")
TREND = ("trendmicro", "apex")          # products whose rows say "Trend Micro agent installed"
_APEX_CSV = {"endpoint": "Endpoint name", "ip address": "IP address", "platform": "OS name", "logon user": "Last logged on user", "connection status": "Connection status"}


class MatchError(Exception):
    http_error = True

    def __init__(self, message, status=400):
        super().__init__(message)
        self.status = status


# ---------------------------------------------------------------- reading the uploaded file
def _cell(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return v
    return "" if v is None else str(v).strip()


def _table(rows):
    """Header row = the first of the first 20 rows that has at least 2 filled cells and nearly as many as the widest of them (title rows above are skipped)."""
    rows = [list(r) for r in rows]
    head = 0
    widest = max(sum(1 for c in x if str(c).strip()) for x in rows[:20])
    for i, r in enumerate(rows[:20]):
        filled = sum(1 for c in r if str(c).strip())
        if filled >= 2 and filled >= 0.6 * widest:
            head = i
            break
    headers = [str(c).strip() or f"Column {j + 1}" for j, c in enumerate(rows[head])]
    body = [(r + [""] * len(headers))[:len(headers)] for r in rows[head + 1:] if any(str(c).strip() for c in r)]
    return headers, body


def _canon(headers, rows):
    """The Apex One 'Security Agent List' CSV (Endpoint, Platform, Logon User...) gets the headings of the other Apex export, so both read as one product."""
    low = {h.strip().lower() for h in headers}
    if {"endpoint", "platform", "logon user", "connection status"} <= low and "endpoint name" not in low:
        headers = [_APEX_CSV.get(h.strip().lower(), h) for h in headers]
    return headers, rows


def read_file(filename, data):
    """-> (headers, rows). Excel (.xlsx/.xlsm; the sheet with the most rows) or CSV."""
    low = filename.lower()
    if low.endswith((".xlsx", ".xlsm")) or (low.endswith(".txt") and data[:2] == b"PK"):
        if data[:2] != b"PK":
            raise MatchError("That file is not an Excel workbook (it is damaged or a different format).")
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheets = []
        for ws in wb.worksheets:
            rows = [[_cell(c) for c in r] for r in ws.iter_rows(values_only=True)]
            if len(rows) > 1 and any(any(str(c).strip() for c in r) for r in rows[:20]):
                sheets.append(_table(rows))
        wb.close()
        if not sheets:
            raise MatchError("The workbook has no rows to read.")
        headers, body = max(sheets, key=lambda t: len(t[1]))
        for h2, b2 in sheets:          # sheets with the very same headings (one per site) are more rows of the same report; repeats are found by the duplicate check
            named = lambda hh: [x for x in hh if not x.startswith("Column ")]      # noqa: E731 - blank trailing headings do not make a different layout
            if named(h2) == named(headers) and b2 is not body:
                body = body + [(r + [""] * len(headers))[:len(headers)] for r in b2]
        return _canon(headers, body)
    if low.endswith((".csv", ".txt")):
        text = data.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        csv.field_size_limit(10_000_000)
        rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), dialect)]
        if len(rows) < 2:
            raise MatchError("The file has no rows to read.")
        return _canon(*_table(rows))
    raise MatchError("Upload an Excel workbook (.xlsx) or a CSV file.")


def _norm(h):
    return re.sub(r"[^a-z0-9]", "", h.lower())


def profile_file(headers, rows=None):
    """What is this file? -> {product, label, confidence, why, mapping}. Decided from the column headings, never from the file name."""
    have = {h.strip().lower(): h for h in headers}
    best, score = None, 0
    for key, p in PRODUCTS.items():
        hits = [s for s in p["sig"] if s in have]
        if len(hits) >= p["min"] and len(hits) > score:
            best, score, why = key, len(hits), hits
    if best:
        p = PRODUCTS[best]
        mapping = {f: have[h.lower()] for f, h in p["map"].items() if h.lower() in have}
        return {"product": best, "label": p["label"], "confidence": round(score / len(p["sig"]), 2), "why": f"matches {score} of the {len(p['sig'])} column headings of a {p['label']} export: " + ", ".join(why[:4]), "mapping": mapping}
    norm = {_norm(h): h for h in headers}
    mapping = {}
    for field, names in GUESS.items():
        for n in names:
            if n in norm:
                mapping[field] = norm[n]
                break
    return {"product": "generic", "label": "Endpoint report", "confidence": 0.3 if mapping.get("name") else 0.0,
            "why": "an unfamiliar layout - the columns were guessed from their headings; please check them" if mapping.get("name") else "no computer-name column could be recognised", "mapping": mapping}


def guess_mapping(headers):
    return profile_file(headers)["mapping"]


def _file_time(rows, col, product):
    """Newest check-in time inside a file - the moment the report was really taken."""
    if "seen" not in col:
        return None
    ts = [t for t in (_seen(r[col["seen"]], product) for r in rows[:50000]) if t]
    return max(ts) if ts else None


def stage(filename, data, user):
    name = re.sub(r"[^\w .()-]", "_", Path(filename or "report.xlsx").name)[:120]
    if not data or len(data) > MAX_BYTES:
        raise MatchError(f"The file is empty or larger than {MAX_BYTES // (1024 * 1024)} MB.")
    headers, rows = read_file(name, data)
    prof = profile_file(headers, rows)
    m = prof["mapping"]
    sid = uuid.uuid4().hex[:12]
    STAGE.mkdir(parents=True, exist_ok=True)
    (STAGE / f"{sid}__{name}").write_bytes(data)
    col = {f: headers.index(m[f]) for f in BASE_FIELDS if m.get(f) in headers}
    prefixes, groups = {}, {}
    if "name" in col:
        for r in rows:
            p = re.match(r"[A-Za-z]{3,4}", str(r[col["name"]]).strip())
            if p:
                prefixes[p.group(0).upper()] = prefixes.get(p.group(0).upper(), 0) + 1
    if "group" in col:
        for r in rows:
            g = str(r[col["group"]]).strip()
            groups[g] = groups.get(g, 0) + 1
    seen = _file_time(rows, col, prof["product"])
    return {"stage_id": sid, "filename": name, "rows": len(rows), "headers": headers, "mapping": m, "product": prof["product"], "label": prof["label"], "confidence": prof["confidence"], "why": prof["why"],
            "sample": [[str(c) for c in r][:12] for r in rows[:5]], "prefixes": sorted(prefixes.items(), key=lambda kv: -kv[1])[:12],
            "groups": sorted(groups.items(), key=lambda kv: -kv[1])[:6], "report_time": seen.strftime("%Y-%m-%d %H:%M") if seen else None}


def _staged(stage_id):
    if not re.fullmatch(r"[0-9a-f]{12}", stage_id or ""):
        raise MatchError("Upload the file again.")
    hits = list(STAGE.glob(f"{stage_id}__*")) if STAGE.exists() else []
    if not hits:
        raise MatchError("That upload is no longer available. Upload the file again.")
    return hits[0].name.split("__", 1)[1], hits[0].read_bytes()


# ---------------------------------------------------------------- understanding one row
def key_of(v):
    """A computer name as a match key: trimmed, upper-case, domain suffix dropped (host.domain.local -> HOST)."""
    s = str(v or "").strip().upper().rstrip("$")
    return s.split(".")[0] if re.search(r"[A-Z]", s) else s


def _when(v):
    if isinstance(v, dt.datetime):
        return v.replace(tzinfo=None) if not v.tzinfo else v.astimezone(dt.timezone.utc).replace(tzinfo=None)
    if isinstance(v, dt.date):
        return dt.datetime(v.year, v.month, v.day)
    s = str(v or "").strip()
    if not s:
        return None
    for f in (lambda: parsedate_to_datetime(s), lambda: dt.datetime.fromisoformat(s.replace("Z", "+00:00")), lambda: dt.datetime.strptime(s[:19], "%Y-%m-%d %H:%M:%S"),
              lambda: dt.datetime.strptime(s[:10], "%d-%m-%Y"), lambda: dt.datetime.strptime(s[:10], "%d/%m/%Y")):
        try:
            d = f()
            return d.astimezone(dt.timezone.utc).replace(tzinfo=None) if d.tzinfo else d
        except Exception:      # noqa: BLE001 - try the next format
            continue
    return None


def _seen(v, product):
    """Trend Micro writes 'Last 24 hours (2026-10-06 10:39:40)': the real time is inside the brackets."""
    m = re.search(r"\((\d{4}-\d\d-\d\d[ T]\d\d:\d\d:\d\d)\)", str(v or ""))
    return _when(m.group(1)) if m else _when(v)


def _os_family(s):
    s = str(s or "").upper()
    for pat, label in ((r"WIN(DOWS)?\s*11", "Windows 11"), (r"WIN(DOWS)?\s*10", "Windows 10"), (r"WIN(DOWS)?\s*(8|7|XP|VISTA)", "Windows 8 or older"), (r"2012|2016|2019|2022|2025|SERVER", "Windows Server"),
                       (r"LINUX|UBUNTU|RHEL|CENTOS|DEBIAN", "Linux"), (r"MAC", "macOS")):
        if re.search(pat, s):
            return label
    return "Other / unknown" if s else ""


def _install_reason(text):
    t = text.lower()
    for pat, label in (("missing one or more required files", "installation package incomplete"), ("unable to download", "agent could not be downloaded (internet / firewall)"), ("network error", "network error during installation"),
                       ("temporary issue", "temporary installation error - retry needed"), ("disk", "not enough disk space"), ("access", "access denied during installation")):
        if pat in t:
            return label
    return re.sub(r"\s+", " ", text)[:70]


def _suite(av, edr):
    """Trend Micro Vision One: which part of the product is on the machine. Anti-malware present (even if switched off) = Full suite; no anti-malware but the
    EDR/XDR sensor on = Sensor only. Judged only from the console's own columns, so a file without them gives no answer ("")."""
    a, x = str(av).strip().lower(), str(edr).strip().lower()
    none = ("", "-", "n/a", "na", "none", "unknown")
    has_av = a not in none and not a.startswith(("not ", "unavailable", "unsupported"))
    has_edr = x not in none and not x.startswith(("disabled", "not ", "unavailable", "unsupported"))
    return "Full suite" if has_av else "Sensor only" if has_edr else "Unclear"


def _entry(r, col, product):
    """One report row -> a clean entry with its own findings: `fails` stop the agent doing its job, `advice` are things to improve."""
    g = lambda f: str(r[col[f]]).strip() if f in col and not isinstance(r[col[f]], (dt.datetime, dt.date)) else ("" if f not in col else r[col[f]])     # noqa: E731
    name = str(r[col["name"]]).strip()
    seen = _seen(r[col["seen"]], product) if "seen" in col else None
    if product == "trendmicro":
        s2 = _when(g("seen2")) if "seen2" in col else None
        seen = max([t for t in (seen, s2) if t], default=None)
    e = {"name": name, "key": key_of(name), "ip": g("ip"), "os": g("os") if product != "trendmicro" else (g("os") + " " + str(g("osver")).replace("10.0 (Build ", "build ").rstrip(")")).strip(), "type": g("type"),
         "group": g("group"), "seen": seen, "user": g("user") if "user" in col else "", "fails": [], "advice": []}
    if product == "trendmicro":
        if str(g("action")):
            e["fails"].append("Agent installation failed: " + _install_reason(str(g("action"))))
        if str(g("conn")) and str(g("conn")).lower() != "communicating":
            e["fails"].append("Agent not communicating")
        av = str(g("av")).lower()
        if "av" in col or "edr" in col:
            e["suite"] = _suite(g("av"), g("edr"))
        if av.startswith("disabled"):
            e["fails"].append("Anti-malware disabled")
        if "pattern outdated" in av:
            e["advice"].append("Anti-malware pattern outdated")
        if str(g("edr")).lower() == "disabled":
            e["advice"].append("EDR (XDR) not enabled")
        ag = str(g("agent")).lower()
        if ag == "olderversion":
            e["advice"].append("Agent version out of date")
        elif ag == "osversionnotsupported":
            e["advice"].append("Windows build not supported by the agent")
        if str(g("runtime")).lower().startswith("requires attention"):
            e["advice"].append("Console flags 'requires attention'")
        e["scanned"] = _when(g("scanned")) if "scanned" in col else None
    elif product == "apex":
        e["conn"] = str(g("conn"))
        if e["conn"].lower() == "offline":
            e["advice"].append("Offline in the Apex console when the report was taken")
    return e


def _generic(label):
    """Group per-machine wording ('Not reported for 118 days') under one heading for counting."""
    return re.sub(r"Not reported for \d+ days", "Agent has not reported recently", re.sub(r"Last scan \d+ days ago", "Not scanned for over 30 days", label))


def _lev(a, b, cap=2):
    """Edit distance, giving up (returns cap + 1) once it must exceed `cap`."""
    if abs(len(a) - len(b)) > cap:
        return cap + 1
    prev = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        cur = [i]
        for j, cb in enumerate(b, 1):
            cur.append(min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + (ca != cb)))
        if min(cur) > cap:
            return cap + 1
        prev = cur
    return prev[-1]


CONFUSABLE = ({"O", "0"}, {"I", "1"}, {"L", "1"}, {"S", "5"}, {"B", "8"})


def _typo(a, b):
    """Is `b` (a report name the register does not know) so probably `a` mistyped that it is worth a person confirming? Only typing slips count:
    O/0-type look-alikes, or one character missing / added. A different digit on its own is not enough - DT140 and DT147 are two real machines."""
    if a[:6] != b[:6] or a == b:
        return None
    if len(a) == len(b):
        diff = [(x, y) for x, y in zip(a, b) if x != y]
        return "name differs only by look-alike characters (O/0, I/1)" if diff and all({x, y} in CONFUSABLE for x, y in diff) else None
    if abs(len(a) - len(b)) == 1 and _lev(a, b, 1) == 1:
        return "name has one character missing or extra"
    return None


def _assets(prefixes, classes):
    like = " OR ".join(["ci_no ILIKE %s"] * len(prefixes))
    return db.query(f"""SELECT ci_no, asset_class, asset_type, make, model, serial_no, hostname, ip_address, asset_status, user_name, location_code, engineer_name
                        FROM asset WHERE is_current = 1 AND record_level = 'ASSET' AND asset_class = ANY(%s) AND ({like}) ORDER BY ci_no""", [classes, *[p + "%" for p in prefixes]])


def _pct(a, b):
    return round(100 * a / b, 1) if b else 0.0


def _ex(items, n=4):
    items = list(items)
    return ", ".join(str(x) for x in items[:n]) + (f" and {len(items) - n} more" if len(items) > n else "")


# ---------------------------------------------------------------- the analysis
def analyse(files, mapping, params, today=None):
    """files: [{name, headers, rows, product}]. All files are read as parts of ONE report (the centre often splits a fleet by group or site)."""
    today = today or dt.datetime.now(dt.timezone.utc).replace(tzinfo=None)
    products = {f["product"] for f in files}
    if len(products) > 1:
        raise MatchError("These files look like different products (" + ", ".join(sorted(PRODUCTS.get(p, {"label": "unknown layout"})["label"] for p in products)) + "). Analyse each product separately.")
    product = products.pop()
    plabel = PRODUCTS[product]["label"] if product in PRODUCTS else "Endpoint report"
    tool = (params.get("tool") or "").strip()[:40] or plabel
    prefixes = [p.strip().upper() for p in re.split(r"[,\s]+", params.get("prefix") or "ANK") if p.strip()] or ["ANK"]
    classes = [c for c in (params.get("classes") or DEFAULT_CLASSES) if c in ALL_CLASSES] or DEFAULT_CLASSES
    active_days, stale_days = int(params.get("active_days") or 7), int(params.get("stale_days") or 30)
    target = float(params.get("target") or 95)
    if not mapping.get("name"):
        raise MatchError("Choose the column that holds the computer name.")
    checks = []          # data checks on the file itself: what was found and how it was handled

    # ---- read every file into clean entries
    raw, blank, per_file, extras_missing, bad_dates, name_fixed = [], 0, [], [], 0, 0
    for f in files:
        h = f["headers"]
        have = {x.strip().lower(): i for i, x in enumerate(h)}
        col = {fld: h.index(mapping[fld]) for fld in BASE_FIELDS if mapping.get(fld) in h}
        if "name" not in col:
            raise MatchError(f"The computer-name column '{mapping['name']}' is not in {f['name']}.")
        for fld, hd in PRODUCTS.get(product, {}).get("extra", {}).items():
            if hd.lower() in have:
                col[fld] = have[hd.lower()]
        if "seen" not in col:
            extras_missing.append(f"{f['name']}: no check-in date column")
        n0 = len(raw)
        for r in f["rows"]:
            nm = str(r[col["name"]]).strip()
            if not nm:
                blank += 1
                continue
            e = _entry(r, col, product)
            e["file"] = f["name"]
            if "seen" in col and r[col["seen"]] and not e["seen"]:
                bad_dates += 1
            if nm != nm.upper() or "." in nm and re.search(r"[A-Za-z]", nm):
                name_fixed += 1
            raw.append(e)
        per_file.append({"name": f["name"], "rows": len(f["rows"]), "used": len(raw) - n0, "product": product, "report_time": _file_time(f["rows"], col, product)})
    if not raw:
        raise MatchError("The file has no usable rows.")

    # ---- the report: one entry per computer (the most recent check-in wins); duplicates and group placement kept for the checks
    by_key, dup_detail, all_by_ip, dup_rows = {}, {}, {}, 0
    for e in raw:
        old = by_key.get(e["key"])
        if old:
            dup_rows += 1
            d = dup_detail.setdefault(e["key"], {"name": e["name"], "rows": [old]})
            d["rows"].append(e)
            if (e["seen"] or dt.datetime.min) < (old["seen"] or dt.datetime.min):
                for f in ("ip", "os", "type", "user", "conn"):          # the older record still knows things the kept one does not
                    if not old.get(f) and e.get(f):
                        old[f] = e[f]
                continue
            for f in ("ip", "os", "type", "user", "conn"):
                if not e.get(f) and old.get(f):
                    e[f] = old[f]
        by_key[e["key"]] = e
    files_of = {}
    for e in raw:
        files_of.setdefault(e["key"], set()).add(e["file"])
    for e in by_key.values():
        e["files"] = ", ".join(sorted(files_of[e["key"]]))
    for e in by_key.values():
        if e["ip"]:
            all_by_ip.setdefault(e["ip"], []).append(e)
    seen_all = [e["seen"] for e in by_key.values() if e["seen"]]
    as_of = max(seen_all) if seen_all else today
    report_age = (today - as_of).days if seen_all else None
    in_scope = {k: e for k, e in by_key.items() if any(k.startswith(p) for p in prefixes)}
    if not in_scope:
        raise MatchError(f"No computer in the report starts with {', '.join(prefixes)}. Check the site prefix, or whether this is the right file.")

    # ---- the register
    assets = _assets(prefixes, classes)
    ci_keys = {key_of(a["ci_no"]): a for a in assets}
    other = {key_of(r["ci_no"]): r for r in db.query("SELECT ci_no, asset_class, asset_type, asset_status FROM asset WHERE is_current = 1 AND (" + " OR ".join(["ci_no ILIKE %s"] * len(prefixes)) + ")",
                                                      [p + "%" for p in prefixes])}
    unmatched = {k: e for k, e in in_scope.items() if k not in ci_keys and k not in other}          # names the register does not know at all

    # ---- pair the not-found machines with report names that probably are them (same IP, or a name that differs by a character or two)
    pair = {}
    taken = set()
    cand = []
    for a in assets:
        k = key_of(a["ci_no"])
        if k in in_scope:
            continue
        ip = (a.get("ip_address") or "").strip()
        for e in all_by_ip.get(ip, []) if ip else []:
            if e["key"] in unmatched and len(all_by_ip[ip]) == 1:
                cand.append((0, k, e["key"], "same IP address"))
        for uk in unmatched:
            why = _typo(k, uk)
            if why:
                cand.append((1, k, uk, why))
    for d, k, uk, why in sorted(cand):
        if k not in pair and uk not in taken:
            pair[k] = (uk, why)
            taken.add(uk)

    out = []
    for a in assets:
        k = key_of(a["ci_no"])
        e = in_scope.get(k)
        row = {"ci": a["ci_no"], "class": a["asset_class"], "type": a["asset_type"], "make": a["make"] or "", "model": a["model"] or "", "serial": a["serial_no"] or "", "status": a["asset_status"] or "",
               "user": a["user_name"] or "", "location": a["location_code"] or "", "engineer": a["engineer_name"] or "UNASSIGNED", "installed": e is not None, "health": "Not installed", "issues": "", "advice": "", "suite": "", "suite_pc": "",
               "r_name": "", "r_files": "", "r_conn": "", "r_ip": "", "r_os": "", "r_os_family": "", "r_group": "", "r_user": "", "r_seen": None, "r_age": None, "reporting": "", "priority": "", "possible": "", "possible_by": "", "action": ""}
        if e:
            age = (as_of - e["seen"]).days if e["seen"] else None
            fails = list(e["fails"])
            row.update(r_name=e["name"], r_files=e["files"], r_conn=e.get("conn", ""), r_ip=e["ip"], r_os=e["os"], r_os_family=_os_family(e["os"]), r_group=e["group"], r_user=e["user"], r_seen=e["seen"].strftime("%Y-%m-%d %H:%M") if e["seen"] else None, r_age=age,
                       reporting="" if age is None else "Active" if age <= active_days else "Stale" if age <= stale_days else "Dormant")
            if row["reporting"] in ("Stale", "Dormant") and not any("not communicating" in x for x in fails):
                fails.append(f"Not reported for {age} days")
            elif row["reporting"] in ("Stale", "Dormant"):
                fails = [f"Not reported for {age} days" if x == "Agent not communicating" else x for x in fails]
            advice = list(e["advice"])
            if product == "trendmicro":
                sc = e.get("scanned")
                if not sc:
                    advice.append("Never scanned")
                elif (as_of - sc).days > 30:
                    advice.append(f"Last scan {(as_of - sc).days} days ago")
            row["suite"] = e.get("suite", "")
            row["health"] = "Needs attention" if fails else "Healthy"
            row["issues"], row["advice"] = "; ".join(fails), "; ".join(advice)
            if fails:
                row["action"] = ("Repair or reinstall the agent: " if any("installation failed" in x or "communicating" in x or "Not reported" in x for x in fails) else "Fix: ") + "; ".join(fails) + "."
                if row["reporting"] == "Dormant":
                    row["action"] += " If the machine is retired, update the register."
        else:
            row["priority"] = PRIORITY_BY_STATUS.get(row["status"], LOW)
            if k in pair:
                uk, why = pair[k]
                hit = unmatched[uk]
                row.update(possible=hit["name"], possible_by=why)
                row["action"] = f"Probably this machine, listed as {hit['name']} ({why}) - confirm, then correct the CI in the register or rename the computer."
            elif row["priority"] == HIGH:
                row["action"] = f"Install the {tool} agent (or restore its connection) - the machine is in use."
            elif row["priority"] == MEDIUM:
                row["action"] = f"Install {tool} when the machine is back on the network / in use."
            else:
                row["action"] = f"Not urgent while {row['status'].replace('_', ' ').lower() or 'not in service'}; install {tool} before it is issued."
        out.append(row)

    # ---- report names the register does not know
    unreg = []
    pair_rev = {v[0]: (k, v[1]) for k, v in pair.items()}
    for k, e in sorted(in_scope.items()):
        if k in ci_keys:
            continue
        o = other.get(k)
        age = (as_of - e["seen"]).days if e["seen"] else None
        if o:
            kind, why, act = "outside", f"Registered, but as {o['asset_class'].title()} / {o['asset_type'].title()} (outside the classes checked)", "No action unless the asset class is wrong in the register."
        elif k in pair_rev:
            ci, how = pair_rev[k]
            kind, why, act = "mismatch", f"Probably register CI {ci} ({how})", f"Confirm it is {ci}; then rename the computer or correct the CI."
        else:
            kind, why, act = "unregistered", "Not in the asset register", "Find the device and add it to the register, or confirm it is not ours / retired."
        unreg.append({"name": e["name"], "ip": e["ip"], "os": e["os"], "type": e["type"], "group": e["group"], "seen": e["seen"].strftime("%Y-%m-%d %H:%M") if e["seen"] else None, "age": age, "why": why, "action": act, "kind": kind, "files": e["files"]})

    # ---- group placement (Trend Micro 'Endpoint group', etc.): is the fleet where it should be?
    group_issues, foreign = [], []
    groups = Counter(e["group"] for e in in_scope.values() if e["group"])
    home = groups.most_common(1)[0][0] if groups else ""
    if home:
        group_issues = [e for e in in_scope.values() if e["group"] and e["group"] != home]
        foreign = [e for e in by_key.values() if e["key"] not in in_scope and e["group"].strip().upper() in prefixes]
    dup_detail = {k: v for k, v in dup_detail.items() if k in in_scope}          # a repeat of another site's computer is not this site's problem
    dups = [{"name": d["name"], "rows": len(d["rows"]), "detail": "; ".join(f"{(r['ip'] or 'no IP')}, last {r['seen']:%d %b}" if r["seen"] else (r["ip"] or "no IP") for r in d["rows"])} for d in sorted(dup_detail.values(), key=lambda d: d["name"])]
    # a machine that has several records but only one recent one is a re-registered machine; several recent ones is a real conflict
    conflicts = [d["name"] for d in dup_detail.values() if sum(1 for r in d["rows"] if r["seen"] and (as_of - r["seen"]).days <= active_days) > 1]

    # ---- the numbers
    total, inst = len(out), sum(1 for r in out if r["installed"])
    miss = [r for r in out if not r["installed"]]
    healthy = sum(1 for r in out if r["health"] == "Healthy")
    attn = sum(1 for r in out if r["health"] == "Needs attention")
    summary = {
        "tool": tool, "product": product, "product_label": plabel, "prefixes": prefixes, "classes": classes, "target": target, "active_days": active_days, "stale_days": stale_days,
        "report_rows": sum(f["rows"] for f in per_file), "report_unique": len(by_key), "report_in_scope": len(in_scope), "assets": total, "installed": inst, "missing": len(miss), "coverage": _pct(inst, total),
        "healthy": healthy, "attention": attn, "healthy_pct": _pct(healthy, total), "installed_healthy_pct": _pct(healthy, inst),
        "missing_high": sum(1 for r in miss if r["priority"] == HIGH), "missing_medium": sum(1 for r in miss if r["priority"] == MEDIUM), "missing_low": sum(1 for r in miss if r["priority"] == LOW),
        "active": sum(1 for r in out if r["reporting"] == "Active"), "stale": sum(1 for r in out if r["reporting"] == "Stale"), "dormant": sum(1 for r in out if r["reporting"] == "Dormant"),
        "unknown_age": sum(1 for r in out if r["installed"] and not r["reporting"]), "name_mismatch": sum(1 for r in miss if r["possible"]),
        "unregistered": sum(1 for u in unreg if u["kind"] == "unregistered"), "outside_scope": sum(1 for u in unreg if u["kind"] == "outside"), "duplicate_names": len(dups), "duplicate_conflicts": len(conflicts),
        "has_seen": bool(seen_all), "has_os": any(e["os"] for e in by_key.values()), "as_of": today.strftime("%Y-%m-%d"), "report_as_of": as_of.strftime("%Y-%m-%d %H:%M"), "report_age_days": report_age,
        "files": [{**f, "report_time": f["report_time"].strftime("%Y-%m-%d %H:%M") if f["report_time"] else None} for f in per_file], "home_group": home,
        "group_issues": len(group_issues), "foreign": len(foreign),
    }
    summary["in_use"] = sum(1 for r in out if r["status"] == "IN_USE")
    summary["in_use_installed"] = sum(1 for r in out if r["status"] == "IN_USE" and r["installed"])
    summary["in_use_healthy"] = sum(1 for r in out if r["status"] == "IN_USE" and r["health"] == "Healthy")
    summary["in_use_coverage"] = _pct(summary["in_use_installed"], summary["in_use"])
    summary["rag"] = "green" if summary["coverage"] >= target else "amber" if summary["coverage"] >= target - 15 else "red"
    ic, ac = Counter(), Counter()
    for r in out:
        for x in filter(None, r["issues"].split("; ")):
            ic[_generic(x)] += 1
        for x in filter(None, r["advice"].split("; ")):
            ac[_generic(x)] += 1
    summary["issue_counts"] = [{"label": k, "n": n} for k, n in ic.most_common()]
    summary["advice_counts"] = [{"label": k, "n": n, "systemic": inst > 0 and n / inst >= 0.6} for k, n in ac.most_common()]

    def by(field, label_fn=lambda v: v or "(none)"):
        g = {}
        for r in out:
            lab = label_fn(r[field])
            x = g.setdefault(lab, {"label": lab, "total": 0, "installed": 0, "missing": 0, "high": 0, "dormant": 0, "attention": 0, "healthy": 0})
            x["total"] += 1
            x["installed"] += r["installed"]
            x["missing"] += not r["installed"]
            x["high"] += (not r["installed"]) and r["priority"] == HIGH
            x["dormant"] += r["reporting"] == "Dormant"
            x["attention"] += r["health"] == "Needs attention"
            x["healthy"] += r["health"] == "Healthy"
        for x in g.values():
            x["coverage"] = _pct(x["installed"], x["total"])
            x["healthy_pct"] = _pct(x["healthy"], x["total"])
        return sorted(g.values(), key=lambda x: (-x["missing"], x["label"]))

    summary["by_class"], summary["by_status"] = by("class"), by("status", lambda v: (v or "(none)").replace("_", " "))
    summary["by_location"], summary["by_engineer"] = by("location"), by("engineer")
    mm = Counter(((r["make"] + " " + r["model"]).strip() or "(model not recorded)") for r in miss)
    summary["gap_models"] = [{"label": k, "n": n} for k, n in mm.most_common(10)]
    fam = Counter(r["r_os_family"] for r in out if r["installed"] and r["r_os_family"])
    summary["os_mix"] = [{"label": k, "n": n} for k, n in fam.most_common()]
    summary["register_quality"] = {"no_engineer": sum(1 for r in out if r["engineer"] == "UNASSIGNED"), "no_user": sum(1 for r in out if r["status"] == "IN_USE" and not r["user"]),
                                   "no_location": sum(1 for r in out if not r["location"])}

    # ---- checks on the file itself, and how each problem was handled
    def chk(title, status, detail, handled):
        checks.append({"check": title, "status": status, "detail": detail, "handled": handled})
    chk("What the file is", "ok" if product != "generic" else "warn", f"{plabel} ({'recognised from its column headings' if product != 'generic' else 'layout not recognised - columns guessed'}). Tool name used in this report: {tool}.",
        "The file name was ignored; the product was identified from the columns.")
    if len(files) > 1:
        chk("Several files, one report", "ok", f"{len(files)} files were combined: " + "; ".join(f"{p['name']} ({p['used']:,} row{'s' if p['used'] != 1 else ''})" for p in per_file) + ".", "Read as parts of one fleet; a computer in more than one part is counted once.")
    chk("Report date", "warn" if report_age is not None and report_age > 3 else "ok", f"Newest check-in in the report: {as_of:%d %b %Y %H:%M}" + (f" ({report_age} days before this analysis)." if report_age else "."),
        "Days silent are counted from the report's own date, not from today." if not report_age else "The report is older than a few days; machines may have changed since. Ask the centre for a fresh export before acting.")
    if blank:
        chk("Rows without a computer name", "warn", f"{blank:,} rows.", "Ignored.")
    if bad_dates:
        chk("Check-in times that could not be read", "warn", f"{bad_dates:,} rows.", "Treated as 'no check-in date'.")
    if name_fixed:
        chk("Names in lower case or with a domain suffix", "info", f"{name_fixed:,} names.", "Normalised before matching (case ignored, domain suffix dropped).")
    if dups:
        chk("Computers listed more than once", "warn", f"{len(dups):,} of this site's names ({sum(d['rows'] - 1 for d in dups):,} extra rows): " + _ex([d["name"] for d in dups]) + ". Usually a re-installed or re-registered machine.", "The most recent record is used. " + (f"{len(conflicts)} names have more than one recent record - ask the centre to remove the stale ones: " + _ex(conflicts) if conflicts else ""))
    if home and group_issues:
        chk(f"Machines outside the usual group '{home}'", "warn", f"{len(group_issues):,} of {len(in_scope):,} {'/'.join(prefixes)} computers sit in other groups: " + ", ".join(f"{g} ({n})" for g, n in Counter(e['group'] for e in group_issues).most_common(3)) + ".",
            "Counted as installed, but they are outside the site's policies. Ask the centre to move them to '" + home + "'.")
    if foreign:
        chk("Other sites' devices in this site's group", "warn", f"{len(foreign):,} devices in group '{home}' have names that are not {'/'.join(prefixes)}: " + _ex([e["name"] for e in foreign]) + ".", "Not counted. Ask the centre whether they belong to this site.")
    ours = {".".join(e["ip"].split(".")[:3]) for e in in_scope.values() if e["ip"].count(".") == 3}          # the network ranges this site's computers use
    dflt = [e["name"] for e in by_key.values() if re.match(r"^(DESKTOP|LAPTOP|WIN|PC)-[A-Z0-9]{5,}$", e["key"]) and e["ip"].count(".") == 3 and ".".join(e["ip"].split(".")[:3]) in ours]
    if dflt:
        chk("Computers with a default Windows name on this site's network", "warn", f"{len(dflt):,}: " + _ex(dflt) + ". They sit in this site's IP ranges but cannot be matched to any CI.",
            "Not counted. Identify each one and rename it to its CI so it can be matched.")
    ipdup = {ip: es for ip, es in all_by_ip.items() if len(es) > 1 and all(e["key"] in in_scope for e in es)}
    if ipdup:
        chk("The same IP address on different computers", "info", f"{len(ipdup):,} address{'es' if len(ipdup) != 1 else ''}, e.g. " + _ex(f"{ip} ({', '.join(e['name'] for e in es[:2])})" for ip, es in ipdup.items()) + ".", "Not used for matching when an address is shared.")
    if extras_missing:
        chk("Columns missing", "warn", "; ".join(extras_missing), "Ages and health could not be worked out for those files.")
    chk("Register-side data", "warn" if summary["register_quality"]["no_engineer"] else "ok", f"{summary['register_quality']['no_engineer']:,} checked machines have no engineer; {summary['register_quality']['no_user']:,} deployed machines have no user recorded.",
        "Machines with no engineer cannot be e-mailed to anyone - assign an engineer in the register." if summary["register_quality"]["no_engineer"] else "")

    # ---- every total must add up before anyone sees it
    recon = [
        ("Rows read from the file(s)", summary["report_rows"], "all parts together"),
        ("  less rows without a name", -blank, ""),
        ("  less repeat rows of the same computer", -dup_rows, "latest record kept"),
        ("= Different computers in the report", len(by_key), ""),
        (f"  of which named {'/'.join(prefixes)}...", len(in_scope), "this site"),
        ("  of which other sites", len(by_key) - len(in_scope), "not counted"),
        ("Register machines checked", total, ", ".join(c.title() for c in classes)),
        ("  found in the report (installed)", inst, ""),
        ("  not found (not installed)", len(miss), ""),
        ("Installed machines: healthy", healthy, ""),
        ("Installed machines: need attention", attn, ""),
    ]
    ok = [summary["report_rows"] - blank - dup_rows == len(by_key), inst + len(miss) == total, healthy + attn == inst, sum(x["total"] for x in summary["by_class"]) == total, sum(x["installed"] for x in summary["by_engineer"]) == inst, len(in_scope) == inst + len(unreg)]
    nx = swobs.cross_check(out, product)
    if product in TREND:
        swobs.suite_seen(out)
        for r in out:
            if not r["installed"] and r["suite_pc"]:
                r["action"] = f"The PC itself reports {r['suite_pc']} but it is not in the console report: do not reinstall - check its connection to the Apex server and its console group."
        sc = Counter(r["suite"] or r["suite_pc"] or "Not determined" for r in out if r["installed"])
        summary["suite_counts"] = {k: sc.get(k, 0) for k in ("Full suite", "Sensor only", "Unclear", "Not determined")}
        summary["suite_disagree"] = sum(1 for r in out if r["suite"] and r["suite_pc"] and r["suite"] != r["suite_pc"])
    if nx:
        summary["xcheck"] = nx
    summary["reconciled"] = all(ok)
    summary["recon"] = [{"label": a, "n": b, "note": c} for a, b, c in recon]
    summary["checks"] = checks
    summary["findings"] = findings(summary)
    return summary, {"assets": out, "unregistered": unreg, "duplicates": dups, "group_issues": [{"name": e["name"], "ip": e["ip"], "group": e["group"], "seen": e["seen"].strftime("%Y-%m-%d %H:%M") if e["seen"] else None} for e in group_issues],
                     "foreign": [{"name": e["name"], "ip": e["ip"], "group": e["group"]} for e in foreign]}


def findings(s):
    """Plain-language findings, most important first, each naming the problem and what to do about it."""
    f, t = [], s["tool"]
    sev = {"green": "Good", "amber": "Needs attention", "red": "Critical"}[s["rag"]]
    f.append({"severity": sev, "finding": f"{t}: {s['installed']:,} of {s['assets']:,} checked machines appear in the centre's report ({s['coverage']}%); the target is {s['target']:g}%. {s['missing']:,} are missing.",
              "action": "Close the gap below, starting with machines in use." if s["missing"] else "Keep the register and the report in step."})
    if s["installed"] and s["attention"]:
        st = "Critical" if s["installed_healthy_pct"] < 50 else "Needs attention"
        top = "; ".join(f"{x['label']} ({x['n']:,})" for x in s["issue_counts"][:3])
        f.append({"severity": st, "finding": f"Being in the report does not mean protected: only {s['healthy']:,} of the {s['installed']:,} machines that appear ({s['installed_healthy_pct']}%) are working properly. Main problems: {top}.",
                  "action": "Treat 'installed' machines with a problem as unprotected. Engineers repair or reinstall the agent (see 'Needs attention')."})
    fail_inst = next((x for x in s["issue_counts"] if x["label"].startswith("Agent installation failed")), None)
    if fail_inst:
        n = sum(x["n"] for x in s["issue_counts"] if x["label"].startswith("Agent installation failed"))
        f.append({"severity": "Needs attention", "finding": f"{n:,} machines are in the console but the agent installation failed ({'; '.join(x['label'].split(': ')[1] for x in s['issue_counts'] if x['label'].startswith('Agent installation failed'))}).",
                  "action": "Ask the centre to re-push a complete installation package; check firewall / internet access for the downloader on those machines."})
    if s["in_use"]:
        f.append({"severity": "Critical" if s["in_use_coverage"] < s["target"] - 15 else "Needs attention" if s["in_use_coverage"] < s["target"] else "Good",
                  "finding": f"Of the {s['in_use']:,} machines deployed to users, {s['in_use_installed']:,} appear in the report ({s['in_use_coverage']}%) and {s['in_use_healthy']:,} are healthy. {s['missing_high']:,} deployed machines have no agent at all.",
                  "action": "Engineers install or repair the agent on these first - they are the real exposure."})
    worst = [x for x in s["by_class"] if x["missing"]][:2]
    if worst and s["missing"]:
        f.append({"severity": "Needs attention", "finding": "Largest gaps by class: " + "; ".join(f"{x['label'].title()} {x['missing']:,} of {x['total']:,} missing ({x['coverage']}% covered)" for x in worst) + ".",
                  "action": "Check whether this class is outside the centre's deployment scope or a roll-out has not reached it."})
    eng = [x for x in s["by_engineer"] if x["missing"] + x["attention"]][:3]
    if eng and len(s["by_engineer"]) > 1:
        f.append({"severity": "Needs attention", "finding": "Most machines to fix sit with: " + "; ".join(f"{x['label'].title()} ({x['missing'] + x['attention']:,})" for x in eng) + ".",
                  "action": "Share each engineer's list (e-mail or the Excel 'By engineer' sheet) and agree a completion date."})
    sc = s.get("suite_counts")
    if sc and sc["Sensor only"]:
        f.append({"severity": "Needs attention", "finding": f"{sc['Sensor only']:,} installed machines have the Trend Micro sensor only (no anti-malware); {sc['Full suite']:,} have the full suite.",
                  "action": "Sensor-only machines are detected but not protected against malware: deploy the full Apex One agent on them (see the field report for the list)."})
    if sc and sc.get("Not determined"):
        f.append({"severity": "Information", "finding": f"{sc['Not determined']:,} machines are in the console but nothing says whether they have the sensor only or the full suite.",
                  "action": "Run ankcheck.bat on those PCs (Inventory match > Software data); the next consolidated report will then show Sensor only or Full suite for them."})
    if s["name_mismatch"]:
        f.append({"severity": "Needs attention", "finding": f"{s['name_mismatch']:,} machines counted as 'not installed' are probably in the report under a slightly different name (same IP, or a name differing by a character or two).",
                  "action": "Confirm each pair, then correct the CI or rename the computer. This raises coverage without installing anything."})
    if s["group_issues"]:
        f.append({"severity": "Needs attention", "finding": f"{s['group_issues']:,} of this site's computers are in the console under another group than '{s['home_group']}' (see the data checks), so they are outside the site's policies.",
                  "action": "Ask the centre to move them into the correct group."})
    if s["foreign"]:
        f.append({"severity": "Information", "finding": f"{s['foreign']:,} devices in this site's console group have names from other sites (for example servers) and were not counted.", "action": "Ask the centre to confirm they belong here."})
    for a in s["advice_counts"][:3]:
        if a["systemic"]:
            f.append({"severity": "Needs attention", "finding": f"{a['n']:,} of {s['installed']:,} installed machines show '{a['label']}'. Because it affects most of the fleet, it is a central/configuration problem rather than a problem on each machine.",
                      "action": "Raise it once with the centre (update source, relay or policy) instead of fixing machines one by one."})
    if s["has_seen"] and s["dormant"]:
        f.append({"severity": "Needs attention", "finding": f"{s['dormant']:,} installed machines have not reported for over {s['stale_days']} days ({s['stale']:,} more for over {s['active_days']} days).",
                  "action": "Locate them; repair the agent, or confirm they are retired and update the register."})
    if s["unregistered"]:
        f.append({"severity": "Needs attention", "finding": f"{s['unregistered']:,} devices report to {t} with names that are not in the asset register.",
                  "action": "Identify each one and add it to the register (or confirm it is not ours). Unregistered devices are outside asset control."})
    if s["has_os"]:
        w10 = next((x["n"] for x in s["os_mix"] if x["label"] == "Windows 10"), 0)
        old = next((x["n"] for x in s["os_mix"] if x["label"] == "Windows 8 or older"), 0)
        if w10 or old:
            f.append({"severity": "Needs attention", "finding": f"{w10:,} reporting machines run Windows 10" + (f" and {old:,} run Windows 8 or older" if old else "") + ".",
                      "action": "Plan upgrades to a supported Windows release; see the OS column in the full list."})
    if s["duplicate_names"]:
        f.append({"severity": "Information", "finding": f"{s['duplicate_names']:,} computer names appear more than once in the report" + (f"; {s['duplicate_conflicts']:,} have more than one recent record" if s["duplicate_conflicts"] else " (re-installed or re-registered machines)") + ".",
                  "action": "Ask the centre to remove the stale duplicate records; the latest check-in was used here."})
    if s["missing_low"]:
        f.append({"severity": "Information", "finding": f"{s['missing_low']:,} of the missing machines are in store, surplus or transferred, so no agent is expected yet.", "action": "No urgent action; install before they are issued to a user."})
    if s["register_quality"]["no_engineer"]:
        f.append({"severity": "Information", "finding": f"{s['register_quality']['no_engineer']:,} checked machines have no engineer in the register, so nobody is responsible for fixing them.", "action": "Assign an engineer to each in the register."})
    if s.get("report_age_days") is not None and s["report_age_days"] > 3:
        f.insert(0, {"severity": "Needs attention", "finding": f"The report is {s['report_age_days']} days old (newest check-in {s['report_as_of']}).", "action": "Get a fresh export from the centre before acting on this."})
    return f


# ---------------------------------------------------------------- run, store, reopen
def _observe(ids, user):
    """The same report also becomes per-machine observations (shown on each asset's Software tab). A file that is not a BigFix/Trend export, or any
    failure here, never stops the match itself."""
    tot = {"recorded": 0, "unchanged": 0, "matched": 0}
    for sid in ids:
        try:
            fname, data = _staged(sid)
            r = swobs.ingest_report(fname, data, user)
        except Exception:      # noqa: BLE001 - observations are a by-product
            continue
        for k in tot:
            tot[k] += r[k]
    return tot


def run(stage_ids, mapping, params, user):
    ids = [stage_ids] if isinstance(stage_ids, str) else list(stage_ids or [])
    if not ids or len(ids) > MAX_FILES:
        raise MatchError(f"Upload between 1 and {MAX_FILES} files.")
    files = []
    for sid in ids:
        fname, data = _staged(sid)
        headers, rows = read_file(fname, data)
        files.append({"name": fname, "headers": headers, "rows": rows, "product": profile_file(headers)["product"]})
    summary, result = analyse(files, mapping or {}, params or {})
    summary["observations"] = _observe(ids, user)
    label = files[0]["name"] if len(files) == 1 else f"{files[0]['name']} + {len(files) - 1} more"
    run_id = uuid.uuid4().hex[:12]
    with db.write() as con:
        con.execute("""INSERT INTO portal_match_run (run_id, username, tool, filename, params, summary, result, published_at, published_by) VALUES (%s,%s,%s,%s,%s::jsonb,%s::jsonb,%s::jsonb, now(), %s)""",
                    (run_id, user, summary["tool"], label, json.dumps(params or {}), json.dumps(summary, default=str), json.dumps(result, default=str), user))      # shared with engineers from the first moment: nothing later can leave it unshared
    return {"run_id": run_id, "filename": label, "summary": summary, **result}


def get(run_id):
    r = db.one("SELECT * FROM portal_match_run WHERE run_id = %s", [run_id])
    if not r:
        raise MatchError("That analysis no longer exists.", 404)
    return {"run_id": r["run_id"], "at": r["at"], "username": r["username"], "filename": r["filename"], "summary": r["summary"], **r["result"]}


def history(limit=15):
    return db.query("SELECT run_id, at, username, tool, filename, published_at, summary->>'coverage' AS coverage, summary->>'assets' AS assets, summary->>'installed' AS installed FROM portal_match_run ORDER BY at DESC LIMIT %s", [limit])


# ---------------------------------------------------------------- each engineer sees their own part of an analysis in the portal
def publish(run_id, user, on=True):
    """Make an analysis visible to engineers (each sees only the machines assigned to them). Reversible; the analysis itself is never removed."""
    if not db.one("SELECT 1 AS x FROM portal_match_run WHERE run_id = %s", [run_id]):
        raise MatchError("That analysis no longer exists.", 404)
    with db.write() as con:
        con.execute("UPDATE portal_match_run SET published_at = CASE WHEN %s THEN now() END, published_by = CASE WHEN %s THEN %s END WHERE run_id = %s", (on, on, user, run_id))
    return {"run_id": run_id, "published": bool(on)}


def published(limit=10):
    return db.query("SELECT run_id, at, tool, filename, published_at, summary->>'report_as_of' AS report_as_of FROM portal_match_run WHERE published_at IS NOT NULL ORDER BY published_at DESC LIMIT %s", [limit])


def notice(user):
    """The sign-in message for an engineer: the newest shared analysis they have not yet opened, with their own figures. None when there is nothing to tell."""
    eng = user.get("engineer_key")
    last = (published(1) or [None])[0]
    if user["role"] == "ADMIN" or not eng or not last or db.one("SELECT 1 AS x FROM portal_match_seen WHERE username = %s AND run_id = %s", [user["username"], last["run_id"]]):
        return None
    sc = scoped(get(last["run_id"]), eng)
    if not sc["assets"]:
        return None
    urgent = [r for r in sc["assets"] if not r["installed"] and r["priority"] == HIGH][:8]
    return {"run_id": last["run_id"], "shared_at": last["published_at"], "summary": sc["summary"],
            "urgent": [{k: r[k] for k in ("ci", "user", "location", "action")} for r in urgent], "needs": bool(sc["summary"]["missing"] or sc["summary"]["attention"])}


def mark_seen(username, run_id):
    with db.write() as con:
        con.execute("INSERT INTO portal_match_seen (username, run_id) VALUES (%s, %s) ON CONFLICT DO NOTHING", (username, run_id))


def engineer_choices(run_id):
    """For an administrator previewing: every engineer who has machines in this analysis."""
    d = get(run_id)
    names = {r["engineer_key"]: r["display_name"] for r in db.query("SELECT engineer_key, display_name FROM portal_engineer")}
    return [{"key": k, "name": names.get(k) or k.title(), "machines": n} for k, n in sorted(Counter(r["engineer"] for r in d["assets"]).items())]


def scoped(d, eng):
    """One engineer's part of an analysis: only their machines, with figures worked out for those machines alone."""
    s = d["summary"]
    w = _work(d, eng)
    mine, miss, attn, named = w["mine"], w["miss"], w["quiet"], w["named"]
    inst = len(mine) - len(miss)
    healthy = sum(1 for r in mine if r["health"] == "Healthy")
    high = [r for r in miss if r["priority"] == HIGH]
    ic = Counter()
    for r in attn:
        for x in filter(None, r["issues"].split("; ")):
            ic[_generic(x)] += 1
    t = s["tool"]
    f = [{"severity": "Good" if not miss and not attn else "Needs attention", "finding": f"{inst:,} of your {len(mine):,} machines appear in the {t} report ({_pct(inst, len(mine))}%); {healthy:,} are working properly.",
          "action": "Nothing to do." if not miss and not attn else "Work through the lists below, urgent machines first."}]
    if high:
        f.append({"severity": "Critical", "finding": f"{len(high):,} machines deployed to users have no agent at all.", "action": "Install the agent, or restore its connection, on these first."})
    if ic:
        f.append({"severity": "Needs attention", "finding": "Installed, but not working properly: " + "; ".join(f"{k} ({n:,})" for k, n in ic.most_common(4)) + ".", "action": "Repair or reinstall the agent; the 'Needs attention' list says what is wrong with each."})
    if named:
        f.append({"severity": "Needs attention", "finding": f"{len(named):,} of your machines are probably in the report under a slightly different name.", "action": "Confirm each pair, then correct the CI or rename the computer."})
    low = [r for r in miss if r["priority"] == LOW]
    if low:
        f.append({"severity": "Information", "finding": f"{len(low):,} of your missing machines are in store, surplus or transferred.", "action": "Not urgent; install the agent before they are issued."})
    by_class = {}
    for r in mine:
        x = by_class.setdefault(r["class"], {"label": r["class"], "total": 0, "installed": 0, "missing": 0, "attention": 0})
        x["total"] += 1
        x["installed"] += r["installed"]
        x["missing"] += not r["installed"]
        x["attention"] += r["health"] == "Needs attention"
    return {"run_id": d["run_id"], "filename": d["filename"], "engineer": eng, "assets": mine,
            "summary": {"tool": t, "report_as_of": s.get("report_as_of"), "target": s["target"], "machines": len(mine), "installed": inst, "missing": len(miss), "healthy": healthy, "attention": len(attn), "high": len(high),
                        "coverage": _pct(inst, len(mine)), "healthy_pct": _pct(healthy, len(mine)), "findings": f, "by_class": sorted(by_class.values(), key=lambda x: -x["missing"])}}


# ---------------------------------------------------------------- files
def _t(title, columns, rows):
    return {"title": title, "columns": columns, "rows": rows}


ASSET_COLS = [("ci", "Asset (CI)"), ("class", "Class"), ("type", "Type"), ("make", "Make"), ("model", "Model"), ("serial", "Serial no"), ("status", "Asset status"), ("user", "User"), ("location", "Location"),
              ("engineer", "Engineer"), ("priority", "Priority"), ("action", "Action"), ("possible", "Probably listed as"), ("possible_by", "Because")]
INSTALLED_COLS = [("ci", "Asset (CI)"), ("class", "Class"), ("make", "Make"), ("model", "Model"), ("status", "Asset status"), ("user", "User"), ("location", "Location"), ("engineer", "Engineer"),
                  ("r_group", "Group"), ("r_ip", "Report IP"), ("r_os", "Report OS"), ("r_seen", "Last report"), ("r_age", "Days silent"), ("health", "Health"), ("issues", "Problems"), ("advice", "Also note"), ("action", "Action")]
GROUP_COLS = [("label", "{}"), ("total", "Machines"), ("installed", "In the report"), ("missing", "Not in the report"), ("high", "Missing, in use"), ("coverage", "Coverage %"), ("attention", "Need attention"), ("healthy_pct", "Healthy %")]


def _group(title, label, rows):
    return _t(title, [(k, label if k == "label" else v) for k, v in GROUP_COLS], rows)


def _yn(rows):
    return [{**r, "installed": "Yes" if r["installed"] else "No"} for r in rows]


def _titled(rows, *keys):
    return [{**r, **{k: (r[k].title() if k in r and isinstance(r[k], str) else r.get(k)) for k in keys}} for r in rows]


SUITE_COLS = [("ci", "Asset (CI)"), ("state", "Sensor only / Full suite"), ("todo", "What to do"), ("class", "Class"), ("model", "Model"), ("serial", "Serial no"), ("user", "User"), ("location", "Location"),
              ("engineer", "Engineer"), ("status", "Asset status"), ("suite", "Console says"), ("suite_pc", "PC check script says"), ("health", "Health"), ("issues", "Problems"), ("r_seen", "Last report"), ("priority", "Priority")]
FIELD_COLS = [("ci", "Asset (CI) - computer name"), ("state", "Sensor only / Full suite"), ("todo", "What to do"), ("user", "User"), ("location", "Location"), ("model", "Model"), ("serial", "Serial no"),
              ("class", "Class"), ("status", "Asset status"), ("priority", "Priority"), ("r_seen", "Last report")]
GUIDE = [("Purpose", "This list shows, for every machine assigned to you, whether Trend Micro is missing, sensor only, or the full suite, and what to do about it."),
         ("Step 1 - find the machine", "Use 'Asset (CI)' (the computer name), then User and Location to walk to the right desk. Serial no and Model confirm you have the right PC. To read a PC's name: right-click This PC > Properties > Device name."),
         ("Step 2 - read the state", "Full suite = Trend Micro is installed with anti-malware: nothing to install. Sensor only = only the monitoring sensor is installed, there is NO anti-malware protection: the full agent must be installed. "
                                     "Not installed = no Trend Micro at all. Not determined = the report did not say; run ankcheck.bat on the PC to find out."),
         ("Step 3 - do the work", "Work top to bottom: High priority (PCs in use) first. Do what 'What to do' says. Afterwards run ankcheck.bat on the PC so the portal records the result."),
         ("Step 4 - if it cannot be done", "PC retired, in store or not found: tell the Team Leader so the register can be corrected. Name differs from the report: the 'What to do' column says so."),
         ("Where results appear", "After the next Trend Micro report is loaded in the portal, fixed machines move to Full suite. Reports > My inventory report always shows your latest list.")]


def _with_state(rows):
    out = []
    for r in rows:
        state = "Not installed" if not r["installed"] else (r["suite"] or r["suite_pc"] or "Not determined")
        if not r["installed"] and r["suite_pc"]:
            todo = f"The PC itself reports {r['suite_pc']}, but {r['ci']} is not in the console report. Do not reinstall: check the PC's connection to the Apex server and its console group."
        elif not r["installed"]:
            todo = f"Install the Trend Micro agent (full suite) on {r['ci']} at {r['location'] or 'its location'}" + (f" - user {r['user'].title()}" if r["user"] else "") + "."
        elif state == "Sensor only":
            todo = "Only the sensor is installed - no anti-malware. Install the full Trend Micro (Apex One) agent."
        elif r["health"] == "Needs attention":
            todo = r["action"] or "Repair the agent."
        elif state == "Not determined":
            todo = "Run ankcheck.bat on this PC to confirm what is installed."
        else:
            todo = "Nothing to do."
        out.append({**r, "state": state, "todo": todo, "user": r["user"].title() if r["user"] else "", "engineer": r["engineer"].title()})
    return out


_ORDER = {"Not installed": 0, "Sensor only": 1, "Not determined": 2, "Full suite": 3}


def _sorted(rows):
    return sorted(rows, key=lambda r: (_ORDER.get(r["state"], 4), {HIGH: 0, MEDIUM: 1, LOW: 2}.get(r["priority"], 3), r["location"], r["ci"]))


def suite_table(rows):
    return _t("All machines - Sensor only / Full suite", SUITE_COLS, _sorted(_with_state(rows)))


def _suite_summary(rows):
    return _t("Sensor only vs Full suite", [("label", "State"), ("n", "Machines")], [{"label": k, "n": n} for k, n in Counter(r["state"] for r in _with_state(rows)).most_common()])


def investigation(d, fmt, user, eng=None):
    """The investigation report: administrators/team leads get everything; an engineer gets the same report for their own machines only."""
    s = d["summary"]
    if eng is None:
        if fmt == "csv":
            return export.render("csv", s["tool"], "", [suite_table(d["assets"])])
        return render(d, fmt, user)
    mine = [r for r in d["assets"] if r["engineer"] == eng]
    sc = scoped(d, eng)
    t = [_t("Key findings", [("severity", "Rating"), ("finding", "Finding"), ("action", "What to do")], sc["summary"]["findings"]), _suite_summary(mine), suite_table(mine)]
    if fmt == "csv":
        t = t[2:]
    m = sc["summary"]
    k = [("Machines assigned", m["machines"]), (f"In the {s['tool']} report", m["installed"]), ("Not installed", m["missing"]), ("Need attention", m["attention"]), ("Coverage (%)", m["coverage"])]
    return export.render(fmt, f"{s['tool']} investigation - {eng.title()}", f"Report: {d['filename']} - issued {dt.date.today():%d %b %Y}", t, k, footer=user["username"])


def field_report(d, fmt, user, eng=None):
    """The light report for the people doing the work: which machine, what state, what to do, and a how-to-use sheet. eng=None: every engineer's machines."""
    s = d["summary"]
    rows = _with_state([r for r in d["assets"] if eng is None or r["engineer"] == eng])
    work = _sorted([r for r in rows if r["todo"] != "Nothing to do."])
    cols = FIELD_COLS if eng else FIELD_COLS[:5] + [("engineer", "Engineer")] + FIELD_COLS[5:]
    t = [_t("Your work list" if eng else "Work list", cols, work), _t("All machines", cols, _sorted(rows)), _t("How to use this report", [("k", "Topic"), ("v", "Guidance")], [{"k": a, "v": b} for a, b in GUIDE])]
    if fmt == "csv":
        return export.render("csv", s["tool"], "", t[:1])
    k = [("Machines", len(rows)), ("To do", len(work)), ("Not installed", sum(r["state"] == "Not installed" for r in rows)), ("Sensor only", sum(r["state"] == "Sensor only" for r in rows)),
         ("Full suite", sum(r["state"] == "Full suite" for r in rows))]
    if fmt == "xlsx":
        t = [t[2], t[0], t[1]]          # the guide first: it is the first thing a new reader needs
    return export.render(fmt, f"{s['tool']} field report" + (f" - {eng.title()}" if eng else ""), f"Report: {d['filename']} - issued {dt.date.today():%d %b %Y}", t, k, footer=user["username"])


CONSOL_COLS = [("endpoint", "Endpoint (computer name)"), ("found", "Found in file"), ("registered", "In asset register"), ("state", "Sensor only / Full suite"), ("todo", "What to do"),
               ("conn", "Console connection"), ("user", "User"), ("location", "Location"), ("engineer", "Engineer"), ("class", "Class"), ("model", "Model"), ("serial", "Serial no"),
               ("os", "Operating system"), ("ip", "IP address"), ("pc", "PC check script says")]
CONSOL_GUIDE = [("What this report is", "Every Trend Micro endpoint of this site from the console file(s), joined into ONE list with repeats removed, and compared with the asset register and with what the check script saw on the PC."),
                ("Sensor only", "Only the Trend Micro sensor is on the PC. It is seen by the console but has NO anti-malware protection. Install the full Apex One agent."),
                ("Full suite", "The full Apex One agent is on the PC (anti-malware plus sensor). Nothing to install."),
                ("Not determined", "The console file does not say which of the two it is and the PC was never checked. Run ankcheck.bat on it (share \\\\10.205.64.203\\ank-helpdesk\\common-for-all\\)."),
                ("Not installed", "The register has the machine but the console report does not. If 'What to do' says the PC itself reports an agent, do not reinstall - check its connection and group."),
                ("Sheets", "Merged list = the working list. Sensor only = fix these first. Not in console = register machines with no agent in the console. Not in register = console names the register does not know. Duplicates = names repeated across the files (the latest record was kept).")]


def _merged_rows(d, eng=None):
    ws = _with_state([r for r in d["assets"] if eng is None or r["engineer"] == eng])
    base = lambda r: {"registered": "Yes", "state": r["state"], "todo": r["todo"], "user": r["user"], "location": r["location"], "engineer": r["engineer"], "class": r["class"], "model": r["model"],   # noqa: E731
                      "serial": r["serial"], "pc": r["suite_pc"]}
    merged = [{**base(r), "endpoint": r["r_name"] or r["ci"], "found": r["r_files"], "conn": r["r_conn"], "os": r["r_os"], "ip": r["r_ip"]} for r in ws if r["installed"]]
    if eng is None:
        merged += [{"endpoint": u["name"], "found": u.get("files", ""), "registered": "No", "state": "Not in register", "todo": u["action"], "conn": "", "user": "", "location": "", "engineer": "",
                    "class": "", "model": "", "serial": "", "os": u["os"], "ip": u["ip"], "pc": ""} for u in d["unregistered"] if u["kind"] != "outside"]
    order = {"Sensor only": 0, "Not determined": 1, "Not in register": 2, "Full suite": 3}
    merged.sort(key=lambda x: (order.get(x["state"], 4), x["endpoint"]))
    absent = [{**base(r), "endpoint": r["ci"], "found": "", "conn": "", "os": "", "ip": ""} for r in ws if not r["installed"]]
    return merged, absent


def consolidated(d, fmt, user, eng=None):
    """The consolidated comparison: all console files merged into one de-duplicated list, set against the register and the PC check script, each machine marked Sensor only /
    Full suite / Not determined with a plain 'what to do'. Administrators and team leads get everything; an engineer only their own machines."""
    s = d["summary"]
    merged, absent = _merged_rows(d, eng)
    sensor = [r for r in merged if r["state"] == "Sensor only"]
    n = Counter(r["state"] for r in merged)
    t = [_t("Merged list", CONSOL_COLS, merged)]
    if fmt != "csv":
        t = [_t("How to use this report", [("k", "Topic"), ("v", "Guidance")], [{"k": a, "v": b} for a, b in CONSOL_GUIDE]), t[0], _t("Sensor only", CONSOL_COLS, sensor),
             _t("Not in console", CONSOL_COLS, absent)]
        if eng is None:
            t += [_t("Not in register", [("name", "Computer name"), ("ip", "IP address"), ("os", "OS"), ("files", "Found in file"), ("why", "Why"), ("action", "What to do")], [u for u in d["unregistered"] if u["kind"] != "outside"]),
                  _t("Duplicates removed", [("name", "Computer name"), ("rows", "Rows in the files"), ("detail", "Records seen")], d["duplicates"])]
    if fmt == "pdf":
        t = [x for x in t if x["title"] not in ("Merged list", "Not in console")]          # key people get the headline and the problem lists; the full lists are in the Excel
    k = [("Endpoints in merged list", len(merged)), ("Full suite", n["Full suite"]), ("Sensor only", n["Sensor only"]), ("Not determined", n["Not determined"]), ("Not in console", len(absent))]
    if eng is None:
        k.append(("Repeat rows removed", sum(x["rows"] - 1 for x in d["duplicates"])))
    if fmt == "csv":
        return export.render("csv", s["tool"], "", t)
    return export.render(fmt, f"{s['tool']} consolidated comparison" + (f" - {eng.title()}" if eng else ""), f"Report: {d['filename']} - issued {dt.date.today():%d %b %Y}", t, k, footer=user["username"])


def latest_trend_run(published_only):
    r = db.one("SELECT run_id FROM portal_match_run WHERE summary->>'product' IN ('trendmicro', 'apex')" + (" AND published_at IS NOT NULL" if published_only else "") + " ORDER BY at DESC LIMIT 1")
    return r["run_id"] if r else None


def tables(d, management=False):
    s, a = d["summary"], d["assets"]
    miss = sorted((r for r in a if not r["installed"]), key=lambda r: ({HIGH: 0, MEDIUM: 1, LOW: 2}[r["priority"]], r["engineer"], r["ci"]))
    attn = sorted((r for r in a if r["health"] == "Needs attention"), key=lambda r: (-(r["r_age"] or 0), r["ci"]))
    about = [{"item": "Tool / product", "value": f"{s['tool']} ({s.get('product_label', '')})"}, {"item": "File(s)", "value": "; ".join(f"{x['name']} - {x['rows']:,} rows" for x in s.get("files", [])) or d["filename"]},
             {"item": "Report date", "value": s.get("report_as_of", "")}, {"item": "Site / classes checked", "value": f"{'/'.join(s['prefixes'])} - {', '.join(c.title() for c in s['classes'])}"},
             {"item": "Method", "value": "A register machine counts as installed when its Asset (CI) equals a computer name in the report (case and domain suffix ignored). An installed machine is 'healthy' only if the agent is communicating, "
                                          "has reported recently, and shows no installation or protection fault."},
             {"item": "Totals reconciled", "value": "Yes - every total below adds up to the file and the register" if s.get("reconciled", True) else "NO - a total did not reconcile; do not circulate this report"}]
    t = [_t("About this report", [("item", "Item"), ("value", "Detail")], about), _t("Key findings", [("severity", "Rating"), ("finding", "Finding"), ("action", "What to do")], s["findings"]),
         _t("Data checks on the report", [("status", "Result"), ("check", "Check"), ("detail", "What was found"), ("handled", "How it was handled")], [{**c, "status": {"ok": "OK", "warn": "Check", "info": "Note"}.get(c["status"], c["status"])} for c in s.get("checks", [])]),
         _group("Coverage by class", "Class", s["by_class"]), _group("Coverage by engineer", "Engineer", s["by_engineer"])]
    if s.get("suite_counts"):
        t.insert(2, _suite_summary(a))
    if management:                  # the PDF is the short version: findings, class, engineer and the worst machines; the Excel pack has everything
        t.append(_t(f"Top {min(25, len(miss))} machines to fix first (deployed, no agent)", [("ci", "Asset (CI)"), ("class", "Class"), ("model", "Model"), ("user", "User"), ("location", "Location"), ("engineer", "Engineer")],
                    [r for r in miss if r["priority"] == HIGH][:25]))
        return t
    t += [_group("Coverage by location", "Location", s["by_location"]), _group("Coverage by asset status", "Asset status", s["by_status"]),
          _t("Problems seen on installed machines", [("label", "Problem"), ("n", "Machines")], s.get("issue_counts", [])), _t("Things to improve (not faults)", [("label", "Note"), ("n", "Machines")], s.get("advice_counts", [])),
          _t("Models with most gaps", [("label", "Make / model"), ("n", "Not installed")], s["gap_models"]),
          _t("Report versus what was seen on the machine", [("ci", "Asset (CI)"), ("observed", "Observed on the machine"), ("xcheck", "Disagreement"), ("engineer", "Engineer")], [r for r in a if r.get("xcheck")]),
          _t("Not installed - action list", ASSET_COLS, miss),
          _t("Installed - needs attention", INSTALLED_COLS, attn),
          _t("Probably the same machine", [("ci", "Asset (CI) in register"), ("possible", "Name in report"), ("possible_by", "Because"), ("engineer", "Engineer"), ("action", "Action")], [r for r in miss if r["possible"]]),
          _t("In report, not in register", [("name", "Name in report"), ("ip", "IP"), ("os", "OS"), ("group", "Group"), ("seen", "Last report"), ("age", "Days silent"), ("why", "Finding"), ("action", "Action")], [u for u in d["unregistered"] if u["kind"] != "mismatch"]),
          _t("In a group other than the usual", [("name", "Computer"), ("ip", "IP"), ("group", "Group"), ("seen", "Last report")], d.get("group_issues", [])),
          _t("Other sites' devices in this group", [("name", "Computer"), ("ip", "IP"), ("group", "Group")], d.get("foreign", [])),
          _t("Duplicate names in report", [("name", "Computer name"), ("rows", "Rows"), ("detail", "Records")], d["duplicates"]),
          _t("How the totals add up", [("label", "Line"), ("n", "Count"), ("note", "Note")], s.get("recon", [])),
          _t("Full match", [("ci", "Asset (CI)"), ("installed", "In the report"), ("health", "Health")] + INSTALLED_COLS[1:13] + [("priority", "Priority")], _yn(a))]
    if s.get("suite_counts"):
        t.append(suite_table(a))
    return t


def kpis(s):
    t = s["tool"]
    return [("Machines checked", s["assets"]), (f"In the {t} report", s["installed"]), ("Not in the report", s["missing"]), ("Coverage (%)", s["coverage"]), ("Target (%)", s["target"]),
            ("Installed and healthy", s["healthy"]), ("Installed, need attention", s["attention"]), ("Healthy of all checked (%)", s["healthy_pct"]),
            ("Deployed, no agent", s["missing_high"]), ("Not in register", s["unregistered"]), ("Probably same machine, name differs", s["name_mismatch"])]


def render(d, fmt, user):
    s = d["summary"]
    ch = [{"title": "Coverage by class (%)", "rows": [{"label": x["label"].title(), "n": x["coverage"]} for x in s["by_class"]]},
          {"title": "Machines to fix (missing or unhealthy), by engineer", "rows": [{"label": x["label"].title(), "n": x["missing"] + x["attention"]} for x in s["by_engineer"] if x["missing"] + x["attention"]][:12]},
          {"title": "Machines without the agent, by location", "rows": [{"label": x["label"], "n": x["missing"]} for x in s["by_location"] if x["missing"]][:12]}]
    if s.get("issue_counts"):
        ch.append({"title": "Problems on installed machines", "rows": [{"label": x["label"][:48], "n": x["n"]} for x in s["issue_counts"][:8]]})
    if s["has_os"] and s["os_mix"]:
        ch.append({"title": "Operating systems on installed machines", "rows": s["os_mix"]})
    title = f"{s['tool']} coverage - {', '.join(s['prefixes'])}"
    sub = f"Report: {d['filename']} (as of {s.get('report_as_of', '')[:10]}) - {s['assets']:,} register machines checked - issued {dt.date.today():%d %b %Y}"
    return export.render(fmt, title, sub, tables(d, management=fmt == "pdf"), kpis(s), ch, footer=user["username"])


def distribute(run_id, user):
    """What happens automatically after every analysis: share it with the engineers in the portal (each sees only their own machines) and e-mail each
    engineer their own list. A mail problem (e-mail switched off, a server error) never loses the analysis; the outcome is stored with it."""
    from . import mailer
    publish(run_id, user["username"])
    try:
        mail = send_to_engineers(run_id, user)
        mail["error"] = None
    except mailer.MailError as e:
        mail = {**send_to_engineers(run_id, user, dry_run=True), "error": str(e)}          # still show who would have been mailed
    with db.write() as con:
        con.execute("UPDATE portal_match_run SET result = result || jsonb_build_object('mail', %s::jsonb) WHERE run_id = %s", (json.dumps(mail, default=str), run_id))
    return mail


# ---------------------------------------------------------------- each engineer's own list, by e-mail
MAIL_RULE = "MATCH_REPORT"      # notify_log.rule_key: a re-send of the same analysis never mails the same engineer twice
UNASSIGNED = "UNASSIGNED"


def _work(d, eng):
    """What one engineer has to do: machines without the agent (most urgent first), silent agents, wrong-name cases. Low-priority gaps (in store) are listed but do not by themselves trigger a mail."""
    mine = [r for r in d["assets"] if r["engineer"] == eng]
    miss = sorted((r for r in mine if not r["installed"] and not r.get("suite_pc")), key=lambda r: ({HIGH: 0, MEDIUM: 1, LOW: 2}[r["priority"]], r["ci"]))
    quiet = sorted((r for r in mine if r["health"] == "Needs attention" or (not r["installed"] and r.get("suite_pc"))), key=lambda r: (-(r["r_age"] or 0), r["ci"]))
    named = [r for r in miss if r["possible"]]
    needs = [r for r in miss if r["priority"] != LOW] + quiet
    return {"mine": mine, "miss": miss, "quiet": quiet, "named": named, "needs": bool(needs)}


def engineer_workbook(d, eng, user):
    return field_report(d, "xlsx", user, eng)[0]


def engineer_plan(d):
    """One line per engineer with something to do: who, how many, which address (or why there is none). Nothing is sent."""
    from . import mailer
    names = {r["engineer_key"]: r["display_name"] for r in db.query("SELECT engineer_key, display_name FROM portal_engineer")}
    sent = {r["event_key"]: r["status"] for r in db.query("SELECT event_key, status FROM notify_log WHERE rule_key = %s AND event_key LIKE %s", [MAIL_RULE, d["run_id"] + ":%"])}
    out = []
    for eng in sorted({r["engineer"] for r in d["assets"]}):
        w = _work(d, eng)
        if not w["miss"] and not w["quiet"]:
            continue
        item = {"engineer": eng, "name": names.get(eng) or eng.title(), "machines": len(w["mine"]), "missing": len(w["miss"]), "high": sum(1 for r in w["miss"] if r["priority"] == HIGH),
                "silent": len(w["quiet"]), "email": None, "status": "READY", "earlier": sent.get(f"{d['run_id']}:{eng}")}
        if eng == UNASSIGNED:
            item["status"] = "NO_ENGINEER"          # nobody to mail: these machines have no engineer in the register - fix that first
        elif not w["needs"]:
            item["status"] = "NOTHING_URGENT"
        else:
            item["email"] = mailer.engineer_email(eng)
            if not item["email"]:
                item["status"] = "NO_ADDRESS"
            elif item["earlier"] == "SENT":
                item["status"] = "ALREADY_SENT"
        out.append(item)
    return out


def send_to_engineers(run_id, user, dry_run=False, only=None):
    """Mail each engineer their own machines (summary in the body, Excel attached). Idempotent per analysis and engineer; a failure for one never stops the rest."""
    from . import mailer
    d = get(run_id)
    cfg = mailer.get_smtp()
    if not cfg["enabled"] and not dry_run:
        raise mailer.MailError("E-mail is not switched on. Ask an administrator to set it up (Administration > E-mail).")
    plan = engineer_plan(d)
    if dry_run:
        return {"mail_enabled": cfg["enabled"], "plan": plan}
    s = d["summary"]
    for item in plan:
        if item["status"] != "READY" or (only and item["engineer"] not in only):
            continue
        w = _work(d, item["engineer"])
        urgent = [r for r in w["miss"] if r["priority"] != LOW][:15]
        subject = f"{s['tool']}: {item['missing']} of your {item['machines']} machines have no agent" + (f", {item['silent']} need attention" if item["silent"] else "") + f" - {dt.date.today():%d %b %Y}"
        intro = (f"Dear {item['name'].title()}, the centre's {s['tool']} report ({d['filename']}) was matched with the asset register. Of the {item['machines']} machines assigned to you, "
                 f"{item['missing']} do not appear in the report ({item['high']} of them deployed to users)" + (f" and {item['silent']} have an agent that is installed but not working properly" if item["silent"] else "") +
                 ". The first ones to fix are below; your full list is in the attached Excel file. Please install or repair the agent, or tell us if the machine is retired or the CI is wrong.")
        body, text = mailer.page(subject, intro, [{**r, "class": r["class"].title()} for r in urgent], [("ci", "Asset (CI)"), ("class", "Class"), ("model", "Model"), ("user", "User"), ("location", "Location"), ("action", "Action")],
                                 footer=f"Automatic message from the ITAM Portal (inventory match, sent by {user['username']})." + (f" {cfg['portal_url']}" if cfg["portal_url"] else ""))
        status, err = "SENT", None
        try:
            fname = f"{export.safe_name(s['tool'])}_{export.safe_name(item['engineer'])}_{s['as_of'].replace('-', '')}.xlsx"
            mailer.send(item["email"], subject, text, body, [(fname, engineer_workbook(d, item["engineer"], user), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")], cfg=cfg)
        except mailer.MailError as e:
            status, err = "FAILED", str(e)[:300]
        with db.write() as con:
            con.execute("""INSERT INTO notify_log (rule_key, event_key, recipient, subject, status, error) VALUES (%s,%s,%s,%s,%s,%s)
                           ON CONFLICT (rule_key, event_key, recipient) DO UPDATE SET status = EXCLUDED.status, error = EXCLUDED.error, at = now(), attempts = notify_log.attempts + 1""",
                        (MAIL_RULE, f"{run_id}:{item['engineer']}", item["email"], subject, status, err))
        item["status"], item["error"] = status, err
    return {"mail_enabled": True, "plan": plan}
