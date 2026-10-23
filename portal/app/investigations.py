"""Investigations: a finished analysis (an Excel workbook, optionally with its PDF summary) kept in the portal so it can be read sheet by sheet on screen
and downloaded again exactly as it was made. The workbook is parsed once on upload (values as shown in Excel, bold and fill colours, merged cells) and stored
in PostgreSQL next to the original file, so backups, restores and release rollbacks cover it. Nothing is ever erased: "remove" only hides an investigation
from the current list. The data lives in the database, never in the repository (a report names staff, users and machines).

Who may see / add: a real administrator or a Team Leader/SI User - never the read-only demo account (its password is public)."""
import datetime as dt
import hashlib
import io
import json
import re
import zipfile

from . import db

MAX_XLSX = 15 * 1024 * 1024
MAX_PDF = 25 * 1024 * 1024
MAX_UNZIPPED = 150 * 1024 * 1024        # a workbook is a zip file: refuse one that expands into something huge (zip bomb)
MAX_SHEETS, MAX_ROWS, MAX_COLS, MAX_CELLS, MAX_TEXT = 20, 20000, 60, 800000, 4000
NAVY = "1F3864"                          # header fill of every table the portal writes; a row of it marks the header of a table sheet

DDL = [
    """CREATE TABLE IF NOT EXISTS portal_investigation (
         inv_id BIGSERIAL PRIMARY KEY, title TEXT NOT NULL, note TEXT NOT NULL DEFAULT '',
         xlsx_name TEXT NOT NULL, xlsx_size INTEGER NOT NULL, xlsx_sha256 TEXT NOT NULL, xlsx BYTEA NOT NULL,
         pdf_name TEXT, pdf_size INTEGER, pdf BYTEA,
         sheet_index JSONB NOT NULL, sheets JSONB NOT NULL,
         uploaded_by TEXT NOT NULL, uploaded_at TIMESTAMPTZ NOT NULL DEFAULT now(), superseded_at TIMESTAMPTZ, superseded_by TEXT)""",
    "CREATE INDEX IF NOT EXISTS ix_investigation_at ON portal_investigation (uploaded_at DESC)",
]


class InvError(Exception):
    http_error = True

    def __init__(self, message, status=400):
        super().__init__(message)
        self.status = status


def can_access(user):
    return not user.get("read_only") and (user["role"] == "ADMIN" and user.get("group", user["role"]) == "ADMIN" or bool(user.get("lead_tools")))


# ---------------------------------------------------------------- reading the workbook
_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _show(v, nf):
    """A cell value as Excel would show it where that matters (percent, thousands, signed, dates); plain numbers stay numbers."""
    if isinstance(v, bool):
        return "Yes" if v else "No"
    if isinstance(v, dt.datetime):
        return v.strftime("%d-%b-%Y %H:%M") if (v.hour or v.minute or v.second) else v.strftime("%d-%b-%Y")
    if isinstance(v, dt.date):
        return v.strftime("%d-%b-%Y")
    if isinstance(v, (int, float)):
        if "%" in nf:
            m = re.search(r"0\.(0+)%", nf)
            return f"{v * 100:.{len(m.group(1)) if m else 0}f}%"
        if nf.startswith("+"):
            return f"{int(round(v)):+,}" if round(v) else "0"
        return int(v) if isinstance(v, float) and v.is_integer() else v
    s = _CTRL.sub("", str(v))
    return s[:MAX_TEXT]


def _style(cell):
    """[bold, fill, font colour] for a cell that has any of them, else None. Only plain RGB colours are kept (theme colours cannot be reproduced)."""
    f = cell.font
    bold = 1 if (f is not None and f.b) else 0
    bg = fg = ""
    fl = cell.fill
    if fl is not None and fl.fill_type == "solid" and fl.fgColor is not None and fl.fgColor.type == "rgb" and isinstance(fl.fgColor.rgb, str):
        bg = fl.fgColor.rgb[-6:].upper()
    if f is not None and f.color is not None and f.color.type == "rgb" and isinstance(f.color.rgb, str) and f.color.rgb[-6:].upper() != "000000":
        fg = f.color.rgb[-6:].upper()
    return [bold, bg, fg] if (bold or bg or fg) else None


def _sheet(ws):
    rows_n, cols_n = min(ws.max_row or 0, MAX_ROWS), min(ws.max_column or 0, MAX_COLS)
    grid, styles = [], {}
    for r, row in enumerate(ws.iter_rows(min_row=1, max_row=rows_n, max_col=cols_n)):
        out = []
        for c, cell in enumerate(row):
            v = None if cell.value is None else _show(cell.value, cell.number_format or "")
            if isinstance(v, str) and v.startswith("="):
                v = None                                    # a formula whose result Excel never stored: nothing to show
            out.append(v)
            st = _style(cell)
            if st:
                styles[f"{r}:{c}"] = st
        grid.append(out)
    last_r = max((i for i, row in enumerate(grid) if any(v not in (None, "") for v in row)), default=-1)
    grid = grid[:last_r + 1]
    last_c = max((j for row in grid for j, v in enumerate(row) if v not in (None, "")), default=-1)
    grid = [row[:last_c + 1] for row in grid]
    styles = {k: v for k, v in styles.items() if int(k.split(":")[0]) <= last_r and int(k.split(":")[1]) <= last_c}
    merges = []
    for m in ws.merged_cells.ranges:
        r1, c1, r2, c2 = m.min_row - 1, m.min_col - 1, min(m.max_row - 1, last_r), min(m.max_col - 1, last_c)
        if r1 <= last_r and c1 <= last_c and (r2 > r1 or c2 > c1):
            merges.append([r1, c1, r2, c2])
    header = next((i for i in range(min(3, len(grid))) if sum(1 for j in range(len(grid[i])) if styles.get(f"{i}:{j}", [0, "", ""])[1] == NAVY) >= 2), None)
    kind = "table" if header is not None and len(grid) > header + 1 and all(m[0] < header for m in merges) else "document"
    return {"name": ws.title, "kind": kind, "header": header if kind == "table" else None, "rows": grid, "styles": styles, "merges": merges}


def parse(data):
    if not data.startswith(b"PK\x03\x04"):
        raise InvError("That is not an Excel workbook (.xlsx).", 415)
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if len(z.infolist()) > 2000 or sum(i.file_size for i in z.infolist()) > MAX_UNZIPPED:
                raise InvError("That workbook is too large to open safely.", 413)
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except InvError:
        raise
    except Exception:
        raise InvError("That file could not be read as an Excel workbook.", 415)
    sheets = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        if len(sheets) >= MAX_SHEETS:
            raise InvError(f"The workbook has more than {MAX_SHEETS} sheets.", 413)
        if (ws.max_row or 0) * (ws.max_column or 0) > MAX_CELLS:
            raise InvError(f"Sheet '{ws.title}' is too large to show on screen.", 413)
        sheets.append(_sheet(ws))
    if not sheets:
        raise InvError("The workbook has no visible sheets.")
    return sheets


# ---------------------------------------------------------------- storing, listing, fetching
def _name(raw, ext):
    base = re.split(r"[\\/]", raw or "")[-1]
    stem = re.sub(r"[^\w.\- ()]", "_", base.rsplit(".", 1)[0]).strip(" .") or "investigation"
    return f"{stem[:100]}.{ext}"


def _check_pdf(data):
    if not data or not data.startswith(b"%PDF-"):
        raise InvError("That is not a PDF file.", 415)
    if len(data) > MAX_PDF:
        raise InvError(f"The PDF is larger than {MAX_PDF // (1024 * 1024)} MB.", 413)


def add(title, note, filename, data, by, pdf_name=None, pdf=None):
    title = re.sub(r"\s+", " ", (title or "").strip())[:120]
    if len(title) < 3:
        raise InvError("Give the investigation a title (at least 3 characters).")
    note = (note or "").strip()[:1000]
    if not data:
        raise InvError("The file is empty.")
    if len(data) > MAX_XLSX:
        raise InvError(f"The workbook is larger than {MAX_XLSX // (1024 * 1024)} MB.", 413)
    sheets = parse(data)
    if pdf is not None:
        _check_pdf(pdf)
    index = [{"name": s["name"], "kind": s["kind"], "rows": max(0, len(s["rows"]) - (s["header"] + 1 if s["header"] is not None else 0)), "cols": max((len(r) for r in s["rows"]), default=0)} for s in sheets]
    fname = _name(filename, "xlsx")
    with db.write() as con:
        row = con.execute("""INSERT INTO portal_investigation (title, note, xlsx_name, xlsx_size, xlsx_sha256, xlsx, pdf_name, pdf_size, pdf, sheet_index, sheets, uploaded_by)
                             VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s::jsonb,%s::jsonb,%s) RETURNING inv_id""",
                          (title, note, fname, len(data), hashlib.sha256(data).hexdigest(), data, _name(pdf_name, "pdf") if pdf else None, len(pdf) if pdf else None, pdf,
                           json.dumps(index), json.dumps(sheets), by)).fetchone()
    return {"inv_id": row[0], "title": title, "sheets": len(sheets), "size": len(data)}


def attach_pdf(inv_id, filename, data, by):
    """The summary PDF that goes with a workbook. Only when there is none yet: a PDF is never overwritten (upload a new investigation instead)."""
    _check_pdf(data)
    with db.write() as con:
        r = con.execute("SELECT pdf_name FROM portal_investigation WHERE inv_id = %s FOR UPDATE", (inv_id,)).fetchone()
        if not r:
            raise InvError("No such investigation.", 404)
        if r[0]:
            raise InvError("This investigation already has a PDF. Upload a new investigation instead; the earlier PDF is kept.", 409)
        con.execute("UPDATE portal_investigation SET pdf_name = %s, pdf_size = %s, pdf = %s WHERE inv_id = %s", (_name(filename, "pdf"), len(data), data, inv_id))
    return {"inv_id": inv_id}


def listing():
    rows = db.query("""SELECT inv_id, title, note, xlsx_name, xlsx_size, pdf_name, pdf_size, sheet_index, uploaded_by, uploaded_at, superseded_at, superseded_by
                       FROM portal_investigation ORDER BY uploaded_at DESC, inv_id DESC""")
    return {"max_xlsx_mb": MAX_XLSX // (1024 * 1024), "max_pdf_mb": MAX_PDF // (1024 * 1024), "items": rows}


def sheet(inv_id, idx):
    r = db.one("SELECT sheets -> %s::int AS s FROM portal_investigation WHERE inv_id = %s", [idx, inv_id])
    if not r or r["s"] is None:
        raise InvError("No such sheet.", 404)
    return r["s"]


def fetch(inv_id, which):
    col = {"xlsx": ("xlsx_name", "xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"), "pdf": ("pdf_name", "pdf", "application/pdf")}.get(which)
    if not col:
        raise InvError("No such file.", 404)
    r = db.one(f"SELECT {col[0]} AS name, {col[1]} AS content FROM portal_investigation WHERE inv_id = %s", [inv_id])
    if not r or r["content"] is None:
        raise InvError("No such file.", 404)
    return r["name"], col[2], bytes(r["content"])


def remove(inv_id, by):
    with db.write() as con:
        n = con.execute("UPDATE portal_investigation SET superseded_at = now(), superseded_by = %s WHERE inv_id = %s AND superseded_at IS NULL", (by, inv_id)).rowcount
    if not n:
        raise InvError("No such current investigation.", 404)
