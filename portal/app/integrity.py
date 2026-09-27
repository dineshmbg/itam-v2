"""Data-integrity page: referential checks on the live database + reconciliation of the masters/ workbooks against it."""
import datetime as dt
import re
from pathlib import Path

from . import config, db

_file_cache = {}


def _n(sql, params=None):
    return db.one(sql, params)["n"]


def checks():
    items = []

    def add(cid, title, detail, bad, total, severity="fail", note=None, sample=None):
        """sample: SQL returning up to a few asset keys (column k) for the offending rows, shown as links on the page."""
        status = "pass" if bad == 0 else ("warn" if severity == "warn" else "fail")
        keys = [r["k"] for r in db.query(sample)] if sample and bad else []
        items.append({"id": cid, "title": title, "detail": detail, "bad": bad, "total": total, "status": status, "note": note, "sample": keys})

    def asset_check(cid, title, detail, where, note=None, extra_from="", severity="warn"):
        """A check over current ASSET rows (not components) - one condition drives the count and the example list."""
        base = f"FROM asset a {extra_from} WHERE a.is_current=1 AND a.record_level='ASSET' AND ({where})"
        add(cid, title, detail, _n(f"SELECT count(*) n {base}"), _n("SELECT count(*) n FROM asset WHERE is_current=1 AND record_level='ASSET'"), severity, note,
            f"SELECT a.asset_key k {base} ORDER BY 1 LIMIT 8")

    total_calls = _n("SELECT count(*) n FROM svc_call WHERE is_current=1")
    add("calls_asset", "Every call points to a known asset", "Calls whose asset (CI) is not in the asset register",
        _n("SELECT count(*) n FROM svc_call c WHERE c.is_current=1 AND (c.asset_key IS NULL OR NOT EXISTS (SELECT 1 FROM asset a WHERE a.asset_key=c.asset_key AND a.is_current=1))"), total_calls)
    total_in = _n("SELECT count(*) n FROM spare_inward WHERE is_current=1")
    add("inward_call", "Every part received belongs to a logged call", "Inward lines whose SR ID is not in the call tracker",
        _n("SELECT count(*) n FROM spare_inward s WHERE s.is_current=1 AND NOT EXISTS (SELECT 1 FROM svc_call c WHERE c.sr_id=s.sr_id AND c.is_current=1)"), total_in, "warn",
        "Usually a call that was logged after the part arrived, or a typing error in the SR ID.")
    total_out = _n("SELECT count(*) n FROM spare_outward WHERE is_current=1")
    add("outward_call", "Every faulty part sent belongs to a logged call", "Outward lines whose SR ID is not in the call tracker",
        _n("SELECT count(*) n FROM spare_outward s WHERE s.is_current=1 AND NOT EXISTS (SELECT 1 FROM svc_call c WHERE c.sr_id=s.sr_id AND c.is_current=1)"), total_out, "warn")
    add("outward_asset", "Every faulty part sent points to a known asset", "Outward lines whose asset (CI) is not in the asset register",
        _n("SELECT count(*) n FROM spare_outward s WHERE s.is_current=1 AND (s.asset_key IS NULL OR NOT EXISTS (SELECT 1 FROM asset a WHERE a.asset_key=s.asset_key AND a.is_current=1))"), total_out)
    add("components", "Every component has its parent asset", "Components whose parent asset is missing",
        _n("SELECT count(*) n FROM asset c WHERE c.is_current=1 AND c.parent_asset_key IS NOT NULL AND NOT EXISTS (SELECT 1 FROM asset p WHERE p.asset_key=c.parent_asset_key AND p.is_current=1)"),
        _n("SELECT count(*) n FROM asset WHERE is_current=1 AND parent_asset_key IS NOT NULL"))
    total_rma = _n("SELECT count(*) n FROM oem_rma WHERE is_current=1")
    add("rma_asset", "OEM RMA cases are linked to an asset", "RMA lines with no linked asset",
        _n("SELECT count(*) n FROM oem_rma WHERE is_current=1 AND asset_key IS NULL"), total_rma, "warn",
        "Early cases were logged without a device serial; link them when the information is found.")
    total_assets = _n("SELECT count(*) n FROM asset WHERE is_current=1 AND record_level='ASSET'")
    add("asset_users", "Asset users exist in the HR master", "Assets assigned to a CPF that is not active in the HR master",
        _n("SELECT count(*) n FROM asset WHERE is_current=1 AND record_level='ASSET' AND user_hr_status <> 'ACTIVE'"), total_assets, "warn",
        "Retired / transferred users, or a wrong CPF on the asset.")
    unknown = _n("""SELECT count(*) n FROM (SELECT engineer_name AS k FROM asset WHERE is_current=1 UNION SELECT engineer FROM svc_call WHERE is_current=1 UNION SELECT pm_done_by FROM asset WHERE is_current=1) x
                    WHERE k IS NOT NULL AND NOT EXISTS (SELECT 1 FROM portal_engineer e WHERE e.engineer_key = x.k)""")
    add("engineer_known", "Every engineer name is registered", "Engineer names in assets / calls that are not in the engineer register", unknown,
        _n("SELECT count(*) n FROM portal_engineer"))
    add("engineer_roster", "Engineers are on the CIPL roster", "Engineers used in the data who are not on the CIPL roster",
        _n("SELECT count(*) n FROM portal_engineer e WHERE e.source = 'DATA_ONLY'"), _n("SELECT count(*) n FROM portal_engineer"), "warn",
        "Names such as visiting or ONGC staff appear in the call tracker but not on the contractor roster.")
    add("call_dates", "Call dates are consistent", "Closed calls without a closed date, or open calls with one",
        _n("SELECT count(*) n FROM svc_call WHERE is_current=1 AND ((call_status='CLOSED' AND closed_date IS NULL) OR (call_status='OPEN' AND closed_date IS NOT NULL))"), total_calls, "warn")
    IPV4 = r"^(\d{1,3}\.){3}\d{1,3}$"
    asset_check("ip_format", "Each IP address field holds one valid address", "Assets whose IP field is not a single IPv4 address (several addresses, text, or a lone dot)",
                f"a.ip_address IS NOT NULL AND a.ip_address !~ '{IPV4}'", "Several addresses were typed into one field. Keep the main one here and put the rest in Remarks until a multi-address field exists.")
    asset_check("ip_dup", "No two assets share an IP address", "Assets whose IP address is also used by another asset",
                "a.ip_address IS NOT NULL AND EXISTS (SELECT 1 FROM asset b WHERE b.is_current=1 AND b.record_level='ASSET' AND b.ip_address=a.ip_address AND b.asset_key<>a.asset_key)",
                "Either a typing error or an address conflict on the network.")
    asset_check("host_dup", "Hostnames are unique", "Assets sharing a hostname with another asset",
                "a.hostname IS NOT NULL AND EXISTS (SELECT 1 FROM asset b WHERE b.is_current=1 AND b.record_level='ASSET' AND b.hostname=a.hostname AND b.asset_key<>a.asset_key)",
                "Labels such as a circuit name or 'STAND' were used as a hostname for several devices.")
    asset_check("install_date", "Installation date is recorded", "Assets with no 'Installed on' date",
                "a.install_date IS NULL", "Without it the age of the fleet, refresh planning and warranty checks are guesswork. Fill it in as each asset is touched - a new asset now defaults to today.")
    asset_check("cover_expired_use", "Assets in use are under cover", "Assets in use whose AMC / warranty has ended",
                "a.asset_status='IN_USE' AND a.cover_status='EXPIRED'", "Renew the cover, or record the real status of the asset.")
    asset_check("amc_removed_rate", "Assets removed from AMC carry no rate", "Assets marked removed from AMC that still have a rate",
                "a.asset_status='REMOVED_FROM_AMC' AND a.rate_value IS NOT NULL", "If the rate is still being billed, the contract has not caught up with the removal.")
    asset_check("amc_no_rate", "Assets under AMC have a rate", "Assets covered by AMC with no rate recorded",
                "a.cover_type='AMC' AND a.rate_value IS NULL AND a.asset_status<>'REMOVED_FROM_AMC'")
    asset_check("user_retired", "Assets are not held by retired users", "In-use assets whose user's retirement date has passed",
                "a.asset_status='IN_USE' AND a.user_retirement_date < current_date", "Reassign the asset or update the user.")
    asset_check("user_retiring", "Users retiring within a year", "In-use assets whose user retires within 12 months",
                "a.asset_status='IN_USE' AND a.user_retirement_date BETWEEN current_date AND current_date + 365", "Plan the hand-back or reassignment before the last working day.")
    asset_check("not_verified", "Assets have been physically checked in the last year", "In-use assets never checked, or last checked over a year ago",
                "a.asset_status='IN_USE' AND COALESCE((SELECT max(v.verified_on) FROM asset_verification v WHERE v.asset_key=a.asset_key), DATE '1900-01-01') < current_date - 365",
                "Open an asset and use Verify to record that it was found where the register says. Engineers can verify their own assets.")
    add("component_parent", "Every component is linked to a device", "Components with no parent device recorded",
        _n("SELECT count(*) n FROM asset WHERE is_current=1 AND record_level='COMPONENT' AND parent_asset_key IS NULL"),
        _n("SELECT count(*) n FROM asset WHERE is_current=1 AND record_level='COMPONENT'"), "warn",
        "SFPs and modules were loaded with a generated key. Link each to the switch or router it sits in, so it can be found when the device moves.",
        "SELECT asset_key k FROM asset WHERE is_current=1 AND record_level='COMPONENT' AND parent_asset_key IS NULL ORDER BY 1 LIMIT 8")
    add("key_unique", "Keys are unique", "Duplicate asset keys among current rows",
        _n("SELECT count(*) n FROM (SELECT asset_key FROM asset WHERE is_current=1 GROUP BY 1 HAVING count(*)>1) x"), total_assets)
    return items


def freshness():
    rows = db.query("""
        SELECT 'Assets' AS name, max(snapshot_date) AS as_of, count(*) AS n FROM asset WHERE is_current=1
        UNION ALL SELECT 'Calls', max(snapshot_date), count(*) FROM svc_call WHERE is_current=1
        UNION ALL SELECT 'Inward', max(snapshot_date), count(*) FROM spare_inward WHERE is_current=1
        UNION ALL SELECT 'Outward', max(snapshot_date), count(*) FROM spare_outward WHERE is_current=1
        UNION ALL SELECT 'OEM RMA', max(snapshot_date), count(*) FROM oem_rma WHERE is_current=1
        UNION ALL SELECT 'HR employees', max(load_date), count(*) FROM employee_history WHERE load_date = (SELECT max(load_date) FROM employee_history)
        UNION ALL SELECT 'CIPL roster', max(snapshot_date), count(*) FROM cipl_employee""")
    today = dt.date.today()
    for r in rows:
        r["age_days"] = (today - r["as_of"]).days if r["as_of"] else None
    return rows


# ---------------------------------------------------------------- masters/ files vs database
_FILE_SPECS = [
    (r"^IT_Asset_Master_.*\.xlsx$", [("ASSET_MASTER", "Assets", "SELECT count(*) n FROM asset WHERE is_current=1")]),
    (r"^Call_Tracker_.*\.xlsx$", [("CALLS", "Calls", "SELECT count(*) n FROM svc_call WHERE is_current=1"),
                                  ("SPARE_INWARD", "Inward", "SELECT count(*) n FROM spare_inward WHERE is_current=1"),
                                  ("SPARE_OUTWARD", "Outward", "SELECT count(*) n FROM spare_outward WHERE is_current=1")]),
    (r"^OEM_RMA_.*\.xlsx$", [("OEM_RMA", "OEM RMA lines", "SELECT count(*) n FROM oem_rma WHERE is_current=1")]),
    (r"^Employee_Master_Upload_\d.*\.xlsx$", [("EMPLOYEE_MASTER", "HR employees", "SELECT count(*) n FROM employee")]),
    (r"^CIPL_Employee_Master_\d.*\.xlsx$", [("EMPLOYEE_MASTER", "CIPL roster", "SELECT count(*) n FROM cipl_employee")]),
]


def _count_rows(path, sheet):
    key = (str(path), path.stat().st_mtime, sheet)
    if key in _file_cache:
        return _file_cache[key]
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet]
        n = sum(1 for r in ws.iter_rows(min_row=2, max_col=3, values_only=True) if any(v is not None for v in r))
    finally:
        wb.close()
    _file_cache[key] = n
    return n


def masters():
    out = []
    d = config.MASTERS_DIR
    if not d.is_dir():
        return out
    for f in sorted(d.glob("*.xlsx")):
        if f.name.startswith("~$"):
            continue
        for pattern, sheets in _FILE_SPECS:
            if re.match(pattern, f.name):
                for sheet, label, sql in sheets:
                    try:
                        in_file = _count_rows(f, sheet)
                        in_db = _n(sql)
                        status = "pass" if in_file == in_db else "fail"
                    except Exception as e:  # unreadable / locked file must not break the page
                        in_file, in_db, status = None, None, "warn"
                        label += f" (unreadable: {type(e).__name__})"
                    out.append({"file": f.name, "sheet": sheet, "label": label, "in_file": in_file, "in_db": in_db, "status": status,
                                "modified": dt.datetime.fromtimestamp(f.stat().st_mtime).isoformat(timespec="seconds"), "size": f.stat().st_size})
    return out


def report():
    c = checks()
    m = masters()
    return {"checks": c, "freshness": freshness(), "masters": m,
            "summary": {"pass": sum(1 for x in c if x["status"] == "pass"), "warn": sum(1 for x in c if x["status"] == "warn"), "fail": sum(1 for x in c if x["status"] == "fail"),
                        "files_ok": sum(1 for x in m if x["status"] == "pass"), "files_total": len(m)}}
