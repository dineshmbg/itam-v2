"""File output: Excel, CSV and PDF from the same table description.

A table is {"title", "columns": [(key, label)], "rows": [dict]}. A dashboard pack adds KPI figures and bar charts.
"""
import csv
import datetime as dt
import io
import re
from decimal import Decimal
from pathlib import Path

from . import config

LOGO = config.PROJECT_ROOT / "ongc-logo.png"
PDF_ROW_LIMIT = 3000


def _plain(v):
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, dt.datetime):
        return v.replace(tzinfo=None) if v.tzinfo else v
    if isinstance(v, (list, dict)):
        return str(v)
    return v


# Stored status codes -> the wording shown on screen (frontend/src/js/ui/badge.js), applied only to these columns when a file is
# written. Stored values are untouched, so an upload with the old codes still loads.
STATUS_LABELS = {
    "asset_status": {"IN_USE": "Deployed", "IN_STORE": "In stock", "NOT_ON_NETWORK": "Offline", "REMOVED_FROM_AMC": "Out of AMC"},
    "cover_status": {"EXPIRING_90D": "Renewal due", "EXPIRED": "Lapsed"},
    "pm_status": {"PENDING": "Scheduled", "DONE": "Completed", "DONE_OUTSIDE_QUARTER": "Completed late"},
    "call_status": {"OPEN": "Raised", "CLOSED": "Resolved"},
}


def shown(key, v):
    m = STATUS_LABELS.get(key)
    return m.get(v, v) if m else v


def safe_name(s):
    return re.sub(r"[^A-Za-z0-9._-]+", "_", s).strip("_")[:80] or "report"


# ---------------------------------------------------------------- CSV
def csv_bytes(table):
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow([label for _, label in table["columns"]])
    for r in table["rows"]:
        w.writerow(["" if r.get(k) is None else (r[k].isoformat() if isinstance(r[k], (dt.date, dt.datetime)) else shown(k, r[k])) for k, _ in table["columns"]])
    return ("﻿" + buf.getvalue()).encode("utf-8")       # BOM so Excel opens the file as UTF-8


# ---------------------------------------------------------------- asset barcode labels
def asset_barcode_svg(value):
    """Compact Code128 barcode (bars only, no text - the caller already shows the value as text) for the asset hover card."""
    from reportlab.graphics import renderSVG
    from reportlab.graphics.barcode import createBarcodeDrawing
    d = createBarcodeDrawing("Code128", value=str(value), barHeight=26, humanReadable=False, quiet=True)
    return renderSVG.drawToString(d)


def asset_labels_pdf(rows):
    """One Code128 barcode sticker per asset - the same Asset (CI) number already on the record, plus the class/type, make/model and
    serial in plain text so a technician can read it without scanning. Rows: export_table("assets", ...)'s rows (already scoped: a
    non-admin only gets the assets assigned to them, same as the CSV export). Sheet: 4 columns x 12 rows of ~45x22mm labels on A4,
    with dashed cut lines - close to a standard sheet of adhesive asset tags."""
    from reportlab.graphics import renderPDF
    from reportlab.graphics.barcode import createBarcodeDrawing
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas as pdfcanvas

    page_w, page_h = A4
    cols, rows_per_page = 4, 12
    lw, lh = 45 * mm, 22 * mm
    margin_x, margin_y = (page_w - cols * lw) / 2, (page_h - rows_per_page * lh) / 2

    buf = io.BytesIO()
    c = pdfcanvas.Canvas(buf, pagesize=A4)
    per_page = cols * rows_per_page
    for i, r in enumerate(rows):
        pos = i % per_page
        if i and pos == 0:
            c.showPage()
        col, row = pos % cols, pos // cols
        x0, y0 = margin_x + col * lw, page_h - margin_y - (row + 1) * lh

        c.setDash(1, 2)
        c.setStrokeColor(colors.HexColor("#b0b0b0"))
        c.rect(x0, y0, lw, lh)
        c.setDash()

        key = str(r.get("asset_key") or "")
        if not key:
            continue
        bc = createBarcodeDrawing("Code128", value=key, barHeight=8 * mm, humanReadable=False, quiet=True)
        scale = min(1.0, (lw - 6 * mm) / bc.width) if bc.width else 1.0
        bw, bh = bc.width * scale, bc.height * scale
        renderPDF.draw(bc, c, x0 + (lw - bw) / 2, y0 + lh - bh - 3 * mm, showBoundary=False)

        c.setFillColor(colors.black)
        c.setFont("Courier-Bold", 8)
        c.drawCentredString(x0 + lw / 2, y0 + lh - bh - 8 * mm, key[:26])
        c.setFont("Helvetica", 5.6)
        line2 = " / ".join(filter(None, [r.get("asset_class"), r.get("asset_type")]))
        c.drawCentredString(x0 + lw / 2, y0 + lh - bh - 12.2 * mm, line2[:46])
        line3 = " · ".join(filter(None, [r.get("make_model"), f"SN {r['serial_no']}" if r.get("serial_no") else None]))
        c.drawCentredString(x0 + lw / 2, y0 + lh - bh - 15.8 * mm, line3[:50])
    c.showPage()
    c.save()
    return buf.getvalue()


# ---------------------------------------------------------------- Excel
def xlsx_bytes(tables, title=None, kpis=None, meta=None, generated_by=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="1F3864")

    # Every workbook gets a SUMMARY sheet (even with no KPIs) so the generation stamp always has a home, not just dashboard exports.
    # Format is fixed: "username - DD Mon YYYY hh:mm AM/PM" - same stamp on screen (this sheet) and when printed (the footer below).
    stamp = dt.datetime.now().strftime("%d %b %Y %I:%M %p")
    generated = f"{generated_by} - {stamp}" if generated_by else stamp
    ws = wb.create_sheet("SUMMARY")
    ws["A1"] = title or "REPORT"
    ws["A1"].font = Font(bold=True, size=14)
    row = 2
    if meta:
        ws.cell(row, 1, meta); row += 1
    ws.cell(row, 1, generated).font = Font(italic=True, color="5F6B7A")
    if kpis:
        for i, (k, v) in enumerate(kpis, start=row + 2):
            ws.cell(i, 1, k).font = Font(bold=True)
            ws.cell(i, 2, _plain(v))
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 18

    def _print_footer(sheet):
        # Native Excel fields: &P/&N recompute to the real page count whenever it's actually printed. Excel has no field code to
        # hide it conditionally on a single page (unlike the PDF path below, which we render ourselves and can suppress outright).
        sheet.oddFooter.left.text = generated
        sheet.oddFooter.center.text = "Page &P of &N"
    _print_footer(ws)

    used = {"SUMMARY"}
    for t in tables:
        name = re.sub(r"[\\/*?:\[\]]", " ", (t.get("title") or "DATA").upper())[:31] or "DATA"
        base, n = name, 2
        while name in used:
            name = f"{base[:28]} {n}"; n += 1
        used.add(name)
        ws = wb.create_sheet(name)
        cols = t["columns"]
        ws.append([label for _, label in cols])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = head_fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
        widths = [len(str(label)) for _, label in cols]
        for r in t["rows"]:
            row = [_plain(shown(k, r.get(k))) for k, _ in cols]
            ws.append(row)
            for i, v in enumerate(row):
                if v is not None and len(str(v)) > widths[i]:
                    widths[i] = min(60, len(str(v)))
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(9, min(60, w + 2))
            for cell in ws.iter_cols(min_col=i, max_col=i, min_row=2, max_row=min(ws.max_row, 20000)):
                for c in cell:
                    if isinstance(c.value, dt.datetime):
                        c.number_format = "DD-MMM-YYYY HH:MM"
                    elif isinstance(c.value, dt.date):
                        c.number_format = "DD-MMM-YYYY"
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        _print_footer(ws)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ---------------------------------------------------------------- PDF
def _txt(v):
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%d %b %Y").upper()
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, (int, float, Decimal)) and not isinstance(v, bool):
        return f"{v:,}" if isinstance(v, int) else f"{float(v):,.1f}"
    return str(v).encode("cp1252", "replace").decode("cp1252")


def pdf_bytes(title, subtitle="", kpis=None, tables=None, charts=None, footer=""):
    from reportlab.graphics.shapes import Drawing, Rect, String
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas as pdfcanvas
    from reportlab.platypus import Image, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    W, H = landscape(A4)
    navy, grey, light = colors.HexColor("#1F3864"), colors.HexColor("#5F6B7A"), colors.HexColor("#F1F3F5")
    st_title = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=16, leading=20, textColor=navy)
    st_sub = ParagraphStyle("s", fontName="Helvetica", fontSize=9, leading=12, textColor=grey)
    st_h = ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=navy, spaceBefore=10, spaceAfter=4)
    st_cell = ParagraphStyle("c", fontName="Helvetica", fontSize=7.5, leading=9, alignment=TA_LEFT)
    st_head = ParagraphStyle("hd", fontName="Helvetica-Bold", fontSize=7.5, leading=9, textColor=colors.white)
    story = [Paragraph(_txt(title).upper(), st_title), Paragraph(_txt(subtitle), st_sub), Spacer(1, 6)]

    if kpis:
        cells = [[Paragraph(f'<font size="7" color="#5F6B7A">{_txt(k).upper()}</font>', st_cell), Paragraph(f'<font size="14"><b>{_txt(v)}</b></font>', st_cell)] for k, v in kpis]
        per_row = 6
        rows = []
        for i in range(0, len(cells), per_row):
            chunk = cells[i:i + per_row]
            chunk += [[Paragraph("", st_cell), Paragraph("", st_cell)]] * (per_row - len(chunk))
            rows.append([Table([[c[0]], [c[1]]], colWidths=[(W - 30 * mm) / per_row - 4]) for c in chunk])
        t = Table(rows, colWidths=[(W - 30 * mm) / per_row] * per_row)
        t.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDE1E6")), ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDE1E6")),
                               ("BACKGROUND", (0, 0), (-1, -1), light), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        story += [t, Spacer(1, 6)]

    for ch in charts or []:
        rows = ch["rows"][:12]
        if not rows:
            continue
        mx = max(float(r["n"] or 0) for r in rows) or 1
        d = Drawing(W - 30 * mm, 12 * len(rows) + 6)
        for i, r in enumerate(rows):
            y = 12 * (len(rows) - 1 - i) + 2
            d.add(String(0, y + 2, _txt(r["label"])[:34].upper(), fontName="Helvetica", fontSize=7))
            d.add(Rect(190, y, max(1, 330 * float(r["n"] or 0) / mx), 8, fillColor=colors.HexColor("#2E75B6"), strokeColor=None))
            d.add(String(190 + max(1, 330 * float(r["n"] or 0) / mx) + 4, y + 2, _txt(r["n"]), fontName="Helvetica", fontSize=7))
        story.append(KeepTogether([Paragraph(_txt(ch["title"]).upper(), st_h), d]))

    for t in tables or []:
        cols = t["columns"]
        rows = t["rows"][:PDF_ROW_LIMIT]
        story.append(Paragraph(_txt(t.get("title") or "").upper(), st_h))
        if not rows:
            story.append(Paragraph("No records.", st_sub))
            continue
        weights = []
        for k, label in cols:
            m = max([len(_txt(label))] + [len(_txt(shown(k, r.get(k)))) for r in rows[:200]])
            weights.append(min(40, max(6, m)))
        total = sum(weights)
        avail = W - 30 * mm
        widths = [avail * w / total for w in weights]
        data = [[Paragraph(_txt(label).upper(), st_head) for _, label in cols]]
        for r in rows:
            data.append([Paragraph(_txt(shown(k, r.get(k))).replace("&", "&amp;").replace("<", "&lt;"), st_cell) for k, _ in cols])
        tb = Table(data, colWidths=widths, repeatRows=1)
        tb.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), navy), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, light]),
                                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#DDE1E6")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3), ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
        story.append(tb)
        if len(t["rows"]) > PDF_ROW_LIMIT:
            story.append(Paragraph(f"Showing the first {PDF_ROW_LIMIT:,} of {len(t['rows']):,} rows. Download Excel or CSV for the complete data.", st_sub))

    # Computed here, not by the caller, so every PDF report carries it regardless of whether a caller remembers to ask. Fixed format:
    # "username - DD Mon YYYY hh:mm AM/PM", printed flush left at the bottom of every page - same stamp `xlsx_bytes` puts on Excel.
    stamp = dt.datetime.now().strftime("%d %b %Y %I:%M %p")
    generated = f"{footer} - {stamp}" if footer else stamp

    def deco(canvas, doc):
        canvas.saveState()
        if LOGO.exists():
            canvas.drawImage(str(LOGO), 15 * mm, H - 16 * mm, height=11 * mm, width=11 * mm * 1.0, preserveAspectRatio=True, mask="auto")
        canvas.setFont("Helvetica-Bold", 8)
        canvas.setFillColor(navy)
        canvas.drawString(30 * mm, H - 11 * mm, "ITAM PORTAL - ANKLESHWAR ASSET")
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(grey)
        canvas.drawString(15 * mm, 8 * mm, _txt(generated))
        canvas.restoreState()

    class _NumberedCanvas(pdfcanvas.Canvas):
        """Defers the page number until save(), once the true page count is known - draws it at all only when there is more than
        one page, per the user's request. (SimpleDocTemplate's normal single-pass build only knows the page it's currently on.)"""
        def __init__(self, *a, **kw):
            pdfcanvas.Canvas.__init__(self, *a, **kw)
            self._saved_page_states = []

        def showPage(self):
            self._saved_page_states.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._saved_page_states)
            for state in self._saved_page_states:
                self.__dict__.update(state)
                if total > 1:
                    self.saveState()
                    self.setFont("Helvetica", 7)
                    self.setFillColor(grey)
                    self.drawCentredString(W / 2, 8 * mm, f"PAGE {self._pageNumber} OF {total}")
                    self.restoreState()
                pdfcanvas.Canvas.showPage(self)
            pdfcanvas.Canvas.save(self)

    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=15 * mm, rightMargin=15 * mm, topMargin=20 * mm, bottomMargin=14 * mm, title=_txt(title), author="ITAM Portal")
    doc.build(story, onFirstPage=deco, onLaterPages=deco, canvasmaker=_NumberedCanvas)
    return out.getvalue()


MIME = {"xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "csv": "text/csv; charset=utf-8", "pdf": "application/pdf"}


def render(fmt, title, subtitle, tables, kpis=None, charts=None, footer=""):
    """-> (bytes, mime, extension). One table for CSV; everything for Excel and PDF."""
    if fmt == "csv":
        return csv_bytes(tables[0]), MIME["csv"], "csv"
    if fmt == "xlsx":
        return xlsx_bytes(tables, title, kpis, subtitle, generated_by=footer), MIME["xlsx"], "xlsx"
    if fmt == "pdf":
        return pdf_bytes(title, subtitle, kpis, tables, charts, footer), MIME["pdf"], "pdf"
    raise ValueError("format must be xlsx, csv or pdf")


__all__ = ["render", "safe_name", "MIME", "Path"]
