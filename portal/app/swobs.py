"""Software / OS observations per asset: what a report, the check script or an engineer saw on a machine, kept as an append-only history.
Sources never overwrite each other (conflicts are shown side by side for the engineer to judge). The inspection date is stamped by the database
at insert - no caller can supply it - and a trigger refuses UPDATE and DELETE, so history cannot be rewritten."""
import datetime as dt
import re

from . import db

SOURCES = ("BIGFIX_REPORT", "TREND_REPORT", "SCRIPT", "MANUAL")
STALE_DAYS, VERY_STALE_DAYS = 30, 90
FIELDS = ("os_edition", "os_build", "os_feature_version", "trend_installed", "trend_version", "trend_status", "bigfix_installed", "bigfix_version",
          "vpn_name", "vpn_version", "vpn_other", "sap_gui_version", "o365_account", "o365_version", "last_user", "connection",
          "bitlocker", "tpm_secureboot", "office_channel", "domain_joined", "last_boot", "browser_version", "local_admins", "remarks",
          # added with the richer check script: service states, how each product was found, the Office sign-in's origin, last Windows / Defender updates
          "os_product_raw", "trend_services", "trend_candidate", "vpn_service_state", "vpn_candidate", "bigfix_service_state", "o365_status", "o365_account_source",
          "os_update_title", "os_update_at", "defender_signature", "defender_signature_at", "chrome_version", "firefox_version",
          # ZENworks agent (check script only)
          "zenworks_installed", "zenworks_version", "zenworks_status", "zenworks_services")
BOOLS = {"trend_installed", "bigfix_installed", "zenworks_installed", "vpn_other", "bitlocker", "tpm_secureboot", "domain_joined"}
INTERNAL = {"apps_count": "INTEGER", "apps_hash": "TEXT", "apps_obs_id": "BIGINT"}      # set by the check-script intake only, never by a person

DDL = [
    f"""CREATE TABLE IF NOT EXISTS portal_sw_obs (
          obs_id BIGSERIAL PRIMARY KEY, asset_key TEXT NOT NULL, source TEXT NOT NULL,
          inspected_at TIMESTAMPTZ NOT NULL DEFAULT now(), data_as_of TIMESTAMPTZ, recorded_by TEXT NOT NULL,
          {', '.join(f + (' BOOLEAN' if f in BOOLS else ' TEXT') for f in FIELDS)})""",
    *[f"ALTER TABLE portal_sw_obs ADD COLUMN IF NOT EXISTS {f} {'BOOLEAN' if f in BOOLS else 'TEXT'}" for f in FIELDS],       # an existing table gains the newer columns
    *[f"ALTER TABLE portal_sw_obs ADD COLUMN IF NOT EXISTS {c} {t}" for c, t in INTERNAL.items()],
    "CREATE INDEX IF NOT EXISTS ix_sw_obs_asset ON portal_sw_obs (asset_key, source, inspected_at DESC)",
    # the installed-software list a script run saw (one row per program), and every Windows update event seen on a machine (never repeated)
    """CREATE TABLE IF NOT EXISTS portal_sw_obs_app (
          app_id BIGSERIAL PRIMARY KEY, obs_id BIGINT NOT NULL REFERENCES portal_sw_obs (obs_id), name TEXT NOT NULL, version TEXT, publisher TEXT, hive TEXT NOT NULL)""",
    "CREATE INDEX IF NOT EXISTS ix_sw_obs_app_obs ON portal_sw_obs_app (obs_id)",
    """CREATE TABLE IF NOT EXISTS portal_sw_update (
          asset_key TEXT NOT NULL, installed_at TIMESTAMPTZ NOT NULL, title TEXT NOT NULL, kb TEXT, first_obs_id BIGINT, PRIMARY KEY (asset_key, installed_at, title))""",
    """CREATE OR REPLACE FUNCTION portal_sw_obs_append_only() RETURNS trigger LANGUAGE plpgsql AS $$
       BEGIN RAISE EXCEPTION 'portal_sw_obs is append-only'; END $$""",
    "DROP TRIGGER IF EXISTS portal_sw_obs_no_change ON portal_sw_obs",
    "CREATE TRIGGER portal_sw_obs_no_change BEFORE UPDATE OR DELETE ON portal_sw_obs FOR EACH ROW EXECUTE FUNCTION portal_sw_obs_append_only()",
    "DROP TRIGGER IF EXISTS portal_sw_obs_app_no_change ON portal_sw_obs_app",
    "CREATE TRIGGER portal_sw_obs_app_no_change BEFORE UPDATE OR DELETE ON portal_sw_obs_app FOR EACH ROW EXECUTE FUNCTION portal_sw_obs_append_only()",
    "DROP TRIGGER IF EXISTS portal_sw_update_no_change ON portal_sw_update",
    "CREATE TRIGGER portal_sw_update_no_change BEFORE UPDATE OR DELETE ON portal_sw_update FOR EACH ROW EXECUTE FUNCTION portal_sw_obs_append_only()",
]


class ObsError(Exception):
    http_error = True

    def __init__(self, message, status=400):
        super().__init__(message)
        self.status = status


def record(asset_key, source, facts, by, data_as_of=None, internal=None, con=None):
    """Add one observation and return its id. `inspected_at` is deliberately not a parameter: the database stamps it.
    `internal` (apps_*) is for the check-script intake only; `con` lets that intake write the observation and its children in one transaction."""
    if source not in SOURCES:
        raise ObsError(f"Unknown source {source!r}.")
    bad = sorted(set(facts) - set(FIELDS))
    if bad:
        raise ObsError(f"Unknown field(s): {', '.join(bad)}.")
    facts = {**facts, **(internal or {})}
    cols = list(facts)
    sql = f"INSERT INTO portal_sw_obs (asset_key, source, recorded_by, data_as_of{''.join(', ' + c for c in cols)}) VALUES (%s,%s,%s,%s{', %s' * len(cols)}) RETURNING obs_id"
    args = [asset_key, source, by, data_as_of] + [facts[c] for c in cols]
    if con is not None:
        row = con.execute(sql, args).fetchone()
    else:
        with db.write() as c:
            row = c.execute(sql, args).fetchone()
    return row[0] if not isinstance(row, dict) else row["obs_id"]


def _stale(r):
    age = (dt.datetime.now(dt.timezone.utc) - (r["data_as_of"] or r["inspected_at"])).days
    return "VERY_STALE" if age > VERY_STALE_DAYS else "STALE" if age > STALE_DAYS else None


def latest_by_source(asset_key):
    """The newest observation from each source, side by side; `stale` is STALE / VERY_STALE by the report's own date when it has one."""
    rows = db.query("SELECT DISTINCT ON (source) * FROM portal_sw_obs WHERE asset_key = %s ORDER BY source, obs_id DESC", [asset_key])
    return [{**r, "stale": _stale(r)} for r in rows]


# ---------------------------------------------------------------- reports (BigFix, Trend Micro) -> observations
SCOPE_CLASSES = ["DESKTOP", "WORKSTATION", "SERVER"]
CI_PREFIX = "ANK"      # every CI of this site starts with ANK; anything else in a report belongs to another site and is discarded
_BIGFIX_OS = re.compile(r"^\s*Win(?P<ed>\d+)\s+(?P<build>\d+\.\d+\.\d+\.\d+)(?:\s*\((?P<feat>[^)]+)\))?")


def _edition(code):
    return f"Windows {code}" if code in ("10", "11") else f"Windows Server {code}"


def _bigfix(g):
    m = _BIGFIX_OS.match(g("os"))
    if not m:
        return {}
    return {"os_edition": _edition(m["ed"]), "os_build": m["build"], "os_feature_version": m["feat"]}


def _trend(g):
    build = re.search(r"Build\s+(\d+)", g("osver"))
    return {"trend_installed": True, "trend_status": g("agent") or None, "connection": g("conn") or None, "last_user": g("user") or None,
            "os_edition": g("os") or None, "os_build": build[1] if build else None, "trend_version": g("unified") or None}


def ingest_report(filename, data, by):
    """Read a BigFix or Trend Micro export (recognised from its columns, like the inventory match does) and record one observation for every row
    whose computer name is one of our desktop/workstation/server assets; names not starting with ANK are discarded unseen. Rows that match nothing are counted and named, not stored. A row whose
    (asset, source, report time) is already recorded is `unchanged`: re-importing a file adds nothing."""
    from . import inventory_match as im
    try:
        headers, rows = im.read_file(filename, data)
    except im.MatchError as e:
        raise ObsError(str(e)) from e
    prof = im.profile_file(headers, rows)
    product = prof["product"]
    if product not in ("bigfix", "trendmicro") or "name" not in prof["mapping"]:
        raise ObsError("That file is not a recognised BigFix or Trend Micro export.")
    have = {h.strip().lower(): i for i, h in enumerate(headers)}
    idx = {f: headers.index(h) for f, h in prof["mapping"].items()}
    idx.update({f: have[h.lower()] for f, h in im.PRODUCTS[product]["extra"].items() if h.lower() in have})
    if "unified agent version" in have:
        idx["unified"] = have["unified agent version"]
    g_row = lambda r: (lambda f: ("" if f not in idx or isinstance(r[idx[f]], (dt.datetime, dt.date)) else str(r[idx[f]]).strip()))      # noqa: E731
    source, build = ("BIGFIX_REPORT", _bigfix) if product == "bigfix" else ("TREND_REPORT", _trend)
    ours = {im.key_of(a["ci_no"]): a["asset_key"] for a in db.query("SELECT asset_key, ci_no FROM asset WHERE is_current = 1 AND record_level = 'ASSET' AND asset_class = ANY(%s)", [SCOPE_CLASSES])}
    latest, ignored = {}, 0
    for r in rows:
        key = im.key_of(r[idx["name"]])
        when = im._seen(r[idx["seen"]], product) if "seen" in idx else None
        if not key.startswith(CI_PREFIX):
            ignored += 1                                 # another site's machine: not ours, not even reported as unmatched
            continue
        if key and (key not in latest or (when or dt.datetime.min) >= (latest[key][1] or dt.datetime.min)):
            latest[key] = (r, when)                      # a machine listed twice: the newest record wins
    known = {(r["asset_key"], r["data_as_of"]) for r in db.query("SELECT asset_key, data_as_of FROM portal_sw_obs WHERE source = %s", [source])}
    out = {"product": product, "rows": len(rows), "matched": 0, "unmatched": 0, "unmatched_names": [], "recorded": 0, "unchanged": 0, "ignored_other_sites": 0}
    for key, (r, when) in latest.items():
        ak = ours.get(key)
        if not ak:
            out["unmatched"] += 1
            out["unmatched_names"].append(str(r[idx["name"]]).strip())
            continue
        out["matched"] += 1
        as_of = when.replace(tzinfo=dt.timezone.utc) if when else None
        if (ak, as_of) in known:
            out["unchanged"] += 1
            continue
        record(ak, source, {k: v for k, v in build(g_row(r)).items() if v is not None}, by, data_as_of=as_of)
        out["recorded"] += 1
    out["ignored_other_sites"] = ignored
    out["unmatched_names"] = out["unmatched_names"][:50]
    return out


# ---------------------------------------------------------------- the asset record's "Software" tab
def _build_tail(v):
    """BigFix says 10.0.19045.6456, Trend only 19045: compare on the Windows build number (third part) so the two can be set side by side."""
    parts = str(v or "").split(".")
    return parts[2] if len(parts) >= 3 else (parts[0] if parts and parts[0].isdigit() else None)


def panel(asset_key, asset_class):
    """What the record window shows: the newest observation per source (with its flags), conflicts between sources for the engineer to judge,
    and the full history. None for classes outside the tracked scope."""
    if asset_class not in SCOPE_CLASSES:
        return None
    latest = latest_by_source(asset_key)
    for r in latest:
        r["flags"] = [f for f in (r["stale"] and ("Report is more than %d days old" % (VERY_STALE_DAYS if r["stale"] == "VERY_STALE" else STALE_DAYS)),
                                  r["os_edition"] == "Windows 10" and "Windows 10 is past end of support") if f]
    conflicts = []
    for label, get in (("Windows edition", lambda r: r["os_edition"]), ("Windows build", lambda r: _build_tail(r["os_build"]))):
        seen = {r["source"]: get(r) for r in latest if get(r)}
        if len(set(seen.values())) > 1:
            conflicts.append({"field": label, "values": seen})
    history = db.query("SELECT * FROM portal_sw_obs WHERE asset_key = %s ORDER BY obs_id DESC LIMIT 100", [asset_key])
    last = db.one("SELECT COALESCE(apps_obs_id, obs_id) AS holder FROM portal_sw_obs WHERE asset_key = %s AND apps_hash IS NOT NULL ORDER BY obs_id DESC LIMIT 1", [asset_key])
    apps = db.query("SELECT name, version, publisher, hive FROM portal_sw_obs_app WHERE obs_id = %s ORDER BY lower(name), hive", [last["holder"]]) if last else []
    updates = db.query("SELECT installed_at, title, kb FROM portal_sw_update WHERE asset_key = %s ORDER BY installed_at DESC LIMIT 30", [asset_key])
    return {"latest": latest, "conflicts": conflicts, "history": history, "apps": apps, "updates": updates}


# ---------------------------------------------------------------- inventory match: report versus what was seen on the machine
SOURCE_LABELS = {"BIGFIX_REPORT": "BigFix report", "TREND_REPORT": "Trend Micro report", "SCRIPT": "Script", "MANUAL": "Entered by hand"}
_PRODUCT_FLAG = {"trendmicro": ("trend_installed", "Trend Micro"), "bigfix": ("bigfix_installed", "BigFix")}


def cross_check(rows, product):
    """Add `observed` / `xcheck` to every inventory-match row whose machine has a Script or hand-entered observation about this product, and say where
    it contradicts the report. Rows with no such observation are left exactly as they were. Returns how many rows contradict."""
    if product not in _PRODUCT_FLAG or not rows:
        return 0
    flag, label = _PRODUCT_FLAG[product]
    obs = db.query(f"""SELECT DISTINCT ON (asset_key, source) asset_key, source, inspected_at, {flag} AS state FROM portal_sw_obs
                       WHERE asset_key = ANY(%s) AND source IN ('SCRIPT', 'MANUAL') AND {flag} IS NOT NULL ORDER BY asset_key, source, obs_id DESC""", [[r["ci"] for r in rows]])
    by = {}
    for o in obs:
        by.setdefault(o["asset_key"], []).append(o)
    n = 0
    for r in rows:
        mine = by.get(r["ci"])
        if not mine:
            continue
        r["observed"] = "; ".join(f"{SOURCE_LABELS[o['source']]} {o['inspected_at']:%d %b %Y}: {label} {'installed' if o['state'] else 'not installed'}" for o in mine)
        bad = [o for o in mine if bool(o["state"]) != bool(r["installed"])]
        if bad:
            o = bad[0]
            r["xcheck"] = (f"The report lists this machine, but {SOURCE_LABELS[o['source']]} ({o['inspected_at']:%d %b %Y}) found {label} not installed." if r["installed"]
                           else f"The report does not list this machine, but {SOURCE_LABELS[o['source']]} ({o['inspected_at']:%d %b %Y}) found {label} installed.")
            n += 1
    return n


def suite_from_services(text):
    """What the check script saw on the PC -> 'Full suite', 'Sensor only' or '' (cannot tell).
    Per Trend Micro's documentation the Apex One Security Agent always runs TmListen (its link to the server), so TmListen says nothing about the sensor.
    The real-time scan service/process (ntrtscan / NTRtScan.exe) means the anti-malware agent is there: Full suite. The Vision One endpoint agent
    (EndpointBasecamp) without it carries only the Endpoint Sensor: Sensor only. Anything else is left undecided rather than guessed."""
    t = (text or "").lower()
    if "ntrtscan" in t:
        return "Full suite"
    if "endpointbasecamp" in t:
        return "Sensor only"
    return ""


def suite_seen(rows):
    """Add `suite_pc` to inventory-match rows: what the check script saw on the PC itself. The real-time scan service (ntrtscan) means the full agent;
    EndpointBasecamp without it means just the sensor (see suite_from_services). Rows the script never checked are left blank."""
    keys = [r["ci"] for r in rows]
    if not keys:
        return
    seen = {o["asset_key"]: o["trend_services"] for o in db.query("""SELECT DISTINCT ON (asset_key) asset_key, trend_services FROM portal_sw_obs
                  WHERE asset_key = ANY(%s) AND source = 'SCRIPT' AND trend_services IS NOT NULL ORDER BY asset_key, obs_id DESC""", [keys])}
    for r in rows:
        t = seen.get(r["ci"])
        if t:
            r["suite_pc"] = suite_from_services(t)


# ---------------------------------------------------------------- the check script (a .bat: reads the PC, installs nothing)
import hmac      # noqa: E402
import secrets   # noqa: E402
from pathlib import Path   # noqa: E402

MAX_CHECK_BYTES = 400_000      # the key lines are tiny; the room is for the update events and the installed-software list that follow them
_TEMPLATE = Path(__file__).with_name("ankcheck.bat.tmpl")
DEFAULT_DROP = "\\\\10.205.64.203\\ank-helpdesk\\common-for-all\\check-results"
_SERVER_NAME = re.compile(r"Windows Server\s+(\d{4}(?:\s+R2)?)", re.I)


_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")      # what a PC sends is not trusted: control characters would break the Excel report


_SECTION = re.compile(r"^\[(UPDATES|SOFTWARE)\]\s*$", re.M)
_SVC_WORD = {"1": "stopped", "2": "starting", "3": "stopping", "4": "running", "5": "continuing", "6": "pausing", "7": "paused"}
_DEFENDER = re.compile(r"Security Intelligence|Defender", re.I)
_OS_UPDATE = re.compile(r"Cumulative Update for Windows|Feature update to Windows|Security Update for Windows|Update for Windows \d", re.I)      # not .NET, MSRT or Defender


def _sections(text):
    """The script's output is KEY=VALUE lines, then optional [UPDATES] (wevtutil text) and [SOFTWARE] (reg query lines) sections."""
    parts = _SECTION.split(text.replace("\x00", ""))
    return parts[0], dict(zip(parts[1::2], parts[2::2]))


def parse_updates(text):
    """wevtutil /f:text events (ID 19, 'Installation Successful') -> [{at, title, kb}], newest first as given. Anything not recognisable is skipped."""
    out = []
    for block in re.split(r"^Event\[\d+\]\s*$", text, flags=re.M):
        d = re.search(r"^\s*Date:\s*(\d{4}-\d\d-\d\dT\d\d:\d\d:\d\d)", block, re.M)
        t = re.search(r"following update:\s*(.+)", block)
        if not (d and t):
            continue
        title = re.sub(r"\s+", " ", t[1]).strip()[:300]
        kb = re.search(r"KB\d{5,8}", title)
        out.append({"at": dt.datetime.fromisoformat(d[1]).replace(tzinfo=dt.timezone.utc), "title": title, "kb": kb[0] if kb else None})
    return out


def parse_software(text):
    """`reg query ... /s | findstr` lines -> sorted unique [(name, version, publisher, hive)]. A key without a DisplayName is not a program."""
    apps, cur, hive = set(), None, ""

    def close():
        if cur and cur.get("n"):
            apps.add((cur["n"][:200], (cur.get("v") or None) and cur["v"][:80], (cur.get("p") or None) and cur["p"][:200], cur["h"]))
    for line in text.splitlines():
        if line.startswith("HKEY_"):
            close()
            hive = "HKCU" if line.startswith("HKEY_CURRENT_USER") else "HKLM (32-bit)" if "WOW6432Node" in line.split("\\Uninstall")[0] else "HKLM"
            cur = {"h": hive}
            continue
        m = re.match(r"^\s*(DisplayName|DisplayVersion|Publisher)\s+REG_\w+\s+(.*)$", line)
        if m and cur is not None:
            cur[{"DisplayName": "n", "DisplayVersion": "v", "Publisher": "p"}[m[1]]] = m[2].strip()
    close()
    return sorted(apps, key=lambda a: (a[0].lower(), a[3], a[1] or "", a[2] or ""))[:2000]


def parse_check(text):
    """The script's text -> {computer, facts, updates, apps}. Strict: it must start as one of ours, name a computer, and stay small."""
    if len(text.encode("utf-8", "replace")) > MAX_CHECK_BYTES:
        raise ObsError("That is not a check-script result (too large).")
    head, sec = _sections(text)
    kv = {}
    for line in head.splitlines():
        k, sep, v = line.partition("=")
        if sep and re.fullmatch(r"[A-Z0-9_]{2,24}", k.strip()):
            kv[k.strip()] = _CTRL.sub("", v).strip()[:200]
    if kv.get("ANKCHECK") != "1":
        raise ObsError("That is not a check-script result.")
    computer = kv.get("COMPUTER", "").upper()
    if not re.fullmatch(r"[A-Z0-9._-]{1,63}", computer):
        raise ObsError("The result does not name a computer.")
    g = kv.get
    facts = {"trend_installed": bool(g("TREND_NAME")), "bigfix_installed": bool(g("BIGFIX_NAME")) or g("BIGFIX_SVC") == "1" or g("BIGFIX_SERVICE_STATE", "") != ""}
    for fld, key in (("trend_version", "TREND_VER"), ("bigfix_version", "BIGFIX_VER"), ("sap_gui_version", "SAP_VER"), ("o365_account", "O365_ACCOUNT"), ("o365_version", "OFFICE_VER"),
                     ("office_channel", "OFFICE_CHANNEL"), ("last_user", "LAST_USER"), ("browser_version", "EDGE_VER"), ("chrome_version", "CHROME_VER"), ("firefox_version", "FIREFOX_VER"), ("os_product_raw", "WIN_PRODUCT_RAW"),
                     ("o365_status", "O365_STATUS"), ("o365_account_source", "O365_ACCOUNT_SOURCE")):
        if g(key):
            facts[fld] = g(key)
    if g("TREND_NAME"):
        facts["trend_status"] = g("TREND_NAME")
    svc = [f"{n} {_SVC_WORD[g(k)]}" for n, k in (("TmListen", "TREND_LISTENER_STATE"), ("ntrtscan", "TREND_SCAN_STATE")) if g(k) in _SVC_WORD]
    if g("TREND_BASECAMP_STATE") in _SVC_WORD:
        svc.append(f"EndpointBasecamp {_SVC_WORD[g('TREND_BASECAMP_STATE')]}")
    if g("TREND_PROCS"):
        svc.append("processes " + g("TREND_PROCS").strip(", "))
    if g("TREND_SVC_SEEN"):
        svc.append("services seen " + g("TREND_SVC_SEEN").strip(", ")[:300])
    if svc:
        facts["trend_services"] = "; ".join(svc)
    if g("ZEN_STATUS"):
        facts["zenworks_installed"] = bool(g("ZEN_NAME"))
        if g("ZEN_NAME"):
            facts["zenworks_status"] = g("ZEN_NAME")
        if g("ZEN_VER"):
            facts["zenworks_version"] = g("ZEN_VER")
        zs = ([f"processes {g('ZEN_PROCS').strip(', ')}"] if g("ZEN_PROCS") else []) + ([f"services seen {g('ZEN_SVC_SEEN').strip(', ')[:300]}"] if g("ZEN_SVC_SEEN") else [])
        if zs:
            facts["zenworks_services"] = "; ".join(zs)
    if g("ZSCALER_SERVICE_STATE") in _SVC_WORD:
        facts["vpn_service_state"] = "ZSAService " + _SVC_WORD[g("ZSCALER_SERVICE_STATE")]
    if g("BIGFIX_SERVICE_STATE") in _SVC_WORD:
        facts["bigfix_service_state"] = "BESClient " + _SVC_WORD[g("BIGFIX_SERVICE_STATE")]
    for fld, name, ver in (("trend_candidate", "TREND_CANDIDATE_NAME", "TREND_CANDIDATE_VER"), ("vpn_candidate", "ZSCALER_CANDIDATE_NAME", "ZSCALER_CANDIDATE_VER")):
        if g(name):
            facts[fld] = f"{g(name)} {g(ver, '')}".strip()
    build = g("WIN_BUILD", "")
    if build.isdigit():
        server = g("WIN_INSTALLTYPE", "").lower() == "server" or "server" in g("WIN_PRODUCT", "").lower()
        m = _SERVER_NAME.search(g("WIN_PRODUCT", ""))
        facts["os_edition"] = (f"Windows Server {m[1]}" if m else "Windows Server") if server else ("Windows 11" if int(build) >= 22000 else "Windows 10")   # ProductName still says 10 on 11
        facts["os_build"] = f"10.0.{build}" + (f".{g('WIN_UBR')}" if g("WIN_UBR", "").isdigit() else "")
        if g("WIN_DISPLAY"):
            facts["os_feature_version"] = g("WIN_DISPLAY")
    other = g("VPN_OTHER", "").strip(" ;")
    if g("ZSCALER_NAME"):
        facts["vpn_name"], facts["vpn_version"] = "Zscaler", g("ZSCALER_VER") or None
    else:
        facts["vpn_name"] = f"{other.split(';')[0].strip()} (not Zscaler)" if other else "None found"
    facts["vpn_other"] = bool(other)
    if g("DOMAIN"):
        facts["domain_joined"] = True
    if g("SECUREBOOT") in ("0", "1"):
        facts["tpm_secureboot"] = g("SECUREBOOT") == "1"
    updates = parse_updates(sec.get("UPDATES", ""))
    os_up = next((u for u in updates if _OS_UPDATE.search(u["title"]) and not _DEFENDER.search(u["title"])), None)
    sig = next((u for u in updates if _DEFENDER.search(u["title"])), None)
    if os_up:
        facts["os_update_title"], facts["os_update_at"] = os_up["title"], os_up["at"].strftime("%Y-%m-%d %H:%M")
    if sig:
        v = re.search(r"Version\s+([\d.]+)", sig["title"])
        facts["defender_signature"], facts["defender_signature_at"] = (v[1] if v else sig["title"]), sig["at"].strftime("%Y-%m-%d %H:%M")
    notes = ([f"Other VPN: {other}"] if other else []) + ([f"ran {g('RUN_AT_RAW')} (PC clock)"] if g("RUN_AT_RAW") else [])
    if notes:
        facts["remarks"] = "; ".join(notes)
    return {"computer": computer, "facts": {k: v for k, v in facts.items() if v is not None}, "updates": updates, "apps": parse_software(sec.get("SOFTWARE", ""))}


def ingest_check(text, by, engineer_key=None):
    """Record one check-script result against the matching desktop/workstation/server asset (name starts with ANK; with `engineer_key`, only one assigned to that engineer): the observation, every Windows
    update event not seen before on that machine, and the installed-software list (stored again only when it changed since the last run)."""
    import hashlib
    c = parse_check(text)
    key = c["computer"].split(".")[0]
    asset = db.one("SELECT asset_key, engineer_name FROM asset WHERE is_current = 1 AND record_level = 'ASSET' AND asset_class = ANY(%s) AND upper(ci_no) = %s", [SCOPE_CLASSES, key]) if key.startswith(CI_PREFIX) else None
    if not asset:
        return {"computer": c["computer"], "matched": False, "recorded": 0, "reason": "Not a desktop, workstation or server in the register."}
    if engineer_key and asset["engineer_name"] != engineer_key:           # an engineer may only upload the PCs assigned to them
        return {"computer": c["computer"], "matched": False, "recorded": 0, "reason": "This PC is not assigned to you."}
    ak, apps = asset["asset_key"], c["apps"]
    internal, store_apps = {}, False
    if apps:
        h = hashlib.sha256(repr(apps).encode()).hexdigest()
        prev = db.one("SELECT obs_id, apps_hash, apps_obs_id FROM portal_sw_obs WHERE asset_key = %s AND apps_hash IS NOT NULL ORDER BY obs_id DESC LIMIT 1", [ak])
        same = bool(prev) and prev["apps_hash"] == h
        internal = {"apps_count": len(apps), "apps_hash": h, "apps_obs_id": (prev["apps_obs_id"] or prev["obs_id"]) if same else None}
        store_apps = not same
    with db.write() as con:
        oid = record(ak, "SCRIPT", c["facts"], by, internal=internal, con=con)
        for u in c["updates"]:
            con.execute("INSERT INTO portal_sw_update (asset_key, installed_at, title, kb, first_obs_id) VALUES (%s,%s,%s,%s,%s) ON CONFLICT DO NOTHING", (ak, u["at"], u["title"], u["kb"], oid))
        if store_apps:
            con.cursor().executemany("INSERT INTO portal_sw_obs_app (obs_id, name, version, publisher, hive) VALUES (%s,%s,%s,%s,%s)", [(oid, *a) for a in apps])
    return {"computer": c["computer"], "matched": True, "recorded": 1, "updates_seen": len(c["updates"]), "software": len(apps)}


def check_token(create=False, by="system", rotate=False):
    """The shared token baked into the script. Stored in portal_setting; it only lets a script add observations for our own machines."""
    row = db.one("SELECT value FROM portal_setting WHERE key = 'sw_check_token'")
    if row and not rotate:
        return row["value"]
    if not (create or rotate):
        return None
    tok = secrets.token_urlsafe(24)
    with db.write() as con:
        con.execute("INSERT INTO portal_setting (key, value, updated_by) VALUES ('sw_check_token', to_jsonb(%s::text), %s) ON CONFLICT (key) DO UPDATE SET value = EXCLUDED.value, updated_at = now(), updated_by = EXCLUDED.updated_by", (tok, by))
    return tok


def token_ok(given):
    real = check_token()
    return bool(real) and bool(given) and hmac.compare_digest(str(real), str(given))


def script_text(portal_url, token, drop=DEFAULT_DROP):
    """The .bat, ready to put on the share: this portal's address and the token are filled in, Windows line endings."""
    body = _TEMPLATE.read_text(encoding="utf-8").replace("__URL__", portal_url.rstrip("/")).replace("__TOKEN__", token).replace("__DROP__", drop)
    return body.replace("\r\n", "\n").replace("\n", "\r\n")



# ---------------------------------------------------------------- manual entry and the Excel bulk template
TEMPLATE_COLUMNS = [("Windows edition", "os_edition"), ("Windows build", "os_build"), ("Windows feature version", "os_feature_version"), ("Trend Micro installed (Yes/No)", "trend_installed"),
                    ("Trend Micro version", "trend_version"), ("BigFix installed (Yes/No)", "bigfix_installed"), ("BigFix version", "bigfix_version"), ("VPN", "vpn_name"), ("VPN version", "vpn_version"),
                    ("Other VPN present (Yes/No)", "vpn_other"), ("SAP GUI version", "sap_gui_version"), ("Microsoft 365 account", "o365_account"), ("Office version", "o365_version"),
                    ("Office update channel", "office_channel"), ("Last logged-on user", "last_user"), ("BitLocker on (Yes/No)", "bitlocker"), ("TPM / Secure Boot on (Yes/No)", "tpm_secureboot"),
                    ("Domain joined (Yes/No)", "domain_joined"), ("Last boot", "last_boot"), ("Edge version", "browser_version"), ("Chrome version", "chrome_version"), ("Firefox version", "firefox_version"), ("Local administrators", "local_admins"), ("Remarks", "remarks")]
_YES, _NO = {"yes", "y", "true", "1"}, {"no", "n", "false", "0"}


def template_headers():
    return ["Asset (CI)"] + [h for h, _ in TEMPLATE_COLUMNS]


def _clean_facts(raw):
    """Typed values for one manual observation: blanks are dropped, Yes/No become booleans, nothing unknown or oversized gets through."""
    facts = {}
    for k, v in (raw or {}).items():
        if k not in FIELDS:
            raise ObsError(f"Unknown field: {k}.")
        if v is None or (isinstance(v, str) and not v.strip()):
            continue
        if k in BOOLS:
            s = str(v).strip().lower()
            if s not in _YES | _NO:
                raise ObsError(f"{k}: answer Yes or No.")
            facts[k] = s in _YES
        else:
            s = str(v).strip()
            if len(s) > 200:
                raise ObsError(f"{k}: too long (200 characters at most).")
            facts[k] = s
    if not facts:
        raise ObsError("Enter at least one value.")
    return facts


def _in_scope_asset(key):
    return db.one("SELECT asset_key, engineer_name FROM asset WHERE is_current = 1 AND record_level = 'ASSET' AND asset_class = ANY(%s) AND upper(asset_key) = upper(%s)", [SCOPE_CLASSES, str(key or "").strip()])


def record_manual(asset_key, raw, by):
    facts = _clean_facts(raw)
    asset = _in_scope_asset(asset_key)
    if not asset:
        raise ObsError("Software is tracked for desktops, workstations and servers only.", 404)
    return record(asset["asset_key"], "MANUAL", facts, by)


def template_bytes():
    import io
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = "Software"
    ws.append(template_headers())
    dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    ws.add_data_validation(dv)
    for i, (h, f) in enumerate(TEMPLATE_COLUMNS, start=2):
        letter = ws.cell(1, i).column_letter
        ws.column_dimensions[letter].width = 22
        if f in BOOLS:
            dv.add(f"{letter}2:{letter}1000")
    ws.column_dimensions["A"].width = 20
    ws.freeze_panes = "B2"
    n = wb.create_sheet("How to fill")
    for line in ("One row per machine. Asset (CI) must be a desktop, workstation or server in the register.", "Leave a cell blank when you did not check it - blank is not the same as No.",
                 "Yes/No columns accept Yes or No only.", "Every row is saved as an observation 'Entered by hand' with the date and your name; the date cannot be set or changed.",
                 "If any row has a problem, nothing from the file is saved: fix the rows listed and upload again."):
        n.append([line])
    n.column_dimensions["A"].width = 110
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def ingest_workbook(data, by, engineer_key=None):
    """Read a filled template. All or nothing: observations cannot be deleted, so one bad row means nothing is saved.
    `engineer_key`: an engineer may only fill in the machines assigned to them (administrators and team leads pass None)."""
    import io
    from openpyxl import load_workbook
    if data[:2] != b"PK":
        raise ObsError("That is not an Excel workbook (.xlsx).")
    try:
        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    except Exception as e:      # noqa: BLE001
        raise ObsError("That Excel file could not be read.") from e
    if not rows or str(rows[0][0] or "").strip() != "Asset (CI)":
        raise ObsError("This is not the Software upload template (the first column must be 'Asset (CI)'). Download the template again.")
    field_of = {h: f for h, f in TEMPLATE_COLUMNS}
    cols = {i: field_of[str(c).strip()] for i, c in enumerate(rows[0]) if str(c or "").strip() in field_of}
    todo, errors = [], []
    for n, r in enumerate(rows[1:], start=2):
        ci = str(r[0] or "").strip()
        if not ci and not any(str(c or "").strip() for c in r):
            continue
        try:
            asset = _in_scope_asset(ci)
            if not asset:
                raise ObsError(f"{ci or '(blank)'} is not a desktop, workstation or server in the register.")
            if engineer_key and asset["engineer_name"] != engineer_key:
                raise ObsError(f"{ci} is not assigned to you.")
            todo.append((asset["asset_key"], _clean_facts({f: r[i] for i, f in cols.items() if i < len(r)})))
        except ObsError as e:
            errors.append(f"Row {n}: {e}")
    if errors or not todo:
        return {"recorded": 0, "errors": errors or ["The file has no rows to save."]}
    for key, facts in todo:
        record(key, "MANUAL", facts, by)
    return {"recorded": len(todo), "errors": []}


# ---------------------------------------------------------------- the software report (Excel): which software and version is on which CI
_SRC = {"BIGFIX_REPORT": "BigFix report", "TREND_REPORT": "Trend Micro report", "SCRIPT": "Check script", "MANUAL": "Entered by hand"}
_MACHINE_COLS = [("ci", "Asset (CI)"), ("class", "Class"), ("engineer", "Engineer"), ("user", "User"), ("source", "Source"), ("inspected", "Inspected"), ("data_as_of", "Data as of"),
                 ("os_edition", "Windows"), ("os_build", "Windows build"), ("os_feature_version", "Feature version"), ("trend_installed", "Trend Micro installed"), ("trend_version", "Trend Micro version"),
                 ("trend_status", "Trend Micro status / name"), ("zenworks_installed", "ZENworks installed"), ("zenworks_version", "ZENworks version"), ("zenworks_status", "ZENworks status / name"),
                 ("bigfix_installed", "BigFix installed"), ("bigfix_version", "BigFix version"), ("vpn_name", "VPN"), ("vpn_version", "VPN version"),
                 ("sap_gui_version", "SAP GUI version"), ("o365_account", "Microsoft 365 account"), ("o365_version", "Office version"), ("browser_version", "Edge version"),
                 ("chrome_version", "Chrome version"), ("firefox_version", "Firefox version"), ("last_user", "Last logged-on user"), ("os_update_title", "Last Windows update"),
                 ("os_update_at", "Last Windows update on"), ("defender_signature", "Defender signature"), ("remarks", "Remarks")]


def _cell(v):
    if isinstance(v, bool):
        return "Yes" if v else "No"
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d %H:%M")
    return _CTRL.sub("", v) if isinstance(v, str) else v


def report_tables(engineer_key=None):
    """Three tables for an Excel report: the newest observation of every source per machine (machines never checked still get a row), the installed-software
    list of each machine's latest script run, and a fleet view of program + version with how many machines have it. `engineer_key` limits it to one engineer's machines."""
    where, args = ("AND engineer_name = %s", [engineer_key]) if engineer_key else ("", [])
    assets = db.query(f"SELECT asset_key, ci_no, asset_class, engineer_name, user_name FROM asset WHERE is_current = 1 AND record_level = 'ASSET' AND asset_class = ANY(%s) {where} ORDER BY ci_no", [SCOPE_CLASSES, *args])
    keys = [a["asset_key"] for a in assets]
    obs = db.query("SELECT DISTINCT ON (asset_key, source) * FROM portal_sw_obs WHERE asset_key = ANY(%s) ORDER BY asset_key, source, obs_id DESC", [keys])
    by_asset = {}
    for o in obs:
        by_asset.setdefault(o["asset_key"], []).append(o)
    machines = []
    for a in assets:
        base = {"ci": a["ci_no"], "class": a["asset_class"], "engineer": a["engineer_name"] or "", "user": a["user_name"] or ""}
        rows = by_asset.get(a["asset_key"]) or [{"source": None}]
        for o in sorted(rows, key=lambda x: SOURCES.index(x["source"]) if x["source"] else 9):
            r = {**base, **{k: _cell(o.get(k)) for k, _ in _MACHINE_COLS[7:]}}
            r["source"] = _SRC.get(o["source"], "Not checked")
            r["inspected"], r["data_as_of"] = _cell(o.get("inspected_at")) or "", _cell(o.get("data_as_of")) or ""
            machines.append(r)
    ci_of = {a["asset_key"]: a["ci_no"] for a in assets}
    apps = db.query("""WITH last AS (SELECT DISTINCT ON (asset_key) asset_key, inspected_at, COALESCE(apps_obs_id, obs_id) AS holder FROM portal_sw_obs
                                     WHERE apps_hash IS NOT NULL AND asset_key = ANY(%s) ORDER BY asset_key, obs_id DESC)
                       SELECT l.asset_key, l.inspected_at, a.name, a.version, a.publisher, a.hive FROM last l JOIN portal_sw_obs_app a ON a.obs_id = l.holder
                       ORDER BY l.asset_key, lower(a.name), a.hive""", [keys])
    soft = [{"ci": ci_of[x["asset_key"]], "program": x["name"], "version": x["version"] or "", "publisher": x["publisher"] or "", "where": x["hive"], "checked": _cell(x["inspected_at"])} for x in apps]
    fleet = {}
    for r in soft:
        fleet.setdefault((r["program"], r["version"], r["publisher"]), set()).add(r["ci"])
    summary = [{"program": p, "version": v, "publisher": pub, "machines": len(c)} for (p, v, pub), c in sorted(fleet.items(), key=lambda kv: (kv[0][0].lower(), kv[0][1]))]
    return [{"title": "Machines", "columns": _MACHINE_COLS, "rows": machines},
            {"title": "Installed software", "columns": [("ci", "Asset (CI)"), ("program", "Program"), ("version", "Version"), ("publisher", "Publisher"), ("where", "Where registered"), ("checked", "Checked")], "rows": soft},
            {"title": "Software and versions", "columns": [("program", "Program"), ("version", "Version"), ("publisher", "Publisher"), ("machines", "Machines")], "rows": summary}]


def report_bytes(engineer_key, by):
    from . import export
    t = report_tables(engineer_key)
    n_checked = len({r["ci"] for r in t[0]["rows"] if r["source"] != "Not checked"})
    n_all = len({r["ci"] for r in t[0]["rows"]})
    meta = "Latest observation per source for each desktop, workstation and server" + (f" assigned to {engineer_key}" if engineer_key else "") + ". Installed software comes from each machine's latest check-script run."
    return export.xlsx_bytes(t, "Software and versions by Asset (CI)", [("Machines in scope", n_all), ("Machines with at least one observation", n_checked), ("Machines never checked", n_all - n_checked),
                                                                  ("Installed-software rows", len(t[1]["rows"]))], meta, generated_by=by)
